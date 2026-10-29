#!/usr/bin/env python3
"""
Negative Screening System — step-by-step build scaffold
=======================================================

Weekly negative / adverse media screening for a target company,
its approved aliases and approved personnel. Prototype by Yap Xiang Yang
(independent student developer; not an employee of the client).

This file is BOTH the program and the build plan. It is written so that Claude Code
(or a human) can implement it one step at a time, proving each step with the
acceptance tests that are already written below before moving on.

Design documents (source of truth — read the referenced section before each step):
    docs/PROJECT_STATE.md               status, decisions D-01..D-18, open items
    docs/SOURCE_ACCESS_BLUEPRINT.md     how each source is accessed / never accessed
    docs/DATA_COLLECTION_ANALYSIS.md    dates, entities, keywords, rules R1-R7,
                                        URL verification, deduplication
    docs/FINAL_ARCHITECTURE.md          pipeline, calendar, CLI, DB schema, report, email
    docs/OPERATIONS_GUIDE.md            setup, testing, Task Scheduler

-------------------------------------------------------------------------------
HOW TO BUILD THIS FILE (instructions for Claude Code)
-------------------------------------------------------------------------------
1.  Check progress:            python negative_screening.py --status
2.  Take the FIRST step whose STEP_STATUS is not "DONE". Never skip ahead.
3.  Read the step's banner below and the doc sections it cites.
4.  Implement ONLY the functions in that step's section. Replace each
    `raise NotImplementedError(...)` with real code. Do not change the function
    signatures or dataclasses without updating every caller and the docs.
5.  Run that step's acceptance tests:
                               python negative_screening.py --self-test S06
    Every test must PASS (not PENDING, not FAIL). Tests marked "needs <package>"
    need `pip install -r requirements.txt` first.
6.  You MAY add more tests for the step (decorate with @selftest("Sxx")). You must
    NOT weaken or delete an existing acceptance test. If a test contradicts the
    docs, stop and ask the user which one is right.
7.  Re-run all earlier steps' tests:  python negative_screening.py --self-test all
8.  Set STEP_STATUS["Sxx"] = "DONE" and tick the phase row in docs/PROJECT_STATE.md.
9.  Stop and report to the user: what was built, test results, anything unresolved.
    Wait for approval before the next step.

Hard rules — apply to every step (docs/PROJECT_STATE.md D-02, FINAL_ARCHITECTURE):
    * No Google Search HTML scraping; no direct scraping of Facebook/Instagram/LinkedIn.
    * No CAPTCHA solving, proxy rotation, fingerprint spoofing, evasive browser
      automation, personal login-cookie scraping, or bypassing access controls.
    * No paid services or subscriptions without the user's approval.
    * No secrets in code, config JSON, logs or reports (environment variables only).
    * Never fabricate or guess URLs, dates or IDs. Unknown stays UNKNOWN.
    * No NLP, sentiment analysis or machine-learning classification of any kind
      (project owner's decision, 29 Sep 2026). Screening is approved name + adverse
      keyword rules only. Do not add it back unless the owner explicitly asks.
    * Keyword matches are screening aids, never proof.
    * A failed source is reported as FAILED, never as "no results".
    * Nothing outside the reporting period goes into the main findings.
    * Never claim complete Internet coverage.
    * Third-party packages are imported INSIDE the step functions that need them,
      so --status and --self-test keep working with the standard library alone.
    * Keep this a single file for now. Section banners map to the modules in
      FINAL_ARCHITECTURE §5 so the file can be split later without redesign.

Step plan (development order; the runtime pipeline order is in run_pipeline()):
    Phase 1  Core framework      S01 models · S02 calendar · S03 config · S04 logging
                                 S05 database · S06 URL normalisation · S07 dedup
    Phase 2  Discovery           S08 collector framework · S09 alert queries
                                 S10 Google Alerts (Facebook / Instagram / LinkedIn /
                                 public web via Google's index) · S11 Company website
                                 S12 SC Malaysia
    Phase 4* Verification        S13 platform detection & URL verification
    Phase 3  Screening           S14 dates · S15 names (entities) · S16 keywords
                                 S17 flagging rules & report sections (rule-based)
    Phase 5  Reporting           S18 Excel report
    Phase 6  Email               S19 email decision & sending
    Phase 7  Scheduling          S20 run lock, Task Scheduler wrapper, end-to-end run
    (*) Platform detection is built before screening because the platform / content
        class (social, public web, regulatory, company) is used by the rules.
    Optional social-media API connectors come only after S20 and client authorisation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import re
import sqlite3
import sys
import tempfile
import traceback
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from enum import Enum
from html.parser import HTMLParser


class StrEnum(str, Enum):
    """String-valued enum that also works on Python 3.10 (enum.StrEnum is 3.11+)."""

    def __str__(self) -> str:
        return str(self.value)

    def __format__(self, spec: str) -> str:
        return format(str(self.value), spec)
from pathlib import Path
from typing import Any, Callable, Iterable, Protocol

APP_VERSION = "0.1.0-scaffold"

# -----------------------------------------------------------------------------
# Build progress. Claude Code updates these values as steps are completed.
# Allowed values: "TODO", "IN PROGRESS", "DONE".
# -----------------------------------------------------------------------------
STEP_STATUS: dict[str, str] = {
    "S01": "DONE",   # data models
    "S02": "DONE",   # reporting-period calendar
    "S03": "DONE",   # configuration files + loading/validation
    "S04": "DONE",   # logging with secret redaction
    "S05": "DONE",   # SQLite schema + upserts
    "S06": "DONE",   # URL decoding / canonicalisation / keys
    "S07": "DONE",   # deduplication
    "S08": "DONE",   # collector framework (isolation, retries, HTTP session)
    "S09": "DONE",   # Google Alerts query generator
    "S10": "DONE",   # Google Alerts collector (IMAP, RSS fallback, public-web page check)
    "S11": "DONE",   # Company website collector
    "S12": "DONE",   # SC Malaysia collector
    "S13": "DONE",   # platform detection + URL classification + verification
    "S14": "DONE",   # date resolution + period filtering
    "S15": "DONE",   # entity matcher
    "S16": "DONE",   # keyword scanner (+ keyword-to-name proximity)
    "S17": "DONE",   # flagging rules R1-R7 + report sections (rule-based, no NLP)
    "S18": "DONE",   # Excel report
    "S19": "DONE",   # email decision + sending
    "S20": "DONE",   # run lock, run_weekly.bat, end-to-end
}

STEP_TITLES: dict[str, str] = {
    "S01": "Data models", "S02": "Reporting-period calendar", "S03": "Configuration",
    "S04": "Logging + secret redaction", "S05": "SQLite database",
    "S06": "URL normalisation", "S07": "Deduplication", "S08": "Collector framework",
    "S09": "Google Alerts query generator", "S10": "Google Alerts collector",
    "S11": "Company website collector", "S12": "SC Malaysia collector",
    "S13": "Platform + URL verification", "S14": "Date resolution + period filter",
    "S15": "Entity matcher", "S16": "Keyword scanner", "S17": "Flagging rules (no NLP)",
    "S18": "Excel report", "S19": "Email", "S20": "Scheduling + end-to-end",
}


def resolve_project_root(frozen: bool, executable: str, source_file: str) -> Path:
    """The folder holding config/, data/, logs/ and reports/. Script: this file's folder.
    Packaged exe (PyInstaller): the exe's folder; an exe still in <project>\\dist\\ (fresh
    build, not yet copied up) uses the project folder above it."""
    if not frozen:
        return Path(source_file).resolve().parent
    exe_dir = Path(executable).resolve().parent
    if not (exe_dir / "config").is_dir() and (exe_dir.parent / "config").is_dir():
        return exe_dir.parent
    return exe_dir


IS_FROZEN = bool(getattr(sys, "frozen", False))
PROJECT_ROOT = resolve_project_root(IS_FROZEN, sys.executable, __file__)
EXE_NAME = "NegativeScreening.exe"


# =============================================================================
# S01 — DATA MODELS                                    (→ core/models.py)
# Docs: DATA_COLLECTION_ANALYSIS §1, FINAL_ARCHITECTURE §6
# Status: DONE. Change only with a matching docs update.
# =============================================================================

class Mode(StrEnum):
    AUTOMATIC = "AUTOMATIC"                 # started by Task Scheduler (--scheduled)
    MANUAL = "MANUAL"                       # started by a person
    MANUAL_TEST_DATE = "MANUAL_TEST_DATE"   # --date override (testing only)


class SourceStatus(StrEnum):
    SUCCESS = "SUCCESS"
    FAILED = "FAILED"
    SKIPPED = "SKIPPED"
    # Report rows only (derive_source_health): the collector ran without error but had nothing to
    # read — no Google Alerts emails arrived — so Google's index (Facebook / Instagram / LinkedIn /
    # public websites) was not screened. Never shown as SUCCESS: a 0 there is not an all-clear (D-25).
    NOT_CHECKED = "NOT CHECKED"


class Platform(StrEnum):
    FACEBOOK = "Facebook"
    INSTAGRAM = "Instagram"
    LINKEDIN = "LinkedIn"
    SC_MALAYSIA = "SC Malaysia"
    COMPANY_WEBSITE = "Company website"
    PUBLIC_WEB = "Public website"       # any URL whose platform cannot be identified from its hostname


class ContentClass(StrEnum):
    SOCIAL = "social"
    PUBLIC_WEB = "public_web"
    REGULATORY = "regulatory"
    COMPANY = "company"


class UrlType(StrEnum):
    SPECIFIC_CONTENT = "SPECIFIC_CONTENT"
    CONTAINER = "CONTAINER"
    NON_EVIDENCE = "NON_EVIDENCE"
    UNKNOWN = "UNKNOWN"


class Verification(StrEnum):
    VERIFIED = "VERIFIED"
    PARTIAL = "PARTIAL"
    REJECTED = "REJECTED"


class DateStatus(StrEnum):
    """How reliable the event / content date is. Never describes when the system found the item."""
    VERIFIED = "VERIFIED"           # stated by the source itself (at the precision it states)
    DERIVED = "DERIVED"
    APPROXIMATE = "APPROXIMATE"     # from a search-result snippet
    FIRST_SEEN = "FIRST_SEEN"       # source gives no event date: listed in the week the system first found it
    UNKNOWN = "UNKNOWN"


class DatePrecision(StrEnum):
    """How precisely the source states the event / content date. A year or month is never
    turned into an exact day."""
    DAY = "DAY"                     # a full date (with a time where the source gives one)
    MONTH = "MONTH"
    YEAR = "YEAR"


class HistoricalStatus(StrEnum):
    """The event / content date compared with the reporting week (S18 'Historical Status')."""
    CURRENT = "CURRENT"             # dated in the reporting week (or its year / month when only that is known)
    HISTORICAL = "HISTORICAL"       # dated before the reporting week: an older record, not a new incident
    DATE_UNKNOWN = "DATE UNKNOWN"   # no reliable event / content date


class Priority(StrEnum):
    CRITICAL = "CRITICAL"
    HIGH = "HIGH"
    MEDIUM = "MEDIUM"
    REVIEW = "REVIEW"
    LOW = "LOW"


PRIORITY_RANK = {Priority.CRITICAL: 5, Priority.HIGH: 4, Priority.MEDIUM: 3,
                 Priority.REVIEW: 2, Priority.LOW: 1}


class ReviewStatus(StrEnum):
    PENDING = "PENDING"
    MANUAL_REVIEW = "MANUAL REVIEW"
    NEEDS_REVIEW = "NEEDS REVIEW"
    CONFIRMED_ADVERSE = "CONFIRMED ADVERSE"   # set by staff, never by the program
    FALSE_POSITIVE = "FALSE POSITIVE"         # set by staff, never by the program


class Section(StrEnum):
    MAIN = "MAIN"                         # Findings sheet
    DATE_UNKNOWN = "DATE_UNKNOWN"         # "Date Unknown — Manual Review" sheet
    UNVERIFIED_URL = "UNVERIFIED_URL"     # "Unverified URLs — Manual Lookup" sheet
    OUT_OF_PERIOD = "OUT_OF_PERIOD"       # dated before the period (optional audit sheet)
    EXCLUDED_FUTURE = "EXCLUDED_FUTURE"   # dated after the period (current week) — excluded
    NOT_REPORTED = "NOT_REPORTED"         # not relevant / not adverse


@dataclass(frozen=True)
class ReportingPeriod:
    reference_date: date
    period_start: date            # Monday
    period_end: date              # Sunday (inclusive, for display)
    start_dt: datetime            # Monday 00:00:00 MYT
    end_dt_exclusive: datetime    # following Monday 00:00:00 MYT

    def label(self) -> str:
        return f"{self.period_start:%d %b %Y} – {self.period_end:%d %b %Y}"

    def contains(self, moment: datetime) -> bool:
        """True if an aware datetime falls inside [start_dt, end_dt_exclusive)."""
        if moment.tzinfo is None:
            raise ValueError("naive datetime: convert to an aware datetime first")
        return self.start_dt <= moment < self.end_dt_exclusive


@dataclass
class RawItem:
    """One result as delivered by a collector. Unknown fields stay None — never guessed."""
    source: str                                   # google_alerts | website | sc_malaysia
    discovery_method: str                         # e.g. "Google Alerts (search index)"
    discovered_url: str                           # after decoding a Google redirect
    title: str = ""
    snippet: str = ""
    discovery_query: str | None = None
    final_url: str | None = None                  # public web / Company website only; never for social
    content_text: str | None = None
    publisher: str | None = None
    alert_received_at: datetime | None = None     # Google Alerts email time — NOT the content date
    source_published_at: datetime | None = None   # date stated by the source itself
    language: str | None = None
    source_message_id: str | None = None          # email Message-ID / feed entry id
    extra: dict = field(default_factory=dict)     # e.g. {"sc_alert_list": True}


@dataclass
class Match:
    entity_id: str
    matched_text: str
    keyword: str | None = None
    category: str | None = None
    priority: str | None = None
    ambiguous: bool = False
    proximity_words: int | None = None     # keyword: word distance to the nearest name mention
    position: int | None = None            # entity: index of the match's first word (tokenize_words)
    look_alike: bool = False               # entity: ambiguous term confirmed only via a web address (R9)


@dataclass
class Candidate:
    """A RawItem plus everything the processing steps add. Each step fills its own fields."""
    raw: RawItem
    # S06
    canonical_url: str = ""
    url_key: str = ""
    # S07
    discovery_queries: set[str] = field(default_factory=set)
    # S13
    platform: Platform | None = None
    content_class: ContentClass | None = None
    content_type: str = "Unknown"
    url_type: UrlType = UrlType.UNKNOWN
    verification: Verification | None = None
    verification_basis: str = ""
    # S14
    content_date: datetime | None = None          # event / content date (start of its precision)
    date_precision: DatePrecision | None = None   # None when content_date is None
    date_status: DateStatus = DateStatus.UNKNOWN
    date_basis: str = ""
    # S15 / S16
    entity_matches: list[Match] = field(default_factory=list)
    keyword_matches: list[Match] = field(default_factory=list)
    # S17
    priority: Priority | None = None
    rules_applied: list[str] = field(default_factory=list)
    review_status: ReviewStatus = ReviewStatus.PENDING
    section: Section | None = None
    notes: list[str] = field(default_factory=list)


@dataclass
class SourceResult:
    source: str
    status: SourceStatus
    items: list[RawItem] = field(default_factory=list)
    error: str | None = None                       # short, secret-free
    warnings: list[str] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)


@dataclass
class Config:
    settings: dict
    entities: dict
    keywords: dict
    email: dict
    config_dir: Path


@dataclass
class RunContext:
    run_id: str
    mode: Mode
    period: ReportingPeriod
    started_at: datetime
    dry_run: bool
    config: Config
    logger: logging.Logger
    db_path: Path
    cli_no_email: bool = False
    cli_send_email: bool = False
    cli_resend_email: bool = False
    only_sources: set[str] | None = None
    validation: bool = False            # staging database + reports/validation/, never emailed
    email_error: str | None = None      # set by send_report_email() on failure
    project_root: Path | None = None    # base for relative paths; None = PROJECT_ROOT


class Collector(Protocol):
    name: str

    def collect(self, ctx: RunContext) -> SourceResult: ...


# =============================================================================
# S02 — REPORTING-PERIOD CALENDAR                      (→ core/calendar.py)
# Docs: FINAL_ARCHITECTURE §3, PROJECT_STATE D-08/D-09
# Status: DONE. This is the ONLY place that calculates reporting periods.
# Never add "exit unless Monday" logic: manual runs must work on any day.
# =============================================================================

def malaysia_tz():
    """Asia/Kuala_Lumpur. On Windows `zoneinfo` needs the `tzdata` package;
    falls back to a fixed UTC+8 offset (Malaysia has no daylight saving)."""
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo("Asia/Kuala_Lumpur")
    except Exception:  # ZoneInfoNotFoundError when tzdata is missing
        return timezone(timedelta(hours=8), name="MYT")


MYT = malaysia_tz()


def today_in_malaysia() -> date:
    return datetime.now(MYT).date()


def now_in_malaysia() -> datetime:
    return datetime.now(MYT)


def calculate_previous_completed_week(reference_date: date) -> ReportingPeriod:
    """Return the immediately preceding COMPLETE Monday–Sunday week (MYT).

    Works for every weekday and never returns the current, incomplete week.
    Mon 21 Sep 2026 -> 14–20 Sep; Fri 25 Sep -> 14–20 Sep; Mon 28 Sep -> 21–27 Sep.
    """
    this_monday = reference_date - timedelta(days=reference_date.weekday())
    period_start = this_monday - timedelta(days=7)
    period_end = this_monday - timedelta(days=1)
    return ReportingPeriod(
        reference_date=reference_date,
        period_start=period_start,
        period_end=period_end,
        start_dt=datetime.combine(period_start, time.min, tzinfo=MYT),
        end_dt_exclusive=datetime.combine(this_monday, time.min, tzinfo=MYT),
    )


# =============================================================================
# S03 — CONFIGURATION                                   (→ core/config.py)
# Docs: DATA_COLLECTION_ANALYSIS §3–4, FINAL_ARCHITECTURE §12
# Goal: `--init-config` writes the four JSON files below (only if missing);
#       `load_config()` reads and validates them. Runtime logic always reads the
#       JSON files — the DEFAULT_* dicts are only templates for first creation.
# Rules: only approved entities are used; secrets are never in these files;
#        a validation error names the file and the field.
# =============================================================================

SAME_ITEM_WINDOW = 1_000_000                 # context_window_words meaning "anywhere in the same item"
_PERSON_ZH_CONTEXT = ["Acme Asset Management", "艾克米投资管理", "艾克米投資管理", "Tan Ah Kow", "Ah Kow Tan"]

# Placeholder targets (the self-tests rely on them). Put the real targets in config/entities.json.
DEFAULT_ENTITIES: dict = {
    "version": "2026-10-05",
    "note": ("Only client-approved names. Add new aliases/personnel only after approval (PROJECT_STATE O-1). "
             "Exact full names only (owner, 5 Oct 2026, D-27): no single word of a name is screened. "
             "The only ambiguous terms are the founder's Chinese name (a common name): they count only when "
             "the same item also names the company or the founder (D-29)."),
    "entities": [
        {
            "id": "acme_company",
            "type": "company",
            "display_name": "Acme Asset Management Sdn Bhd",
            "names": ["Acme Asset Management Sdn Bhd", "Acme Asset Management"],
            "variants": ["Acme Asset Management Sdn. Bhd."],
            "aliases": ["艾克米投资管理", "艾克米投資管理"],   # client-approved Chinese name (Simplified, Traditional)
            "ambiguous_terms": [],                # D-27: the full company name only, never one word of it
            "alert_names": ["Acme Asset Management Sdn Bhd", "Acme Asset Management",
                            "艾克米投资管理", "艾克米投資管理"],
            "approved": True,
        },
        {
            "id": "tan_ah_kow",
            "type": "person",
            "display_name": "Tan Ah Kow",
            "role": "Founder",
            "names": ["Tan Ah Kow"],
            "variants": ["Ah Kow Tan"],
            "aliases": [],
            # D-29: the founder's Chinese name is a common name, so it counts only when the same item also
            # names the company or the founder (context window = the whole item). Never used as an alert query.
            "ambiguous_terms": [{"term": form, "requires_any_context": _PERSON_ZH_CONTEXT,
                                 "context_window_words": SAME_ITEM_WINDOW} for form in ("陈亚九", "陳亞九")],
            "alert_names": ["Tan Ah Kow"],
            "approved": True,
        },
    ],
    # Client-approved links only; never inferred from text (S15 approved_relationships).
    "relationships": [
        {"subject": "tan_ah_kow", "type": "associated_with", "label": "Founder of",
         "object": "acme_company", "approved": True},
    ],
}


# Chinese forms of the English keywords (owner, 5 Oct 2026, D-28): Simplified first, then Traditional
# (incl. common Taiwan / Hong Kong wording). Stored as each keyword's "variants", so a Chinese form has
# the same category, priority and ambiguity as its English term. Chinese has no spaces: a form matches
# as an exact character sequence and proximity is counted in characters (tokenize_words).
ZH_KEYWORD_FORMS: dict[str, list[str]] = {
    "money laundering": ["洗钱", "洗錢"],
    "fraud": ["诈骗", "欺诈", "詐騙", "欺詐"],
    "embezzlement": ["侵吞", "挪用公款"],
    "tax evasion": ["逃税", "逃稅"],
    "bribery": ["贿赂", "受贿", "行贿", "賄賂", "受賄", "行賄"],
    "corruption": ["贪污", "腐败", "貪污", "貪汙", "腐敗"],
    "kickback": ["回扣"],
    "insider trading": ["内幕交易", "內幕交易"],
    "market manipulation": ["操纵市场", "操纵股价", "操縱市場", "操縱股價"],
    "ponzi": ["庞氏骗局", "龐氏騙局"],
    "pyramid scheme": ["金字塔计划", "传销", "金字塔計劃", "傳銷"],
    "forgery": ["伪造", "偽造"],
    "counterfeit": ["假冒"],
    "shell company": ["空壳公司", "空殼公司"],
    "misappropriation": ["挪用", "盗用", "盜用"],
    "sanctions": ["制裁"],
    "enforcement action": ["执法行动", "執法行動"],
    "regulatory penalty": ["监管处罚", "監管處罰"],
    "prosecution": ["起诉", "起訴"],
    "subpoena": ["传票", "傳票"],
    "blacklisted": ["黑名单", "黑名單"],
    "asset freeze": ["冻结资产", "凍結資產"],
    "cease and desist": ["停止令"],
    "seizure": ["查封", "扣押"],
    "fined": ["罚款", "罰款"],
    "banned": ["被禁"],
    "disqualified": ["取消资格", "取消資格"],
    "arrested": ["被捕"],
    "charged": ["被控"],
    "convicted": ["定罪"],
    "guilty": ["有罪"],
    "sentenced": ["判刑", "判处", "判處"],
    "bankrupt": ["破产", "破產"],
    "insolvent": ["无力偿债", "無力償債"],
    "data breach": ["数据泄露", "數據泄露", "數據洩露", "資料外洩"],
    "cyber attack": ["网络攻击", "網絡攻擊", "網路攻擊"],
    "cybercrime": ["网络犯罪", "網絡犯罪", "網路犯罪"],
    "hacking": ["黑客", "駭客"],
    "ransomware": ["勒索软件", "勒索軟件", "勒索軟體"],
    "identity theft": ["身份盗窃", "身份盜竊"],
    "phishing": ["网络钓鱼", "網絡釣魚", "網路釣魚"],
}


def _kw(terms: Iterable[str], ambiguous: Iterable[str] = ()) -> list[dict]:
    amb = set(ambiguous)
    return [{"term": t, "ambiguous": t in amb,
             **({"variants": list(ZH_KEYWORD_FORMS[t])} if t in ZH_KEYWORD_FORMS else {})} for t in terms]


DEFAULT_KEYWORDS: dict = {
    "version": "2026-10-05",
    "note": ("Client-supplied dictionary, spelling corrected ('rrested' -> 'arrested'). "
             "Only Sanctions/Regulatory priority was set by the client; others are proposed "
             "defaults pending confirmation (PROJECT_STATE O-3). Chinese forms (Simplified and "
             "Traditional) are each keyword's 'variants' (owner, 5 Oct 2026, D-28)."),
    "categories": [
        {"id": "financial_crime", "name": "Financial Crime", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["money laundering", "aml", "ctf", "fraud", "embezzlement", "tax evasion"],
                         ambiguous=["aml", "ctf"])},
        {"id": "bribery_corruption", "name": "Bribery / Corruption", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["bribery", "corruption", "kickback"])},
        {"id": "market_misconduct", "name": "Market Misconduct", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["insider trading", "market manipulation"])},
        {"id": "investment_schemes", "name": "Investment Schemes", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["ponzi", "pyramid scheme"])},
        {"id": "forgery_counterfeit", "name": "Forgery / Counterfeit", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["forgery", "counterfeit"])},
        {"id": "corporate_misconduct", "name": "Corporate / Financial Misconduct", "priority": "HIGH",
         "client_confirmed": False,
         "keywords": _kw(["structuring", "shell company", "misappropriation"], ambiguous=["structuring"])},
        {"id": "sanctions_regulatory", "name": "Sanctions / Regulatory Enforcement", "priority": "CRITICAL",
         "client_confirmed": True,
         "keywords": _kw(["sanctions", "ofac", "enforcement action", "regulatory penalty", "indictment",
                          "prosecution", "subpoena", "blacklisted", "debarred", "asset freeze",
                          "cease and desist", "seizure", "fined", "banned", "disqualified", "arrested",
                          "charged", "convicted", "guilty", "sentenced", "fugitive", "wanted by",
                          "extradition"],
                         ambiguous=["sanctions", "seizure", "fined", "banned", "charged", "guilty",
                                    "wanted by"])},
        {"id": "financial_distress", "name": "Financial Distress", "priority": "HIGH", "client_confirmed": False,
         "keywords": _kw(["bankrupt", "insolvent", "receivership"])},
        {"id": "cybersecurity", "name": "Cybersecurity", "priority": "MEDIUM", "client_confirmed": False,
         "keywords": _kw(["data breach", "cyber attack", "cybercrime", "hacking", "ransomware",
                          "identity theft", "phishing"], ambiguous=["hacking", "phishing"])},
    ],
}

DEFAULT_SETTINGS: dict = {
    "timezone": "Asia/Kuala_Lumpur",
    "paths": {"data_dir": "data", "reports_dir": "reports", "logs_dir": "logs",
              "database": "data/screening.sqlite", "dry_run_database": "data/screening_dryrun.sqlite",
              "validation_database": "data/screening_validation.sqlite"},
    "sources": {
        "google_alerts": {"enabled": True, "mode": "imap", "rss_feeds": [],
                          "fetch_public_web_pages": True,
                          "note": ("mode: imap (default) or rss (fallback). Facebook / Instagram / "
                                   "LinkedIn result pages are never fetched; other public web result "
                                   "pages may be fetched once, politely, to confirm names, keywords "
                                   "and dates.")},
        "website": {"enabled": True, "base_url": "https://acme.example.com/",
                    "feed_url": "https://acme.example.com/feed/",
                    "comments_feed_url": "https://acme.example.com/comments/feed/",
                    "max_feed_pages": 5, "api_endpoint": None, "listing_pages": [],
                    "note": ("Site inspected 29 Sep 2026 (WordPress): the RSS feed is used "
                             "(SOURCE_ACCESS_BLUEPRINT §4, option 1); robots.txt allows it.")},
        "sc_malaysia": {"enabled": True,
                        "investor_alert_list_url": "https://www.sc.com.my/investor-alert-list",
                        "announcement_urls": ["https://www.sc.com.my/resources/media/media-release"],
                        "fetch_announcement_pages": True,
                        "note": ("Pages confirmed 29 Sep 2026 (robots.txt allows all). Both pages embed "
                                 "their full list as data, which is read instead of the HTML layout.")},
    },
    "http": {"timeout_seconds": 30, "max_retries": 3, "backoff_seconds": 2,
             "delay_between_requests_seconds": 2,
             "user_agent": "NegativeScreeningPrototype/0.1 (+contact: configured operator)"},
    "screening": {"keyword_name_window_words": 30,
                  "note": ("Rule-based only: approved name + adverse keyword. NLP is excluded from "
                           "this project by the owner's decision (29 Sep 2026).")},
    "report": {"include_out_of_period_sheet": False},
    "email_defaults": {"send_on_automatic": True, "send_on_manual": True, "manual_email_cooldown_minutes": 10},
    "run": {"max_run_minutes": 120},
}

DEFAULT_EMAIL: dict = {
    "note": ("No passwords here. They are saved in the app's Setup (Windows Credential Manager) or set "
             "as NS_SMTP_PASSWORD / NS_IMAP_PASSWORD environment variables."),
    "smtp": {"host": "smtp.gmail.com", "port": 587, "starttls": True,
             "username": "operator@example.com"},
    "sender": "operator@example.com",
    "recipients": ["operator@example.com"],
    "subject_template": "[Negative Screening] {period} — {critical} critical, {review} for review{failed_suffix}",
    "alerts_mailbox": {"imap_host": "imap.gmail.com", "imap_port": 993,
                       "username": "operator@example.com", "label": "NS-Alerts",
                       "sender_filter": "googlealerts-noreply@google.com"},
}

CONFIG_FILES = {"settings.json": DEFAULT_SETTINGS, "entities.json": DEFAULT_ENTITIES,
                "adverse_keywords.json": DEFAULT_KEYWORDS, "email.json": DEFAULT_EMAIL}

SECRET_ENV_VARS = ("NS_SMTP_PASSWORD", "NS_IMAP_PASSWORD")
NO_ALERT_EMAILS = ("no Google Alerts emails arrived in the {label!r} label since {since}: nothing from Google's "
                   "index could be read. Create the alerts (app: Setup → Google Alerts) delivered to {mailbox}, "
                   "and a Gmail filter that applies the label")
IMAP_NOT_CONFIGURED = ("Google Alerts email access is not configured: no mailbox password "
                       "(NS_IMAP_PASSWORD) in the app's Setup or the environment")
SMTP_NOT_CONFIGURED = ("email sending is not configured: no sending-account password "
                       "(NS_SMTP_PASSWORD) in the app's Setup or the environment")
SECRET_LABELS = {"NS_IMAP_PASSWORD": "Google Alerts mailbox (IMAP) app password",
                 "NS_SMTP_PASSWORD": "report sending account (SMTP) app password"}
CREDENTIAL_TARGET_PREFIX = "NegativeScreening/"


class WindowsCredentialStore:
    """Secrets in Windows Credential Manager (generic credentials of the current Windows
    user, protected by Windows). Staff save them from the app's Setup dialog, so no
    command line is needed; they never go into source, JSON, logs, reports or Git."""

    def _api(self):
        import ctypes
        from ctypes import wintypes

        class CREDENTIALW(ctypes.Structure):
            _fields_ = [("Flags", wintypes.DWORD), ("Type", wintypes.DWORD),
                        ("TargetName", wintypes.LPWSTR), ("Comment", wintypes.LPWSTR),
                        ("LastWritten", wintypes.FILETIME), ("CredentialBlobSize", wintypes.DWORD),
                        ("CredentialBlob", ctypes.POINTER(ctypes.c_ubyte)), ("Persist", wintypes.DWORD),
                        ("AttributeCount", wintypes.DWORD), ("Attributes", ctypes.c_void_p),
                        ("TargetAlias", wintypes.LPWSTR), ("UserName", wintypes.LPWSTR)]
        advapi = ctypes.WinDLL("advapi32", use_last_error=True)
        return ctypes, CREDENTIALW, advapi

    def read(self, name: str) -> str | None:
        ctypes, CREDENTIALW, advapi = self._api()
        pcred = ctypes.POINTER(CREDENTIALW)()
        if not advapi.CredReadW(CREDENTIAL_TARGET_PREFIX + name, 1, 0, ctypes.byref(pcred)):
            return None                                  # not saved (ERROR_NOT_FOUND)
        try:
            cred = pcred.contents
            blob = ctypes.string_at(cred.CredentialBlob, cred.CredentialBlobSize)
            return blob.decode("utf-16-le") or None
        finally:
            advapi.CredFree(pcred)

    def write(self, name: str, value: str) -> None:
        ctypes, CREDENTIALW, advapi = self._api()
        blob = value.encode("utf-16-le")
        buffer = (ctypes.c_ubyte * len(blob)).from_buffer_copy(blob)
        cred = CREDENTIALW(Type=1, TargetName=CREDENTIAL_TARGET_PREFIX + name,
                           Comment="Negative Screening: " + SECRET_LABELS.get(name, name),
                           CredentialBlobSize=len(blob), CredentialBlob=buffer,
                           Persist=2, UserName=name)      # 2 = CRED_PERSIST_LOCAL_MACHINE
        if not advapi.CredWriteW(ctypes.byref(cred), 0):
            raise OSError(ctypes.get_last_error(), "Windows Credential Manager refused the password")

    def delete(self, name: str) -> None:
        ctypes, _, advapi = self._api()
        advapi.CredDeleteW(CREDENTIAL_TARGET_PREFIX + name, 1, 0)


class MemorySecretStore:
    """Stand-in store for the self-tests: the machine's saved passwords are never touched."""

    def __init__(self, values: dict[str, str] | None = None) -> None:
        self.values = dict(values or {})

    def read(self, name: str) -> str | None:
        return self.values.get(name)

    def write(self, name: str, value: str) -> None:
        self.values[name] = value

    def delete(self, name: str) -> None:
        self.values.pop(name, None)


SECRET_STORE: WindowsCredentialStore | MemorySecretStore | None = (
    WindowsCredentialStore() if os.name == "nt" else None)


def get_secret(name: str) -> str:
    """A secret by name: the environment variable if set (developer / admin override),
    else the password saved in Windows Credential Manager by the app's Setup; '' if neither."""
    value = os.environ.get(name, "")
    if value or SECRET_STORE is None:
        return value
    try:
        return SECRET_STORE.read(name) or ""
    except OSError:
        return ""


def secret_source(name: str) -> str | None:
    """'environment variable', 'Windows Credential Manager' or None — never the value."""
    if os.environ.get(name):
        return "environment variable"
    if SECRET_STORE is not None and get_secret(name):
        return "Windows Credential Manager"
    return None


class ConfigError(Exception):
    """Raised with a message naming the file and the field."""


SECRET_KEY_MARKERS = ("password", "passwd", "secret")   # any config key containing these is refused
SOURCE_NAMES = ("google_alerts", "website", "sc_malaysia")


def init_config(config_dir: Path, overwrite: bool = False) -> list[Path]:
    """S03: write CONFIG_FILES into config_dir as pretty JSON (UTF-8). Skip files that
    already exist unless overwrite=True. Return the paths actually written."""
    config_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for name, template in CONFIG_FILES.items():
        path = config_dir / name
        if path.exists() and not overwrite:
            continue
        path.write_text(json.dumps(template, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        written.append(path)
    return written


def _is_text(v) -> bool:
    return isinstance(v, str) and v.strip() != ""


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_int(v) -> bool:
    return isinstance(v, int) and not isinstance(v, bool)


def _is_email_address(v) -> bool:
    return isinstance(v, str) and re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", v) is not None


def _check(ok: bool, file: str, field: str, problem: str) -> None:
    if not ok:
        raise ConfigError(f"{file}: {field}: {problem}")


def _field(data: dict, dotted: str, file: str):
    """Value at a dotted path, or ConfigError naming the missing field."""
    node = data
    for part in dotted.split("."):
        _check(isinstance(node, dict) and part in node, file, dotted, "missing")
        node = node[part]
    return node


def _read_config_json(path: Path) -> dict:
    name = path.name
    if not path.is_file():
        raise ConfigError(f"{name}: file not found in {path.parent} (create it with --init-config)")
    try:
        # utf-8-sig: Notepad and PowerShell may save UTF-8 with a byte-order mark
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except UnicodeDecodeError as e:
        raise ConfigError(f"{name}: not valid UTF-8 ({e.reason})") from None
    except json.JSONDecodeError as e:
        raise ConfigError(f"{name}: invalid JSON at line {e.lineno}, column {e.colno}: {e.msg}") from None
    _check(isinstance(data, dict), name, "(top level)", "must be a JSON object")
    return data


def _secret_key_paths(node, path: str = "") -> list[str]:
    """Dotted paths of keys that look like secrets. Values are never returned."""
    found: list[str] = []
    if isinstance(node, dict):
        for key, value in node.items():
            p = f"{path}.{key}" if path else str(key)
            if any(m in str(key).lower() for m in SECRET_KEY_MARKERS):
                found.append(p)
            found.extend(_secret_key_paths(value, p))
    elif isinstance(node, list):
        for i, value in enumerate(node):
            found.extend(_secret_key_paths(value, f"{path}[{i}]"))
    return found


def _validate_entities(data: dict, f: str = "entities.json") -> None:
    _check(_is_text(data.get("version")), f, "version", "must be a non-empty string")
    entities = _field(data, "entities", f)
    _check(isinstance(entities, list) and len(entities) > 0, f, "entities", "must be a non-empty list")
    seen_ids: set[str] = set()
    for i, e in enumerate(entities):
        at = f"entities[{i}]"
        _check(isinstance(e, dict), f, at, "must be an object")
        _check(_is_text(e.get("id")), f, f"{at}.id", "must be a non-empty string")
        _check(e["id"] not in seen_ids, f, f"{at}.id", f"duplicate id {e['id']!r}")
        seen_ids.add(e["id"])
        _check(_is_text(e.get("display_name")), f, f"{at}.display_name", "must be a non-empty string")
        names = e.get("names")
        _check(isinstance(names, list) and len(names) > 0 and all(_is_text(n) for n in names),
               f, f"{at}.names", "must be a non-empty list of non-empty strings")
        _check(isinstance(e.get("approved"), bool), f, f"{at}.approved", "must be true or false")
        for key in ("variants", "aliases", "alert_names"):
            if key in e:
                _check(isinstance(e[key], list) and all(_is_text(n) for n in e[key]),
                       f, f"{at}.{key}", "must be a list of non-empty strings")
        for form in e.get("variants", []) + e.get("aliases", []):
            _check(bool(tokenize_words(form)), f, f"{at}: {form!r}",
                   "has no letters, digits or Chinese characters, so it can never match")
        amb = e.get("ambiguous_terms", [])
        _check(isinstance(amb, list), f, f"{at}.ambiguous_terms", "must be a list")
        for j, t in enumerate(amb):
            tat = f"{at}.ambiguous_terms[{j}]"
            _check(isinstance(t, dict), f, tat, "must be an object")
            _check(_is_text(t.get("term")), f, f"{tat}.term", "must be a non-empty string")
            ctx = t.get("requires_any_context")
            _check(isinstance(ctx, list) and len(ctx) > 0 and all(_is_text(c) for c in ctx),
                   f, f"{tat}.requires_any_context", "must be a non-empty list of non-empty strings")
            win = t.get("context_window_words")
            _check(_is_int(win) and win > 0, f, f"{tat}.context_window_words", "must be a positive integer")
    _check(any(e["approved"] is True for e in entities), f, "entities",
           "no entity has \"approved\": true — nothing would be screened")
    links = data.get("relationships", [])
    _check(isinstance(links, list), f, "relationships", "must be a list")
    for i, r in enumerate(links):
        at = f"relationships[{i}]"
        _check(isinstance(r, dict), f, at, "must be an object")
        for key in ("subject", "object"):
            _check(r.get(key) in seen_ids, f, f"{at}.{key}", f"must be the id of an entity above, not {r.get(key)!r}")
        _check(r["subject"] != r["object"], f, at, "subject and object must be different entities")
        for key in ("type", "label"):
            _check(_is_text(r.get(key)), f, f"{at}.{key}", "must be a non-empty string")
        _check(isinstance(r.get("approved"), bool), f, f"{at}.approved", "must be true or false")


def _validate_keywords(data: dict, f: str = "adverse_keywords.json") -> None:
    _check(_is_text(data.get("version")), f, "version", "must be a non-empty string")
    categories = _field(data, "categories", f)
    _check(isinstance(categories, list) and bool(categories), f, "categories", "must be a non-empty list")
    priorities = [p.value for p in Priority]
    seen_ids: set[str] = set()
    for i, c in enumerate(categories):
        at = f"categories[{i}]"
        _check(isinstance(c, dict), f, at, "must be an object")
        _check(_is_text(c.get("id")), f, f"{at}.id", "must be a non-empty string")
        _check(c["id"] not in seen_ids, f, f"{at}.id", f"duplicate id {c['id']!r}")
        seen_ids.add(c["id"])
        _check(_is_text(c.get("name")), f, f"{at}.name", "must be a non-empty string")
        _check(c.get("priority") in priorities, f, f"{at}.priority", f"must be one of {', '.join(priorities)}")
        kws = c.get("keywords")
        _check(isinstance(kws, list) and bool(kws), f, f"{at}.keywords", "must be a non-empty list")
        for j, k in enumerate(kws):
            kat = f"{at}.keywords[{j}]"
            _check(isinstance(k, dict), f, kat, "must be an object like {\"term\": ..., \"ambiguous\": false}")
            _check(_is_text(k.get("term")), f, f"{kat}.term", "must be a non-empty string")
            _check(isinstance(k.get("ambiguous"), bool), f, f"{kat}.ambiguous", "must be true or false")
            if "variants" in k:
                _check(isinstance(k["variants"], list) and all(_is_text(v) for v in k["variants"]),
                       f, f"{kat}.variants", "must be a list of non-empty strings")


_SETTINGS_RULES: list[tuple[str, Callable[[object], bool], str]] = [
    ("timezone", lambda v: v == "Asia/Kuala_Lumpur", 'must be "Asia/Kuala_Lumpur" (the only supported time zone)'),
    ("paths.data_dir", _is_text, "must be a non-empty string"),
    ("paths.reports_dir", _is_text, "must be a non-empty string"),
    ("paths.logs_dir", _is_text, "must be a non-empty string"),
    ("paths.database", _is_text, "must be a non-empty string"),
    ("paths.dry_run_database", _is_text, "must be a non-empty string"),
    ("sources", lambda v: isinstance(v, dict), "must be an object"),
    ("http.timeout_seconds", lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0, "must be a number > 0"),
    ("http.max_retries", lambda v: isinstance(v, int) and not isinstance(v, bool) and v >= 0, "must be an integer >= 0"),
    ("http.backoff_seconds", lambda v: _is_number(v) and isinstance(v, (int, float)) and v >= 0, "must be a number >= 0"),
    ("http.delay_between_requests_seconds", lambda v: _is_number(v) and isinstance(v, (int, float)) and v >= 0, "must be a number >= 0"),
    ("http.user_agent", _is_text, "must be a non-empty string"),
    ("screening.keyword_name_window_words", lambda v: _is_int(v) and isinstance(v, int) and v > 0, "must be a positive integer"),
    ("report.include_out_of_period_sheet", lambda v: isinstance(v, bool), "must be true or false"),
    ("email_defaults.send_on_automatic", lambda v: isinstance(v, bool), "must be true or false"),
    ("email_defaults.send_on_manual", lambda v: isinstance(v, bool), "must be true or false"),
    ("run.max_run_minutes", lambda v: isinstance(v, int) and not isinstance(v, bool) and v > 0, "must be a positive integer"),
]


def _validate_settings(data: dict, f: str = "settings.json") -> None:
    _check("nlp" not in data, f, "nlp",
           "NLP is excluded from this project by the owner's decision; remove this section")
    for dotted, ok, problem in _SETTINGS_RULES:
        _check(ok(_field(data, dotted, f)), f, dotted, problem)
    cooldown = data["email_defaults"].get("manual_email_cooldown_minutes")
    if cooldown is not None:            # absent: manual runs email a week only once, like scheduled runs
        _check(_is_int(cooldown) and isinstance(cooldown, int) and 0 <= cooldown <= 1440, f,
               "email_defaults.manual_email_cooldown_minutes", "must be a whole number of minutes, 0-1440")
    for name in SOURCE_NAMES:
        if name in data["sources"]:     # an absent source is simply disabled
            src = data["sources"][name]
            _check(isinstance(src, dict), f, f"sources.{name}", "must be an object")
            _check(isinstance(src.get("enabled"), bool), f, f"sources.{name}.enabled", "must be true or false")
    ga = data["sources"].get("google_alerts", {})
    if "mode" in ga:
        _check(ga["mode"] in ("imap", "rss"), f, "sources.google_alerts.mode", 'must be "imap" or "rss"')
    if "fetch_public_web_pages" in ga:
        _check(isinstance(ga["fetch_public_web_pages"], bool), f,
               "sources.google_alerts.fetch_public_web_pages", "must be true or false")
    site = data["sources"].get("website", {})
    for key in ("feed_url", "comments_feed_url"):
        if site.get(key) is not None:
            _check(isinstance(site[key], str) and site[key].startswith(("http://", "https://")), f,
                   f"sources.website.{key}", "must be an http(s) URL or null")
    if "max_feed_pages" in site:
        _check(_is_int(site["max_feed_pages"]) and site["max_feed_pages"] > 0, f,
               "sources.website.max_feed_pages", "must be a positive integer")


def _validate_email(data: dict, settings: dict, f: str = "email.json") -> None:
    _check(_is_email_address(_field(data, "sender", f)), f, "sender", "must be an email address")
    rcpt = _field(data, "recipients", f)
    _check(isinstance(rcpt, list) and len(rcpt) > 0 and all(_is_email_address(r) for r in rcpt),
           f, "recipients", "must be a non-empty list of email addresses")
    _check(_is_text(_field(data, "subject_template", f)), f, "subject_template", "must be a non-empty string")
    _check(_is_text(_field(data, "smtp.host", f)), f, "smtp.host", "must be a non-empty string")
    port = _field(data, "smtp.port", f)
    _check(isinstance(port, int) and not isinstance(port, bool) and 0 < port < 65536,
            f, "smtp.port", "must be a port number (1-65535)")
    ga = settings["sources"].get("google_alerts", {})
    if ga.get("enabled") and ga.get("mode", "imap") == "imap":
        for dotted in ("alerts_mailbox.imap_host", "alerts_mailbox.username",
                       "alerts_mailbox.label", "alerts_mailbox.sender_filter"):
            _check(_is_text(_field(data, dotted, f)), f, dotted,
                   "must be a non-empty string (needed because Google Alerts IMAP mode is enabled)")
        port = _field(data, "alerts_mailbox.imap_port", f)
        _check(isinstance(port, int) and not isinstance(port, bool) and 0 < port < 65536,
             f, "alerts_mailbox.imap_port", "must be a port number (1-65535)")


def load_config(config_dir: Path) -> Config:
    """S03: read the four JSON files, validate, and return Config.
    Validation (raise ConfigError naming file + field):
      - every file present and valid JSON
      - entities: each entity has id, display_name, non-empty names, approved (bool)
      - keywords: each category has id, name, priority in Priority, keywords list of {term, ambiguous}
      - email: sender and recipients present; no key named like 'password' anywhere
    Entities with approved != True are REMOVED from the returned config."""
    files = {name: _read_config_json(config_dir / name) for name in CONFIG_FILES}
    for name, data in files.items():     # secrets are checked first, in every file
        leaked = _secret_key_paths(data)
        _check(not leaked, name, ", ".join(leaked),
               "secrets are not allowed in config files; save them in the app's Setup (Windows "
               "Credential Manager) or as environment variables (" + ", ".join(SECRET_ENV_VARS) + ")")
    settings, entities = files["settings.json"], files["entities.json"]
    keywords, email = files["adverse_keywords.json"], files["email.json"]
    _validate_settings(settings)
    _validate_entities(entities)
    _validate_keywords(keywords)
    _validate_email(email, settings)
    approved = {**entities, "entities": [e for e in entities["entities"] if e["approved"] is True]}
    return Config(settings=settings, entities=approved, keywords=keywords, email=email,
                  config_dir=config_dir)


# =============================================================================
# S04 — LOGGING WITH SECRET REDACTION                   (→ core/logging_setup.py)
# Docs: FINAL_ARCHITECTURE §11
# Goal: rotating logs/negative_screening.log + logs/run_<run_id>.log + console.
#       A filter replaces the values of SECRET_ENV_VARS (and 'password=...' pairs)
#       with '***' in every record. Never log full email bodies.
# =============================================================================

LOGGER_NAME = "negative_screening"
REDACTED = "***"
# key=value / key: value pairs whose value is a secret, e.g. password=x, "api_key": "x", token: x
_SECRET_PAIR_RE = re.compile(
    r"(?i)\b(password|passwd|pwd|secret|client[_-]?secret|token|access[_-]?token|api[_-]?key)"
    r"([\"']?\s*[=:]\s*)(\"[^\"]*\"|'[^']*'|[^\s,;&\"']+)")
_BEARER_RE = re.compile(r"(?i)\b(bearer)\s+[A-Za-z0-9._~+/=-]+")


def redact(text: str) -> str:
    """Remove secret values (environment variable or Credential Manager) and
    secret-looking pairs from text."""
    for name in SECRET_ENV_VARS:
        value = get_secret(name)
        # Gmail shows app passwords as 'abcd efgh ijkl mnop'; redact with and without spaces
        for form in {value, value.replace(" ", "")}:
            if form:
                text = text.replace(form, REDACTED)
    text = _SECRET_PAIR_RE.sub(lambda m: m.group(1) + m.group(2) + REDACTED, text)
    return _BEARER_RE.sub(lambda m: m.group(1) + " " + REDACTED, text)


class RedactingFilter(logging.Filter):
    """S04: replace secret values in record.msg/args with '***'."""

    def filter(self, record: logging.LogRecord) -> bool:
        record.msg = redact(record.getMessage())   # merge args first, so secrets in args are caught
        record.args = None
        if record.exc_info and not record.exc_text:
            record.exc_text = logging.Formatter().formatException(record.exc_info)
        if record.exc_text:
            record.exc_text = redact(record.exc_text)
        if record.stack_info:
            record.stack_info = redact(record.stack_info)
        return True


class _MytFormatter(logging.Formatter):
    """Log timestamps in Malaysia time with the offset, whatever the machine's zone."""

    def formatTime(self, record: logging.LogRecord, datefmt: str | None = None) -> str:  # noqa: ARG002
        return datetime.fromtimestamp(record.created, MYT).isoformat(timespec="seconds")


def setup_logging(logs_dir: Path, run_id: str, verbose: bool = False) -> logging.Logger:
    """S04: return the 'negative_screening' logger with file + console handlers and
    RedactingFilter attached to every handler. Idempotent (no duplicate handlers)."""
    from logging.handlers import RotatingFileHandler

    logs_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger(LOGGER_NAME)
    for h in list(logger.handlers):      # idempotent: replace the handlers of an earlier call
        logger.removeHandler(h)
        h.close()
    logger.setLevel(logging.DEBUG)
    logger.propagate = False

    # A title with e.g. Chinese characters must not break logging on a cp1252 console
    # (Task Scheduler redirects output); unencodable characters are escaped instead.
    reconfigure = getattr(sys.stderr, "reconfigure", None)
    if reconfigure is not None:
        reconfigure(errors="backslashreplace")
    handlers: list[logging.Handler] = [
        RotatingFileHandler(logs_dir / "negative_screening.log", maxBytes=5_000_000,
                            backupCount=5, encoding="utf-8"),
        logging.FileHandler(logs_dir / f"run_{run_id}.log", encoding="utf-8"),
    ]
    if sys.stderr is not None:           # none in the packaged app and under Task Scheduler
        console = logging.StreamHandler(sys.stderr)
        console.setLevel(logging.DEBUG if verbose else logging.INFO)
        handlers.append(console)
    formatter = _MytFormatter("%(asctime)s %(levelname)-7s %(message)s")
    for h in handlers:
        h.setFormatter(formatter)
        h.addFilter(RedactingFilter())
        logger.addHandler(h)
    return logger


# =============================================================================
# S05 — SQLITE DATABASE                                 (→ database/database.py)
# Docs: FINAL_ARCHITECTURE §7 (schema is authoritative there — copy it exactly)
# Goal: create schema; upsert findings by url_key; upsert finding_periods by
#       (finding_id, period_start) WITHOUT overwriting a review_status that staff
#       have changed to CONFIRMED ADVERSE / FALSE POSITIVE; append discovery_hits;
#       record runs, source_status, classifications, email_log.
# =============================================================================

SCHEMA_SQL = """
-- One row per execution
CREATE TABLE IF NOT EXISTS runs (
    run_id            TEXT PRIMARY KEY,           -- e.g. 20260922T103012+0800
    started_at        TEXT NOT NULL,              -- ISO 8601, MYT
    finished_at       TEXT,
    execution_mode    TEXT NOT NULL,              -- AUTOMATIC | MANUAL | MANUAL_TEST_DATE
    reference_date    TEXT NOT NULL,
    period_start      TEXT NOT NULL,
    period_end        TEXT NOT NULL,
    dry_run           INTEGER NOT NULL DEFAULT 0,
    exit_code         INTEGER,
    report_path       TEXT,
    app_version       TEXT,
    config_versions   TEXT                        -- JSON: entities/keywords/settings versions
);

-- Source health per run
CREATE TABLE IF NOT EXISTS source_status (
    run_id            TEXT NOT NULL REFERENCES runs(run_id),
    source            TEXT NOT NULL,              -- Google Alerts, Facebook, Instagram, LinkedIn, Public website, Company website, SC Malaysia
    status            TEXT NOT NULL,              -- SUCCESS | FAILED | SKIPPED | NOT CHECKED
    retrieved         INTEGER DEFAULT 0,
    filtered_out      INTEGER DEFAULT 0,
    flagged           INTEGER DEFAULT 0,
    duplicates        INTEGER DEFAULT 0,
    rejected          INTEGER DEFAULT 0,
    error             TEXT,
    PRIMARY KEY (run_id, source)
);

-- One row per unique piece of content, ever
CREATE TABLE IF NOT EXISTS findings (
    finding_id          INTEGER PRIMARY KEY AUTOINCREMENT,
    url_key             TEXT NOT NULL UNIQUE,     -- sha256 of canonical URL / platform key
    canonical_url       TEXT NOT NULL,
    discovered_url      TEXT NOT NULL,
    final_url           TEXT,
    platform            TEXT NOT NULL,
    content_class       TEXT NOT NULL,            -- social | public_web | regulatory | company
    content_type        TEXT,
    url_type            TEXT,                     -- SPECIFIC_CONTENT | CONTAINER | NON_EVIDENCE | UNKNOWN
    url_verification    TEXT NOT NULL,            -- VERIFIED | PARTIAL | REJECTED
    verification_basis  TEXT,
    title               TEXT,
    snippet             TEXT,
    publisher           TEXT,
    content_hash        TEXT,                     -- when content legitimately available
    language            TEXT,
    source_published_at TEXT,
    alert_received_at   TEXT,
    content_date        TEXT,                     -- event / content date, never the discovery date
    date_status         TEXT NOT NULL,            -- VERIFIED | DERIVED | APPROXIMATE | FIRST_SEEN | UNKNOWN
    date_basis          TEXT,
    first_discovered_at TEXT NOT NULL,            -- when this system first found the item (Discovery Date)
    last_seen_at        TEXT NOT NULL,
    date_precision      TEXT                      -- DAY | MONTH | YEAR; NULL when content_date is NULL
);
CREATE INDEX IF NOT EXISTS idx_findings_content_hash ON findings(content_hash);

-- Entity/keyword matches for a finding
CREATE TABLE IF NOT EXISTS finding_matches (
    finding_id        INTEGER NOT NULL REFERENCES findings(finding_id),
    entity_id         TEXT NOT NULL,
    matched_text      TEXT NOT NULL,
    keyword           TEXT,                       -- NULL for entity-only matches
    keyword_category  TEXT,
    keyword_priority  TEXT,
    ambiguous         INTEGER DEFAULT 0,
    proximity_words   INTEGER,
    UNIQUE (finding_id, entity_id, matched_text, keyword)
);

-- Which reporting week a finding belongs to, and its review state for that week
CREATE TABLE IF NOT EXISTS finding_periods (
    finding_id        INTEGER NOT NULL REFERENCES findings(finding_id),
    period_start      TEXT NOT NULL,
    period_end        TEXT NOT NULL,
    report_section    TEXT NOT NULL,              -- MAIN | DATE_UNKNOWN | UNVERIFIED_URL | OUT_OF_PERIOD | NOT_REPORTED
    priority          TEXT,
    rules_applied     TEXT,                       -- e.g. "R2,R5"
    review_status     TEXT NOT NULL DEFAULT 'PENDING',   -- PENDING | MANUAL REVIEW | NEEDS REVIEW | CONFIRMED ADVERSE | FALSE POSITIVE
    reviewer_note     TEXT,
    first_run_id      TEXT NOT NULL,
    last_run_id       TEXT NOT NULL,
    PRIMARY KEY (finding_id, period_start)
);

-- Every time a finding was surfaced (audit of discovery queries)
CREATE TABLE IF NOT EXISTS discovery_hits (
    hit_id            INTEGER PRIMARY KEY AUTOINCREMENT,
    finding_id        INTEGER NOT NULL REFERENCES findings(finding_id),
    run_id            TEXT NOT NULL REFERENCES runs(run_id),
    source            TEXT NOT NULL,
    discovery_method  TEXT NOT NULL,
    discovery_query   TEXT,
    source_message_id TEXT,                       -- alert email Message-ID / feed entry id
    seen_at           TEXT NOT NULL,
    UNIQUE (finding_id, run_id, source, discovery_query)
);

-- Email sends, used to prevent duplicate emails for a period
CREATE TABLE IF NOT EXISTS email_log (
    email_id          INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id            TEXT NOT NULL REFERENCES runs(run_id),
    period_start      TEXT NOT NULL,
    sent_at           TEXT NOT NULL,
    recipient_count   INTEGER NOT NULL,
    status            TEXT NOT NULL,              -- SENT | FAILED | SKIPPED
    reason            TEXT                        -- e.g. "already sent for period", "manual mode default"
);
"""

STAFF_REVIEW_DECISIONS = (ReviewStatus.CONFIRMED_ADVERSE.value, ReviewStatus.FALSE_POSITIVE.value)


def _iso(moment: datetime | None) -> str | None:
    """Aware datetime -> ISO 8601 text in Malaysia time (None stays None)."""
    return moment.astimezone(MYT).isoformat() if moment is not None else None


def db_connect(db_path: Path) -> sqlite3.Connection:
    """S05: open connection (foreign keys ON, row_factory=sqlite3.Row), create parent
    dir, apply SCHEMA_SQL idempotently (CREATE TABLE IF NOT EXISTS)."""
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA_SQL)
    _migrate_date_precision(conn)
    return conn


_LEGACY_SC_YEAR_RE = re.compile(r"the SC list gives only the year (\d{4})\b")


def _migrate_date_precision(conn: sqlite3.Connection) -> None:
    """Add findings.date_precision to a database created before it existed (once). Earlier, a
    FIRST_SEEN row's content_date held the start of the week the item was first found — a
    discovery date, not the event date. Such rows get the year the SC list stated (kept in
    their date_basis) as a YEAR-precision VERIFIED date, or no event date at all."""
    if "date_precision" in {r[1] for r in conn.execute("PRAGMA table_info(findings)")}:
        return
    with conn:
        conn.execute("ALTER TABLE findings ADD COLUMN date_precision TEXT")
        conn.execute("UPDATE findings SET date_precision = ? WHERE content_date IS NOT NULL"
                     " AND date_status <> ?", (DatePrecision.DAY.value, DateStatus.FIRST_SEEN.value))
        legacy = conn.execute("SELECT finding_id, date_basis FROM findings WHERE date_status = ?",
                              (DateStatus.FIRST_SEEN.value,)).fetchall()
        for finding_id, basis in legacy:
            m = _LEGACY_SC_YEAR_RE.search(basis or "")
            if m:
                conn.execute("UPDATE findings SET content_date = ?, date_precision = ?, date_status = ?"
                             " WHERE finding_id = ?",
                             (_iso(datetime(int(m.group(1)), 1, 1, tzinfo=MYT)), DatePrecision.YEAR.value,
                              DateStatus.VERIFIED.value, finding_id))
            else:
                conn.execute("UPDATE findings SET content_date = NULL WHERE finding_id = ?", (finding_id,))


def db_record_run_start(conn: sqlite3.Connection, ctx: RunContext) -> None:
    """S05: insert into runs."""
    versions = {"entities": ctx.config.entities.get("version"),
                "keywords": ctx.config.keywords.get("version"),
                "settings": ctx.config.settings.get("version")}
    with conn:
        conn.execute(
            "INSERT INTO runs (run_id, started_at, execution_mode, reference_date, period_start,"
            " period_end, dry_run, app_version, config_versions) VALUES (?,?,?,?,?,?,?,?,?)",
            (ctx.run_id, _iso(ctx.started_at), str(ctx.mode), ctx.period.reference_date.isoformat(),
             ctx.period.period_start.isoformat(), ctx.period.period_end.isoformat(),
             int(ctx.dry_run), APP_VERSION, json.dumps(versions)))


def db_record_run_end(conn: sqlite3.Connection, run_id: str, exit_code: int,
                      report_path: Path | None) -> None:
    """S05: update runs.finished_at / exit_code / report_path."""
    with conn:
        conn.execute("UPDATE runs SET finished_at = ?, exit_code = ?, report_path = ? WHERE run_id = ?",
                     (_iso(now_in_malaysia()), exit_code,
                      str(report_path) if report_path is not None else None, run_id))


def db_record_source_status(conn: sqlite3.Connection, run_id: str,
                            rows: list[dict]) -> None:
    """S05: insert one source_status row per dict
    (keys: source, status, retrieved, filtered_out, flagged, duplicates, rejected, error)."""
    with conn:
        for r in rows:
            conn.execute(
                "INSERT OR REPLACE INTO source_status (run_id, source, status, retrieved, filtered_out,"
                " flagged, duplicates, rejected, error) VALUES (?,?,?,?,?,?,?,?,?)",
                (run_id, r["source"], str(r["status"]), r.get("retrieved", 0), r.get("filtered_out", 0),
                 r.get("flagged", 0), r.get("duplicates", 0), r.get("rejected", 0), r.get("error")))


def db_upsert_candidate(conn: sqlite3.Connection, c: Candidate, ctx: RunContext) -> int:
    """S05: upsert findings (by url_key; update last_seen_at), finding_matches,
    finding_periods (by finding_id+
    period_start; keep staff review decisions), and one discovery_hits row per query.
    Return finding_id. Must be safe to call repeatedly for the same content."""
    missing = [name for name, value in (("url_key", c.url_key), ("platform", c.platform),
                                        ("content_class", c.content_class),
                                        ("verification", c.verification), ("section", c.section))
               if not value]
    if missing:
        raise ValueError(f"candidate not fully processed, missing: {', '.join(missing)}")
    raw, now = c.raw, _iso(now_in_malaysia())
    chash = content_hash(raw.content_text) if raw.content_text else None
    with conn:
        # Latest screening results overwrite; source facts are kept when a later run lacks them.
        conn.execute(
            """INSERT INTO findings (url_key, canonical_url, discovered_url, final_url, platform,
                   content_class, content_type, url_type, url_verification, verification_basis,
                   title, snippet, publisher, content_hash, language, source_published_at,
                   alert_received_at, content_date, date_precision, date_status, date_basis,
                   first_discovered_at, last_seen_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(url_key) DO UPDATE SET
                   canonical_url = excluded.canonical_url,
                   discovered_url = excluded.discovered_url,
                   final_url = COALESCE(excluded.final_url, findings.final_url),
                   platform = excluded.platform,
                   content_class = excluded.content_class,
                   content_type = excluded.content_type,
                   url_type = excluded.url_type,
                   url_verification = excluded.url_verification,
                   verification_basis = excluded.verification_basis,
                   title = COALESCE(NULLIF(excluded.title, ''), findings.title),
                   snippet = COALESCE(NULLIF(excluded.snippet, ''), findings.snippet),
                   publisher = COALESCE(excluded.publisher, findings.publisher),
                   content_hash = COALESCE(excluded.content_hash, findings.content_hash),
                   language = COALESCE(excluded.language, findings.language),
                   source_published_at = COALESCE(excluded.source_published_at, findings.source_published_at),
                   alert_received_at = COALESCE(excluded.alert_received_at, findings.alert_received_at),
                   content_date = excluded.content_date,
                   date_precision = excluded.date_precision,
                   date_status = excluded.date_status,
                   date_basis = excluded.date_basis,
                   last_seen_at = excluded.last_seen_at""",
            (c.url_key, c.canonical_url, raw.discovered_url, raw.final_url, str(c.platform),
             str(c.content_class), c.content_type, str(c.url_type), str(c.verification),
             c.verification_basis, raw.title, raw.snippet, raw.publisher, chash, raw.language,
             _iso(raw.source_published_at), _iso(raw.alert_received_at), _iso(c.content_date),
             str(c.date_precision) if c.date_precision else None, str(c.date_status), c.date_basis, now, now))
        finding_id = conn.execute("SELECT finding_id FROM findings WHERE url_key = ?",
                                  (c.url_key,)).fetchone()[0]

        # Matches describe the latest screening of this content: replace, never accumulate.
        conn.execute("DELETE FROM finding_matches WHERE finding_id = ?", (finding_id,))
        for m in c.entity_matches + c.keyword_matches:
            conn.execute(
                "INSERT OR IGNORE INTO finding_matches (finding_id, entity_id, matched_text, keyword,"
                " keyword_category, keyword_priority, ambiguous, proximity_words) VALUES (?,?,?,?,?,?,?,?)",
                (finding_id, m.entity_id, m.matched_text, m.keyword, m.category, m.priority,
                 int(m.ambiguous), m.proximity_words))

        # A review decision made by staff is never overwritten by a re-run.
        conn.execute(
            """INSERT INTO finding_periods (finding_id, period_start, period_end, report_section,
                   priority, rules_applied, review_status, first_run_id, last_run_id)
               VALUES (?,?,?,?,?,?,?,?,?)
               ON CONFLICT(finding_id, period_start) DO UPDATE SET
                   period_end = excluded.period_end,
                   report_section = excluded.report_section,
                   priority = excluded.priority,
                   rules_applied = excluded.rules_applied,
                   review_status = CASE WHEN finding_periods.review_status IN (?, ?)
                                        THEN finding_periods.review_status
                                        ELSE excluded.review_status END,
                   last_run_id = excluded.last_run_id""",
            (finding_id, ctx.period.period_start.isoformat(), ctx.period.period_end.isoformat(),
             str(c.section), str(c.priority) if c.priority else None, ",".join(c.rules_applied),
             str(c.review_status), ctx.run_id, ctx.run_id, *STAFF_REVIEW_DECISIONS))

        # SQLite treats NULLs as distinct in UNIQUE, so check "no query" hits with IS.
        for query in sorted(c.discovery_queries) or [None]:
            seen = conn.execute(
                "SELECT 1 FROM discovery_hits WHERE finding_id = ? AND run_id = ? AND source = ?"
                " AND discovery_query IS ?", (finding_id, ctx.run_id, raw.source, query)).fetchone()
            if not seen:
                conn.execute(
                    "INSERT INTO discovery_hits (finding_id, run_id, source, discovery_method,"
                    " discovery_query, source_message_id, seen_at) VALUES (?,?,?,?,?,?,?)",
                    (finding_id, ctx.run_id, raw.source, raw.discovery_method, query,
                     raw.source_message_id, now))
    return finding_id


def db_findings_for_period(conn: sqlite3.Connection, period: ReportingPeriod) -> list[sqlite3.Row]:
    """S05: all findings joined with finding_periods for the period (input to S18)."""
    return conn.execute(
        """SELECT f.*, fp.period_start, fp.period_end, fp.report_section, fp.priority,
                  fp.rules_applied, fp.review_status, fp.reviewer_note,
                  (SELECT group_concat(DISTINCT m.entity_id) FROM finding_matches m
                    WHERE m.finding_id = f.finding_id AND m.keyword IS NULL) AS entity_ids,
                  (SELECT group_concat(pair, char(31)) FROM
                     (SELECT DISTINCT m.entity_id || char(30) || m.matched_text AS pair FROM finding_matches m
                       WHERE m.finding_id = f.finding_id AND m.keyword IS NULL)) AS matched_names,
                  (SELECT group_concat(DISTINCT CASE WHEN m.matched_text IS NOT NULL
                                                     AND lower(m.matched_text) != lower(m.keyword)
                                                THEN m.keyword || ' (' || m.matched_text || ')'
                                                ELSE m.keyword END) FROM finding_matches m
                    WHERE m.finding_id = f.finding_id AND m.keyword IS NOT NULL) AS keywords,
                  (SELECT group_concat(DISTINCT m.keyword_category) FROM finding_matches m
                    WHERE m.finding_id = f.finding_id AND m.keyword IS NOT NULL) AS keyword_categories,
                  (SELECT group_concat(DISTINCT h.discovery_method) FROM discovery_hits h
                    WHERE h.finding_id = f.finding_id) AS discovery_methods,
                  (SELECT group_concat(DISTINCT h.discovery_query) FROM discovery_hits h
                    WHERE h.finding_id = f.finding_id) AS discovery_queries,
                  (SELECT COUNT(DISTINCT h.discovery_query) FROM discovery_hits h
                    WHERE h.finding_id = f.finding_id) AS query_count
           FROM findings f JOIN finding_periods fp ON fp.finding_id = f.finding_id
           WHERE fp.period_start = ?
           ORDER BY CASE fp.priority WHEN 'CRITICAL' THEN 5 WHEN 'HIGH' THEN 4 WHEN 'MEDIUM' THEN 3
                                     WHEN 'REVIEW' THEN 2 WHEN 'LOW' THEN 1 ELSE 0 END DESC,
                    f.content_date, f.finding_id""",
        (period.period_start.isoformat(),)).fetchall()


REPORTED_SECTIONS = (Section.MAIN.value, Section.DATE_UNKNOWN.value, Section.UNVERIFIED_URL.value)


def db_first_reported_period(conn: sqlite3.Connection, finding_id: int, period: ReportingPeriod) -> date | None:
    """S05: start of the earliest EARLIER week whose report listed this finding on a reviewed sheet
    (Findings, Date Unknown or Unverified URLs), else None. An item that is out of period now but was
    already listed then (an SC Investor Alert List entry, listed in the week first found) is not
    counted 'for review' again (D-24)."""
    row = conn.execute(
        "SELECT MIN(period_start) FROM finding_periods WHERE finding_id = ? AND period_start < ?"
        f" AND report_section IN ({','.join('?' * len(REPORTED_SECTIONS))})",
        (finding_id, period.period_start.isoformat(), *REPORTED_SECTIONS)).fetchone()
    return date.fromisoformat(row[0]) if row and row[0] else None


def db_minutes_since_last_email(conn: sqlite3.Connection, now: datetime) -> float | None:
    """S05: minutes since the newest SENT email of any week; None if none was ever sent."""
    row = conn.execute("SELECT MAX(sent_at) FROM email_log WHERE status = 'SENT'").fetchone()
    if not row or not row[0]:
        return None
    try:
        return max((now - datetime.fromisoformat(row[0])).total_seconds() / 60, 0.0)
    except (TypeError, ValueError):
        return None


def db_email_already_sent(conn: sqlite3.Connection, period: ReportingPeriod) -> bool:
    """S05: True if email_log has status='SENT' for period_start."""
    row = conn.execute("SELECT 1 FROM email_log WHERE period_start = ? AND status = 'SENT' LIMIT 1",
                       (period.period_start.isoformat(),)).fetchone()
    return row is not None


def db_record_email(conn: sqlite3.Connection, run_id: str, period: ReportingPeriod,
                    status: str, recipient_count: int, reason: str | None) -> None:
    """S05: insert into email_log."""
    with conn:
        conn.execute("INSERT INTO email_log (run_id, period_start, sent_at, recipient_count, status, reason)"
                     " VALUES (?,?,?,?,?,?)",
                     (run_id, period.period_start.isoformat(), _iso(now_in_malaysia()),
                      recipient_count, status, reason))


# =============================================================================
# S06 — URL NORMALISATION                              (→ processing/normalizer.py)
# Docs: DATA_COLLECTION_ANALYSIS §9.1, SOURCE_ACCESS_BLUEPRINT §2.2
# Goal: decode Google redirect links WITHOUT fetching them; canonicalise URLs;
#       derive platform content keys; hash into url_key. Never visit social URLs.
# =============================================================================

TRACKING_PARAMS = {"fbclid", "gclid", "igsh", "igshid", "__cft__", "__tn__", "ref", "refid",
                   "trk", "rcm", "mibextid", "sfnsn", "si"}          # plus any utm_*
IDENTITY_PARAMS = {"story_fbid", "id", "fbid", "v", "comment_id", "p"}
HOST_PREFIXES_TO_DROP = ("www.", "m.", "mobile.", "web.", "l.")


_GOOGLE_HOST_RE = re.compile(r"(www\.)?google\.[a-z]{2,3}(\.[a-z]{2})?")
_DEFAULT_PORTS = {"http": 80, "https": 443}


def collapse_whitespace(text: str | None) -> str:
    return " ".join((text or "").split())


def decode_google_redirect(url: str) -> str:
    """S06: if url is a google.com/url?...&url=<target> (or q=<target>) wrapper, return the
    unquoted target; otherwise return url unchanged. Pure string work — no network."""
    from urllib.parse import parse_qs, urlsplit

    try:
        parts = urlsplit(url.strip())
    except ValueError:
        return url
    host = (parts.hostname or "").lower()
    if not _GOOGLE_HOST_RE.fullmatch(host) or parts.path != "/url":
        return url
    params = parse_qs(parts.query)
    for name in ("url", "q"):
        target = (params.get(name) or [""])[0].strip()
        if target.lower().startswith(("http://", "https://")):
            return target
    return url


def _host_without_prefixes(host: str) -> str:
    changed = True
    while changed:
        changed = False
        for prefix in HOST_PREFIXES_TO_DROP:
            rest = host[len(prefix):]
            if host.startswith(prefix) and "." in rest:     # never reduce 'web.dev' to 'dev'
                host, changed = rest, True
    if host.endswith(".linkedin.com"):                      # my.linkedin.com, uk.linkedin.com, ...
        host = "linkedin.com"
    return host


def canonicalize_url(url: str) -> str:
    """S06: lower-case scheme+host; drop HOST_PREFIXES_TO_DROP; LinkedIn country
    subdomains (e.g. my.linkedin.com) -> linkedin.com; remove fragment and trailing slash;
    drop utm_* and TRACKING_PARAMS; keep other params sorted by name. Path case is kept.
    Exception: a '#comment-<n>' fragment is kept, because it identifies a comment (S11)."""
    from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

    url = url.strip()
    try:
        parts = urlsplit(url)
        port = parts.port
    except ValueError:            # e.g. an invalid port number
        return url
    if not parts.scheme or not parts.hostname:
        return url
    scheme = parts.scheme.lower()
    netloc = _host_without_prefixes(parts.hostname.lower())
    if port is not None and port != _DEFAULT_PORTS.get(scheme):
        netloc += f":{port}"
    path = parts.path.rstrip("/")
    params = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
              if not k.lower().startswith("utm_") and k.lower() not in TRACKING_PARAMS]
    query = urlencode(sorted(params, key=lambda kv: kv[0]))
    fragment = parts.fragment if re.fullmatch(r"comment-\d+", parts.fragment) else ""
    return urlunsplit((scheme, netloc, path, query, fragment))


def platform_content_key(canonical_url: str) -> str | None:
    """S06: same content, different URL shapes -> one key:
      LinkedIn  /posts/...-activity-<id>-...  and  /feed/update/urn:li:activity:<id>
                -> 'linkedin:activity:<id>'
      Instagram /p/<code>/, /reel/<code>/, /reels/<code>/, /tv/<code>/ -> 'instagram:media:<code>'
      Facebook  permalink.php / story.php ?story_fbid=<s>&id=<p>        -> 'facebook:story:<p>:<s>'
                (append ':comment:<cid>' when comment_id present)
    Return None when no platform key applies."""
    from urllib.parse import parse_qs, urlsplit

    parts = urlsplit(canonical_url)
    host, path = (parts.hostname or "").lower(), parts.path
    if host == "linkedin.com":
        m = (re.match(r"/posts/[^/]*-activity-(\d+)", path)
             or re.match(r"/feed/update/urn:li:activity:(\d+)", path))
        return f"linkedin:activity:{m.group(1)}" if m else None
    if host == "instagram.com":
        m = re.match(r"/(?:[^/]+/)?(?:p|reel|reels|tv)/([A-Za-z0-9_-]+)", path)
        return f"instagram:media:{m.group(1)}" if m else None
    if host in ("facebook.com", "fb.com") and path in ("/permalink.php", "/story.php"):
        q = {k: v[0] for k, v in parse_qs(parts.query).items()}
        if q.get("story_fbid") and q.get("id"):
            key = f"facebook:story:{q['id']}:{q['story_fbid']}"
            return key + (f":comment:{q['comment_id']}" if q.get("comment_id") else "")
    return None


def make_url_key(canonical_url: str) -> str:
    """S06: sha256 hex of platform_content_key(url) if not None, else of canonical_url."""
    basis = platform_content_key(canonical_url) or canonical_url
    return hashlib.sha256(basis.encode("utf-8")).hexdigest()


def content_key_url_key(content_key: str) -> str:
    """url_key for an item a collector identifies itself (e.g. one SC alert-list entry)."""
    return hashlib.sha256(f"content:{content_key}".encode("utf-8")).hexdigest()


def normalize_item(item: RawItem) -> Candidate:
    """S06: build a Candidate: decode redirect, canonicalise, set canonical_url, url_key,
    discovery_queries={item.discovery_query} (if any). Also collapse whitespace in title/snippet.
    When the collector supplies extra['content_key'] (several items on one page, e.g. SC Investor
    Alert List entries), that key identifies the item instead of the shared page URL."""
    item.discovered_url = decode_google_redirect(item.discovered_url.strip())
    item.title, item.snippet = collapse_whitespace(item.title), collapse_whitespace(item.snippet)
    canonical = canonicalize_url(item.discovered_url)
    if item.extra.get("content_key"):
        url_key = content_key_url_key(item.extra["content_key"])
    elif canonical:
        url_key = make_url_key(canonical)
    else:   # no URL at all: key on the text so unrelated URL-less items are never merged
        url_key = hashlib.sha256(f"no-url:{item.source}:{item.title}:{item.snippet}".encode()).hexdigest()
    return Candidate(raw=item, canonical_url=canonical, url_key=url_key,
                     discovery_queries={item.discovery_query} if item.discovery_query else set())


# =============================================================================
# S07 — DEDUPLICATION                                   (→ processing/deduplicator.py)
# Docs: DATA_COLLECTION_ANALYSIS §9.2–9.3
# Goal: merge candidates with the same url_key (union discovery_queries, keep the
#       richest fields); flag probable duplicates by content hash or
#       publisher + normalised title + same content date.
# =============================================================================

def content_hash(text: str | None) -> str | None:
    """S07: sha256 of lower-cased, whitespace-collapsed text; None for empty/None."""
    normalised = collapse_whitespace(text).lower()
    return hashlib.sha256(normalised.encode("utf-8")).hexdigest() if normalised else None


def _merge_into(kept: Candidate, dup: Candidate) -> None:
    """Fold a duplicate into the kept candidate, keeping the richest value of each field."""
    kept.discovery_queries |= dup.discovery_queries
    for attr in ("title", "snippet", "content_text"):          # longer text is richer
        if len(getattr(dup.raw, attr) or "") > len(getattr(kept.raw, attr) or ""):
            setattr(kept.raw, attr, getattr(dup.raw, attr))
    for attr in ("source_published_at", "final_url", "publisher", "language",
                 "alert_received_at", "source_message_id"):   # fill what is missing
        if getattr(kept.raw, attr) is None:
            setattr(kept.raw, attr, getattr(dup.raw, attr))
    kept.raw.extra = {**dup.raw.extra, **kept.raw.extra}
    if kept.content_date is None and dup.content_date is not None:
        kept.content_date, kept.date_status, kept.date_basis = dup.content_date, dup.date_status, dup.date_basis
        kept.date_precision = dup.date_precision
    kept.notes += [n for n in dup.notes if n not in kept.notes]


def dedupe_in_run(candidates: list[Candidate]) -> tuple[list[Candidate], int]:
    """S07: return (unique candidates in first-seen order, number merged).
    Merge rules: same url_key -> one candidate; union discovery_queries; keep non-empty
    title/snippet/content_text/source_published_at from whichever has them."""
    unique: list[Candidate] = []
    by_key: dict[str, Candidate] = {}
    for c in candidates:
        if c.url_key and c.url_key in by_key:
            _merge_into(by_key[c.url_key], c)
            continue
        unique.append(c)
        if c.url_key:
            by_key[c.url_key] = c
    return unique, len(candidates) - len(unique)


def _publisher_key(c: Candidate) -> str:
    """The site an item comes from: its stated publisher, else the URL host."""
    from urllib.parse import urlsplit
    return (c.raw.publisher or urlsplit(c.canonical_url).hostname or "").strip().lower()


def _title_key(c: Candidate) -> str:
    return " ".join(tokenize_words(c.raw.title))


def _add_note(c: Candidate, note: str) -> None:
    if note not in c.notes:
        c.notes.append(note)


def mark_probable_duplicates(candidates: list[Candidate]) -> None:
    """S07: after dates are resolved (S14): same content_hash, or same publisher +
    normalised title + same content date -> merge into the first and add a Note.
    Different publishers are NOT merged; add 'also published by …' to Notes instead.
    Social posts are never merged here: Google Alerts titles them with the page or profile
    name, so two different posts share publisher + title + day. The same post under two URL
    shapes is already one item through its url_key (S06 platform_content_key)."""
    kept: list[Candidate] = []
    by_hash: dict[tuple[str, str], Candidate] = {}       # (publisher, hash)
    by_title: dict[tuple[str, str, date], Candidate] = {}  # (publisher, title, content date)
    for c in candidates:
        if c.content_class == ContentClass.SOCIAL:
            kept.append(c)
            continue
        pub, h, title = _publisher_key(c), content_hash(c.raw.content_text), _title_key(c)
        day = c.content_date.astimezone(MYT).date() if c.content_date else None
        target, reason = None, ""
        if h and (pub, h) in by_hash:
            target, reason = by_hash[(pub, h)], "same text"
        elif title and day and (pub, title, day) in by_title:
            target, reason = by_title[(pub, title, day)], "same publisher, title and date"
        if target is not None:
            _merge_into(target, c)
            _add_note(target, f"Probable duplicate merged ({reason}): {c.canonical_url}")
            continue
        kept.append(c)
        if h:
            by_hash[(pub, h)] = c
        if title and day:
            by_title[(pub, title, day)] = c

    # Syndicated copies (same text, or same title and date, on different sites) stay separate.
    groups: dict[tuple, list[Candidate]] = {}
    for c in kept:
        if c.content_class == ContentClass.SOCIAL:
            continue
        h, title = content_hash(c.raw.content_text), _title_key(c)
        if h:
            groups.setdefault(("hash", h), []).append(c)
        if title and c.content_date:
            groups.setdefault(("title", title, c.content_date.astimezone(MYT).date()), []).append(c)
    for members in groups.values():
        for c in members:
            for other in members:
                if other is not c:
                    _add_note(c, f"Also published by {_publisher_key(other) or 'unknown site'}: "
                                 f"{other.canonical_url}")
    candidates[:] = kept


# =============================================================================
# S08 — COLLECTOR FRAMEWORK                            (→ collectors/base.py)
# Docs: FINAL_ARCHITECTURE §6, §10; SOURCE_ACCESS_BLUEPRINT §1, §7
# Goal: shared polite HTTP session (User-Agent, timeout, retries with back-off on
#       timeouts/429/5xx, no retry on 401/403, delay between requests, robots.txt
#       check for direct-retrieval sites); run each collector in isolation so one
#       failure never stops the others; derive per-platform health rows.
# =============================================================================

MAX_RETRY_AFTER_SECONDS = 60


def is_social_url(url: str) -> bool:
    """True for Facebook / Instagram / LinkedIn URLs, which are never fetched."""
    from urllib.parse import urlsplit
    try:
        host = (urlsplit(url).hostname or "").lower()
    except ValueError:
        return False
    return any(host == s or host.endswith("." + s) for s in SOCIAL_HOST_SUFFIXES)


def build_http_session(settings: dict):
    """S08: return a requests.Session configured from settings['http'].
    (import requests inside this function)."""
    import requests

    session = requests.Session()
    session.headers["User-Agent"] = settings["http"]["user_agent"]
    return session


def _wait_between_requests(session, delay: float) -> None:
    """One request at a time, at least `delay` seconds apart (per session)."""
    from time import monotonic, sleep
    state = session.__dict__
    last = state.get("_ns_last_request_at")
    if last is not None and monotonic() - last < delay:
        sleep(delay - (monotonic() - last))
    state["_ns_last_request_at"] = monotonic()


def _robots_allows(session, url: str, http: dict, logger: logging.Logger) -> bool:
    """robots.txt check, cached per site for the life of the session. A missing robots.txt
    (4xx) allows everything; 401/403 or an unreachable/5xx robots.txt means 'do not fetch'."""
    from urllib.parse import urlsplit
    from urllib.robotparser import RobotFileParser

    parts = urlsplit(url)
    site = f"{parts.scheme}://{parts.netloc}"
    cache = session.__dict__.setdefault("_ns_robots_cache", {})
    if site not in cache:
        parser = RobotFileParser()
        _wait_between_requests(session, http["delay_between_requests_seconds"])
        resp = None
        try:
            resp = session.get(f"{site}/robots.txt", timeout=http["timeout_seconds"])
            code = resp.status_code
        except Exception as e:  # noqa: BLE001 — any network problem: be conservative
            logger.warning("robots.txt unreachable for %s (%s): not fetching", site, type(e).__name__)
            code = None
        if code == 200 and resp is not None:
            parser.parse(resp.text.splitlines())
        elif code in (401, 403) or code is None or code >= 500:
            parser.parse(["User-agent: *", "Disallow: /"])
        else:
            parser.parse(["User-agent: *", "Disallow:"])
        cache[site] = parser
    return cache[site].can_fetch(http["user_agent"], url)


MAX_REDIRECTS = 5
REDIRECT_CODES = (301, 302, 303, 307, 308)


class HttpStatusError(ConnectionError):
    """A final HTTP error status (e.g. 404) that is not worth retrying."""

    def __init__(self, status_code: int, message: str):
        super().__init__(message)
        self.status_code = status_code


def polite_get(session, url: str, settings: dict, logger: logging.Logger, *,
               check_robots: bool = True, **kwargs):
    """S08: GET with timeout, retries/back-off, per-request delay and robots.txt respect.
    Raise a clear exception after the final failure (caught by run_collector_safely).
    Redirects are followed here, one hop at a time, so every hop is checked again:
    a page that redirects to Facebook / Instagram / LinkedIn is never fetched."""
    from urllib.parse import urljoin

    http = settings["http"]
    kwargs.pop("allow_redirects", None)
    for _hop in range(MAX_REDIRECTS + 1):
        if is_social_url(url):
            raise PermissionError(f"refusing to fetch a social-media URL: {url}")
        if check_robots and not _robots_allows(session, url, http, logger):
            raise PermissionError(f"robots.txt does not allow fetching {url}")
        resp = _get_with_retries(session, url, http, logger, allow_redirects=False, **kwargs)
        location = resp.headers.get("Location")
        if resp.status_code not in REDIRECT_CODES or not location:
            return resp
        url = urljoin(url, location)
        logger.debug("redirected to %s", url)
    raise ConnectionError(f"more than {MAX_REDIRECTS} redirects: {url}")


def _get_with_retries(session, url: str, http: dict, logger: logging.Logger, **kwargs):
    """One GET with retries: back-off on timeouts, connection errors, 429 and 5xx;
    no retry on 401/403 (access problems) or other 4xx."""
    from time import sleep

    retryable: tuple[type[BaseException], ...] = (ConnectionError, TimeoutError)
    try:
        import requests
        retryable += (requests.ConnectionError, requests.Timeout)
    except ModuleNotFoundError:
        pass
    retries = http["max_retries"]
    problem, server_wait = "", 0
    for attempt in range(retries + 1):
        if attempt:
            wait = max(http["backoff_seconds"] * 2 ** (attempt - 1), server_wait)
            logger.info("retry %d/%d for %s in %.1fs (%s)", attempt, retries, url, wait, problem)
            sleep(wait)
        _wait_between_requests(session, http["delay_between_requests_seconds"])
        try:
            resp = session.get(url, timeout=http["timeout_seconds"], **kwargs)
        except retryable as e:
            problem, server_wait = type(e).__name__, 0
            continue
        code = resp.status_code
        if code in (401, 403):
            raise PermissionError(f"HTTP {code} (access denied) for {url}")
        if code == 429 or code >= 500:
            retry_after = str(resp.headers.get("Retry-After", ""))
            problem = f"HTTP {code}"
            server_wait = min(int(retry_after), MAX_RETRY_AFTER_SECONDS) if retry_after.isdigit() else 0
            continue
        if code >= 400:
            raise HttpStatusError(code, f"HTTP {code} for {url}")
        return resp
    raise ConnectionError(f"{problem} after {retries} retries: {url}")


def _short_error(e: BaseException, limit: int = 300) -> str:
    text = redact(f"{type(e).__name__}: {e}")
    return text if len(text) <= limit else text[:limit - 3] + "..."


def run_collector_safely(collector: Collector, ctx: RunContext) -> SourceResult:
    """S08: call collector.collect(ctx) inside try/except. Any exception ->
    SourceResult(status=FAILED, error=<short message without secrets>). Log start,
    finish, status, item count and duration."""
    from time import monotonic

    log, name = ctx.logger, getattr(collector, "name", type(collector).__name__)
    log.info("source started: %s", name)
    started = monotonic()
    try:
        result = collector.collect(ctx)
        if not isinstance(result, SourceResult):
            raise TypeError(f"collector returned {type(result).__name__}, not SourceResult")
    except Exception as e:  # noqa: BLE001 — one failing source must never stop the others
        result = SourceResult(name, SourceStatus.FAILED, error=_short_error(e))
        log.error("source failed: %s — %s", name, result.error)
        log.debug("traceback for %s:\n%s", name, traceback.format_exc())
    for w in result.warnings:
        log.warning("source warning: %s — %s", name, w)
    log.info("source completed: %s status=%s items=%d duration=%.1fs",
             name, result.status, len(result.items), monotonic() - started)
    return result


GOOGLE_ALERTS_PLATFORMS = (Platform.FACEBOOK, Platform.INSTAGRAM, Platform.LINKEDIN, Platform.PUBLIC_WEB)


def _health_counts(cands: list[Candidate]) -> dict[str, int]:
    return {"filtered_out": sum(1 for c in cands if not c.entity_matches),
            "flagged": sum(1 for c in cands if c.priority is not None),
            "rejected": sum(1 for c in cands if c.verification == Verification.REJECTED)}


def derive_source_health(results: list[SourceResult], candidates: list[Candidate]) -> list[dict]:
    """S08 (uses S13 platform fields): one row per report source:
    'Google Alerts', 'Facebook', 'Instagram', 'LinkedIn', 'Public website', 'Company website',
    'SC Malaysia'. Facebook / Instagram / LinkedIn / Public website are all discovered through
    Google Alerts, so they inherit its status (FAILED if it failed, with error 'Google Alerts
    unavailable'); their 'retrieved' = count of candidates on that platform. 'Company website'
    follows the website collector, 'SC Malaysia' the SC collector (SKIPPED when disabled).
    A FAILED row must never show as SUCCESS with 0 results, and neither may Google Alerts and
    its four platforms when no alert emails arrived: they show NOT CHECKED (D-25)."""
    by_source = {r.source: r for r in results}

    def collector_row(label: str, source: str) -> dict:
        r = by_source.get(source)
        if r is None:                   # disabled, or left out with --sources
            return {"source": label, "status": str(SourceStatus.SKIPPED), "retrieved": 0,
                    "filtered_out": 0, "flagged": 0, "duplicates": 0, "rejected": 0,
                    "error": "not run (disabled or not selected)"}
        cands = [c for c in candidates if c.raw.source == source]
        error = r.error
        if r.status == SourceStatus.SUCCESS and r.warnings:
            error = "Warnings: " + "; ".join(r.warnings)
        return {"source": label, "status": str(r.status), "retrieved": len(r.items),
                "duplicates": max(len(r.items) - len(cands), 0) if r.status == SourceStatus.SUCCESS else 0,
                "error": error, **_health_counts(cands)}

    alerts = collector_row("Google Alerts", "google_alerts")
    no_alert_emails = (alerts["status"] == SourceStatus.SUCCESS and alerts["retrieved"] == 0
                       and "no Google Alerts emails arrived" in (alerts["error"] or ""))
    if no_alert_emails:         # ran, but there was nothing to read: not an all-clear
        alerts["status"] = str(SourceStatus.NOT_CHECKED)
    rows = [alerts]
    for platform in GOOGLE_ALERTS_PLATFORMS:
        cands = [c for c in candidates if c.raw.source == "google_alerts" and c.platform == platform]
        error = alerts["error"]
        if alerts["status"] == SourceStatus.FAILED:
            error = "Google Alerts unavailable"
        elif alerts["status"] == SourceStatus.SKIPPED:
            error = "Google Alerts not run"
        elif no_alert_emails:
            error = "No Google Alerts emails arrived in the period — not checked through Google's index"
        elif alerts["status"] == SourceStatus.SUCCESS:
            error = None
        rows.append({"source": str(platform), "status": alerts["status"], "retrieved": len(cands),
                     "duplicates": 0, "error": error, **_health_counts(cands)})
    rows.append(collector_row("Company website", "website"))
    rows.append(collector_row("SC Malaysia", "sc_malaysia"))
    return rows


# =============================================================================
# S09 — GOOGLE ALERTS QUERY GENERATOR                  (→ collectors/google_alerts.py)
# Docs: SOURCE_ACCESS_BLUEPRINT §2.3
# Goal: `--print-alert-queries` prints the Tier 1 + Tier 2 alert list generated
#       from entities.json + adverse_keywords.json. Never a single-word name query.
#       Split categories so no single query has more than max_terms keywords.
# =============================================================================

ALERT_SCOPES = ("site:facebook.com", "site:instagram.com", "site:linkedin.com", "")


def generate_alert_queries(entities: dict, keywords: dict, max_terms: int = 8) -> list[str]:
    """S09: Tier 1 = every entity alert_name x ALERT_SCOPES -> '<scope> "<name>"'
    (no scope prefix for the whole-web alert).
    Tier 2 = the SHORTEST alert_name of each entity x each keyword-category chunk of at most
    max_terms keywords -> '"<name>" (kw1 OR "kw two" OR ...)'; multi-word keywords quoted.
    Return a de-duplicated list, Tier 1 first. Keep the total well under 100."""
    queries: list[str] = []

    def add(q: str) -> None:
        if q not in queries:
            queries.append(q)

    def quoted(term: str) -> str:
        term = term.replace('"', "").strip()
        return f'"{term}"' if " " in term else term

    alert_names: list[list[str]] = []
    for e in entities.get("entities", []):
        if e.get("approved") is not True:            # alerts only for approved names
            continue
        ambiguous = {t["term"].strip().lower() for t in e.get("ambiguous_terms", [])}
        names = [n.replace('"', "").strip() for n in (e.get("alert_names") or e["names"])]
        names = [n for n in names if n and n.lower() not in ambiguous]   # never a single-word name
        if names:
            alert_names.append(names)

    for names in alert_names:                           # Tier 1: entity x scope
        for name in names:
            for scope in ALERT_SCOPES:
                add(f'{scope} "{name}"'.strip())

    for names in alert_names:                           # Tier 2: entity x keyword-category chunk
        latin = [n for n in names if re.search(r"[A-Za-z]", n)] or names   # English keywords need an English name
        name = min(latin, key=len)
        for category in keywords.get("categories", []):
            terms = [k["term"] for k in category["keywords"]]
            chunks = -(-len(terms) // max_terms)        # ceil: fewest queries, evenly sized
            size = -(-len(terms) // chunks) if chunks else 0
            for i in range(0, len(terms), size or 1):
                add(f'"{name}" (' + " OR ".join(quoted(t) for t in terms[i:i + size]) + ")")
    return queries


# =============================================================================
# S10 — GOOGLE ALERTS COLLECTOR                        (→ collectors/google_alerts.py)
# Docs: SOURCE_ACCESS_BLUEPRINT §2.1–2.4, DATA_COLLECTION_ANALYSIS §2
# Goal: read (never delete/modify) alert emails from the configured label, received
#       from period.start_dt until now; parse each result into a RawItem.
#       IMAP password from NS_IMAP_PASSWORD. RSS mode reads settings feeds instead.
# Save one or two REAL alert emails (with personal data removed) as test fixtures
# once available and add tests for them; the synthetic sample below is a stand-in.
# =============================================================================

ALERTS_METHOD = "Google Alerts (search index)"
_ALERT_SUBJECT_RE = re.compile(r"^\s*Google Alerts?\s*[-–—:]\s*(.+?)\s*$", re.IGNORECASE)
_BLOCK_TAGS = {"div", "p", "td", "tr", "table", "li", "ul", "ol", "br", "h1", "h2", "h3", "h4",
               "h5", "h6", "section", "article", "header", "footer", "blockquote"}
MONTH_ABBR = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def parse_iso_datetime(text: str | None) -> tuple[datetime | None, bool]:
    """ISO 8601 text -> (aware datetime, timezone_was_assumed). A value without a time zone
    (e.g. '2026-09-18') is read as Malaysia time and flagged as assumed. Unparseable -> (None, False).
    A week without a day ('2022-W10') is a date range, not a date, so it is not read as one."""
    text = (text or "").strip()
    if not text or re.fullmatch(r"\d{4}-?W\d{2}", text, re.IGNORECASE):
        return None, False
    text = re.sub(r"[zZ]$", "+00:00", text)
    text = re.sub(r"([+-]\d{2})(\d{2})$", r"\1:\2", text)          # +0800 -> +08:00 (Python 3.10)
    text = re.sub(r"(\.\d{6})\d+", r"\1", text)                    # at most microseconds
    try:
        value = datetime.fromisoformat(text)
    except ValueError:
        return None, False
    if value.tzinfo is None:
        return value.replace(tzinfo=MYT), True
    return value, False


def _query_from_subject(subject: str) -> str | None:
    m = _ALERT_SUBJECT_RE.match(subject or "")
    return m.group(1) if m else (collapse_whitespace(subject) or None)


class _TextCollector(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data):
        self.parts.append(data)


def html_to_text(html_text: str | None) -> str:
    """Plain text of an HTML fragment (tags removed, entities decoded, whitespace collapsed)."""
    parser = _TextCollector()
    parser.feed(html_text or "")
    parser.close()
    return collapse_whitespace(" ".join(parser.parts))


def _alert_html_results(html_text: str) -> list[dict]:
    """Result links of an alert email: [{url, title, publisher, snippet}]. Only google.com/url
    wrappers are results; any other link (share, flag, footer) ends the current result."""

    class Parser(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.results: list[dict] = []
            self.current: dict | None = None
            self.in_result_link = self.in_other_link = self.closed = False
            self.block: list[str] = []
            self.open_blocks: list[tuple[str, str | None]] = []   # (tag, itemprop in force)

        def finish_block(self):
            text = collapse_whitespace("".join(self.block))
            prop = self.open_blocks[-1][1] if self.open_blocks else None
            if text and self.current is not None and not self.closed:
                self.current["blocks"].append((text, prop))
            self.block = []

        def handle_starttag(self, tag, attrs):
            a = dict(attrs)
            if tag == "a":
                self.finish_block()
                href = (a.get("href") or "").strip()
                target = decode_google_redirect(href)
                if href and target != href:
                    if self.current is None or self.current["url"] != target:
                        self.current = {"url": target, "title": [], "blocks": []}
                        self.results.append(self.current)
                        self.closed = False
                    self.in_result_link = True
                else:
                    self.closed, self.in_other_link = True, True
            elif tag == "br":
                self.finish_block()
            elif tag in _BLOCK_TAGS:
                self.finish_block()
                inherited = self.open_blocks[-1][1] if self.open_blocks else None
                prop = a.get("itemprop")
                self.open_blocks.append((tag, prop if prop in ("publisher", "description") else inherited))

        def handle_endtag(self, tag):
            if tag == "a":
                self.in_result_link = self.in_other_link = False
            elif tag in _BLOCK_TAGS:
                self.finish_block()
                if any(t == tag for t, _ in self.open_blocks):
                    while self.open_blocks and self.open_blocks.pop()[0] != tag:
                        pass

        def handle_data(self, data):
            if self.in_result_link and self.current is not None:
                self.current["title"].append(data)
            elif not self.in_other_link:
                self.block.append(data)

    p = Parser()
    p.feed(html_text)
    p.finish_block()
    p.close()
    out = []
    for r in p.results:
        blocks = r["blocks"]
        publisher = next((t for t, prop in blocks if prop == "publisher"), None)
        described = [t for t, prop in blocks if prop == "description"]
        rest = [t for t, prop in blocks if prop is None]
        if publisher is None and not described and len(rest) >= 2 and len(rest[0]) <= 80:
            publisher, rest = rest[0], rest[1:]         # layout: publisher line, then snippet
        out.append({"url": r["url"], "title": collapse_whitespace("".join(r["title"])),
                    "publisher": publisher, "snippet": " ".join(described or rest)})
    return out


def parse_alert_email(raw_message: bytes, received_at: datetime | None = None) -> list[RawItem]:
    """S10: parse one Google Alerts email (MIME, HTML part). For each result link:
      discovered_url = decode_google_redirect(href) (only google.com/url wrappers are results;
                       ignore footer/unsubscribe/alerts-management links),
      title = link text, snippet = the text block that follows, publisher if shown,
      discovery_query = query from the subject 'Google Alert - <query>',
      alert_received_at = received_at (or the Date header), source_message_id = Message-ID,
      source='google_alerts', discovery_method='Google Alerts (search index)'.
    source_published_at stays None (alerts do not state it). Use html.parser from the
    standard library or BeautifulSoup (import inside)."""
    from email import policy
    from email.parser import BytesParser
    from email.utils import parsedate_to_datetime

    msg = BytesParser(policy=policy.default).parsebytes(raw_message)
    body = msg.get_body(preferencelist=("html",))
    if body is None:
        raise ValueError("alert email has no HTML part")
    if received_at is None and msg["Date"]:
        try:
            received_at = parsedate_to_datetime(str(msg["Date"]))
        except (TypeError, ValueError):
            received_at = None
        if received_at is not None and received_at.tzinfo is None:
            received_at = received_at.replace(tzinfo=timezone.utc)     # RFC 5322 '-0000'
    query = _query_from_subject(str(msg["Subject"] or ""))
    message_id = str(msg["Message-ID"]).strip() if msg["Message-ID"] else None
    return [RawItem(source="google_alerts", discovery_method=ALERTS_METHOD, discovered_url=r["url"],
                    title=r["title"], snippet=r["snippet"], discovery_query=query,
                    publisher=r["publisher"], alert_received_at=received_at,
                    source_message_id=message_id)
            for r in _alert_html_results(body.get_content())]


def parse_alert_feed(xml_text: str) -> list[RawItem]:
    """Google Alerts RSS (Atom) fallback. An entry's 'published' time is when Google alerted,
    so it becomes alert_received_at, never the content date."""
    import xml.etree.ElementTree as ET

    ns = {"a": "http://www.w3.org/2005/Atom"}
    root = ET.fromstring(xml_text)
    query = _query_from_subject(html_to_text(root.findtext("a:title", "", ns)))
    items = []
    for entry in root.findall("a:entry", ns):
        link = entry.find("a:link", ns)
        href = (link.get("href") or "") if link is not None else ""
        url = decode_google_redirect(href)
        if not url:
            continue
        alerted, _ = parse_iso_datetime(entry.findtext("a:published", "", ns)
                                        or entry.findtext("a:updated", "", ns))
        items.append(RawItem(source="google_alerts", discovery_method=ALERTS_METHOD, discovered_url=url,
                             title=html_to_text(entry.findtext("a:title", "", ns)),
                             snippet=html_to_text(entry.findtext("a:content", "", ns)),
                             discovery_query=query, alert_received_at=alerted,
                             source_message_id=entry.findtext("a:id", None, ns)))
    return items


SOCIAL_HOST_SUFFIXES = ("facebook.com", "fb.com", "fb.watch", "instagram.com", "linkedin.com")
MAX_PAGE_TEXT_CHARS = 200_000


def _read_page(html_text: str) -> tuple[str, str | None]:
    """(visible text, publication date text) of an HTML page. Date order: article:published_time
    meta, JSON-LD datePublished, first <time datetime>."""
    hidden = {"script", "style", "noscript", "template", "svg", "head"}

    class Parser(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.text: list[str] = []
            self.hidden_depth, self.in_json_ld = 0, False
            self.meta_date = self.time_date = None
            self.json_ld: list[str] = []

        def handle_starttag(self, tag, attrs):
            a = {k: (v or "") for k, v in attrs}
            if tag == "meta" and (a.get("property") or a.get("name")) == "article:published_time":
                self.meta_date = self.meta_date or a.get("content")
            elif tag == "time" and a.get("datetime"):
                self.time_date = self.time_date or a["datetime"]
            if tag == "script" and a.get("type", "").lower() == "application/ld+json":
                self.in_json_ld = True
                self.json_ld.append("")
            if tag == "body":
                self.hidden_depth = 0          # browsers close an unclosed <head> here too
            elif tag in hidden:
                self.hidden_depth += 1
            elif tag in _BLOCK_TAGS:
                self.text.append(" ")

        def handle_endtag(self, tag):
            if tag in hidden and self.hidden_depth:
                self.hidden_depth -= 1
            if tag == "script":
                self.in_json_ld = False

        def handle_data(self, data):
            if self.in_json_ld:
                self.json_ld[-1] += data
            elif not self.hidden_depth:
                self.text.append(data)

    p = Parser()
    p.feed(html_text)
    p.close()
    ld_date = None
    for block in p.json_ld:
        try:
            ld_date = _find_key(json.loads(block), "datePublished")
        except ValueError:
            continue
        if ld_date:
            break
    date_text = p.meta_date or ld_date or p.time_date
    return collapse_whitespace(" ".join(p.text))[:MAX_PAGE_TEXT_CHARS], date_text


def _find_key(node, key: str):
    """First string value of `key` anywhere in parsed JSON (JSON-LD may nest it in @graph)."""
    if isinstance(node, dict):
        if isinstance(node.get(key), str):
            return node[key]
        node = list(node.values())
    if isinstance(node, list):
        for child in node:
            found = _find_key(child, key)
            if found:
                return found
    return None


def enrich_public_web_item(item: RawItem, session, settings: dict, logger: logging.Logger) -> RawItem:
    """S10: for a Google Alerts result whose host is NOT a social platform (SOCIAL_HOST_SUFFIXES),
    and only when settings.sources.google_alerts.fetch_public_web_pages is true: fetch the page
    once with polite_get (robots.txt respected) and fill content_text (visible text),
    source_published_at (article:published_time / JSON-LD datePublished / <time datetime>)
    and final_url. Facebook / Instagram / LinkedIn URLs are NEVER fetched. A failed fetch
    leaves the item unchanged (logged as a warning)."""
    url = item.discovered_url
    if (is_social_url(url) or not url.lower().startswith(("http://", "https://"))
            or not settings["sources"]["google_alerts"].get("fetch_public_web_pages", False)):
        return item
    try:
        resp = polite_get(session, url, settings, logger)
    except Exception as e:  # noqa: BLE001 — blocked, paywalled or down: keep the alert title/snippet
        logger.warning("public web page not read: %s (%s)", url, _short_error(e, 150))
        return item
    content_type = str(resp.headers.get("Content-Type", "text/html")).lower()
    if "html" not in content_type:
        logger.info("public web page skipped (not HTML: %s): %s", content_type, url)
        return item
    text, date_text = _read_page(resp.text)
    published, assumed = parse_iso_datetime(date_text)
    item.content_text = text or None
    item.final_url = getattr(resp, "url", None) or url
    if published is not None:
        item.source_published_at = published
        if assumed:
            item.extra["published_timezone_assumed"] = "Asia/Kuala_Lumpur"
    return item


def _imap_quote(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def _imap_date(d: date) -> str:
    """IMAP SEARCH date, e.g. 13-Sep-2026 (English month names whatever the locale)."""
    return f"{d.day:02d}-{MONTH_ABBR[d.month - 1]}-{d.year}"


def _parse_internaldate(meta: bytes) -> datetime | None:
    m = re.search(rb'INTERNALDATE "\s*(\d{1,2})-([A-Za-z]{3})-(\d{4}) (\d{2}):(\d{2}):(\d{2}) ([+-])(\d{2})(\d{2})"',
                  meta or b"")
    if not m:
        return None
    day, mon, year, hh, mm, ss, sign, oh, om = (g.decode() for g in m.groups())
    offset = timedelta(hours=int(oh), minutes=int(om)) * (-1 if sign == "-" else 1)
    return datetime(int(year), MONTH_ABBR.index(mon.title()) + 1, int(day), int(hh), int(mm), int(ss),
                    tzinfo=timezone(offset))


class GoogleAlertsCollector:
    name = "google_alerts"

    def __init__(self, imap_factory=None, session=None):
        self.imap_factory = imap_factory      # tests pass a fake; default imaplib.IMAP4_SSL
        self.session = session                # tests pass a fake; default build_http_session()

    def collect(self, ctx: RunContext) -> SourceResult:
        """S10: IMAP (imaplib, SSL, readonly select of the label, SEARCH FROM sender SINCE
        date) or RSS (Atom parsed with the standard library's xml.etree) per settings; then
        enrich_public_web_item() for non-social results. stats: emails_read, results, pages_fetched."""
        settings = ctx.config.settings
        ga = settings["sources"]["google_alerts"]
        warnings: list[str] = []
        if ga.get("mode", "imap") == "rss":
            items, stats = self._collect_rss(ctx, warnings)
        else:
            items, stats = self._collect_imap(ctx, warnings)
        items = [i for i in items if i.alert_received_at is None or i.alert_received_at >= ctx.period.start_dt]
        stats["results"] = len(items)
        stats["pages_fetched"] = self._enrich(items, ctx, warnings) if ga.get("fetch_public_web_pages") else 0
        return SourceResult(self.name, SourceStatus.SUCCESS, items=items, warnings=warnings, stats=stats)

    def _collect_imap(self, ctx: RunContext, warnings: list[str]) -> tuple[list[RawItem], dict]:
        import imaplib

        box, log = ctx.config.email["alerts_mailbox"], ctx.logger
        password = get_secret("NS_IMAP_PASSWORD")
        if not password:
            raise RuntimeError(IMAP_NOT_CONFIGURED)
        factory = self.imap_factory or imaplib.IMAP4_SSL
        imap = factory(box["imap_host"], box["imap_port"], timeout=ctx.config.settings["http"]["timeout_seconds"])
        items: list[RawItem] = []
        emails = 0
        try:
            imap.login(box["username"], password)
            typ, _ = imap.select(_imap_quote(box["label"]), readonly=True)     # read-only: nothing changes
            if typ != "OK":
                raise RuntimeError(f"alerts label {box['label']!r} not found in the mailbox")
            since = ctx.period.start_dt.date() - timedelta(days=1)            # IMAP SINCE is date-only
            typ, data = imap.search(None, "FROM", _imap_quote(box["sender_filter"]), "SINCE", _imap_date(since))
            if typ != "OK":
                raise RuntimeError("mailbox search failed")
            for msg_id in (data[0] or b"").decode("ascii", "replace").split():
                typ, parts = imap.fetch(msg_id, "(INTERNALDATE BODY.PEEK[])")    # PEEK: stays unread
                raw = next((p for p in parts or [] if isinstance(p, tuple)), None) if typ == "OK" else None
                if raw is None:
                    warnings.append(f"alert email {msg_id} could not be fetched")
                    continue
                emails += 1
                try:
                    items.extend(parse_alert_email(raw[1], _parse_internaldate(raw[0])))
                except Exception as e:  # noqa: BLE001 — one bad email must not fail the source
                    warnings.append(f"alert email {msg_id} skipped: {_short_error(e, 120)}")
            log.info("google alerts: %d emails read, %d results", emails, len(items))
            if emails == 0:
                warnings.append(NO_ALERT_EMAILS.format(label=box["label"], since=f"{since:%d %b %Y}",
                                                       mailbox=box["username"]))
        finally:
            try:
                imap.logout()
            except Exception:  # noqa: BLE001
                pass
        return items, {"emails_read": emails}

    def _collect_rss(self, ctx: RunContext, warnings: list[str]) -> tuple[list[RawItem], dict]:
        settings = ctx.config.settings
        feeds = settings["sources"]["google_alerts"].get("rss_feeds") or []
        if not feeds:
            raise ValueError("Google Alerts RSS mode is selected but sources.google_alerts.rss_feeds is empty")
        session = self.session or build_http_session(settings)
        items: list[RawItem] = []
        ok = 0
        for feed in feeds:
            try:
                # The owner's own alert feeds are made for feed readers; robots.txt governs crawlers.
                resp = polite_get(session, feed, settings, ctx.logger, check_robots=False)
                items.extend(parse_alert_feed(resp.text))
                ok += 1
            except Exception as e:  # noqa: BLE001
                warnings.append(f"alert feed failed: {_short_error(e, 150)}")
        if not ok:
            raise ConnectionError(f"all {len(feeds)} Google Alerts feeds failed")
        return items, {"feeds_read": ok}

    def _enrich(self, items: list[RawItem], ctx: RunContext, warnings: list[str]) -> int:
        """Read each public web result page once (cached by URL); never social pages."""
        settings = ctx.config.settings
        session = self.session or build_http_session(settings)
        read: dict[str, RawItem] = {}
        unread = 0
        for item in items:
            if is_social_url(item.discovered_url):
                continue
            done = read.get(item.discovered_url)
            if done is None:
                enrich_public_web_item(item, session, settings, ctx.logger)
                read[item.discovered_url] = item
                unread += item.content_text is None
            else:
                item.content_text, item.final_url = done.content_text, done.final_url
                item.source_published_at = done.source_published_at
                item.extra.update(done.extra)
        if unread:
            warnings.append(f"{unread} public web page(s) could not be read (blocked, robots.txt, "
                            "not HTML or unavailable); their alert title and snippet are used instead")
        return len(read) - unread


# =============================================================================
# S11 — COMPANY WEBSITE COLLECTOR                      (→ collectors/website.py)
# Docs: SOURCE_ACCESS_BLUEPRINT §4
# Goal: preferred order feed -> public endpoint -> requests+BeautifulSoup listing pages.
#       Collect title, publication date, URL, text/summary, public comments, source type.
#       First inspect the site once and record the chosen path in settings.json.
#       No acceptance test yet: after inspecting the site, save a trimmed copy of one
#       listing page as a fixture and add @selftest("S11") parser tests BEFORE coding.
# =============================================================================

# Site inspected 29 Sep 2026: WordPress; robots.txt only disallows /wp-admin/; RSS feed at /feed/
# (10 newest posts, older pages via ?paged=N, 404 past the last page); site-wide comments feed at
# /comments/feed/. The feed is the preferred path (SOURCE_ACCESS_BLUEPRINT §4, option 1).
WEBSITE_METHOD = "Direct public retrieval"
_CONTENT_NS = "{http://purl.org/rss/1.0/modules/content/}"
_PAGE_DATE_RE = re.compile(r'<p class="date">(.*?)</p>', re.IGNORECASE | re.DOTALL)


def parse_wordpress_feed(xml_text: str, comments: bool = False) -> list[RawItem]:
    """S11: items of a WordPress RSS 2.0 feed (posts, or comments when comments=True).
    The item pubDate (when the company posted it) is the publication date; a date printed in the
    post itself (e.g. the date of a re-posted article) is kept in extra['page_date_text']."""
    import xml.etree.ElementTree as ET
    from email.utils import parsedate_to_datetime

    root = ET.fromstring(xml_text)
    channel = root.find("channel")
    if root.tag != "rss" or channel is None:
        raise ValueError("not an RSS feed")
    publisher = collapse_whitespace(channel.findtext("title")) or None
    items = []
    for entry in channel.findall("item"):
        link = (entry.findtext("link") or "").strip()
        if not link:
            continue
        try:
            published = parsedate_to_datetime(entry.findtext("pubDate") or "")
            if published.tzinfo is None:
                published = published.replace(tzinfo=timezone.utc)
        except (TypeError, ValueError):
            published = None
        description = entry.findtext("description") or ""
        body = entry.findtext(f"{_CONTENT_NS}encoded") or description
        extra: dict = {}
        page_date = _PAGE_DATE_RE.search(body)
        if page_date and html_to_text(page_date.group(1)):
            extra["page_date_text"] = html_to_text(page_date.group(1))
        categories = [collapse_whitespace(c.text) for c in entry.findall("category") if c.text]
        if categories:
            extra["categories"] = categories
        if comments:
            extra["is_comment"] = True
        items.append(RawItem(source="website", discovery_method=WEBSITE_METHOD, discovered_url=link,
                             title=collapse_whitespace(entry.findtext("title")),
                             snippet=html_to_text(description)[:500], content_text=html_to_text(body) or None,
                             publisher=publisher, source_published_at=published,
                             source_message_id=(entry.findtext("guid") or "").strip() or None, extra=extra))
    return items


class WebsiteCollector:
    name = "website"

    def __init__(self, session=None):
        self.session = session                # tests pass a fake; default build_http_session()

    def collect(self, ctx: RunContext) -> SourceResult:
        """S11: return RawItems (source='website', discovery_method='Direct public retrieval')
        for items published in or near the period; set source_published_at from the page."""
        settings = ctx.config.settings
        site = settings["sources"]["website"]
        if not site.get("feed_url"):
            raise ValueError("sources.website.feed_url is not set (the feed is the retrieval path "
                             "implemented; see SOURCE_ACCESS_BLUEPRINT §4)")
        session = self.session or build_http_session(settings)
        warnings: list[str] = []
        posts, pages = self._read_feed(session, site["feed_url"], ctx, warnings, comments=False)
        comments: list[RawItem] = []
        if site.get("comments_feed_url"):
            try:
                comments, more = self._read_feed(session, site["comments_feed_url"], ctx, warnings, comments=True)
                pages += more
            except Exception as e:  # noqa: BLE001 — comments are a sub-part: warn, do not fail
                warnings.append(f"comments feed not read: {_short_error(e, 150)}")
        for item in posts + comments:
            item.discovery_query = site["comments_feed_url"] if item.extra.get("is_comment") else site["feed_url"]
        return SourceResult(self.name, SourceStatus.SUCCESS, items=posts + comments, warnings=warnings,
                            stats={"pages_fetched": pages, "posts": len(posts), "comments": len(comments)})

    @staticmethod
    def _read_feed(session, feed_url: str, ctx: RunContext, warnings: list[str],
                   comments: bool) -> tuple[list[RawItem], int]:
        """Newest-first feed pages until a post older than the period start appears. Items
        from the current (incomplete) week are kept: S14 excludes them from this report."""
        settings, start = ctx.config.settings, ctx.period.start_dt
        max_pages = settings["sources"]["website"].get("max_feed_pages", 5)
        items: list[RawItem] = []
        for page in range(1, max_pages + 1):
            url = feed_url if page == 1 else f"{feed_url}{'&' if '?' in feed_url else '?'}paged={page}"
            try:
                resp = polite_get(session, url, settings, ctx.logger)
            except HttpStatusError as e:
                if page > 1 and e.status_code == 404:     # WordPress: past the last page
                    return items, page - 1
                raise
            page_items = parse_wordpress_feed(resp.text, comments)
            items += [i for i in page_items if i.source_published_at is None or i.source_published_at >= start]
            if not page_items or any(i.source_published_at and i.source_published_at < start for i in page_items):
                return items, page
        warnings.append(f"stopped after {max_pages} feed pages of {feed_url}; older posts in the period may be missing")
        return items, max_pages


# =============================================================================
# S12 — SC MALAYSIA COLLECTOR                           (→ collectors/securities_commission.py)
# Docs: SOURCE_ACCESS_BLUEPRINT §5, DATA_COLLECTION_ANALYSIS §7 (R1, R6)
# Goal: search the Investor Alert List for entity names every run; read announcement
#       pages for the period. Mark list hits with extra={'sc_alert_list': True}.
#       Confirm exact public URLs first and store them in settings.json.
# =============================================================================

# Pages confirmed 29 Sep 2026; robots.txt allows everything. Both pages embed their complete data
# in the HTML as JSON: the Investor Alert List in `$X.PMD` (entries keyed by a stable id, with
# name / address / website / date / remark / bankInfo, where `date` is only a YEAR), and the
# media releases in `$X.PPG` (records with 'Page Name', 'Page Description', 'Page Address' and
# 'Display Date 2' such as '17 September 2026'). The JSON is read instead of the page layout.
SC_METHOD = "Direct public retrieval"
SC_PUBLISHER = "Securities Commission Malaysia"
# An alert-list entry's `url` is a file SC publishes for that entry (e.g. a fake-certificate
# image). The item's URL stays the list page (a CONTAINER); this file is shown, labelled, in
# the snippet. The entry's `website` is the alerted party's own site, never the item URL.
SC_ITEM_FILE_LABEL = "SC item file: "
SC_BASE_URL = "https://www.sc.com.my/"
MONTH_NAMES = ("january", "february", "march", "april", "may", "june", "july", "august",
               "september", "october", "november", "december")


def _embedded_json(html_text: str, variable: str) -> dict | None:
    """The object assigned to `$X.<variable> = {...}` in a page script, or None."""
    m = re.search(r"\$X\." + re.escape(variable) + r"\s*=\s*", html_text)
    if not m:
        return None
    try:
        value, _ = json.JSONDecoder().raw_decode(html_text, m.end())
    except ValueError:
        return None
    return value if isinstance(value, dict) else None


def _mentions_any(text: str, names: list[str]) -> str | None:
    """First name found in text as whole words, ignoring case and punctuation."""
    haystack = f" {' '.join(tokenize_words(text))} "
    for name in names:
        needle = " ".join(tokenize_words(name))
        if needle and f" {needle} " in haystack:
            return name
    return None


def _day_month_year(text: str) -> date | None:
    """'17 September 2026', '5 Sept 2026' or '15/09/2026' -> date; anything else -> None."""
    text = collapse_whitespace(text)
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if m:
        day, month, year = (int(g) for g in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2}) ([A-Za-z]+),? (\d{4})", text)
        if not m:
            return None
        word = m.group(2).lower()          # full name or an abbreviation of 3+ letters
        month = next((i + 1 for i, name in enumerate(MONTH_NAMES) if len(word) >= 3 and name.startswith(word)), 0)
        day, year = int(m.group(1)), int(m.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _sc_alert_item(name: str, text: str, page_url: str, content_key: str, when: date | None,
                   extra: dict) -> RawItem:
    return RawItem(source="sc_malaysia", discovery_method=SC_METHOD, discovered_url=page_url,
                   title=f"SC Investor Alert List: {name}", snippet=text[:500], content_text=text,
                   publisher=SC_PUBLISHER,
                   source_published_at=datetime.combine(when, time.min, tzinfo=MYT) if when else None,
                   source_message_id=content_key,
                   extra={"sc_alert_list": True, "content_key": content_key, **extra})


def sc_alert_content_key(entry: dict) -> str:
    """Identity of an Investor Alert List entry across runs: its name and website(s), as
    words. The embedded data's own keys are random and change on every page load, so they
    cannot be used (an entry would look new — and CRITICAL — every week)."""
    name = " ".join(ascii_words(str(entry.get("name") or "")))
    website = " ".join(ascii_words(str(entry.get("website") or "")))
    return f"sc-alert-list:{name}" + (f" | {website}" if website else "")


def parse_sc_alert_list(html: str, entity_names: list[str], page_url: str) -> list[RawItem]:
    """S12: return one RawItem per alert-list row whose text contains an entity name
    (case-insensitive). source='sc_malaysia', extra['sc_alert_list']=True,
    source_published_at = the row's 'date added' if present, else None.
    Reads the embedded `$X.PMD` data of the real page; falls back to an HTML table.
    The real list only gives a year, kept in extra['sc_listed_year'] (not a date)."""
    data = _embedded_json(html, "PMD")
    entries = [(key, e) for group in (data or {}).values() if isinstance(group, dict)
               for key, e in group.items() if isinstance(e, dict) and "name" in e]
    if entries:
        items = []
        for _, e in entries:
            item_file = str(e.get("url") or "").strip()     # SC's own file for this entry, when it has one
            parts = [str(e.get("name") or ""), SC_ITEM_FILE_LABEL + item_file if item_file else ""]
            parts += [str(e.get(f) or "") for f in ("address", "website", "remark", "bankInfo")]
            text = collapse_whitespace(" | ".join(p for p in parts if p))
            if _mentions_any(text, entity_names):
                year = str(e.get("date") or "").strip()
                items.append(_sc_alert_item(collapse_whitespace(e["name"]), text, page_url,
                                            sc_alert_content_key(e), None,
                                            {"sc_listed_year": year} if year else {}))
        return items
    return _parse_sc_alert_table(html, entity_names, page_url)


def _parse_sc_alert_table(html_text: str, entity_names: list[str], page_url: str) -> list[RawItem]:
    class Rows(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.rows: list[list[tuple[str, str]]] = []   # [(cell tag, text)]
            self.cell: list[str] | None = None
            self.cell_tag = ""

        def handle_starttag(self, tag, attrs):
            if tag == "tr":
                self.rows.append([])
            elif tag in ("td", "th") and self.rows:
                self.cell, self.cell_tag = [], tag

        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.cell is not None:
                self.rows[-1].append((self.cell_tag, collapse_whitespace("".join(self.cell))))
                self.cell = None

        def handle_data(self, data):
            if self.cell is not None:
                self.cell.append(data)

    p = Rows()
    p.feed(html_text)
    p.close()
    header = next(([t.lower() for _, t in row] for row in p.rows if row and all(k == "th" for k, _ in row)), [])
    date_col = next((i for i, h in enumerate(header) if "date" in h), None)
    items = []
    for row in p.rows:
        if not row or all(k == "th" for k, _ in row):
            continue
        cells = [t for _, t in row]
        text = " | ".join(c for c in cells if c)
        if not _mentions_any(text, entity_names):
            continue
        when = _day_month_year(cells[date_col]) if date_col is not None and date_col < len(cells) else None
        items.append(_sc_alert_item(cells[0], text, page_url,
                                    "sc-alert-list:" + " ".join(ascii_words(cells[0])), when, {}))
    return items


def parse_sc_media_releases(html_text: str, page_url: str) -> list[RawItem]:
    """Media releases embedded in the listing page (`$X.PPG`). The 'Display Date 2' shown on
    the SC page is the publication date (day precision, Malaysia time)."""
    from urllib.parse import urljoin

    data = _embedded_json(html_text, "PPG")
    records = [r for group in (data or {}).values() if isinstance(group, dict)
               for page in (group.get("pages") or {}).values() if isinstance(page, list)
               for r in page if isinstance(r, dict) and r.get("Page Address") and "Display Date 2" in r]
    if not records:
        raise ValueError(f"SC page layout changed: no embedded release list found on {page_url}")
    items = []
    for r in records:
        day = _day_month_year(str(r.get("Display Date 2") or ""))
        items.append(RawItem(source="sc_malaysia", discovery_method=SC_METHOD,
                             discovered_url=urljoin(SC_BASE_URL, str(r["Page Address"]).lstrip("/")),
                             title=collapse_whitespace(str(r.get("Page Name") or "")),
                             snippet=collapse_whitespace(str(r.get("Page Description") or "")),
                             publisher=SC_PUBLISHER,
                             source_published_at=datetime.combine(day, time.min, tzinfo=MYT) if day else None,
                             discovery_query=page_url, source_message_id=r.get("xDataId") or None,
                             extra={"date_precision": "day"} if day else {}))
    return items


def entity_screening_terms(entities: dict) -> list[str]:
    """Every name, variant and ambiguous term of the approved entities. Used to pre-select rows
    from long lists; the entity matcher (S15) makes the real decision later."""
    terms: list[str] = []
    for e in entities.get("entities", []):
        if e.get("approved") is True:
            terms += entity_name_forms(e)
            terms += [t["term"] for t in e.get("ambiguous_terms", [])]
    return list(dict.fromkeys(t for t in terms if t.strip()))


def _first_seen_period(ctx: RunContext, content_key: str) -> date:
    """Period start of the first run that found this item (from the database), else this run's."""
    try:
        if ctx.db_path.exists():
            conn = sqlite3.connect(ctx.db_path)
            try:
                row = conn.execute(
                    "SELECT MIN(fp.period_start) FROM findings f JOIN finding_periods fp"
                    " ON fp.finding_id = f.finding_id WHERE f.url_key = ?",
                    (content_key_url_key(content_key),)).fetchone()
            finally:
                conn.close()
            if row and row[0]:
                return min(date.fromisoformat(row[0]), ctx.period.period_start)
    except (sqlite3.Error, ValueError) as e:
        ctx.logger.warning("first-seen lookup failed (%s); using this run's period", type(e).__name__)
    return ctx.period.period_start


class SecuritiesCommissionCollector:
    name = "sc_malaysia"

    def __init__(self, session=None):
        self.session = session                # tests pass a fake; default build_http_session()

    def collect(self, ctx: RunContext) -> SourceResult:
        """S12: fetch configured pages with polite_get; parse; return RawItems."""
        settings, log = ctx.config.settings, ctx.logger
        sc = settings["sources"]["sc_malaysia"]
        alert_url, release_urls = sc.get("investor_alert_list_url"), sc.get("announcement_urls") or []
        if not alert_url:
            raise ValueError("sources.sc_malaysia.investor_alert_list_url is not set")
        session = self.session or build_http_session(settings)
        warnings: list[str] = []

        # 1. Investor Alert List — the main path: a failure here fails the source.
        page = polite_get(session, alert_url, settings, log)
        if _embedded_json(page.text, "PMD") is None and "<table" not in page.text.lower():
            raise ValueError(f"SC page layout changed: no alert list data found on {alert_url}")
        hits = parse_sc_alert_list(page.text, entity_screening_terms(ctx.config.entities), alert_url)
        for item in hits:
            item.final_url = getattr(page, "url", None) or alert_url      # the page the program read
            if item.source_published_at is None:        # no full date: listed in the week first found
                item.extra["first_seen_period_start"] = _first_seen_period(ctx, item.extra["content_key"]).isoformat()
            item.discovery_query = alert_url

        # 2. Media releases dated from the period start (current-week ones are kept; S14 excludes them).
        releases: list[RawItem] = []
        pages_read = 1
        for url in release_urls:
            try:
                listing = polite_get(session, url, settings, log)
                releases += [r for r in parse_sc_media_releases(listing.text, url)
                             if r.source_published_at is None or r.source_published_at >= ctx.period.start_dt]
                pages_read += 1
            except Exception as e:  # noqa: BLE001 — a sub-part: warn, keep the alert-list result
                warnings.append(f"announcements not read from {url}: {_short_error(e, 150)}")
        if sc.get("fetch_announcement_pages", True):
            for item in releases:
                try:
                    resp = polite_get(session, item.discovered_url, settings, log)
                    item.content_text = _read_page(resp.text)[0] or None
                    item.final_url = getattr(resp, "url", None) or item.discovered_url
                    pages_read += 1
                except Exception as e:  # noqa: BLE001
                    warnings.append(f"release page not read ({_short_error(e, 120)}); its summary is used")
        return SourceResult(self.name, SourceStatus.SUCCESS, items=hits + releases, warnings=warnings,
                            stats={"alert_list_hits": len(hits), "releases": len(releases),
                                   "pages_fetched": pages_read})


def build_collectors(ctx: RunContext) -> list[Collector]:
    """Enabled collectors in fixed order, filtered by --sources. (Wiring only.)"""
    all_collectors: list[Collector] = [GoogleAlertsCollector(), WebsiteCollector(),
                                       SecuritiesCommissionCollector()]
    enabled = ctx.config.settings.get("sources", {})
    chosen = [c for c in all_collectors if enabled.get(c.name, {}).get("enabled", False)]
    if ctx.only_sources is not None:
        chosen = [c for c in chosen if c.name in ctx.only_sources]
    return chosen


# =============================================================================
# S13 — PLATFORM DETECTION & URL VERIFICATION          (→ processing/url_validator.py)
# Docs: DATA_COLLECTION_ANALYSIS §8, SOURCE_ACCESS_BLUEPRINT §3 (pattern tables)
# Goal: platform from hostname ONLY (never invented from the query); content class;
#       content type + url_type from path patterns; VERIFIED / PARTIAL / REJECTED with
#       a verification_basis. Never build or guess a URL.
# =============================================================================

PLATFORM_HOSTS: tuple[tuple[Platform, tuple[str, ...]], ...] = (
    (Platform.FACEBOOK, ("facebook.com", "fb.com", "fb.watch")),
    (Platform.INSTAGRAM, ("instagram.com",)),
    (Platform.LINKEDIN, ("linkedin.com",)),
    (Platform.SC_MALAYSIA, ("sc.com.my",)),
    (Platform.COMPANY_WEBSITE, ("acme.example.com",)),
)
# Shorthand used only by the URL pattern tables below.
_S, _C, _N, _U = UrlType.SPECIFIC_CONTENT, UrlType.CONTAINER, UrlType.NON_EVIDENCE, UrlType.UNKNOWN
IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".bmp")


def _url_parts(url: str) -> tuple[str, list[str], dict[str, str], str]:
    """(host, path segments, first value of each query parameter, fragment)."""
    from urllib.parse import parse_qs, urlsplit
    try:
        p = urlsplit(url)
        host = (p.hostname or "").lower()
    except ValueError:
        return "", [], {}, ""
    return host, [s for s in p.path.split("/") if s], {k: v[0] for k, v in parse_qs(p.query).items()}, p.fragment


def detect_platform(canonical_url: str, source: str) -> Platform:
    """S13: from the URL hostname only: facebook.com / fb.com / fb.watch -> Facebook;
    instagram.com -> Instagram; linkedin.com incl. country subdomains -> LinkedIn;
    sc.com.my -> SC Malaysia; the company's own domain -> Company website; ANY other host ->
    Public website. Never inferred from the alert query (a site:facebook.com alert that
    returns another domain is Public website). `source` is kept for logging only."""
    host = _url_parts(canonical_url)[0]
    for platform, domains in PLATFORM_HOSTS:
        if any(host == d or host.endswith("." + d) for d in domains):   # never 'notfacebook.com'
            return platform
    return Platform.PUBLIC_WEB


def content_class_for(platform: Platform, raw: RawItem) -> ContentClass:
    """S13: Facebook/Instagram/LinkedIn -> SOCIAL; SC Malaysia -> REGULATORY;
    Company website -> COMPANY (comments on it -> SOCIAL if raw.extra['is_comment']);
    Public website -> PUBLIC_WEB."""
    if platform in (Platform.FACEBOOK, Platform.INSTAGRAM, Platform.LINKEDIN):
        return ContentClass.SOCIAL
    if platform == Platform.SC_MALAYSIA:
        return ContentClass.REGULATORY
    if platform == Platform.COMPANY_WEBSITE:
        return ContentClass.SOCIAL if raw.extra.get("is_comment") else ContentClass.COMPANY
    return ContentClass.PUBLIC_WEB


def _classify_facebook(host: str, seg: list[str], q: dict) -> tuple[str, UrlType]:
    first = seg[0].lower() if seg else ""
    if host == "fb.watch" and seg:
        return "Video / reel", _S
    if not seg:
        return "Home page", _N
    if first in ("search", "hashtag", "login", "login.php", "recover", "checkpoint", "signup"):
        return "Search / hashtag / login", _N
    if first == "share":
        return "Share link (not resolved)", _U
    if first in ("permalink.php", "story.php"):
        if q.get("story_fbid"):
            return ("Comment on a post", _S) if q.get("comment_id") else ("Post", _S)
        return "Unknown", _U
    if first in ("photo", "photo.php") and q.get("fbid"):
        return "Photo", _S
    if first == "watch":
        return ("Video / reel", _S) if q.get("v") else ("Video home", _N)
    if first == "reel" and len(seg) > 1:
        return "Video / reel", _S
    if first == "profile.php":
        return "Page / profile", _C
    if first == "groups":
        if len(seg) >= 4 and seg[2] in ("posts", "permalink"):
            return "Group post", _S
        return ("Group home", _C) if len(seg) == 2 else ("Unknown", _U)
    if len(seg) >= 3 and seg[1] == "posts":
        return ("Comment on a post", _S) if q.get("comment_id") else ("Post", _S)
    if len(seg) >= 3 and seg[1] == "videos":
        return "Video / reel", _S
    if len(seg) == 2 and seg[1] == "reviews":
        return "Page reviews tab (not a specific review)", _C
    if len(seg) == 1 or (len(seg) == 2 and seg[1] == "about"):
        return "Page / profile", _C
    return "Unknown", _U


def _classify_instagram(seg: list[str]) -> tuple[str, UrlType]:
    kinds = {"p": "Post", "reel": "Reel", "reels": "Reel", "tv": "Video (IGTV)"}
    if not seg:
        return "Home page", _N
    first = seg[0].lower()
    if first == "explore":
        return ("Hashtag page", _N) if len(seg) > 1 and seg[1] == "tags" else ("Explore page", _N)
    if first in ("accounts", "direct"):
        return "Login / account page", _N
    if first == "stories":
        return "Story (ephemeral)", _S
    for i in (0, 1):                                    # /p/<code> or /<user>/p/<code>
        if len(seg) > i + 1 and seg[i].lower() in kinds:
            return kinds[seg[i].lower()], _S
    return ("Profile", _C) if len(seg) == 1 else ("Unknown", _U)


def _classify_linkedin(seg: list[str]) -> tuple[str, UrlType]:
    if not seg:
        return "Home page", _N
    first = seg[0].lower()
    if first == "posts" and len(seg) > 1:
        return "Individual post", _S
    if first == "feed" and len(seg) > 2 and seg[1] == "update" and seg[2].startswith("urn:li:"):
        return "Individual post", _S
    if first == "pulse" and len(seg) > 1:
        return "Article (LinkedIn Pulse)", _S
    if first in ("company", "showcase") and len(seg) > 1:
        return "Company page", _C
    if first == "in" and len(seg) > 1:
        return "Profile", _C
    if first in ("jobs", "school", "search", "login", "authwall", "signup", "uas", "checkpoint"):
        return "Other — not evidence", _N
    return "Unknown", _U


def _classify_web(platform: Platform, seg: list[str], q: dict, fragment: str) -> tuple[str, UrlType]:
    last = seg[-1].lower() if seg else ""
    if "search" in (s.lower() for s in seg) or ({"q", "s", "query", "search"} & set(q)):
        return "Search page", _N
    if not seg:
        return "Home page", _N
    if last.endswith(IMAGE_EXTENSIONS):
        return "Image", _N
    if last.endswith(".pdf"):
        return "Document (PDF)", _S
    if platform == Platform.SC_MALAYSIA:
        if seg == ["investor-alert-list"]:
            return "Investor Alert List", _C
        if seg[:3] == ["resources", "media", "media-release"]:
            return ("Media release", _S) if len(seg) > 3 else ("Media release listing", _C)
        return "SC page", _U
    if platform == Platform.COMPANY_WEBSITE and fragment.startswith("comment-"):
        return "Comment", _S
    if seg[0].lower() in ("tag", "tags", "category", "categories", "topic", "topics", "author", "page", "feed"):
        return "Listing / tag page", _C
    return ("Article / post", _S) if platform == Platform.COMPANY_WEBSITE else ("Web page", _S)


def classify_url(canonical_url: str, platform: Platform) -> tuple[str, UrlType]:
    """S13: (content_type, url_type) from the pattern tables. Examples:
    Instagram /p/x -> ('Post', SPECIFIC_CONTENT); /explore/tags/x -> ('Hashtag page', NON_EVIDENCE);
    Facebook /<page> -> ('Page / profile', CONTAINER); LinkedIn /company/x -> ('Company page', CONTAINER);
    LinkedIn /posts/..-activity-.. -> ('Individual post', SPECIFIC_CONTENT).
    URL structure is a signal, not proof: verify_url() decides what counts as evidence."""
    host, seg, q, fragment = _url_parts(canonical_url)
    if not host:
        return "Unknown", _U
    if platform == Platform.FACEBOOK:
        return _classify_facebook(host, seg, q)
    if platform == Platform.INSTAGRAM:
        return _classify_instagram(seg)
    if platform == Platform.LINKEDIN:
        return _classify_linkedin(seg)
    return _classify_web(platform, seg, q, fragment)


def verify_url(c: Candidate, entity_found_in_snippet_or_title: bool,
               entity_found_in_fetched_text: bool | None) -> tuple[Verification, str]:
    """S13: returns (status, basis).
    REJECTED: url_type NON_EVIDENCE, or decode failed/empty URL.
    VERIFIED: fetched page text contains the entity (basis 'page text match'); or SOCIAL +
              SPECIFIC_CONTENT + entity in title/snippet (basis 'search-index snippet').
    PARTIAL:  everything else relevant (containers, share links, specific URLs whose snippet
              lacks the entity, fetched page where the text could not be re-located).
    Only judges the URL it is given: it never changes, builds or guesses a URL."""
    host, seg, _, _ = _url_parts(c.canonical_url)
    if not c.canonical_url or not host:
        return Verification.REJECTED, "no usable URL (missing or could not be decoded)"
    if _GOOGLE_HOST_RE.fullmatch(host) and seg == ["url"]:
        return Verification.REJECTED, "decode failed: still a Google redirect link"
    if c.url_type == UrlType.NON_EVIDENCE:
        return Verification.REJECTED, f"url pattern: {c.content_type.lower()} is not evidence"
    if entity_found_in_fetched_text:
        keywords = [m.keyword for m in c.keyword_matches if m.keyword]
        if keywords and not _mentions_any(c.raw.content_text or "", keywords):
            return Verification.PARTIAL, "page text match for the name, keyword not re-located in the page"
        return Verification.VERIFIED, "page text match"
    if entity_found_in_fetched_text is False:
        return Verification.PARTIAL, "page read, but the matched text could not be re-located"
    if c.content_class == ContentClass.SOCIAL and c.url_type == UrlType.SPECIFIC_CONTENT:
        if entity_found_in_snippet_or_title:
            return Verification.VERIFIED, "search-index snippet"
        return Verification.PARTIAL, "specific URL, but the name is not in the search-index snippet"
    if c.url_type == UrlType.CONTAINER:
        return Verification.PARTIAL, f"container URL ({c.content_type.lower()}), not the specific item"
    if c.url_type == UrlType.UNKNOWN:
        return Verification.PARTIAL, f"url pattern only ({c.content_type.lower()})"
    return Verification.PARTIAL, "search-index snippet only; page not read"


# =============================================================================
# S14 — DATE RESOLUTION & PERIOD FILTER                 (→ processing/date_filter.py)
# Docs: DATA_COLLECTION_ANALYSIS §2
# Goal: resolve content_date/date_status/date_basis. NEVER use alert_received_at as the
#       content date. Convert everything to MYT, compare with period.contains().
# =============================================================================

_RELATIVE_DATE_RE = re.compile(r"(\d{1,2}|an?|one)\s+(min(?:ute)?s?|hours?|hrs?|days?|weeks?)\s+ago\b", re.I)
_UNIT_DELTA = {"m": timedelta(minutes=1), "h": timedelta(hours=1), "d": timedelta(days=1), "w": timedelta(weeks=1)}
_MONTH_WORD = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"
_SNIPPET_DATE_RES = (
    re.compile(rf"(?P<day>\d{{1,2}})\s+(?P<month>{_MONTH_WORD}),?\s+(?P<year>\d{{4}})\b", re.I),    # 15 Sep 2026
    re.compile(rf"(?P<month>{_MONTH_WORD})\s+(?P<day>\d{{1,2}}),?\s+(?P<year>\d{{4}})\b", re.I),    # Sep 15, 2026
    re.compile(r"(?P<year>\d{4})-(?P<mon>\d{2})-(?P<day>\d{2})\b"),                               # 2026-09-15
)


def _snippet_date(snippet: str, received: datetime | None) -> tuple[datetime | None, str]:
    """(date, basis) from the START of an alert snippet: '3 days ago …' (relative to when the
    alert was received) or 'Sep 15, 2026 …'. Dates later in the text are usually event dates
    inside the content, not the content's own date, so they are ignored."""
    head = collapse_whitespace(snippet)       # patterns are matched at the start only
    m = _RELATIVE_DATE_RE.match(head)
    if m and received is not None:
        count = 1 if m.group(1).lower() in ("a", "an", "one") else int(m.group(1))
        when = received.astimezone(MYT) - count * _UNIT_DELTA[m.group(2)[0].lower()]
        return when, f"alert snippet '{m.group(0)}' counted back from when the alert was received"
    for pattern in _SNIPPET_DATE_RES:
        m = pattern.match(head)
        if not m:
            continue
        parts = m.groupdict()
        month = int(parts["mon"]) if parts.get("mon") else next(
            i + 1 for i, name in enumerate(MONTH_NAMES) if name.startswith(parts["month"].lower().rstrip(".")[:3]))
        try:
            day = date(int(parts["year"]), month, int(parts["day"]))
        except ValueError:
            return None, ""
        when = datetime.combine(day, time.min, tzinfo=MYT)
        if received is not None and when > received.astimezone(MYT) + timedelta(days=1):
            return None, ""               # a date after the alert arrived cannot be the content date
        return when, f"date '{m.group(0)}' at the start of the alert snippet"
    return None, ""


def _source_date_basis(c: Candidate) -> str:
    raw = c.raw
    if raw.source == "website":
        basis = "Company website feed publication date"
    elif raw.source == "sc_malaysia":
        basis = "date added on the SC Investor Alert List" if raw.extra.get("sc_alert_list") else "date shown on the SC page"
    elif raw.content_text:
        basis = "publication date in the page's own metadata"
    else:
        basis = "publication date stated by the source"
    if raw.extra.get("published_timezone_assumed"):
        basis += " (no time zone given; read as Malaysia time)"
    return basis


def _first_seen_week(c: Candidate) -> ReportingPeriod | None:
    """The reporting week of the run that first found an item whose source gives no full date
    (raw.extra['first_seen_period_start'], set by S12), else None."""
    first_seen = c.raw.extra.get("first_seen_period_start")
    if not first_seen:
        return None
    return calculate_previous_completed_week(date.fromisoformat(first_seen) + timedelta(days=7))


def resolve_content_date(c: Candidate) -> None:
    """S14: priority 1–2 source_published_at -> VERIFIED; 3 LinkedIn activity-ID date ->
    DERIVED (only after the decoding is validated; until then skip); 4 explicit or relative
    date in the snippet ('3 days ago' relative to alert_received_at, '15 Sep 2026') ->
    APPROXIMATE; otherwise UNKNOWN. (Public web pages fetched in S10 supply source_published_at.)
    content_date is always the event / content date, at the precision the source states it
    (date_precision) — never the date the system found the item. SC list entries without a
    full date: the year the list states -> VERIFIED, YEAR precision (1 Jan of that year is only
    the start of the range, never shown as a day); no year -> FIRST_SEEN with no content date.
    Both are listed in the week of the run that first found them (period_position), so they are
    reported in that week only. alert_received_at is never used as the content date."""
    raw = c.raw
    if raw.source_published_at is not None:
        c.content_date, c.date_precision = raw.source_published_at.astimezone(MYT), DatePrecision.DAY
        c.date_status, c.date_basis = DateStatus.VERIFIED, _source_date_basis(c)
        return
    # Priority 3 (LinkedIn activity-ID dates) is deliberately not used until the decoding is
    # validated against real posts (DATA_COLLECTION_ANALYSIS §2.1).
    when, basis = _snippet_date(raw.snippet, raw.alert_received_at)
    if when is not None:
        c.content_date, c.date_status, c.date_basis = when, DateStatus.APPROXIMATE, basis
        c.date_precision = DatePrecision.DAY
        return
    week = _first_seen_week(c)
    if week is not None:
        listed = f"listed in the week the system first found it ({week.label()})"
        year = str(raw.extra.get("sc_listed_year") or "").strip()
        if re.fullmatch(r"\d{4}", year):
            c.content_date, c.date_precision = datetime(int(year), 1, 1, tzinfo=MYT), DatePrecision.YEAR
            c.date_status = DateStatus.VERIFIED
            c.date_basis = f"the SC list gives only the year {year} (no day or month); {listed}"
        else:
            c.content_date, c.date_precision, c.date_status = None, None, DateStatus.FIRST_SEEN
            c.date_basis = f"the source gives no date; {listed}"
        return
    c.content_date, c.date_precision, c.date_status = None, None, DateStatus.UNKNOWN
    c.date_basis = "no date stated by the source" + (
        " (the alert email's arrival time is not the content date)" if raw.alert_received_at else "")


def _date_range_end(when: datetime, precision: DatePrecision | str | None) -> datetime:
    """Exclusive end of the time a date of this precision stands for: the next year / month
    for a YEAR / MONTH date; the moment itself for a full date."""
    when = when.astimezone(MYT)
    if precision == DatePrecision.YEAR:
        return datetime(when.year + 1, 1, 1, tzinfo=MYT)
    if precision == DatePrecision.MONTH:
        return datetime(when.year + when.month // 12, when.month % 12 + 1, 1, tzinfo=MYT)
    return when


def period_position(c: Candidate, period: ReportingPeriod) -> str:
    """S14: 'IN' | 'BEFORE' | 'AFTER' | 'UNKNOWN' using the MYT content_date. An item whose
    source gives no full date but which records the week it was first found (SC list) is placed
    by that week, so it is reported once, not every week. Any other year / month-only date is
    longer than a week, so if it overlaps the period it is UNKNOWN (manual review), never IN."""
    partial = c.date_precision in (DatePrecision.YEAR, DatePrecision.MONTH)
    first_seen = _first_seen_week(c)
    if first_seen is not None and (c.content_date is None or partial):
        when = first_seen.start_dt
    elif c.content_date is None:
        return "UNKNOWN"
    elif partial:
        start = c.content_date.astimezone(MYT)
        if _date_range_end(start, c.date_precision) <= period.start_dt:
            return "BEFORE"
        return "AFTER" if start >= period.end_dt_exclusive else "UNKNOWN"
    else:
        when = c.content_date.astimezone(MYT)
    if when < period.start_dt:
        return "BEFORE"
    if when >= period.end_dt_exclusive:
        return "AFTER"
    return "IN"


def historical_status(content_date: datetime | None, precision: DatePrecision | str | None,
                      period: ReportingPeriod) -> HistoricalStatus:
    """CURRENT / HISTORICAL / DATE UNKNOWN for an event / content date and the reporting week.
    HISTORICAL only when the whole date (all of its year or month, if that is all the source
    states) is before the week; a year / month that includes the week cannot be shown to be older."""
    if content_date is None:
        return HistoricalStatus.DATE_UNKNOWN
    if precision in (DatePrecision.YEAR, DatePrecision.MONTH):
        old = _date_range_end(content_date, precision) <= period.start_dt
    else:
        old = content_date.astimezone(MYT) < period.start_dt
    return HistoricalStatus.HISTORICAL if old else HistoricalStatus.CURRENT


def record_age_text(content_date: datetime | None, precision: DatePrecision | str | None,
                    period: ReportingPeriod) -> str:
    """Age of a HISTORICAL event / content date at the end of the reporting week, e.g. 'approx.
    4 years' for a year-only 2022 date in a 2026 week, '12 days' for a full date. Never more
    precise than the date; '' when the item is not HISTORICAL."""
    if content_date is None or historical_status(content_date, precision, period) != HistoricalStatus.HISTORICAL:
        return ""
    when, end = content_date.astimezone(MYT).date(), period.period_end
    if precision == DatePrecision.YEAR:
        years = end.year - when.year
        return f"approx. {years} year{'s' if years != 1 else ''}"
    months = (end.year - when.year) * 12 + end.month - when.month
    if precision != DatePrecision.MONTH:
        days = (end - when).days
        if days < 60:
            return f"{days} day{'s' if days != 1 else ''}"
        if end.day < when.day:                  # complete months only
            months -= 1
    if months < 24:
        return f"approx. {months} month{'s' if months != 1 else ''}"
    return f"approx. {(months + 6) // 12} years"            # to the nearest year


# =============================================================================
# S15 — ENTITY MATCHER                                   (→ processing/entity_matcher.py)
# Docs: DATA_COLLECTION_ANALYSIS §3
# Goal: whole-word, case-insensitive, punctuation-tolerant matching of names/variants;
#       ambiguous terms only with context inside the window. Search title + snippet +
#       content_text + URL path. Only approved entities (load_config already filtered).
# =============================================================================

# Chinese characters (CJK Unified Ideographs, Extension A, Compatibility, Extension B).
_CJK_CHARS = "㐀-䶿一-鿿豈-﫿\U00020000-\U0002a6df"
_TOKEN_RE = re.compile(rf"[{_CJK_CHARS}]|[0-9a-z]+")


def tokenize_words(text: str) -> list[str]:
    """Shared word tokenizer for entity positions and keyword proximity (S15, S16).
    Latin letters/digits form words; each Chinese character is one token, because Chinese
    has no spaces between words. A Chinese alias therefore matches as an exact character
    sequence — no segmentation, translation or other NLP. Proximity in Chinese text is
    counted in characters."""
    return _TOKEN_RE.findall((text or "").lower())


# Web addresses inside text: an email address, 'https://…', 'www.…' or a bare host with a common
# top-level domain ('example-fund.com'). Used only to keep an AMBIGUOUS term (one word of a name) from being confirmed by
# the words of a web address — see web_address_positions().
_URL_TLDS = ("com|net|org|info|biz|io|co|my|sg|hk|cn|tw|asia|xyz|top|vip|app|site|online|club|ltd|"
             "me|cc|us|uk|ai|gov|edu")
_URL_SPAN_RE = re.compile(
    r"[0-9a-z._%+-]+@(?:[0-9a-z-]+\.)+[a-z]{2,}"                     # email address
    rf"|(?:https?://|www\.)[^\s<>\"'|{_CJK_CHARS}]+"
    rf"|(?<![0-9a-z.@-])(?:[0-9a-z](?:[0-9a-z-]*[0-9a-z])?\.)+(?:{_URL_TLDS})(?![0-9a-z-])"
    rf"(?:/[^\s<>\"'|{_CJK_CHARS}]*)?")


def web_address_positions(text: str) -> set[int]:
    """Indexes (in tokenize_words(text)) of the words that are part of a web address, e.g. the
    'example' and 'fund' of 'https://www.example-fund.com/'. A domain name is not a mention:
    an SC Investor Alert List entry for an unrelated scheme with a look-alike domain was matched
    to the company because one word of its name and the context word 'fund' both came from
    that one address (report for 28 Sep – 4 Oct 2026)."""
    lowered = (text or "").lower()
    spans = [m.span() for m in _URL_SPAN_RE.finditer(lowered)]
    if not spans:
        return set()
    return {i for i, m in enumerate(_TOKEN_RE.finditer(lowered))
            if any(start <= m.start() < end for start, end in spans)}


def ascii_words(text: str) -> list[str]:
    """The original ASCII-only word split. Kept for identity keys stored in SQLite (SC alert-list
    entries) so that existing entries keep their keys and are not reported again as new."""
    return re.findall(r"[0-9a-z]+", (text or "").lower())


MAX_MENTIONS_PER_ENTITY = 100


def find_phrase(tokens: list[str], phrase: list[str]) -> list[int]:
    """Start index of every occurrence of a word sequence in a word list."""
    n = len(phrase)
    if not n:
        return []
    return [i for i in range(len(tokens) - n + 1) if tokens[i] == phrase[0] and tokens[i:i + n] == phrase]


def _word_gap(i: int, n: int, j: int, m: int) -> int:
    """Words between span [i, i+n) and span [j, j+m); 0 when they touch or overlap."""
    if j >= i + n:
        return j - (i + n - 1)
    if i >= j + m:
        return i - (j + m - 1)
    return 0


def entity_name_forms(entity: dict) -> list[str]:
    """Every approved way of writing one entity: names, variants and aliases (e.g. a Chinese
    name). Only what entities.json lists — nothing is inferred or translated."""
    return list(entity.get("names", [])) + list(entity.get("variants", [])) + list(entity.get("aliases", []))


LOOK_ALIKE_NOTE = "only in a web address — possible look-alike"


def _entity_mentions(tokens: list[str], entity: dict,
                     web_words: set[int] | None = None) -> list[tuple[int, str, bool]]:
    """(position, matched text, look_alike) of every mention of one entity in a word list: names
    and variants anywhere (longest form wins where forms overlap); ambiguous terms only with one
    of their context phrases inside the context window.
    web_words: positions that belong to a web address (web_address_positions). An ambiguous term
    confirmed only through a web address — the term itself sits in an address such as
    a look-alike domain, or its only context word does — is a LOOK-ALIKE mention (look_alike=True,
    reported at REVIEW priority by R9, never CRITICAL; PROJECT_STATE O-12 / D-26). The company's own
    domain listed as a context is a normal mention. Full names, variants and
    aliases are never look-alikes."""
    web_words = web_words or set()

    def in_web(j: int, m: int) -> bool:
        return any(k in web_words for k in range(j, j + m))

    spans: list[tuple[int, int, str, bool]] = []    # (start, length, matched text, look_alike)

    def free(i: int, n: int) -> bool:
        return not any(s < i + n and i < s + length for s, length, _, _ in spans)

    forms = list(dict.fromkeys(entity_name_forms(entity)))
    for form in sorted(forms, key=lambda f: -len(tokenize_words(f))):
        words = tokenize_words(form)
        spans += [(i, len(words), form, False) for i in find_phrase(tokens, words) if free(i, len(words))]
    for amb in entity.get("ambiguous_terms", []):
        words, n, window = tokenize_words(amb["term"]), len(tokenize_words(amb["term"])), amb["context_window_words"]
        contexts = [(ctx, j, len(cw)) for ctx in amb["requires_any_context"]
                    for cw in [tokenize_words(ctx)] for j in find_phrase(tokens, cw)]
        for i in find_phrase(tokens, words):
            if not free(i, n):
                continue
            term_in_web = in_web(i, n)
            strong = next((ctx for ctx, j, m in contexts if "." in ctx and j <= i and i + n <= j + m), None)
            if strong is None and not term_in_web:
                strong = next((ctx for ctx, j, m in contexts if _word_gap(i, n, j, m) <= window
                               and ("." in ctx or not in_web(j, m))), None)
            if strong is not None:
                spans.append((i, n, f"{amb['term']} (with context '{strong}')", False))
                continue
            weak = next((ctx for ctx, j, m in contexts if _word_gap(i, n, j, m) <= window), None)
            if weak is not None:
                spans.append((i, n, f"{amb['term']} ({LOOK_ALIKE_NOTE}; context '{weak}')", True))
    return sorted((s, text, look) for s, _, text, look in spans)[:MAX_MENTIONS_PER_ENTITY]


def match_entities(text: str, url: str, entities: dict) -> list[Match]:
    """S15: return Match(entity_id, matched_text) per distinct entity found in `text`
    or in the path of `url` (e.g. /company/<company-name>; '-' treated as space).
    Set Match.position = index in tokenize_words(text) of the match's first word
    (None when the match is only in the URL).
    One Match is returned per MENTION (in text order), so that keyword proximity can use the
    nearest mention (DATA_COLLECTION_ANALYSIS §6); {m.entity_id for m in ...} gives the entities.
    Only entities with "approved": true are used."""
    from urllib.parse import unquote, urlsplit

    tokens = tokenize_words(text or "")
    web_words = web_address_positions(text or "")
    try:
        path_tokens = tokenize_words(unquote(urlsplit(url).path)) if url else []
    except ValueError:
        path_tokens = []
    path_words = set(range(len(path_tokens)))       # every word of a URL path is part of a web address
    matches: list[Match] = []
    for e in entities.get("entities", []):
        if e.get("approved") is not True:
            continue
        mentions = _entity_mentions(tokens, e, web_words)
        if mentions:
            matches += [Match(e["id"], matched, position=pos, look_alike=look) for pos, matched, look in mentions]
        else:
            in_url = _entity_mentions(path_tokens, e, path_words)
            if in_url:
                matches.append(Match(e["id"], f"{in_url[0][1]} (in URL)", position=None, look_alike=in_url[0][2]))
    return matches


def approved_relationships(entities: dict) -> list[dict]:
    """The client-approved links between approved entities, exactly as listed under
    entities.json "relationships" (e.g. founder associated_with company,
    label 'Founder of'). Nothing is inferred from text, names or proximity."""
    ids = {e["id"] for e in entities.get("entities", []) if e.get("approved") is True}
    return [r for r in entities.get("relationships", [])
            if r.get("approved") is True and r.get("subject") in ids and r.get("object") in ids]


def linked_relationships(entity_ids: set[str], relationships: list[dict]) -> list[dict]:
    """Approved relationships whose subject AND object are both named in one item. A
    relationship never adds an entity the item does not name, and naming both does not prove
    that the item is about their relationship — such findings go to human review (R8)."""
    return [r for r in relationships if r["subject"] in entity_ids and r["object"] in entity_ids]


def describe_entity_link(entity_ids: set[str], entities: dict) -> dict[str, str]:
    """S18 report fields for the entities an item names: Finding Type, Matched Entity,
    Associated Entity, Relationship. Rule-based: from the matched ids and the approved list."""
    by_id = {e["id"]: e for e in entities.get("entities", [])}
    name = lambda i: by_id.get(i, {}).get("display_name") or i      # noqa: E731
    matched = ", ".join(name(i) for i in sorted(entity_ids))
    links = linked_relationships(entity_ids, approved_relationships(entities))
    if links:
        return {"type": "Relationship-linked finding", "matched": matched,
                "associated": ", ".join(dict.fromkeys(name(r["object"]) for r in links)),
                "relationship": "; ".join(
                    f"{name(r['subject'])} → {r['label']} → {name(r['object'])} (approved relationship). "
                    f"Both are named in this item; whether the keyword concerns {name(r['subject'])} "
                    "needs human review." for r in links)}
    if len(entity_ids) == 1:
        e = by_id.get(next(iter(entity_ids)), {})
        kind = "Company" if e.get("type") == "company" else (e.get("role") or "Person")
        return {"type": f"{kind} finding", "matched": matched, "associated": "", "relationship": ""}
    if entity_ids:
        return {"type": "Several entities — no approved relationship", "matched": matched,
                "associated": "", "relationship": ""}
    return {"type": "", "matched": "", "associated": "", "relationship": ""}


# =============================================================================
# S16 — KEYWORD SCANNER                                  (→ processing/keyword_scanner.py)
# Docs: DATA_COLLECTION_ANALYSIS §4
# Goal: whole-word/phrase, case-insensitive matching ('aml' not in 'Hamlet', 'fined' not
#       in 'refined'); record category, priority, ambiguous flag and word distance to the
#       nearest entity mention (proximity_words).
# =============================================================================

def scan_keywords(text: str, keywords: dict, entity_positions: list[int | None] | None = None) -> list[Match]:
    """S16: return one Match per keyword found (entity_id left '' — caller fills it).
    proximity_words = smallest |keyword word index - entity position| over entity_positions,
    using tokenize_words(); None when entity_positions is empty/None.
    Whole words / phrases only; a keyword's optional "variants" (e.g. 'fraudulent' for 'fraud')
    count as that keyword. Results are in the order the keywords first appear in the text."""
    tokens = tokenize_words(text or "")
    positions = [p for p in (entity_positions or []) if p is not None]
    found: list[tuple[int, Match]] = []
    for category in keywords.get("categories", []):
        for kw in category["keywords"]:
            hits: list[tuple[int, str]] = []                  # (word index, form found)
            for form in [kw["term"]] + list(kw.get("variants", [])):
                hits += [(i, form) for i in find_phrase(tokens, tokenize_words(form))]
            if not hits:
                continue
            hits.sort()
            proximity = (min(abs(i - p) for i, _ in hits for p in positions) if positions else None)
            found.append((hits[0][0], Match(entity_id="", matched_text=hits[0][1], keyword=kw["term"],
                                            category=category["name"], priority=category["priority"],
                                            ambiguous=bool(kw.get("ambiguous")), proximity_words=proximity)))
    return [m for _, m in sorted(found, key=lambda pm: pm[0])]


# =============================================================================
# S17 — FLAGGING RULES & REPORT SECTIONS                 (→ processing/flagging.py)
# Docs: DATA_COLLECTION_ANALYSIS §6–7 (rules R1–R7), §2.2, §8.3
# Goal: RULE-BASED ONLY (approved name + adverse keyword). No NLP / machine learning —
#       excluded by the project owner's decision. Set priority, rules_applied,
#       review_status; then the report section. Highest priority wins.
#       The program never sets CONFIRMED ADVERSE / FALSE POSITIVE.
# =============================================================================

def apply_flagging_rules(c: Candidate, window_words: int, relationships: list[dict] | None = None) -> None:
    """S17: implement R1–R9 exactly as documented:
    R1 SC Investor Alert List hit (raw.extra['sc_alert_list']) + entity -> CRITICAL / MANUAL REVIEW
    R2 entity + unambiguous CRITICAL keyword NEAR the name -> CRITICAL / MANUAL REVIEW
    R3 entity + ambiguous CRITICAL keyword NEAR the name -> HIGH / NEEDS REVIEW
    R4 entity + unambiguous non-critical keyword NEAR the name -> the category's priority / PENDING
    R5 entity + any keyword that is NOT near the name (D-30, owner 5 Oct 2026: an article that
       quotes the founder and mentions 'fraud' about someone else is not a HIGH finding), or any
       ambiguous non-critical keyword -> REVIEW / NEEDS REVIEW
    R6 SC Malaysia regulatory page (content_class REGULATORY) naming an entity, no keyword
       -> HIGH / MANUAL REVIEW
    R7 anything else (name without keyword, keyword without name) -> priority None (not reported)
    R8 a flagged item names both ends of an approved relationship (approved_relationships) ->
       priority unchanged, review at least NEEDS REVIEW: the program cannot tell which of them
       the keyword concerns, so a person decides.
    R9 a flagged item whose ONLY entity mentions are look-alikes (Match.look_alike: an ambiguous
       term confirmed only through a web address, e.g. an SC list entry using a look-alike domain)
       -> priority capped at REVIEW, review NEEDS REVIEW: possibly an impersonation, not news
       about the entity (PROJECT_STATE O-12 / D-26). R9 never makes an item reported by itself.
    'Near' = proximity_words is None or <= window_words. Review status follows the strongest
    rule that fired (MANUAL REVIEW > NEEDS REVIEW > PENDING). Record every rule in rules_applied."""
    fired: list[tuple[str, Priority, ReviewStatus]] = []
    if c.entity_matches:
        if c.raw.extra.get("sc_alert_list"):
            fired.append(("R1", Priority.CRITICAL, ReviewStatus.MANUAL_REVIEW))
        for k in c.keyword_matches:
            critical = k.priority == Priority.CRITICAL
            near = k.proximity_words is None or k.proximity_words <= window_words
            if critical and not k.ambiguous and near:
                fired.append(("R2", Priority.CRITICAL, ReviewStatus.MANUAL_REVIEW))
            elif critical and near:
                fired.append(("R3", Priority.HIGH, ReviewStatus.NEEDS_REVIEW))
            elif not k.ambiguous and near:
                fired.append(("R4", Priority(k.priority), ReviewStatus.PENDING))
            else:                                   # far from the name (D-30), or ambiguous non-critical
                fired.append(("R5", Priority.REVIEW, ReviewStatus.NEEDS_REVIEW))
        if c.content_class == ContentClass.REGULATORY and not c.keyword_matches:
            fired.append(("R6", Priority.HIGH, ReviewStatus.MANUAL_REVIEW))
    review_rank = {ReviewStatus.PENDING: 0, ReviewStatus.NEEDS_REVIEW: 1, ReviewStatus.MANUAL_REVIEW: 2}
    if not fired:                                   # R7: name without keyword, or keyword without name
        c.priority, c.review_status, c.rules_applied = None, ReviewStatus.PENDING, ["R7"]
        return
    c.priority = max((p for _, p, _ in fired), key=lambda p: PRIORITY_RANK[p])
    if linked_relationships({m.entity_id for m in c.entity_matches}, relationships or []):
        fired.append(("R8", c.priority, ReviewStatus.NEEDS_REVIEW))     # attribution needs a person
    c.review_status = max((r for _, _, r in fired), key=lambda r: review_rank[r])
    if all(m.look_alike for m in c.entity_matches):     # R9: name only via a web address
        if PRIORITY_RANK[c.priority] > PRIORITY_RANK[Priority.REVIEW]:
            c.priority = Priority.REVIEW
        c.review_status = ReviewStatus.NEEDS_REVIEW
        fired.append(("R9", c.priority, ReviewStatus.NEEDS_REVIEW))
    c.rules_applied = sorted({rule for rule, _, _ in fired})


def assign_section(c: Candidate, period: ReportingPeriod) -> Section:
    """S17: no entity match or priority None -> NOT_REPORTED;
    date AFTER -> EXCLUDED_FUTURE; BEFORE -> OUT_OF_PERIOD;
    REJECTED -> UNVERIFIED_URL; UNKNOWN date -> DATE_UNKNOWN (review MANUAL REVIEW);
    otherwise MAIN."""
    if not c.entity_matches or c.priority is None:
        return Section.NOT_REPORTED
    position = period_position(c, period)
    if position == "AFTER":
        return Section.EXCLUDED_FUTURE
    if position == "BEFORE":
        return Section.OUT_OF_PERIOD
    if c.verification == Verification.REJECTED:
        return Section.UNVERIFIED_URL
    if position == "UNKNOWN":
        c.review_status = ReviewStatus.MANUAL_REVIEW
        return Section.DATE_UNKNOWN
    return Section.MAIN


# =============================================================================
# S18 — EXCEL REPORT                                      (→ reporting/excel_report.py)
# Docs: FINAL_ARCHITECTURE §8 (sheets, columns, summary contents)
# Goal: build reports/Negative_Screening_<period_start>.xlsx FROM SQLITE ONLY; archive
#       an existing file for the same week to reports/archive/ first; dry runs go to
#       reports/dry_run/. Summary must include the coverage statement and source health.
# =============================================================================

COVERAGE_STATEMENT = ("This report covers content surfaced by the sources listed below during the "
                      "reporting period. It is not a complete record of all Internet mentions.")
SHEET_SUMMARY = "Summary"
SHEET_FINDINGS = "Findings"
SHEET_DATE_UNKNOWN = "Date Unknown — Manual Review"
SHEET_UNVERIFIED = "Unverified URLs — Manual Lookup"
SHEET_OUT_OF_PERIOD = "Out of Period"
FINDINGS_COLUMNS = ["Report Week", "Event / Content Date", "Date Status", "Historical Status", "Record Age",
                    "Discovery Date", "Platform", "Content Type", "Finding Type",
                    "Matched Entity", "Matched Name / Alias", "Associated Entity", "Relationship",
                    "Keyword", "Keyword Category", "Priority", "Title / Comment / Snippet",
                    "URL", "URL Type", "URL Verification", "Discovered URL", "Fetched URL",
                    "Discovery Method", "Manual Review Status", "Notes"]


MODE_LABELS = {Mode.AUTOMATIC: "Automatic", Mode.MANUAL: "Manual",
               Mode.MANUAL_TEST_DATE: "Manual (test reference date)"}
FAILED_SOURCE_NOTE = ("Absence of findings from this source does not mean there is no adverse content: "
                      "the source could not be checked.")
VALIDATION_NOTE = ("Yes — VALIDATION DATA from the staging database (data/screening_validation.sqlite), "
                   "not production results; not emailed (an email preview is saved beside this report)")
SCREENING_NOTE = "Keyword matches are screening aids, not proof. Every finding needs human verification."
COUNT_LABELS = [
    ("discovered", "Total discovered candidates"),
    ("duplicates_merged", "Duplicates merged"),
    ("relevant", "Relevant (entity-matched) candidates"),
    ("potentially_adverse", "Potentially adverse items"),
    ("critical", "Critical items"),
    ("manual_review", "Manual-review items"),
    ("date_unknown", "Items with unknown date"),
    ("unverified_urls", "Unverified URLs"),
    ("historical", "Reported items dated before the period (HISTORICAL — older records, not new incidents)"),
    ("dated_before_period", "Surfaced this week but dated before the period (listed on the Out of Period sheet)"),
    ("already_reported", "Found again but already reported in an earlier week (not repeated, not counted for review)"),
]
DATE_FIELD_NOTES = [        # Summary 'Dates in this report': what each date column of the findings sheets means
    ("Report Week", "The Monday–Sunday week screened by this run. An item is listed in the week it was found or "
                    "listed; that does not mean it happened in this week."),
    ("Event / Content Date", "When the underlying incident or source record is dated, as the source states it and "
                             "only as precisely as the source states it ('2022' when only the year is given). "
                             "UNKNOWN when no reliable date is available. Never the date the system found the item."),
    ("Date Status", "How reliable and precise the Event / Content Date is. VERIFIED: stated by the source "
                    "('(year only)' / '(month only)' when the source gives no day). APPROXIMATE: from a search-result "
                    "snippet. FIRST_SEEN: the source gives no event date, so the item is listed in the week the "
                    "system first found it (see Discovery Date). UNKNOWN: no reliable date."),
    ("Historical Status", "CURRENT: dated in the reporting week (or in its year / month when only that is known). "
                          "HISTORICAL: dated before the reporting week — an older record found or listed this week, "
                          "not a new incident. DATE UNKNOWN: no reliable Event / Content Date; check the source."),
    ("Record Age", "For HISTORICAL items: how long before the end of the reporting week the Event / Content Date "
                   "is (approximate when only the year or month is known)."),
    ("Discovery Date", "When this system first found the item (its first screening run). Not the event date."),
]
SHEET_FOR_SECTION = {Section.MAIN: SHEET_FINDINGS, Section.DATE_UNKNOWN: SHEET_DATE_UNKNOWN,
                     Section.UNVERIFIED_URL: SHEET_UNVERIFIED, Section.OUT_OF_PERIOD: SHEET_OUT_OF_PERIOD}
COLUMN_WIDTHS = {"Report Week": 24, "Event / Content Date": 17, "Date Status": 20, "Historical Status": 15,
                 "Record Age": 15, "Discovery Date": 17, "Platform": 16,
                 "Content Type": 16, "Finding Type": 22, "Matched Entity": 30, "Matched Name / Alias": 30,
                 "Associated Entity": 30, "Relationship": 50, "Keyword": 22, "Keyword Category": 26,
                 "Priority": 10, "Title / Comment / Snippet": 60, "URL": 50, "URL Type": 18,
                 "Discovered URL": 45, "Fetched URL": 45,
                 "URL Verification": 16, "Discovery Method": 32, "Manual Review Status": 20,
                 "Notes": 60}
EXCEL_CELL_LIMIT = 32767                     # longest text an Excel cell can hold
CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # not allowed in xlsx XML


def screened_targets(entities: dict) -> list[dict]:
    """The approved targets this run screened: display name, kind (company / person and
    role) and every name and variant searched. Shown in the report, the email and the app."""
    targets = []
    for e in entities.get("entities", []):
        if e.get("approved") is not True:
            continue
        kind = str(e.get("type") or "").capitalize() or "Target"
        if e.get("role"):
            kind += f" ({e['role']})"
        names = list(dict.fromkeys(entity_name_forms(e)))
        targets.append({"name": e.get("display_name") or e["id"], "kind": kind, "names": names})
    return targets


def targets_text(entities: dict) -> str:
    """'<company> (Company); <founder> (Person (Founder))'."""
    return "; ".join(f"{t['name']} ({t['kind']})" for t in screened_targets(entities)) or "none approved"


def collector_rows(source_health: list[dict]) -> list[dict]:
    """The health rows of the sources actually run (Google Alerts, Company website,
    SC Malaysia). Facebook / Instagram / LinkedIn / Public website rows only repeat
    Google Alerts' status, so one mailbox failure counts as one failed source, not five."""
    derived = {str(p) for p in GOOGLE_ALERTS_PLATFORMS}
    return [h for h in source_health if h.get("source") not in derived]


def failed_collectors(source_health: list[dict]) -> list[dict]:
    return [h for h in collector_rows(source_health) if str(h.get("status")) == SourceStatus.FAILED.value]


def report_path_for(period: ReportingPeriod, reports_dir: Path, dry_run: bool,
                    validation: bool = False) -> Path:
    """S18: reports/Negative_Screening_YYYY-MM-DD.xlsx (period_start); dry_run -> reports/dry_run/;
    validation -> reports/validation/."""
    folder = reports_dir / "dry_run" if dry_run else reports_dir / "validation" if validation else reports_dir
    return folder / f"Negative_Screening_{period.period_start.isoformat()}.xlsx"


def _archive_existing_report(path: Path) -> Path | None:
    """Move an earlier report for the same week to <folder>/archive/, named with the time
    it was generated (its modification time, MYT). Returns the archive path, or None."""
    if not path.exists():
        return None
    generated = datetime.fromtimestamp(path.stat().st_mtime, tz=MYT)
    archive_dir = path.parent / "archive"
    archive_dir.mkdir(parents=True, exist_ok=True)
    target = archive_dir / f"{path.stem}_{generated:%Y%m%dT%H%M%S}{path.suffix}"
    number = 1
    while target.exists():                   # two runs in the same second
        number += 1
        target = archive_dir / f"{path.stem}_{generated:%Y%m%dT%H%M%S}_{number}{path.suffix}"
    path.replace(target)
    return target


def _excel_text(value):
    """Make one value safe for a cell: strip characters xlsx cannot store, cut to Excel's
    cell limit. Non-text values are returned unchanged."""
    if not isinstance(value, str):
        return value
    value = CONTROL_CHARS_RE.sub("", value)
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[:EXCEL_CELL_LIMIT - 1] + "…"
    return value


def _put(ws, row: int, column: int, value):
    """Write one cell. Text is always stored as text, so a snippet starting with '='
    can never become a formula."""
    cell = ws.cell(row=row, column=column, value=_excel_text(value))
    if isinstance(value, str) and cell.data_type == "f":
        cell.data_type = "s"
    return cell


def _report_date(iso_text: str | None) -> str:
    """ISO datetime from SQLite -> 'YYYY-MM-DD HH:MM' in Malaysia time ('YYYY-MM-DD' at midnight)."""
    if not iso_text:
        return ""
    when = datetime.fromisoformat(iso_text)
    if when.tzinfo is None:
        return when.strftime("%Y-%m-%d %H:%M")
    when = when.astimezone(MYT)
    if when.hour == 0 and when.minute == 0:
        return when.strftime("%Y-%m-%d")
    return when.strftime("%Y-%m-%d %H:%M")


def _stored_datetime(iso_text: str | None) -> datetime | None:
    """ISO datetime from SQLite -> aware datetime (MYT if stored without a zone); None stays None."""
    if not iso_text:
        return None
    when = datetime.fromisoformat(iso_text)
    return when if when.tzinfo is not None else when.replace(tzinfo=MYT)


def event_date_text(iso_text: str | None, precision: str | None) -> str:
    """S18 'Event / Content Date': only as precise as the source — '2022' (YEAR), '2022-03'
    (MONTH), '2026-09-16' / '2026-09-16 09:30' (a full date); 'UNKNOWN' when there is none."""
    when = _stored_datetime(iso_text)
    if when is None:
        return "UNKNOWN"
    when = when.astimezone(MYT)
    if precision == DatePrecision.YEAR:
        return f"{when:%Y}"
    if precision == DatePrecision.MONTH:
        return f"{when:%Y-%m}"
    return _report_date(iso_text)


def date_status_text(status: str | None, precision: str | None) -> str:
    """S18 'Date Status': reliability, plus '(year only)' / '(month only)' when the source gives no day."""
    suffix = {DatePrecision.YEAR.value: " (year only)", DatePrecision.MONTH.value: " (month only)"}
    return f"{status or DateStatus.UNKNOWN.value}{suffix.get(precision or '', '')}"


def _split_list(text: str | None) -> list[str]:
    """group_concat() result -> sorted list without blanks."""
    if not text:
        return []
    return sorted(part for part in text.split(",") if part)


def report_url(discovered_url: str, final_url: str | None, url_verification: str,
               verification_basis: str | None) -> str:
    """S18: the URL shown in the report. The final_url (after redirects) of a page that was
    read and matched — VERIFIED by 'page text match' — is preferred, unless the redirect
    went up to a less specific page (the home page or a parent section of the discovered
    URL). Otherwise the discovered URL. Both are URLs the source gave; none is built."""
    if (not final_url or final_url == discovered_url
            or url_verification != Verification.VERIFIED or verification_basis != "page text match"
            or not final_url.lower().startswith(("http://", "https://")) or is_social_url(final_url)):
        return discovered_url
    final_seg, discovered_seg = _url_parts(final_url)[1], _url_parts(discovered_url)[1]
    if not final_seg or (len(final_seg) < len(discovered_seg) and discovered_seg[:len(final_seg)] == final_seg):
        return discovered_url
    return final_url


NO_LINK = "No link: "
NO_LINK_EMPTY = NO_LINK + "the source gave no web address for this item"
NOT_FETCHED_SOCIAL = NO_LINK + "not fetched — social media pages are never opened by the program"
NOT_FETCHED_OTHER = NO_LINK + "not fetched — link, title and snippet as delivered by the source"
NOT_FETCHED_SWITCHED_OFF = (NO_LINK + "not fetched — reading public web pages is switched off "
                            "(sources.google_alerts.fetch_public_web_pages)")
NOT_FETCHED_UNREAD = (NO_LINK + "not fetched — the page could not be read (blocked, robots.txt, "
                      "not HTML or unavailable; see the run log)")
NO_ITEMS_IN_SHEET = (NO_LINK + "no items for this section in the reporting period — see "
                     "'Links in this report' on the Summary sheet for each source")


def is_web_link(value: object) -> bool:
    """True for an http(s) address, the only values the report makes clickable."""
    return isinstance(value, str) and value.strip().lower().startswith(("http://", "https://"))


def link_or_reason(url: str | None) -> str:
    """S18 URL / Discovered URL columns: the address, or 'No link: <reason>'."""
    url = (url or "").strip()
    if not url:
        return NO_LINK_EMPTY
    if not is_web_link(url):
        return f"{NO_LINK}'{url}' is not a web address (http:// or https://)"
    return url


def fetched_url_text(final_url: str | None, discovered_url: str, *, via_alerts: bool = False,
                     fetch_public_web: bool = True) -> str:
    """S18 'Fetched URL' column: the page the program actually downloaded (after redirects),
    or 'No link: <why it was not downloaded>'. Never a guessed URL."""
    if final_url:
        return link_or_reason(final_url)
    if not is_web_link(discovered_url):
        return NO_LINK + "not fetched — the item has no web address to open"
    if is_social_url(discovered_url):
        return NOT_FETCHED_SOCIAL
    if via_alerts:          # a Google Alerts result on the public web: read only when switched on
        return NOT_FETCHED_UNREAD if fetch_public_web else NOT_FETCHED_SWITCHED_OFF
    return NOT_FETCHED_OTHER


def _source_platforms(label: str) -> set[str]:
    """Platforms whose findings come from a source-health row: Google Alerts covers the four
    platforms it discovers; every other row is its own platform."""
    if label == "Google Alerts":
        return {str(p) for p in GOOGLE_ALERTS_PLATFORMS}
    return {label}


def _week_starting(monday: date) -> ReportingPeriod:
    """The Monday–Sunday reporting week that starts on `monday`."""
    return calculate_previous_completed_week(monday + timedelta(days=7))


def source_link_status(h: dict, linked: int, listed: int, out_of_period_sheet: bool,
                       reported_earlier: list[date] | None = None) -> str:
    """S18 Summary 'Links in this report': how many clickable links one source put in the
    report, or 'No link: <reason>' worked out from its health row (first matching reason).
    reported_earlier: start of the earlier week in which each flagged out-of-period item of this
    source was already listed (D-24)."""
    if linked:
        return f"{linked} link(s) in this report"
    if listed:
        return (f"{NO_LINK}{listed} item(s) listed, but none has a web address — "
                "see the URL column for each item's reason")
    status, error = str(h.get("status", "")), (h.get("error") or "").strip()
    retrieved, flagged = h.get("retrieved") or 0, h.get("flagged") or 0
    filtered_out = h.get("filtered_out") or 0
    unique = max(retrieved - (h.get("duplicates") or 0), 0)     # filtered_out counts merged items
    if status == SourceStatus.FAILED.value:
        return f"{NO_LINK}the source failed — {error or 'no error message was recorded'}"
    if status == SourceStatus.SKIPPED.value:
        return f"{NO_LINK}the source was not run — {error or 'disabled or not selected'}"
    if retrieved == 0:
        return f"{NO_LINK}nothing was retrieved — {error}" if error else f"{NO_LINK}nothing was retrieved in this run"
    if filtered_out >= unique:
        return f"{NO_LINK}{retrieved} item(s) retrieved, but none named a screened target"
    if flagged == 0:
        return (f"{NO_LINK}{retrieved} item(s) retrieved and {unique - filtered_out} named a screened "
                "target, but none matched the adverse-keyword rules")
    if reported_earlier:
        weeks = ", ".join(email_period_text(_week_starting(w)) for w in sorted(set(reported_earlier)))
        return (f"{NO_LINK}{len(reported_earlier)} flagged item(s) found again, already reported in the "
                f"report for {weeks} — not repeated and not counted for review this week; "
                "see that week's report for the link")
    where = "" if out_of_period_sheet else " (the Out of Period sheet is switched off: report.include_out_of_period_sheet)"
    return (f"{NO_LINK}{flagged} item(s) flagged, but none is in this report's sheets — "
            f"usually dated outside the reporting period{where}")


def matched_names_text(pairs: str | None, entity_names: dict[str, str]) -> str:
    """S18 'Matched Name / Alias': the exact name, variant or alias that matched (e.g. a Chinese alias),
    for audit. From db_findings_for_period()'s matched_names ('<entity_id>\\x1e<text>' joined
    by '\\x1f'). With several entities, each is prefixed by its display name."""
    by_entity: dict[str, list[str]] = {}
    for pair in sorted((pairs or "").split("\x1f")):
        entity_id, sep, text = pair.partition("\x1e")
        if sep and text:
            by_entity.setdefault(entity_id, []).append(text)
    if len(by_entity) == 1:
        return "; ".join(next(iter(by_entity.values())))
    return "; ".join(f"{entity_names.get(e, e)}: {', '.join(texts)}" for e, texts in sorted(by_entity.items()))


def _finding_row(row: sqlite3.Row, period: ReportingPeriod, entity_names: dict[str, str],
                 fetch_public_web: bool = True, entities: dict | None = None) -> list:
    """One sqlite row from db_findings_for_period() -> the values for FINDINGS_COLUMNS.
    Link columns hold an address or 'No link: <reason>'."""
    link = describe_entity_link(set(_split_list(row["entity_ids"])), entities or {
        "entities": [{"id": i, "display_name": n} for i, n in entity_names.items()]})
    url = report_url(row["discovered_url"] or "", row["final_url"], row["url_verification"],
                     row["verification_basis"])
    via_alerts = ALERTS_METHOD in _split_list(row["discovery_methods"])

    title = row["title"] or ""
    snippet = row["snippet"] or ""
    if title and snippet and snippet != title:
        text = f"{title} — {snippet}"
    else:
        text = title or snippet

    methods = ", ".join(_split_list(row["discovery_methods"]))
    queries = row["query_count"] or 0
    if queries == 1:
        methods += " (1 query)"
    elif queries > 1:
        methods += f" ({queries} queries)"

    notes = []
    if row["rules_applied"]:
        notes.append(f"Rules: {row['rules_applied']}")
    if row["verification_basis"]:
        notes.append(f"URL: {row['verification_basis']}")
    if row["date_basis"]:
        notes.append(f"Date: {row['date_basis']}")
    if row["discovery_queries"]:
        notes.append("Found by: " + "; ".join(_split_list(row["discovery_queries"])))
    if row["reviewer_note"]:
        notes.append(f"Reviewer: {row['reviewer_note']}")

    event_date, precision = _stored_datetime(row["content_date"]), row["date_precision"]
    return [
        f"{period.period_start.isoformat()} – {period.period_end.isoformat()}",
        event_date_text(row["content_date"], precision),
        date_status_text(row["date_status"], precision),
        str(historical_status(event_date, precision, period)),
        record_age_text(event_date, precision, period),
        _report_date(row["first_discovered_at"]),
        row["platform"],
        row["content_type"],
        link["type"],
        link["matched"],
        matched_names_text(row["matched_names"], entity_names),
        link["associated"],
        link["relationship"],
        ", ".join(_split_list(row["keywords"])),
        ", ".join(_split_list(row["keyword_categories"])),
        row["priority"] or "",
        text,
        link_or_reason(url),
        row["url_type"],
        row["url_verification"],
        link_or_reason(row["discovered_url"]),
        fetched_url_text(row["final_url"], row["discovered_url"] or "", via_alerts=via_alerts,
                         fetch_public_web=fetch_public_web),
        methods,
        row["review_status"],
        " | ".join(notes),
    ]


def _setup_print(ws, heading: str, landscape: bool = True, repeat_header_row: bool = False) -> None:
    """S18: printable sheets — A4, fitted to one page wide (as many pages tall as needed),
    header row repeated on every page, report heading and page x of y on each page."""
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.7, bottom=0.6, header=0.3, footer=0.3)
    if repeat_header_row:
        ws.print_title_rows = "1:1"
    ws.oddHeader.center.text = heading.replace("&", "&&")
    ws.oddHeader.center.size = 9
    ws.oddFooter.left.text = "Printed &D &T"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8


def _write_findings_sheet(wb, title: str, rows: list[list], styles: dict) -> None:
    ws = wb.create_sheet(title)
    for col, name in enumerate(FINDINGS_COLUMNS, start=1):
        cell = _put(ws, 1, col, name)
        cell.font, cell.fill = styles["header_font"], styles["header_fill"]
        ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(name, 18)
    link_cols = [FINDINGS_COLUMNS.index(c) + 1 for c in ("URL", "Discovered URL", "Fetched URL")]
    priority_col = FINDINGS_COLUMNS.index("Priority") + 1
    historical_col = FINDINGS_COLUMNS.index("Historical Status") + 1
    for r, values in enumerate(rows, start=2):
        critical = values[priority_col - 1] == Priority.CRITICAL.value
        for col, value in enumerate(values, start=1):
            cell = _put(ws, r, col, value)
            cell.alignment = styles["wrap"]
            if critical:
                cell.fill = styles["critical_fill"]
        if values[historical_col - 1] == HistoricalStatus.HISTORICAL.value:   # stands out even on a CRITICAL row
            cell = ws.cell(row=r, column=historical_col)
            cell.fill, cell.font = styles["historical_fill"], styles["bold"]
        for col in link_cols:
            url = values[col - 1]
            link = ws.cell(row=r, column=col)
            if is_web_link(url):
                link.hyperlink = url.strip()
                link.font = styles["link_font"]
            elif isinstance(url, str) and url.startswith(NO_LINK):
                link.font = styles["no_link_font"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(FINDINGS_COLUMNS)).column_letter}{max(len(rows) + 1, 1)}"
    if not rows:
        _put(ws, 3, 1, NO_ITEMS_IN_SHEET).font = styles["no_link_font"]


def _historical_note(reported_rows: list[list]) -> str:
    """Summary 'Historical findings': whether the reported items are new in this week or older
    records found this week. reported_rows: the Findings, Date Unknown and Unverified URL rows."""
    status_col = FINDINGS_COLUMNS.index("Historical Status")
    date_col = FINDINGS_COLUMNS.index("Event / Content Date")
    old = [v for v in reported_rows if v[status_col] == HistoricalStatus.HISTORICAL.value]
    if not old:
        return "None: no reported item is dated before this reporting week."
    oldest = min(str(v[date_col]) for v in old)
    return (f"{len(old)} of {len(reported_rows)} reported item(s) are HISTORICAL (oldest Event / Content Date: "
            f"{oldest}). They were found or listed during this run, but the underlying record is older: they are "
            "NOT new incidents in this reporting week. They are still reported for manual review — see the "
            "Historical Status and Record Age columns.")


def _write_summary_sheet(ws, ctx: RunContext, source_health: list[dict], counts: dict[str, int],
                         generated: datetime, styles: dict, report_rows: list[list],
                         out_of_period_sheet: bool, reported_rows: list[list] | None = None,
                         reported_earlier: dict[str, list[date]] | None = None) -> None:
    """report_rows: every findings row written to this workbook (for 'Links in this report').
    reported_rows: the Findings, Date Unknown and Unverified URL rows (for 'Dates in this report');
    defaults to report_rows. reported_earlier: platform -> weeks whose reports already listed an
    out-of-period item of this run (not repeated here)."""
    reported_earlier = reported_earlier or {}
    reported_rows = report_rows if reported_rows is None else reported_rows
    period = ctx.period
    row = 1

    def line(label: str, value: object = "") -> None:      # numbers stay numbers in Excel
        nonlocal row
        _put(ws, row, 1, label).font = styles["bold"]
        _put(ws, row, 2, value)
        row += 1

    _put(ws, row, 1, "Negative Screening Report").font = styles["title_font"]
    row += 2
    line("Report Generated", f"{generated:%Y-%m-%d %H:%M:%S} (Asia/Kuala_Lumpur)")
    line("Reporting Period", f"{period.period_start:%A %d %b %Y} – {period.period_end:%A %d %b %Y} "
                             f"(Asia/Kuala_Lumpur)")
    line("Execution Mode", MODE_LABELS.get(ctx.mode, str(ctx.mode)))
    line("Reference Date", f"{period.reference_date:%A %d %b %Y}")
    line("Run ID", ctx.run_id)
    if ctx.dry_run:
        line("Dry Run", "Yes — preview only, not emailed")
    if ctx.validation:
        line("Validation Run", VALIDATION_NOTE)
    row += 1
    line("Coverage", COVERAGE_STATEMENT)
    line("Screening", SCREENING_NOTE)
    row += 1

    # Event date vs discovery date: staff must not read an old record as a new incident.
    _put(ws, row, 1, "Dates in this report").font = styles["heading_font"]
    row += 1
    status_col = FINDINGS_COLUMNS.index("Historical Status")
    by_status = {s.value: sum(1 for v in reported_rows if v[status_col] == s.value) for s in HistoricalStatus}
    line("Reported items by Historical Status", " · ".join(f"{n} {s}" for s, n in by_status.items()))
    line("Historical findings", _historical_note(reported_rows))
    if by_status[HistoricalStatus.HISTORICAL.value]:
        for col in (1, 2):
            ws.cell(row=row - 1, column=col).fill = styles["historical_fill"]
    ws.cell(row=row - 1, column=2).alignment = styles["wrap"]
    for label, note in DATE_FIELD_NOTES:
        line(label, note)
        ws.cell(row=row - 1, column=2).alignment = styles["wrap"]
    row += 1

    _put(ws, row, 1, "Screened targets").font = styles["heading_font"]
    row += 1
    for target in screened_targets(ctx.config.entities):
        line(target["name"], f"{target['kind']} — searched as: " + ", ".join(target["names"]))
        ws.cell(row=row - 1, column=2).alignment = styles["wrap"]
    names = {e["id"]: e.get("display_name") or e["id"] for e in ctx.config.entities.get("entities", [])}
    for r in approved_relationships(ctx.config.entities):
        line("Approved relationship", f"{names[r['subject']]} → {r['label']} → {names[r['object']]} "
                                      f"({r['type']}). Used only when an item names both.")
    row += 1

    _put(ws, row, 1, "Source health").font = styles["heading_font"]
    row += 1
    headers = ["Source", "Status", "Retrieved", "Filtered out", "Flagged", "Duplicates", "Rejected",
               "Error / note"]
    for col, name in enumerate(headers, start=1):
        cell = _put(ws, row, col, name)
        cell.font, cell.fill = styles["header_font"], styles["header_fill"]
    row += 1
    for h in source_health:
        status = str(h.get("status", ""))
        note = h.get("error") or ""
        unchecked = status in (SourceStatus.FAILED.value, SourceStatus.NOT_CHECKED.value)
        if unchecked:
            note = f"{note} — {FAILED_SOURCE_NOTE}" if note else FAILED_SOURCE_NOTE
        values = [h.get("source", ""), status, h.get("retrieved", 0), h.get("filtered_out", 0),
                  h.get("flagged", 0), h.get("duplicates", 0), h.get("rejected", 0), note]
        for col, value in enumerate(values, start=1):
            cell = _put(ws, row, col, value)
            if unchecked:
                cell.fill = styles["failed_fill"]
        ws.cell(row=row, column=len(headers)).alignment = styles["wrap"]
        row += 1
    row += 1

    statuses = [str(h.get("status", "")) for h in collector_rows(source_health)]
    line("Sources attempted", sum(1 for s in statuses if s != SourceStatus.SKIPPED.value))
    line("Sources successful", statuses.count(SourceStatus.SUCCESS.value))
    line("Sources not checked (ran, but had nothing to read)", statuses.count(SourceStatus.NOT_CHECKED.value))
    line("Sources failed", statuses.count(SourceStatus.FAILED.value))
    row += 1

    # Where the clickable links come from, and 'No link: <reason>' for a source with none.
    _put(ws, row, 1, "Links in this report").font = styles["heading_font"]
    row += 1
    url_col, platform_col = FINDINGS_COLUMNS.index("URL"), FINDINGS_COLUMNS.index("Platform")
    linked_rows = [v for v in report_rows if is_web_link(v[url_col])]
    line("Clickable links (URL column)", len(linked_rows) if linked_rows else
         NO_LINK + ("no items in this report's sheets — see each source below" if not report_rows
                    else "no item has a web address — see the URL column for each item's reason"))
    for h in source_health:
        platforms = _source_platforms(str(h.get("source", "")))
        listed = [v for v in report_rows if v[platform_col] in platforms]
        earlier = [w for p in platforms for w in reported_earlier.get(p, [])]
        status = source_link_status(h, sum(1 for v in listed if is_web_link(v[url_col])), len(listed),
                                    out_of_period_sheet, earlier)
        line(str(h.get("source", "")), status)
        cell = ws.cell(row=row - 1, column=2)
        cell.alignment = styles["wrap"]
        if status.startswith(NO_LINK):
            cell.font = styles["no_link_font"]
    row += 1

    _put(ws, row, 1, "Counts").font = styles["heading_font"]
    row += 1
    for key, label in COUNT_LABELS:
        line(label, counts.get(key, "n/a"))
    row += 1

    _put(ws, row, 1, "Configuration versions").font = styles["heading_font"]
    row += 1
    line("Entities", ctx.config.entities.get("version") or "not set")
    line("Keywords", ctx.config.keywords.get("version") or "not set")
    line("Settings", ctx.config.settings.get("version") or "not set")
    line("Application", APP_VERSION)

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 60
    for letter in "CDEFG":
        ws.column_dimensions[letter].width = 12
    ws.column_dimensions["H"].width = 70


def generate_excel_report(conn: sqlite3.Connection, ctx: RunContext,
                          source_health: list[dict], counts: dict[str, int]) -> Path:
    """S18: openpyxl (import inside). reports_dir = settings.paths.reports_dir, resolved
    against PROJECT_ROOT unless absolute. Summary: Report Generated, Reporting Period,
    Execution Mode, Reference Date, COVERAGE_STATEMENT, source health table (FAILED rows
    carry the 'absence of findings…' note), counts, config versions.
    Findings / Date Unknown / Unverified sheets with FINDINGS_COLUMNS, frozen header,
    autofilter, hyperlinks, CRITICAL highlighted. Return the saved path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    styles = {
        "bold": Font(bold=True),
        "title_font": Font(bold=True, size=14),
        "heading_font": Font(bold=True, size=12),
        "header_font": Font(bold=True, color="FFFFFF"),
        "header_fill": PatternFill("solid", fgColor="1F4E78"),
        "critical_fill": PatternFill("solid", fgColor="F8CBAD"),
        "failed_fill": PatternFill("solid", fgColor="FFE699"),
        "historical_fill": PatternFill("solid", fgColor="DDEBF7"),
        "link_font": Font(color="0563C1", underline="single"),
        "no_link_font": Font(italic=True, color="808080"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
    }

    reports_dir = Path(ctx.config.settings["paths"]["reports_dir"])
    if not reports_dir.is_absolute():
        reports_dir = (ctx.project_root or PROJECT_ROOT) / reports_dir
    path = report_path_for(ctx.period, reports_dir, ctx.dry_run, ctx.validation)
    path.parent.mkdir(parents=True, exist_ok=True)

    entity_names = {}
    for e in ctx.config.entities.get("entities", []):
        entity_names[e["id"]] = e.get("display_name") or e["id"]

    # Group the period's findings by report section (SQLite is the only input).
    settings = ctx.config.settings
    fetch_public_web = bool(settings.get("sources", {}).get("google_alerts", {}).get("fetch_public_web_pages"))
    audit_sheet = bool(settings.get("report", {}).get("include_out_of_period_sheet"))
    sheet_rows: dict[str, list[list]] = {title: [] for title in SHEET_FOR_SECTION.values()}
    reported_earlier: dict[str, list[date]] = {}        # platform -> earlier weeks that already listed an item
    platform_col, notes_col = FINDINGS_COLUMNS.index("Platform"), FINDINGS_COLUMNS.index("Notes")
    for row in db_findings_for_period(conn, ctx.period):
        title = SHEET_FOR_SECTION.get(row["report_section"])
        if title is None:
            continue
        values = _finding_row(row, ctx.period, entity_names, fetch_public_web, ctx.config.entities)
        if row["report_section"] == Section.OUT_OF_PERIOD.value:
            # D-24: an item already listed in an earlier week's report is not repeated (its review
            # happens from that report); one never listed before is always shown, so nothing that
            # needs review is hidden.
            earlier = db_first_reported_period(conn, row["finding_id"], ctx.period)
            if earlier is not None:
                reported_earlier.setdefault(str(values[platform_col]), []).append(earlier)
                if not audit_sheet:
                    continue
                note = f"Already reported in the report for {email_period_text(_week_starting(earlier))}."
                values[notes_col] = f"{note} {values[notes_col] or ''}".strip()
        sheet_rows[title].append(values)
    out_of_period_sheet = audit_sheet or bool(sheet_rows[SHEET_OUT_OF_PERIOD])
    sheets = [SHEET_FINDINGS, SHEET_DATE_UNKNOWN, SHEET_UNVERIFIED] + ([SHEET_OUT_OF_PERIOD] if out_of_period_sheet else [])

    generated = now_in_malaysia()
    wb = Workbook()
    summary = wb.active
    if summary is None:
        summary = wb.create_sheet()
    summary.title = SHEET_SUMMARY
    _write_summary_sheet(summary, ctx, source_health, counts, generated, styles,
                         [v for title in sheets for v in sheet_rows[title]], out_of_period_sheet,
                         [v for title in (SHEET_FINDINGS, SHEET_DATE_UNKNOWN, SHEET_UNVERIFIED)
                          for v in sheet_rows[title]], reported_earlier)
    for title in sheets:
        _write_findings_sheet(wb, title, sheet_rows[title], styles)
        ws, url_col = wb[title], FINDINGS_COLUMNS.index("URL") + 1
        for r in range(2, len(sheet_rows[title]) + 2):      # URL trace (run log): what Excel holds
            cell = ws.cell(row=r, column=url_col)
            ctx.logger.debug("url trace [excel] sheet=%r row=%d url=%s hyperlink=%s", title, r, cell.value,
                             cell.hyperlink.target if cell.hyperlink else "(none)")
    week = f"{ctx.period.period_start:%d %b %Y} – {ctx.period.period_end:%d %b %Y}"
    for ws in wb.worksheets:
        _setup_print(ws, f"Negative Screening Report · {week} · {ws.title}",
                     repeat_header_row=ws.title != SHEET_SUMMARY)

    archived = _archive_existing_report(path)
    if archived is not None:
        ctx.logger.info("previous report for this week archived: %s", archived)
    wb.save(path)
    return path


# =============================================================================
# S19 — EMAIL                                              (→ reporting/email_report.py)
# Docs: FINAL_ARCHITECTURE §9
# Goal: pure decision function + SMTP send with attachment. Password only from
#       NS_SMTP_PASSWORD. Failure is logged and recorded; the local report is kept.
# =============================================================================

def decide_email(*, mode: Mode, dry_run: bool, no_email: bool, send_email: bool,
                 resend_email: bool, already_sent: bool, email_defaults: dict,
                 validation: bool = False, minutes_since_last_email: float | None = None) -> tuple[bool, str]:
    """S19: return (send?, reason). Order:
    dry_run or validation or MANUAL_TEST_DATE or no_email -> False
    resend_email -> True
    MANUAL with email_defaults['manual_email_cooldown_minutes'] set -> False while the last
        SENT email (any week) is younger than the cooldown ('cooldown: …'); otherwise the
        week may be emailed again (a manual scan is a deliberate request for the report)
    otherwise (AUTOMATIC, or no cooldown configured): already_sent -> False ('already sent for period')
    send_email -> True
    AUTOMATIC -> email_defaults['send_on_automatic']
    MANUAL -> email_defaults['send_on_manual'] (both true by default)"""
    if dry_run:
        return False, "dry run"
    if validation:
        return False, "validation run (staging database)"
    if mode == Mode.MANUAL_TEST_DATE:
        return False, "test reference date (--date) run"
    if no_email:
        return False, "--no-email"
    if resend_email:
        return True, "--resend-email"
    cooldown = email_defaults.get("manual_email_cooldown_minutes")
    if mode == Mode.MANUAL and cooldown is not None:
        if minutes_since_last_email is not None and minutes_since_last_email < cooldown:
            return False, (f"cooldown: an email was sent {int(minutes_since_last_email)} min ago; manual scans "
                           f"email at most once every {cooldown} min")
    elif already_sent:
        return False, "already sent for period"
    if send_email:
        return True, "--send-email"
    if mode == Mode.AUTOMATIC:
        if email_defaults.get("send_on_automatic"):
            return True, "automatic mode default"
        return False, "automatic mode: send_on_automatic is false"
    if email_defaults.get("send_on_manual"):
        return True, "manual mode default"
    return False, "manual mode: send_on_manual is false"


XLSX_MIME = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def email_period_text(period: ReportingPeriod) -> str:
    """'14–20 Sep 2026', '28 Sep – 4 Oct 2026' or '28 Dec 2026 – 3 Jan 2027'."""
    start, end = period.period_start, period.period_end
    if start.year != end.year:
        return f"{start.day} {start:%b %Y} – {end.day} {end:%b %Y}"
    if start.month != end.month:
        return f"{start.day} {start:%b} – {end.day} {end:%b %Y}"
    return f"{start.day}–{end.day} {end:%b %Y}"


def build_email_subject(template: str, period: ReportingPeriod, source_health: list[dict],
                        counts: dict[str, int]) -> str:
    failed = len(failed_collectors(source_health))
    failed_suffix = ""
    if failed == 1:
        failed_suffix = " — 1 source failed"
    elif failed > 1:
        failed_suffix = f" — {failed} sources failed"
    if any(str(h.get("status")) == SourceStatus.NOT_CHECKED.value and h.get("source") == "Google Alerts"
           for h in source_health):     # D-25: a quiet week must not read as an all-clear
        failed_suffix += " — social media & news NOT CHECKED (no Google Alerts emails)"
    return template.format(period=email_period_text(period), critical=counts.get("critical", 0),
                           review=counts.get("manual_review", 0), failed_suffix=failed_suffix)


def build_email_body(ctx: RunContext, source_health: list[dict], counts: dict[str, int],
                     report_name: str) -> str:
    period = ctx.period
    lines = [
        "Negative Screening — weekly report",
        "",
        f"Reporting period: {period.period_start:%A %d %b %Y} – {period.period_end:%A %d %b %Y} "
        f"(Asia/Kuala_Lumpur)",
        f"Execution mode:   {MODE_LABELS.get(ctx.mode, str(ctx.mode))}",
        f"Reference date:   {period.reference_date:%A %d %b %Y}",
        f"Screened targets: {targets_text(ctx.config.entities)}",
        "",
        "Source health:",
    ]
    for h in source_health:
        status = str(h.get("status", ""))
        line = (f"  - {h.get('source', '')}: {status} "
                f"(retrieved {h.get('retrieved', 0)}, flagged {h.get('flagged', 0)})")
        if h.get("error"):
            line += f" — {h['error']}"
        lines.append(line)
        if status in (SourceStatus.FAILED.value, SourceStatus.NOT_CHECKED.value):
            lines.append(f"      {FAILED_SOURCE_NOTE}")
    lines += ["", "Key counts:"]
    for key, label in COUNT_LABELS:
        lines.append(f"  - {label}: {counts.get(key, 'n/a')}")
    lines += [
        "",
        COVERAGE_STATEMENT,
        SCREENING_NOTE,
        "",
        f"The full report is attached ({report_name}) and saved on the screening machine.",
        "",
        f"— Negative Screening prototype v{APP_VERSION}, run {ctx.run_id}",
    ]
    return "\n".join(lines) + "\n"


def build_report_message(report_path: Path, ctx: RunContext, source_health: list[dict],
                         counts: dict[str, int]):
    """S19: the report email (subject from the template, plain-text body, xlsx attached).
    Raises OSError if the report cannot be read."""
    from email.message import EmailMessage
    from email.utils import formatdate, make_msgid

    cfg = ctx.config.email
    msg = EmailMessage()
    msg["Subject"] = build_email_subject(cfg["subject_template"], ctx.period, source_health, counts)
    msg["From"] = cfg["sender"]
    msg["To"] = ", ".join(cfg["recipients"])
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=cfg["sender"].split("@")[-1])
    msg.set_content(build_email_body(ctx, source_health, counts, report_path.name))
    msg.add_attachment(report_path.read_bytes(), maintype=XLSX_MIME[0], subtype=XLSX_MIME[1],
                       filename=report_path.name)
    return msg


def _smtp_error_text(e: Exception) -> str:
    import smtplib
    if isinstance(e, smtplib.SMTPAuthenticationError):
        return f"the mail server rejected the sending account's user name or app password ({_short_error(e)})"
    return _short_error(e)


def send_report_email(report_path: Path, ctx: RunContext, source_health: list[dict],
                      counts: dict[str, int]) -> tuple[str, int]:
    """S19: build subject from template, body (period, mode, health, counts, coverage
    statement), attach the xlsx, send via smtplib STARTTLS. Return (status, recipient_count).
    On failure ctx.email_error holds the reason (never the password)."""
    import smtplib
    import ssl

    log = ctx.logger
    cfg = ctx.config.email
    smtp = cfg["smtp"]
    recipients = list(cfg["recipients"])
    password = get_secret("NS_SMTP_PASSWORD")
    if not password:
        ctx.email_error = SMTP_NOT_CONFIGURED
        log.error("email not sent: %s", SMTP_NOT_CONFIGURED)
        return "FAILED", 0
    try:
        msg = build_report_message(report_path, ctx, source_health, counts)
    except OSError as e:
        ctx.email_error = f"report file could not be read: {_short_error(e)}"
        log.error("email not sent: %s", ctx.email_error)
        return "FAILED", 0

    timeout = ctx.config.settings.get("http", {}).get("timeout_seconds", 30)
    try:
        with smtplib.SMTP(smtp["host"], smtp["port"], timeout=timeout) as server:
            if smtp.get("starttls", True):
                server.starttls(context=ssl.create_default_context())
            server.login(smtp.get("username") or cfg["sender"], password)
            refused = server.send_message(msg, from_addr=cfg["sender"], to_addrs=recipients)
    except (smtplib.SMTPException, OSError) as e:
        ctx.email_error = _smtp_error_text(e)
        log.error("email not sent: %s", ctx.email_error)
        return "FAILED", 0

    delivered = len(recipients) - len(refused)
    for address, reason in refused.items():
        log.warning("recipient refused by the mail server: %s (%s)", address, reason)
    log.info("email sent to %d of %d recipient(s)", delivered, len(recipients))
    return "SENT", delivered


def write_email_preview(report_path: Path, ctx: RunContext, source_health: list[dict],
                        counts: dict[str, int]) -> Path:
    """Validation runs: save the exact email a normal run would send (with the report
    attached) as <report>_email_preview.eml next to the report. Nothing is sent."""
    path = report_path.with_name(report_path.stem + "_email_preview.eml")
    path.write_bytes(bytes(build_report_message(report_path, ctx, source_health, counts)))
    return path


def check_smtp_account(email_cfg: dict, password: str, timeout: int = 30,
                       send_test_to: list[str] | None = None, smtp_factory=None) -> tuple[bool, str]:
    """S19 (Setup): log in to the sending account with STARTTLS; optionally send a short
    test message (no attachment). Returns (ok, message for staff); never raises."""
    import smtplib
    import ssl
    from email.message import EmailMessage

    smtp = email_cfg["smtp"]
    if not password:
        return False, "No password entered."
    try:
        with (smtp_factory or smtplib.SMTP)(smtp["host"], smtp["port"], timeout=timeout) as server:
            if smtp.get("starttls", True):
                server.starttls(context=ssl.create_default_context())
            server.login(smtp.get("username") or email_cfg["sender"], password)
            if not send_test_to:
                return True, f"Signed in to {smtp['host']} as {smtp.get('username') or email_cfg['sender']}."
            msg = EmailMessage()
            msg["Subject"] = "[Negative Screening] test email"
            msg["From"] = email_cfg["sender"]
            msg["To"] = ", ".join(send_test_to)
            msg.set_content("This is a test from the Negative Screening application. "
                            "Weekly reports will be sent from this account.\n")
            refused = server.send_message(msg, from_addr=email_cfg["sender"], to_addrs=send_test_to)
            sent = len(send_test_to) - len(refused)
            return sent > 0, f"Test email sent to {sent} of {len(send_test_to)} recipient(s)."
    except (smtplib.SMTPException, OSError) as e:
        return False, "Could not sign in / send: " + redact(_smtp_error_text(e))


def check_imap_account(email_cfg: dict, password: str, timeout: int = 30,
                       imap_factory=None) -> tuple[bool, str]:
    """S10/S19 (Setup): log in to the Google Alerts mailbox read-only and check the label.
    Returns (ok, message for staff); never raises."""
    import imaplib

    box = email_cfg["alerts_mailbox"]
    if not password:
        return False, "No password entered."
    try:
        imap = (imap_factory or imaplib.IMAP4_SSL)(box["imap_host"], box["imap_port"], timeout=timeout)
        try:
            imap.login(box["username"], password)
            typ, _ = imap.select(_imap_quote(box["label"]), readonly=True)
            if typ != "OK":
                return False, (f"Signed in, but the label {box['label']!r} was not found. Create it and a "
                               "filter for Google Alerts emails (OPERATIONS_GUIDE §2).")
            signed_in = f"Signed in to {box['imap_host']} as {box['username']}; label {box['label']!r} found."
            typ, data = imap.search(None, "FROM", _imap_quote(box["sender_filter"]))
            total = len((data[0] or b"").split()) if typ == "OK" and data else 0
            since = date.today() - timedelta(days=14)
            typ, data = imap.search(None, "FROM", _imap_quote(box["sender_filter"]), "SINCE", _imap_date(since))
            recent = len((data[0] or b"").split()) if typ == "OK" and data else 0
            if total == 0:
                return True, (signed_in + "\n\nBut it holds NO Google Alerts emails yet, so nothing from Google's "
                              "index can be read. Create the alerts (Setup → Google Alerts) and the Gmail filter.")
            return True, f"{signed_in}\n{total} Google Alerts email(s) in the label, {recent} in the last 14 days."
        finally:
            try:
                imap.logout()
            except Exception:  # noqa: BLE001 — already disconnected
                pass
    except (imaplib.IMAP4.error, OSError) as e:
        return False, "Could not sign in: " + redact(_short_error(e))


# =============================================================================
# S20 — SCHEDULING SUPPORT & END-TO-END                  (→ main.py / run_weekly.bat)
# Docs: FINAL_ARCHITECTURE §4, §10; OPERATIONS_GUIDE §6
# Goal: lock file to prevent overlapping runs (stale lock cleared after
#       settings.run.max_run_minutes); `--write-scheduler-files` creates run_weekly.bat;
#       then one full dry run and one real run for a past week, reviewed with the user.
# =============================================================================

RUN_WEEKLY_BAT = """@echo off
rem Negative Screening - weekly automatic run. Written by --write-scheduler-files.
rem Changes to the project folder first so config, data, logs and reports always resolve.
cd /d "%~dp0"
if exist "NegativeScreening.exe" (
    "NegativeScreening.exe" --scheduled
) else (
    "{python}" negative_screening.py --scheduled
)
exit /b %ERRORLEVEL%
"""
VENV_PYTHON = Path(".venv") / "Scripts" / "python.exe"   # relative to the project folder

_RUN_LOCK_TOKENS: dict[str, str] = {}   # lock path -> token this process wrote into it


def _lock_key(lock_path: Path) -> str:
    return os.path.normcase(str(Path(lock_path).resolve()))


def _read_run_lock(lock_path: Path) -> dict:
    """The lock's JSON content; {} if it cannot be read or parsed."""
    try:
        data = json.loads(Path(lock_path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def _lock_age_minutes(path: Path) -> float:
    """From the file's modification time: the lock is written once, when it is taken."""
    return (datetime.now(timezone.utc).timestamp() - path.stat().st_mtime) / 60


def _move_lock_aside(lock_path: Path) -> Path | None:
    """Rename lock_path to a unique side name (atomic: of two processes only one can move
    it). None if it was already gone."""
    aside = lock_path.with_name(f"{lock_path.name}.{os.urandom(6).hex()}")
    try:
        os.rename(lock_path, aside)
    except FileNotFoundError:
        return None
    return aside


def _put_lock_back(aside: Path, lock_path: Path) -> None:
    """Undo _move_lock_aside() for a lock that turned out to belong to another run."""
    try:
        os.rename(aside, lock_path)
    except OSError:   # the path was taken again meanwhile: that run holds the lock now
        aside.unlink(missing_ok=True)


def _clear_stale_run_lock(lock_path: Path, max_age_minutes: int) -> bool:
    """True if lock_path is gone or was stale and has been removed; False if it is fresh
    (another run is active) or could not be moved."""
    log = logging.getLogger(LOGGER_NAME)
    try:
        if _lock_age_minutes(lock_path) <= max_age_minutes:
            return False
    except FileNotFoundError:
        return True
    stale = _read_run_lock(lock_path)
    try:
        aside = _move_lock_aside(lock_path)
    except OSError as e:   # e.g. still open by its owner on Windows
        log.warning("stale run lock %s could not be removed: %s", lock_path, _short_error(e))
        return False
    if aside is None:
        return True
    # Another process may have replaced the stale lock with its own between the checks above.
    if _lock_age_minutes(aside) <= max_age_minutes:
        _put_lock_back(aside, lock_path)
        return False
    aside.unlink(missing_ok=True)
    log.warning("stale run lock replaced (older than %d minutes; pid %s on %s, taken %s)",
                max_age_minutes, stale.get("pid", "?"), stale.get("host", "?"), stale.get("created", "?"))
    return True


def acquire_run_lock(lock_path: Path, max_age_minutes: int) -> bool:
    """S20: create lock_path exclusively (write pid + timestamp). If it exists and is
    older than max_age_minutes, replace it (log a warning) and return True; if fresh,
    return False."""
    import socket
    lock_path = Path(lock_path)
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    token = os.urandom(16).hex()
    info = {"pid": os.getpid(), "host": socket.gethostname(), "token": token,
            "created": now_in_malaysia().isoformat(timespec="seconds")}
    for _ in range(3):   # a stale lock cleared -> try the exclusive create again
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            if not _clear_stale_run_lock(lock_path, max_age_minutes):
                return False
            continue
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(info, f)
        _RUN_LOCK_TOKENS[_lock_key(lock_path)] = token
        return True
    return False


def release_run_lock(lock_path: Path) -> None:
    """S20: remove the lock file if it is ours."""
    lock_path = Path(lock_path)
    token = _RUN_LOCK_TOKENS.pop(_lock_key(lock_path), None)
    if token is None or _read_run_lock(lock_path).get("token") != token:
        if token is not None:
            logging.getLogger(LOGGER_NAME).warning(
                "run lock %s now belongs to another run (ours was replaced as stale); left in place", lock_path)
        return
    aside = _move_lock_aside(lock_path)
    if aside is None:
        return
    if _read_run_lock(aside).get("token") == token:
        aside.unlink(missing_ok=True)
    else:   # replaced between the check and the move: not ours to delete
        _put_lock_back(aside, lock_path)


def write_scheduler_files(project_root: Path) -> Path:
    """S20: write RUN_WEEKLY_BAT to project_root/run_weekly.bat (CRLF line endings).
    It runs the packaged NegativeScreening.exe when present. Python fallback: the project's
    .venv if it exists (relative, so the folder can move), else the interpreter running
    this command (the .venv path when this is the packaged exe)."""
    project_root = Path(project_root)
    if (project_root / VENV_PYTHON).is_file() or IS_FROZEN:
        python = "\\".join(VENV_PYTHON.parts)
    else:
        python = sys.executable
    path = project_root / "run_weekly.bat"
    path.write_text(RUN_WEEKLY_BAT.replace("{python}", python), encoding="utf-8", newline="\r\n")
    return path


SCHEDULED_TASK_NAME = "Negative Screening - Weekly"
SCHEDULED_TASK_DEFAULT_TIME = "09:00"

_TASK_XML = """<?xml version="1.0" encoding="UTF-16"?>
<Task version="1.2" xmlns="http://schemas.microsoft.com/windows/2004/02/mit/task">
  <RegistrationInfo>
    <Description>Negative Screening: weekly adverse-media screening of the previous complete Monday-Sunday week. Emails the Excel report; logs in {logs}.</Description>
    <URI>\\{name}</URI>
  </RegistrationInfo>
  <Triggers>
    <CalendarTrigger>
      <StartBoundary>{start}</StartBoundary>
      <Enabled>true</Enabled>
      <ScheduleByWeek>
        <DaysOfWeek><Monday /></DaysOfWeek>
        <WeeksInterval>1</WeeksInterval>
      </ScheduleByWeek>
    </CalendarTrigger>
  </Triggers>
  <Principals>
    <Principal id="Author">
      <UserId>{user}</UserId>
      <LogonType>InteractiveToken</LogonType>
      <RunLevel>LeastPrivilege</RunLevel>
    </Principal>
  </Principals>
  <Settings>
    <MultipleInstancesPolicy>IgnoreNew</MultipleInstancesPolicy>
    <DisallowStartIfOnBatteries>false</DisallowStartIfOnBatteries>
    <StopIfGoingOnBatteries>false</StopIfGoingOnBatteries>
    <StartWhenAvailable>true</StartWhenAvailable>
    <RunOnlyIfNetworkAvailable>false</RunOnlyIfNetworkAvailable>
    <IdleSettings>
      <StopOnIdleEnd>false</StopOnIdleEnd>
      <RestartOnIdle>false</RestartOnIdle>
    </IdleSettings>
    <AllowStartOnDemand>true</AllowStartOnDemand>
    <Enabled>true</Enabled>
    <Hidden>false</Hidden>
    <RunOnlyIfIdle>false</RunOnlyIfIdle>
    <WakeToRun>false</WakeToRun>
    <ExecutionTimeLimit>PT{limit_minutes}M</ExecutionTimeLimit>
    <Priority>7</Priority>
  </Settings>
  <Actions Context="Author">
    <Exec>
      <Command>{command}</Command>
      <Arguments>{arguments}</Arguments>
      <WorkingDirectory>{workdir}</WorkingDirectory>
    </Exec>
  </Actions>
</Task>
"""


def scheduled_launch_command(project_root: Path) -> tuple[Path, str]:
    """S20: what the weekly task starts. The packaged exe (no console window) with
    --scheduled when it is in the project folder; otherwise run_weekly.bat (developer
    setup without an exe), written on the spot."""
    exe = Path(project_root) / EXE_NAME
    if exe.is_file():
        return exe, "--scheduled"
    if IS_FROZEN:                        # the exe runs from somewhere else, e.g. dist\
        return Path(sys.executable).resolve(), "--scheduled"
    return write_scheduler_files(project_root), ""


def next_weekday_at(today: date, weekday: int, at: str) -> datetime:
    """The first date on/after today with that weekday (0 = Monday), at HH:MM (naive local)."""
    hour, minute = (int(x) for x in at.split(":"))
    day = today + timedelta(days=(weekday - today.weekday()) % 7)
    return datetime.combine(day, time(hour, minute))


def scheduled_task_xml(command: Path, arguments: str, workdir: Path, logs_dir: Path, user: str,
                       start: datetime, limit_minutes: int = 120,
                       name: str = SCHEDULED_TASK_NAME) -> str:
    """S20: Task Scheduler definition — weekly on Monday at start's time; Start When
    Available (a missed Monday runs as soon as the computer is on); no second instance
    while one runs; stopped after limit_minutes; runs as the signed-in user, who owns
    the saved passwords. No boot/logon trigger."""
    from xml.sax.saxutils import escape
    return _TASK_XML.format(name=escape(name), start=start.strftime("%Y-%m-%dT%H:%M:%S"),
                            user=escape(user), limit_minutes=int(limit_minutes),
                            command=escape(str(command)), arguments=escape(arguments),
                            workdir=escape(str(workdir)), logs=escape(str(logs_dir)))


def _current_windows_user() -> str:
    domain, user = os.environ.get("USERDOMAIN", ""), os.environ.get("USERNAME", "")
    return f"{domain}\\{user}" if domain else user


def _run_schtasks(args: list[str]) -> tuple[int, str]:
    import subprocess
    try:
        done = subprocess.run(["schtasks", *args], capture_output=True, text=True, timeout=60,
                              creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.SubprocessError) as e:
        return 1, _short_error(e)
    return done.returncode, ((done.stdout or "") + (done.stderr or "")).strip()


def install_scheduled_task(project_root: Path, at: str = SCHEDULED_TASK_DEFAULT_TIME,
                           limit_minutes: int = 120, runner=_run_schtasks) -> tuple[bool, str]:
    """S20: create (or replace) the weekly Monday task for the current Windows user.
    Returns (ok, message); ok only when schtasks reports success."""
    if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", at):
        return False, f"time must be HH:MM (24-hour), got {at!r}"
    project_root = Path(project_root).resolve()
    command, arguments = scheduled_launch_command(project_root)
    xml = scheduled_task_xml(command, arguments, project_root, project_root / "logs",
                             _current_windows_user(), next_weekday_at(date.today(), 0, at), limit_minutes)
    xml_path = project_root / "data" / "scheduled_task.xml"
    xml_path.parent.mkdir(parents=True, exist_ok=True)
    xml_path.write_text(xml, encoding="utf-16")
    code, output = runner(["/Create", "/TN", SCHEDULED_TASK_NAME, "/XML", str(xml_path), "/F"])
    if code != 0:
        return False, f"Task Scheduler refused the task: {output or f'exit code {code}'}"
    return True, (f"Installed '{SCHEDULED_TASK_NAME}': every Monday at {at} (runs as soon as possible "
                  f"if the computer was off), starting {command.name} {arguments}".rstrip() + ".")


def uninstall_scheduled_task(runner=_run_schtasks) -> tuple[bool, str]:
    """S20: delete the weekly task. Not installed counts as success (nothing to remove)."""
    if query_scheduled_task(runner) is None:
        return True, f"'{SCHEDULED_TASK_NAME}' is not installed; nothing to remove."
    code, output = runner(["/Delete", "/TN", SCHEDULED_TASK_NAME, "/F"])
    if code != 0:
        return False, f"Task Scheduler could not remove the task: {output or f'exit code {code}'}"
    return True, f"Removed '{SCHEDULED_TASK_NAME}'."


def query_scheduled_task(runner=_run_schtasks) -> dict | None:
    """S20: {'next_run', 'last_run', 'last_result', 'status', 'command'} of the weekly task,
    or None if it is not installed."""
    import csv
    import io
    code, output = runner(["/Query", "/TN", SCHEDULED_TASK_NAME, "/FO", "CSV", "/V"])
    if code != 0 or not output:
        return None
    rows = list(csv.reader(io.StringIO(output)))
    if len(rows) < 2:
        return None
    row = dict(zip(rows[0], rows[1]))
    return {"next_run": row.get("Next Run Time", ""), "last_run": row.get("Last Run Time", ""),
            "last_result": row.get("Last Result", ""), "status": row.get("Status", ""),
            "command": row.get("Task To Run", "")}


# =============================================================================
# PIPELINE — runtime order (same for automatic and manual runs)   (→ core/pipeline.py)
# Docs: FINAL_ARCHITECTURE §1–2. This wiring is designed now; it works once the
# steps it calls are DONE. Do not add a separate manual-mode path.
# =============================================================================

# S03..S20: main() takes the S20 run lock around run_pipeline(), so a real or dry run
# must not start before S20 is DONE.
PIPELINE_REQUIRED_STEPS = [f"S{n:02d}" for n in range(3, 21)]


def pipeline_ready() -> list[str]:
    """Steps that must be DONE before a real or dry run. Empty list = ready."""
    return [s for s in PIPELINE_REQUIRED_STEPS if STEP_STATUS.get(s) != "DONE"]


def run_pipeline(ctx: RunContext) -> int:
    """Return exit code: 0 all sources OK, 1 completed with source/email failure, 2 fatal."""
    log = ctx.logger
    log.info("run started: id=%s mode=%s reference=%s version=%s",
             ctx.run_id, ctx.mode, ctx.period.reference_date, APP_VERSION)
    log.info("reporting period: %s", ctx.period.label())
    settings = ctx.config.settings
    conn = db_connect(ctx.db_path)
    db_record_run_start(conn, ctx)

    # 1. Collection — each source isolated (S08, S10–S12)
    results = [run_collector_safely(c, ctx) for c in build_collectors(ctx)]
    raw_items = [item for r in results if r.status == SourceStatus.SUCCESS for item in r.items]
    for item in raw_items:      # URL trace (run log): what each source supplied
        log.debug("url trace [collected] source=%s url=%s title=%r", item.source,
                  item.discovered_url or "(no URL supplied)", (item.title or "")[:80])

    # 2. Normalise + in-run dedup (S06, S07)
    candidates = [normalize_item(i) for i in raw_items]
    candidates, merged = dedupe_in_run(candidates)

    # 3. Platform from the URL, content class, URL type (S13)
    for c in candidates:
        c.platform = detect_platform(c.canonical_url, c.raw.source)
        c.content_class = content_class_for(c.platform, c.raw)
        c.content_type, c.url_type = classify_url(c.canonical_url, c.platform)

    # 4. Dates (S14) then probable-duplicate merge that needs dates (S07)
    for c in candidates:
        resolve_content_date(c)
    mark_probable_duplicates(candidates)
    for c in candidates:
        for note in c.notes:
            if note.startswith("Probable duplicate merged"):
                log.info("url trace [merged] %s — kept as %s", note, c.raw.discovered_url)

    # 5. Name + adverse keyword screening (S15, S16) — rule-based, no NLP
    entities, keywords = ctx.config.entities, ctx.config.keywords
    window = settings["screening"]["keyword_name_window_words"]
    for c in candidates:
        text = " ".join(filter(None, [c.raw.title, c.raw.snippet, c.raw.content_text]))
        c.entity_matches = match_entities(text, c.canonical_url, entities)
        if c.entity_matches:
            positions: list[int | None] = [
                m.position for m in c.entity_matches if m.position is not None
            ]
            c.keyword_matches = scan_keywords(text, keywords, positions)

    # 6. URL verification, flagging rules, report section (S13, S17)
    for c in candidates:
        in_snippet = bool(match_entities(f"{c.raw.title} {c.raw.snippet}", "", entities))
        in_text = (bool(match_entities(c.raw.content_text, "", entities))
                   if c.raw.content_text else None)
        c.verification, c.verification_basis = verify_url(c, in_snippet, in_text)
        apply_flagging_rules(c, window, approved_relationships(entities))
        c.section = assign_section(c, ctx.period)

    # 7. Persist (S05) — upserts make re-runs safe
    reported_earlier: set[int] = set()      # id() of out-of-period items already listed in an earlier week
    for c in candidates:
        finding_id = db_upsert_candidate(conn, c, ctx)
        if c.section == Section.OUT_OF_PERIOD and db_first_reported_period(conn, finding_id, ctx.period):
            reported_earlier.add(id(c))
        log.debug("url trace [stored] finding=%s url=%s final_url=%s platform=%s section=%s priority=%s review=%s",
                  finding_id, c.raw.discovered_url or "(none)", c.raw.final_url or "(not fetched)",
                  c.platform, c.section, c.priority, c.review_status)
    health = derive_source_health(results, candidates)
    db_record_source_status(conn, ctx.run_id, health)

    counts = {
        "discovered": len(raw_items),
        "duplicates_merged": merged,
        "relevant": sum(1 for c in candidates if c.entity_matches),
        "potentially_adverse": sum(1 for c in candidates if c.section == Section.MAIN
                                   and c.priority not in (None, Priority.LOW)),
        "critical": sum(1 for c in candidates if c.priority == Priority.CRITICAL
                        and c.section in (Section.MAIN, Section.DATE_UNKNOWN)),
        # 'for review' = review items shown in THIS report's sheets (an out-of-period item already
        # listed in an earlier week is not repeated, so it is not counted again — D-24)
        "manual_review": sum(1 for c in candidates if c.review_status in
                             (ReviewStatus.MANUAL_REVIEW, ReviewStatus.NEEDS_REVIEW)
                             and c.section not in (Section.NOT_REPORTED, Section.EXCLUDED_FUTURE)
                             and id(c) not in reported_earlier),
        "date_unknown": sum(1 for c in candidates if c.section == Section.DATE_UNKNOWN),
        "unverified_urls": sum(1 for c in candidates if c.section == Section.UNVERIFIED_URL),
        "historical": sum(1 for c in candidates
                          if c.section in (Section.MAIN, Section.DATE_UNKNOWN, Section.UNVERIFIED_URL)
                          and historical_status(c.content_date, c.date_precision, ctx.period)
                          == HistoricalStatus.HISTORICAL),
        "dated_before_period": sum(1 for c in candidates if c.section == Section.OUT_OF_PERIOD
                                   and id(c) not in reported_earlier),
        "already_reported": len(reported_earlier),
    }
    log.info("counts: %s", counts)

    # 8. Report from SQLite (S18)
    report_path = generate_excel_report(conn, ctx, health, counts)
    log.info("report generated: %s", report_path)

    # 9. Email (S19)
    exit_code = 1 if any(r.status == SourceStatus.FAILED for r in results) else 0
    decision: dict[str, Any] = dict(
        mode=ctx.mode, dry_run=ctx.dry_run, no_email=ctx.cli_no_email,
        send_email=ctx.cli_send_email, resend_email=ctx.cli_resend_email,
        already_sent=db_email_already_sent(conn, ctx.period),
        email_defaults=settings["email_defaults"],
        minutes_since_last_email=db_minutes_since_last_email(conn, now_in_malaysia()))
    send, reason = decide_email(**decision, validation=ctx.validation)
    if ctx.validation:
        would_send, would_reason = decide_email(**decision)
        preview = write_email_preview(report_path, ctx, health, counts)
        log.info("validation: a normal run would %s the email (%s); preview saved, nothing sent: %s",
                 "send" if would_send else "not send", would_reason, preview)
    if send:
        status, n = send_report_email(report_path, ctx, health, counts)
        db_record_email(conn, ctx.run_id, ctx.period, status, n, ctx.email_error)
        log.info("email status: %s", status if status == "SENT" else f"{status} ({ctx.email_error})")
        if status != "SENT":
            exit_code = 1
    else:
        db_record_email(conn, ctx.run_id, ctx.period, "SKIPPED", 0, reason)
        log.info("email status: skipped (%s)", reason)

    db_record_run_end(conn, ctx.run_id, exit_code, report_path)
    conn.close()
    log.info("run completed: exit code %s", exit_code)
    return exit_code


# =============================================================================
# CLI                                                                  (→ main.py)
# Docs: FINAL_ARCHITECTURE §4. Designed now; run commands work once steps are DONE.
# =============================================================================

def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="negative_screening.py",
        description="Weekly negative / adverse media screening. Default: manual scan of the "
                    "previous complete Monday–Sunday week (Asia/Kuala_Lumpur).")
    run = p.add_argument_group("run options")
    run.add_argument("--scheduled", action="store_true",
                     help="automatic run (used only by Windows Task Scheduler)")
    run.add_argument("--date", metavar="YYYY-MM-DD",
                     help="TESTING ONLY: behave as if run on this date (email suppressed)")
    run.add_argument("--dry-run", action="store_true",
                     help="full processing, separate database, report in reports/dry_run/, no email")
    run.add_argument("--validate", action="store_true",
                     help="safe validation run: staging database data/screening_validation.sqlite, "
                          "report in reports/validation/ marked VALIDATION, email preview saved, never sent")
    run.add_argument("--no-email", action="store_true", help="do not send email")
    run.add_argument("--send-email", action="store_true",
                     help="send the email even if the mode's default is off (not for an already-sent week)")
    run.add_argument("--resend-email", action="store_true",
                     help="email even if this period was already emailed")
    run.add_argument("--sources", metavar="LIST",
                     help="comma-separated subset: google_alerts,website,sc_malaysia")
    run.add_argument("--verbose", action="store_true", help="debug logging")
    dev = p.add_argument_group("build / setup commands")
    dev.add_argument("--status", action="store_true", help="show build progress per step")
    dev.add_argument("--self-test", metavar="STEP", help="run acceptance tests: S01..S20 or 'all'")
    dev.add_argument("--show-period", action="store_true",
                     help="print the reporting period (respects --date) and exit")
    dev.add_argument("--init-config", action="store_true", help="create config/*.json if missing")
    dev.add_argument("--print-alert-queries", action="store_true",
                     help="print Google Alerts queries generated from the config")
    dev.add_argument("--write-scheduler-files", action="store_true", help="create run_weekly.bat")
    app = p.add_argument_group("staff application and Windows Task Scheduler")
    app.add_argument("--gui", action="store_true",
                     help="open the staff window (the default for NegativeScreening.exe without options)")
    app.add_argument("--install-task", action="store_true",
                     help=f"install the weekly task '{SCHEDULED_TASK_NAME}' (Mondays, --task-time)")
    app.add_argument("--uninstall-task", action="store_true", help="remove the weekly task")
    app.add_argument("--task-status", action="store_true", help="show the weekly task's next/last run")
    app.add_argument("--task-time", metavar="HH:MM", default=SCHEDULED_TASK_DEFAULT_TIME,
                     help=f"Monday start time for --install-task (default {SCHEDULED_TASK_DEFAULT_TIME})")
    return p


def validate_args(args: argparse.Namespace) -> str | None:
    """Return an error message for invalid combinations, else None."""
    if args.scheduled and args.date:
        return "--date is for testing only and cannot be combined with --scheduled"
    if args.no_email and (args.send_email or args.resend_email):
        return "--no-email cannot be combined with --send-email / --resend-email"
    if args.validate and args.dry_run:
        return "--validate and --dry-run are separate safe modes; use one"
    if args.validate and (args.send_email or args.resend_email):
        return "--validate never sends email (it saves a preview); remove --send-email / --resend-email"
    if args.date:
        try:
            date.fromisoformat(args.date)
        except ValueError:
            return f"--date must be YYYY-MM-DD, got {args.date!r}"
    if args.sources:
        valid = {"google_alerts", "website", "sc_malaysia"}
        bad = {s.strip() for s in args.sources.split(",")} - valid
        if bad:
            return f"unknown source(s) in --sources: {', '.join(sorted(bad))}"
    return None


def resolve_mode_and_reference(args: argparse.Namespace) -> tuple[Mode, date]:
    if args.date:
        return Mode.MANUAL_TEST_DATE, date.fromisoformat(args.date)
    return (Mode.AUTOMATIC if args.scheduled else Mode.MANUAL), today_in_malaysia()


def _load_config_for_cli(config_dir: Path) -> Config | None:
    """load_config() for the CLI: an invalid config is a fatal error (exit code 2)
    reported as one clear line, not a traceback. Runs before logging exists."""
    try:
        return load_config(config_dir)
    except ConfigError as e:
        print(f"Configuration error: {e}", file=sys.stderr)
        return None


PROJECT_DIRS = ("config", "data", "logs")
REPORT_SUBDIRS = ("archive", "dry_run")


def ensure_project_dirs(project_root: Path) -> None:
    """Create the local folders the application uses under the project folder."""
    for d in PROJECT_DIRS:
        (Path(project_root) / d).mkdir(parents=True, exist_ok=True)


def ensure_report_dirs(project_root: Path, cfg: Config) -> Path:
    """Create the reports folder (settings.paths.reports_dir, e.g. C:\\Negative_Screening_Reports)
    with its archive\\ and dry_run\\ subfolders; return the reports folder. Raises OSError."""
    folder = reports_folder(project_root, cfg)
    for sub in REPORT_SUBDIRS:
        (folder / sub).mkdir(parents=True, exist_ok=True)
    return folder


@dataclass
class RunRequest:
    """One screening run, as asked for by the CLI, the scheduled task or the staff window."""
    mode: Mode
    reference: date
    dry_run: bool = False
    validation: bool = False
    no_email: bool = False
    send_email: bool = False
    resend_email: bool = False
    sources: set[str] | None = None
    verbose: bool = False


@dataclass
class RunOutcome:
    exit_code: int                     # 0 ok · 1 source or email failed · 2 fatal · 3 another run active
    message: str                       # one line for staff
    run_id: str | None = None
    db_path: Path | None = None
    report_path: Path | None = None


RUN_MESSAGES = {0: "Completed: every source was checked.",
                1: "Completed with problems: a source or the email failed (see Last run).",
                2: "Stopped by an error; no complete report (see the run log).",
                3: "Another screening run is in progress; try again when it has finished."}


def run_database_path(project_root: Path, paths: dict, dry_run: bool, validation: bool) -> Path:
    key = "dry_run_database" if dry_run else "validation_database" if validation else "database"
    return Path(project_root) / paths.get(key, DEFAULT_SETTINGS["paths"][key])


def unique_run_id(db_path: Path, started: datetime) -> str:
    """e.g. 20260922T103012+0800; a second run started in the same second (one after the
    other — the lock stops overlaps) gets -2, -3 … instead of colliding in the runs table."""
    base = started.strftime("%Y%m%dT%H%M%S%z")
    if not Path(db_path).is_file():
        return base
    try:
        conn = sqlite3.connect(db_path)
        try:
            taken = {r[0] for r in conn.execute("SELECT run_id FROM runs WHERE run_id LIKE ?", (base + "%",))}
        finally:
            conn.close()
    except sqlite3.Error:
        return base
    run_id, n = base, 1
    while run_id in taken:
        n += 1
        run_id = f"{base}-{n}"
    return run_id


def execute_run(req: RunRequest, project_root: Path | None = None,
                log_handler: logging.Handler | None = None) -> RunOutcome:
    """Everything after the arguments: readiness gate, config, logging, run lock, pipeline.
    Same path for the CLI, the scheduled task (--scheduled) and the staff window."""
    project_root = Path(project_root or PROJECT_ROOT)
    missing = pipeline_ready()
    if missing:
        return RunOutcome(2, "Pipeline not ready. Steps still to build: " + ", ".join(missing))
    ensure_project_dirs(project_root)
    period = calculate_previous_completed_week(req.reference)
    started = now_in_malaysia()
    run_id = started.strftime("%Y%m%dT%H%M%S%z")
    try:
        cfg = load_config(project_root / "config")
    except ConfigError as e:
        logger = setup_logging(project_root / "logs", run_id, verbose=req.verbose)
        if log_handler is not None:
            log_handler.addFilter(RedactingFilter())
            logger.addHandler(log_handler)
        logger.error("configuration error: %s", e)
        return RunOutcome(2, f"Configuration error: {e}", run_id)
    paths = cfg.settings["paths"]
    db_path = run_database_path(project_root, paths, req.dry_run, req.validation)
    run_id = unique_run_id(db_path, started)
    logger = setup_logging(project_root / paths["logs_dir"], run_id, verbose=req.verbose)
    if log_handler is not None:
        log_handler.addFilter(RedactingFilter())
        logger.addHandler(log_handler)
    try:
        ensure_report_dirs(project_root, cfg)
    except OSError as e:      # the report save tries again and records the failure
        logger.warning("reports folder not created: %s", _short_error(e, 200))
    ctx = RunContext(
        run_id=run_id, mode=req.mode, period=period, started_at=started, dry_run=req.dry_run,
        config=cfg, logger=logger, db_path=db_path, cli_no_email=req.no_email,
        cli_send_email=req.send_email, cli_resend_email=req.resend_email,
        only_sources=req.sources, validation=req.validation, project_root=project_root)
    if req.validation:
        logger.info("VALIDATION RUN: staging database %s — not production data, nothing is emailed", db_path)

    lock = project_root / paths["data_dir"] / "run.lock"
    if not acquire_run_lock(lock, cfg.settings["run"]["max_run_minutes"]):
        logger.error("another run is in progress (lock: %s)", lock)
        return RunOutcome(3, RUN_MESSAGES[3], run_id, db_path)
    try:
        code = run_pipeline(ctx)
    except Exception:
        logger.error("fatal error:\n%s", traceback.format_exc())
        code = 2
    finally:
        release_run_lock(lock)
    report = None
    try:
        conn = sqlite3.connect(db_path)
        try:
            row = conn.execute("SELECT report_path FROM runs WHERE run_id = ?", (run_id,)).fetchone()
        finally:
            conn.close()
        report = Path(row[0]) if row and row[0] else None
    except sqlite3.Error:
        pass
    return RunOutcome(code, RUN_MESSAGES.get(code, f"exit code {code}"), run_id, db_path, report)


def _attach_console_if_packaged() -> None:
    """The packaged exe has no console of its own (staff never see one). For command-line
    use from Command Prompt, write to that window; otherwise discard output safely."""
    if sys.stdout is not None and sys.stderr is not None:
        return
    attached = False
    if os.name == "nt":
        import ctypes
        attached = bool(ctypes.windll.kernel32.AttachConsole(-1))    # ATTACH_PARENT_PROCESS
    target = "CONOUT$" if attached else os.devnull
    if sys.stdout is None:
        sys.stdout = open(target, "w", encoding="utf-8", errors="backslashreplace")  # noqa: SIM115
    if sys.stderr is None:
        sys.stderr = open(target, "w", encoding="utf-8", errors="backslashreplace")  # noqa: SIM115


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if IS_FROZEN and not argv:
        return run_gui()                   # staff double-click NegativeScreening.exe
    _attach_console_if_packaged()
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    error = validate_args(args)
    if error:
        parser.error(error)   # exits with code 2

    if args.gui:
        return run_gui()
    if args.status:
        print_status()
        return 0
    if args.self_test:
        return run_self_tests(args.self_test)

    mode, reference = resolve_mode_and_reference(args)
    period = calculate_previous_completed_week(reference)
    if args.show_period:
        print(f"Execution mode:   {mode}")
        print(f"Reference date:   {reference:%A, %d %B %Y}")
        print(f"Reporting period: {period.period_start:%A %d %b %Y} 00:00:00 – "
              f"{period.period_end:%A %d %b %Y} 23:59:59 (Asia/Kuala_Lumpur)")
        return 0

    config_dir = PROJECT_ROOT / "config"
    if args.init_config:
        written = init_config(config_dir)
        print("created: " + (", ".join(p.name for p in written) if written else "nothing (all exist)"))
        return 0
    if args.write_scheduler_files:
        print(f"written: {write_scheduler_files(PROJECT_ROOT)}")
        return 0
    if args.install_task or args.uninstall_task or args.task_status:
        return _task_command(args)
    if args.print_alert_queries:
        cfg = _load_config_for_cli(config_dir)
        if cfg is None:
            return 2
        for q in generate_alert_queries(cfg.entities, cfg.keywords):
            print(q)
        return 0

    outcome = execute_run(RunRequest(
        mode=mode, reference=reference, dry_run=args.dry_run, validation=args.validate,
        no_email=args.no_email, send_email=args.send_email, resend_email=args.resend_email,
        sources={s.strip() for s in args.sources.split(",")} if args.sources else None,
        verbose=args.verbose))
    if outcome.run_id is None:                  # stopped before logging existed
        print(outcome.message)
        print("Run `python negative_screening.py --status` for details.")
    return outcome.exit_code


def _task_command(args: argparse.Namespace) -> int:
    if args.install_task:
        limit = 120
        try:
            limit = load_config(PROJECT_ROOT / "config").settings["run"]["max_run_minutes"]
        except ConfigError:
            pass
        ok, message = install_scheduled_task(PROJECT_ROOT, args.task_time, limit)
    elif args.uninstall_task:
        ok, message = uninstall_scheduled_task()
    else:
        info = query_scheduled_task()
        ok = info is not None
        message = (f"'{SCHEDULED_TASK_NAME}': {info['status']}; next run {info['next_run']}; last run "
                   f"{info['last_run']} (result {info['last_result']}); runs {info['command']}"
                   if info else f"'{SCHEDULED_TASK_NAME}' is not installed.")
    print(message)
    return 0 if ok else 1


# =============================================================================
# STAFF APPLICATION (Tkinter, standard library)                      (→ app/gui.py)
# Docs: README "Staff application", OPERATIONS_GUIDE §4, §6
# One window: Run Weekly Screening (previous complete Monday–Sunday week, any day,
# report emailed by default), Open Latest Report, Open Reports Folder, Last run status.
# Setup menu: passwords (Windows Credential Manager) and the weekly Task Scheduler task.
# Developer options (--date, --dry-run, --validate, --sources) are not shown to staff.
# =============================================================================

MODE_TEXT = {"AUTOMATIC": "Automatic (weekly schedule)", "MANUAL": "Manual (started by staff)",
             "MANUAL_TEST_DATE": "Test reference date (developer)"}


def read_last_run(db_path: Path) -> dict | None:
    """The newest run in a database (read-only) with its source health, its email row and
    the latest SENT email for its period; None if there is no database or no run yet."""
    from urllib.request import pathname2url
    db_path = Path(db_path)
    if not db_path.is_file():
        return None
    try:
        conn = sqlite3.connect(f"file:{pathname2url(str(db_path.resolve()))}?mode=ro", uri=True)
    except sqlite3.Error:
        return None
    conn.row_factory = sqlite3.Row
    try:
        run = conn.execute("SELECT * FROM runs ORDER BY started_at DESC, rowid DESC LIMIT 1").fetchone()
        if run is None:
            return None
        sources = [dict(r) for r in conn.execute(
            "SELECT * FROM source_status WHERE run_id = ? ORDER BY rowid", (run["run_id"],))]
        email = conn.execute("SELECT * FROM email_log WHERE run_id = ? ORDER BY email_id DESC LIMIT 1",
                             (run["run_id"],)).fetchone()
        sent = conn.execute("SELECT sent_at, recipient_count FROM email_log WHERE period_start = ?"
                            " AND status = 'SENT' ORDER BY sent_at DESC LIMIT 1",
                            (run["period_start"],)).fetchone()
        findings = [tuple(r) for r in conn.execute(
            "SELECT report_section, priority, COUNT(*) FROM finding_periods WHERE period_start = ?"
            " GROUP BY report_section, priority", (run["period_start"],))]
    except sqlite3.Error:
        return None
    finally:
        conn.close()
    return {"run": dict(run), "sources": sources, "email": dict(email) if email else None,
            "period_sent": dict(sent) if sent else None, "findings": findings}


def friendly_source_note(status: str, error: str | None) -> str:
    """Source-health note for staff: why a source was not checked, never a traceback."""
    error = error or ""
    if status in (SourceStatus.SUCCESS.value, SourceStatus.NOT_CHECKED.value):
        if "No Google Alerts emails arrived" in error:      # Facebook / Instagram / LinkedIn / Public website
            return "Not checked — no Google Alerts emails arrived (see Google Alerts)."
        if "no Google Alerts emails arrived" in error:
            return ("Checked the mailbox, but NO Google Alerts emails arrived — nothing from Google's index was "
                    "read. Create the alerts: Setup → Google Alerts.")
        return "Checked" + (" — " + error.removeprefix("Warnings: ")[:160] if error.startswith("Warnings:") else "")
    if status == SourceStatus.SKIPPED.value:
        return "Skipped (disabled)"
    if "NS_IMAP_PASSWORD" in error:
        return ("NOT CHECKED — Google Alerts email access has not been configured. "
                "Open Setup → Email accounts and passwords.")
    if "Google Alerts unavailable" in error:
        return "NOT CHECKED — found only through Google Alerts, which could not be read."
    return "NOT CHECKED — " + (error.split(" — ")[0][:160] or "the source could not be read")


def friendly_email_status(email: dict | None, period_sent: dict | None) -> str:
    if email is None:
        return "No email step recorded (the run stopped early)."
    status, reason = email.get("status"), email.get("reason") or ""
    if status == "SENT":
        return f"Sent to {email.get('recipient_count', 0)} recipient(s), {_short_time(email.get('sent_at'))}."
    if status == "FAILED":
        if "not configured" in reason:
            return ("NOT SENT — email sending has not been configured (Setup → Email accounts and "
                    "passwords). The report is saved locally.")
        return f"NOT SENT — {reason or 'see the run log'}. The report is saved locally."
    if reason == "already sent for period":
        when = _short_time(period_sent.get("sent_at")) if period_sent else "earlier"
        return f"Not sent again: this week's report was already emailed ({when})."
    if reason == "--no-email":
        return "Not sent (run without email)."
    if reason.startswith("cooldown:"):
        return ("Not sent: " + reason.removeprefix("cooldown: ")
                + ". The report is saved; run again later to email it.")
    return f"Not sent ({reason or 'skipped'})."


def _short_time(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return f"{datetime.fromisoformat(iso):%a %d %b %Y %H:%M}"
    except ValueError:
        return iso


def describe_last_run(info: dict | None) -> dict:
    """Plain-language fields for the Last run panel."""
    if not info:
        return {"when": "No screening has been run yet.", "mode": "", "period": "", "result": "",
                "email": "", "report": "", "report_path": None, "sources": [], "alerts_not_configured": False,
                "findings": "", "alerts_no_emails": False}
    run = info["run"]
    code = run.get("exit_code")
    if code is None:
        result = "Did not finish (stopped, or still running)."
    else:
        result = {0: "Success — every source was checked.",
                  1: "Completed with problems — see the sources and email below.",
                  2: "Failed — see the run log in the logs folder.",
                  3: "Not started — another run was in progress."}.get(code, f"Exit code {code}")
    if code == 0 and any("no Google Alerts emails arrived" in (s.get("error") or "") for s in info["sources"]):
        result = ("Completed — but no Google Alerts emails arrived, so Facebook, Instagram, LinkedIn and news "
                  "in Google's index were not checked (Setup → Google Alerts).")
    start, end = date.fromisoformat(run["period_start"]), date.fromisoformat(run["period_end"])
    report_path = Path(run["report_path"]) if run.get("report_path") else None
    report = str(report_path) if report_path else "No report was produced."
    if report_path and not report_path.exists():
        report += "  (file no longer exists)"
    sources = [(s["source"], s["status"], s.get("retrieved") or 0, s.get("flagged") or 0,
                friendly_source_note(s["status"], s.get("error"))) for s in info["sources"]]
    return {
        "when": _short_time(run.get("started_at")),
        "mode": MODE_TEXT.get(run.get("execution_mode"), run.get("execution_mode", "")),
        "period": f"{start:%A %d %b %Y} – {end:%A %d %b %Y}",
        "result": result,
        "email": friendly_email_status(info["email"], info["period_sent"]),
        "report": report,
        "report_path": report_path,
        "sources": sources,
        "alerts_not_configured": any("NS_IMAP_PASSWORD" in (s.get("error") or "") for s in info["sources"]),
        "alerts_no_emails": any("no Google Alerts emails arrived" in (s.get("error") or "") for s in info["sources"]),
        "findings": describe_findings(info.get("findings", [])),
    }


def describe_findings(rows: list[tuple]) -> str:
    """'2 in Findings (1 critical, 1 high) · 1 date unknown · 0 unverified URLs' for the period."""
    main: dict[str, int] = {}
    other = {"DATE_UNKNOWN": 0, "UNVERIFIED_URL": 0}
    for section, priority, n in rows:
        if section == Section.MAIN.value:
            main[priority or "LOW"] = main.get(priority or "LOW", 0) + n
        elif section in other:
            other[section] += n
    total = sum(main.values())
    if not total and not any(other.values()):
        return "No reportable findings in this period from the sources that were checked."
    order = [p.value for p in (Priority.CRITICAL, Priority.HIGH, Priority.MEDIUM, Priority.REVIEW, Priority.LOW)]
    detail = ", ".join(f"{main[p]} {p.lower()}" for p in order if main.get(p))
    return (f"{total} in Findings" + (f" ({detail})" if detail else "")
            + f" · {other['DATE_UNKNOWN']} date unknown · {other['UNVERIFIED_URL']} unverified URLs")


def reports_folder(project_root: Path, cfg: Config | None) -> Path:
    folder = Path(cfg.settings["paths"]["reports_dir"] if cfg else "reports")
    return folder if folder.is_absolute() else Path(project_root) / folder


def latest_report(folder: Path, info: dict | None = None) -> Path | None:
    """The last production run's report if it still exists, else the newest
    Negative_Screening_*.xlsx directly in the reports folder (not archive/, dry_run/, validation/)."""
    if info and info["run"].get("report_path") and not info["run"].get("dry_run"):
        path = Path(info["run"]["report_path"])
        if path.is_file():
            return path
    files = [f for f in Path(folder).glob("Negative_Screening_*.xlsx") if f.is_file()]
    return max(files, key=lambda f: f.stat().st_mtime) if files else None


PRINT_COLUMNS = ["Priority", "Event / Content Date", "Historical Status", "Platform", "Matched Entity",
                 "Relationship", "Keyword",
                 "Title / Comment / Snippet", "URL", "Fetched URL", "URL Verification", "Manual Review Status"]
PRINT_SECTIONS = [(Section.MAIN.value, "Findings"), (Section.DATE_UNKNOWN.value, "Date unknown — manual review"),
                  (Section.UNVERIFIED_URL.value, "Unverified URLs — manual lookup")]


def build_print_view(db_path: Path, cfg: Config, info: dict) -> str:
    """A printable HTML page of the last run's report (same SQLite data as the Excel file):
    screened targets, period, result, source health, email status and the reportable
    findings with their links. A4 landscape; the browser prints it or saves it as PDF.
    Every value from the web is HTML-escaped."""
    from html import escape
    run = info["run"]
    period = calculate_previous_completed_week(date.fromisoformat(run["reference_date"]))
    d = describe_last_run(info)
    names = {e["id"]: e.get("display_name") or e["id"] for e in cfg.entities.get("entities", [])}
    fetch_public_web = bool(cfg.settings.get("sources", {}).get("google_alerts", {}).get("fetch_public_web_pages"))
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        by_section: dict[str, list[list]] = {s: [] for s, _ in PRINT_SECTIONS}
        for row in db_findings_for_period(conn, period):
            if row["report_section"] in by_section:
                by_section[row["report_section"]].append(_finding_row(row, period, names, fetch_public_web, cfg.entities))
    finally:
        conn.close()
    idx = [FINDINGS_COLUMNS.index(c) for c in PRINT_COLUMNS]

    def cell(name: str, value) -> str:
        text = escape(str(value if value is not None else ""))
        if name in ("URL", "Fetched URL") and is_web_link(value):
            return f'<td class="url"><a href="{text}">{text}</a></td>'
        return f'<td class="{"crit" if text == "CRITICAL" else ""}">{text}</td>'

    targets = "".join(f"<li><b>{escape(x['name'])}</b> — {escape(x['kind'])}; searched as: "
                      f"{escape(', '.join(x['names']))}</li>" for x in screened_targets(cfg.entities))
    health = "".join(
        f'<tr class="{"failed" if s == "FAILED" else ""}"><td>{escape(src)}</td><td>{escape(s)}</td>'
        f"<td>{found}</td><td>{flagged}</td><td>{escape(note)}</td></tr>"
        for src, s, found, flagged, note in d["sources"])
    sections = []
    for section, title in PRINT_SECTIONS:
        rows = by_section[section]
        body = "".join("<tr>" + "".join(cell(PRINT_COLUMNS[i], r[j]) for i, j in enumerate(idx)) + "</tr>"
                       for r in rows) or f'<tr><td colspan="{len(PRINT_COLUMNS)}">None in this period.</td></tr>'
        head = "".join(f"<th>{escape(c)}</th>" for c in PRINT_COLUMNS)
        sections.append(f"<h2>{escape(title)} ({len(rows)})</h2><table class=\"findings\"><thead><tr>{head}"
                        f"</tr></thead><tbody>{body}</tbody></table>")
    report = escape(d["report"])
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>Negative Screening Report {period.period_start:%d %b %Y} – {period.period_end:%d %b %Y}</title>
<style>
  @page {{ size: A4 landscape; margin: 12mm; }}
  body {{ font: 10pt "Segoe UI", Arial, sans-serif; color: #111; margin: 16px; }}
  h1 {{ font-size: 16pt; margin: 0 0 4px; }}
  h2 {{ font-size: 12pt; margin: 18px 0 6px; break-after: avoid; page-break-after: avoid; }}
  table {{ border-collapse: collapse; width: 100%; }} th, td {{ border: 1px solid #999; padding: 3px 5px;
  vertical-align: top; text-align: left; }} th {{ background: #1F4E78; color: #fff; }}
  thead {{ display: table-header-group; }} tr {{ page-break-inside: avoid; }}
  td.url {{ word-break: break-all; max-width: 260px; }} td.crit {{ background: #F8CBAD; font-weight: bold; }}
  tr.failed td {{ background: #FFE699; }} .meta td:first-child {{ font-weight: bold; width: 170px; }}
  .note {{ font-size: 9pt; color: #333; }} .noprint {{ margin-bottom: 12px; }}
  @media print {{ .noprint {{ display: none; }} body {{ margin: 0; }} a {{ color: #111; }} }}
</style></head><body>
<div class="noprint"><button onclick="window.print()">Print / Save as PDF</button></div>
<h1>Negative Screening Report</h1>
<table class="meta">
<tr><td>Reporting period</td><td>{period.period_start:%A %d %b %Y} – {period.period_end:%A %d %b %Y} (Asia/Kuala_Lumpur)</td></tr>
<tr><td>Run</td><td>{escape(d["when"])} · {escape(d["mode"])} · run {escape(run["run_id"])}</td></tr>
<tr><td>Result</td><td>{escape(d["result"])}</td></tr>
<tr><td>Findings</td><td>{escape(d["findings"])}</td></tr>
<tr><td>Email</td><td>{escape(d["email"])}</td></tr>
<tr><td>Excel report</td><td>{report}</td></tr>
</table>
<h2>Screened targets</h2><ul>{targets}</ul>
<h2>Source health</h2>
<table><thead><tr><th>Source</th><th>Status</th><th>Found</th><th>Flagged</th><th>Note</th></tr></thead>
<tbody>{health}</tbody></table>
{"".join(sections)}
<p class="note">{escape(COVERAGE_STATEMENT)} {escape(SCREENING_NOTE)} A FAILED source could not be checked: the
absence of findings from it does not mean there is no adverse content.</p>
<p class="note">Negative Screening v{APP_VERSION} · printed from the local database; the Excel report holds every column.</p>
<script>window.addEventListener("load", function () {{ setTimeout(function () {{ window.print(); }}, 400); }});</script>
</body></html>
"""


def write_print_view(project_root: Path, cfg: Config, db_path: Path) -> Path | None:
    """Write reports/print/Negative_Screening_<week>.html for the last run; None if no run yet."""
    info = read_last_run(db_path)
    if info is None:
        return None
    folder = reports_folder(project_root, cfg) / "print"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"Negative_Screening_{info['run']['period_start']}.html"
    path.write_text(build_print_view(db_path, cfg, info), encoding="utf-8")
    return path


def open_path(path: Path, opener=None) -> None:
    """Open a file or folder with its Windows program (Excel, Explorer)."""
    (opener or os.startfile)(str(path))


def staff_run_request(today: date, no_email: bool = False, resend_email: bool = False) -> RunRequest:
    """'Run Weekly Screening': a manual run for the previous complete week of today (any
    weekday), emailed by default (settings.email_defaults.send_on_manual), at most once per
    email_defaults.manual_email_cooldown_minutes."""
    return RunRequest(mode=Mode.MANUAL, reference=today, no_email=no_email, resend_email=resend_email)


def normalise_app_password(value: str) -> str:
    """Gmail shows app passwords as 'abcd efgh ijkl mnop'; the spaces are not part of it."""
    value = value.strip()
    if re.fullmatch(r"[A-Za-z]{4}( [A-Za-z]{4}){3}", value):
        return value.replace(" ", "")
    return value


class _QueueLogHandler(logging.Handler):
    def __init__(self, q) -> None:
        super().__init__(logging.INFO)
        self.q = q
        self.setFormatter(logging.Formatter("%(message)s"))

    def emit(self, record: logging.LogRecord) -> None:
        try:
            self.q.put(("log", record.levelname, self.format(record)))
        except Exception:  # noqa: BLE001 — logging must never break the run
            self.handleError(record)


class StaffApp:
    """The staff window. Long work runs in a background thread; the window polls a queue."""

    STAGES = [("source started: google_alerts", "Checking Google Alerts (Facebook, Instagram, LinkedIn, news)…"),
              ("source started: website", "Checking the Company website…"),
              ("source started: sc_malaysia", "Checking Securities Commission Malaysia…"),
              ("counts:", "Screening the results…"),
              ("report generated", "Excel report saved. Emailing…"),
              ("email status", "Finishing…")]

    def __init__(self, root, project_root: Path, task_query=None, today=None) -> None:
        import queue
        import tkinter as tk
        from tkinter import ttk
        self.tk, self.ttk, self.root = tk, ttk, root
        self.project_root = Path(project_root)
        self.queue: queue.Queue = queue.Queue()
        self.task_query = task_query if task_query is not None else query_scheduled_task
        self.today = today or today_in_malaysia
        self.running = False
        ensure_project_dirs(self.project_root)
        self.cfg, self.config_error, created = None, None, []
        config_dir = self.project_root / "config"
        if any(not (config_dir / name).exists() for name in CONFIG_FILES):
            created = init_config(config_dir)
        try:
            self.cfg = load_config(config_dir)
        except ConfigError as e:
            self.config_error = str(e)
        if self.cfg is not None:
            try:
                ensure_report_dirs(self.project_root, self.cfg)
            except OSError:     # a run tries again and reports the problem
                pass

        root.title("Negative Screening")
        root.geometry("940x720")
        root.minsize(780, 600)
        root.protocol("WM_DELETE_WINDOW", self.close)
        style = ttk.Style(root)
        style.configure("Run.TButton", font=("Segoe UI", 11, "bold"), padding=(18, 10))
        style.configure("Big.TButton", padding=(12, 10))
        style.configure("Title.TLabel", font=("Segoe UI", 16, "bold"))
        style.configure("Notice.TLabel", background="#FFF4CE", foreground="#5C4400", padding=8)
        style.configure("Key.TLabel", font=("Segoe UI", 9, "bold"))
        self._build_menu()
        self._build_body()
        if created:
            self._activity("Configuration files were created in the config folder: "
                           + ", ".join(f.name for f in created) + ". Review names and email recipients.")
        self.refresh_status()
        root.after(150, self._poll)

    # ---- layout ------------------------------------------------------------
    def _build_menu(self) -> None:
        tk = self.tk
        bar = tk.Menu(self.root)
        setup = tk.Menu(bar, tearoff=False)
        setup.add_command(label="Email accounts and passwords…", command=self.email_setup)
        setup.add_command(label="Google Alerts (search-index discovery)…", command=self.alerts_setup)
        setup.add_command(label="Weekly schedule (Task Scheduler)…", command=self.schedule_setup)
        setup.add_separator()
        setup.add_command(label="Open logs folder", command=lambda: self._open(self.project_root / "logs"))
        setup.add_command(label="Open configuration folder",
                          command=lambda: self._open(self.project_root / "config"))
        bar.add_cascade(label="Setup", menu=setup)
        more = tk.Menu(bar, tearoff=False)
        more.add_command(label="Run without email", command=lambda: self.run_weekly(no_email=True))
        more.add_command(label="Run again and re-send this week's email",
                         command=lambda: self.run_weekly(resend_email=True))
        bar.add_cascade(label="More", menu=more)
        help_menu = tk.Menu(bar, tearoff=False)
        help_menu.add_command(label="About this report", command=self.about)
        bar.add_cascade(label="Help", menu=help_menu)
        self.root.config(menu=bar)

    def _build_body(self) -> None:
        tk, ttk = self.tk, self.ttk
        outer = ttk.Frame(self.root, padding=16)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text="Negative Screening", style="Title.TLabel").pack(anchor="w")
        self.period_var = tk.StringVar()
        ttk.Label(outer, textvariable=self.period_var, wraplength=880).pack(anchor="w", pady=(2, 8))
        self.notice_var = tk.StringVar()
        self.notice = ttk.Label(outer, textvariable=self.notice_var, style="Notice.TLabel", wraplength=880)

        buttons = ttk.Frame(outer)
        buttons.pack(fill="x", pady=(0, 8))
        self.run_button = ttk.Button(buttons, text="Run Weekly Screening", style="Run.TButton",
                                     command=self.run_weekly)
        self.run_button.pack(side="left")
        self.latest_button = ttk.Button(buttons, text="Open Latest Report", style="Big.TButton",
                                        command=self.open_latest)
        self.latest_button.pack(side="left", padx=(10, 0))
        self.folder_button = ttk.Button(buttons, text="Open Reports Folder", style="Big.TButton",
                                        command=self.open_reports_folder)
        self.folder_button.pack(side="left", padx=(10, 0))
        self.print_button = ttk.Button(buttons, text="Print Report", style="Big.TButton",
                                       command=self.print_report)
        self.print_button.pack(side="left", padx=(10, 0))

        progress_row = ttk.Frame(outer)
        progress_row.pack(fill="x", pady=(0, 10))
        self.progress = ttk.Progressbar(progress_row, mode="determinate", value=0, length=220)  # empty when idle
        self.progress.pack(side="left")
        self.status_var = tk.StringVar(value="Ready.")
        ttk.Label(progress_row, textvariable=self.status_var, wraplength=640).pack(side="left", padx=10)

        last = ttk.LabelFrame(outer, text="Last run", padding=10)
        last.pack(fill="both", expand=True)
        self.fields = {}
        rows = [("targets", "Screened targets"), ("when", "Run time"), ("mode", "Started by"),
                ("period", "Reporting period"), ("result", "Result"), ("findings", "Findings"), ("email", "Email"),
                ("report", "Report file"), ("schedule", "Weekly schedule")]
        for r, (key, label) in enumerate(rows):
            ttk.Label(last, text=label, style="Key.TLabel").grid(row=r, column=0, sticky="nw", padx=(0, 12), pady=1)
            var = self.fields[key] = tk.StringVar()
            widget = ttk.Label(last, textvariable=var, wraplength=700, justify="left")
            widget.grid(row=r, column=1, sticky="w", pady=1)
            if key == "report":
                widget.configure(foreground="#0563C1", cursor="hand2")
                widget.bind("<Button-1>", lambda _e: self.open_latest())
        last.columnconfigure(1, weight=1)
        columns = ("source", "status", "found", "flagged", "note")
        self.tree = ttk.Treeview(last, columns=columns, show="headings", height=7)
        for col, title, width in [("source", "Source", 130), ("status", "Status", 80), ("found", "Found", 60),
                                  ("flagged", "Flagged", 60), ("note", "Note", 520)]:
            self.tree.heading(col, text=title)
            self.tree.column(col, width=width, anchor="w", stretch=(col == "note"))
        self.tree.tag_configure("failed", background="#FFE699")
        self.tree.grid(row=len(rows), column=0, columnspan=2, sticky="nsew", pady=(8, 0))
        last.rowconfigure(len(rows), weight=1)

        activity = ttk.LabelFrame(outer, text="Activity", padding=6)
        activity.pack(fill="both", expand=False, pady=(10, 0))
        self.log_text = tk.Text(activity, height=7, wrap="word", state="disabled", font=("Consolas", 9))
        scroll = ttk.Scrollbar(activity, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

    # ---- status ------------------------------------------------------------
    def production_db(self) -> Path:
        paths = self.cfg.settings["paths"] if self.cfg else DEFAULT_SETTINGS["paths"]
        return run_database_path(self.project_root, paths, dry_run=False, validation=False)

    def refresh_status(self) -> None:
        period = calculate_previous_completed_week(self.today())
        self.period_var.set(f"Run Weekly Screening checks the previous complete week: "
                            f"{period.period_start:%A %d %b %Y} – {period.period_end:%A %d %b %Y} "
                            f"(Malaysia time). The Excel report is saved in the reports folder and emailed.")
        self.info = read_last_run(self.production_db())
        d = describe_last_run(self.info)
        for key in ("when", "mode", "period", "result", "findings", "email", "report"):
            self.fields[key].set(d[key])
        self.fields["targets"].set(
            "\n".join(f"{x['name']} — {x['kind']}; searched as: {', '.join(x['names'])}"
                      for x in screened_targets(self.cfg.entities)) if self.cfg else "Configuration not loaded.")
        task = None
        try:
            task = self.task_query()
        except Exception:  # noqa: BLE001 — status display only
            task = None
        self.fields["schedule"].set(
            f"Installed — next run {task['next_run']} ({task['status']})" if task else
            "Not installed (Setup → Weekly schedule)")
        self.tree.delete(*self.tree.get_children())
        for source, status, found, flagged, note in d["sources"]:
            self.tree.insert("", "end", values=(source, status, found, flagged, note),
                             tags=("failed",) if status == "FAILED" else ())
        notes = []
        if self.config_error:
            notes.append(f"Configuration problem: {self.config_error}")
        if not get_secret("NS_IMAP_PASSWORD"):
            notes.append("Google Alerts email access has not been configured, so Facebook, Instagram, "
                         "LinkedIn and news results cannot be checked.")
        if not get_secret("NS_SMTP_PASSWORD"):
            notes.append("Email sending has not been configured, so reports are saved but not emailed.")
        if notes:
            notes.append("Open Setup → Email accounts and passwords.")
        if d["alerts_no_emails"]:
            notes.append("No Google Alerts emails arrived in the last run's week, so nothing from Google's index "
                         "(Facebook, Instagram, LinkedIn, news) was read. Open Setup → Google Alerts.")
        if notes:
            self.notice_var.set("  ".join(notes))
            self.notice.pack(fill="x", pady=(0, 10), before=self.run_button.master)
        else:
            self.notice.pack_forget()

    def _activity(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    # ---- actions -----------------------------------------------------------
    def _open(self, path: Path) -> None:
        from tkinter import messagebox
        try:
            open_path(path)
        except OSError as e:
            messagebox.showerror("Negative Screening", f"Could not open {path}:\n{e}", parent=self.root)

    def open_latest(self) -> None:
        from tkinter import messagebox
        path = latest_report(reports_folder(self.project_root, self.cfg), getattr(self, "info", None))
        if path is None:
            messagebox.showinfo("Negative Screening", "There is no report yet. Click Run Weekly Screening.",
                                parent=self.root)
            return
        self._open(path)

    def print_report(self) -> None:
        """Printable page of the last run (browser print dialog: printer or Save as PDF)."""
        import webbrowser
        from tkinter import messagebox
        try:
            path = write_print_view(self.project_root, self.cfg, self.production_db()) if self.cfg else None
        except (OSError, sqlite3.Error) as e:
            messagebox.showerror("Negative Screening", f"Could not prepare the printable report:\n{e}",
                                 parent=self.root)
            return
        if path is None:
            messagebox.showinfo("Negative Screening", "There is no report yet. Click Run Weekly Screening.",
                                parent=self.root)
            return
        self._activity(f"Printable report: {path}")
        webbrowser.open(path.as_uri())

    def alerts_setup(self) -> None:
        """Guide for creating the Google Alerts that bring in links from Google's index."""
        import webbrowser
        tk, ttk = self.tk, self.ttk
        from tkinter import messagebox
        if self.cfg is None:
            messagebox.showerror("Negative Screening", f"Configuration problem:\n{self.config_error}",
                                 parent=self.root)
            return
        box = self.cfg.email["alerts_mailbox"]
        queries = generate_alert_queries(self.cfg.entities, self.cfg.keywords)
        win = tk.Toplevel(self.root)
        win.title("Google Alerts — search-index discovery")
        win.transient(self.root)
        frame = ttk.Frame(win, padding=16)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, wraplength=700, justify="left", text=(
            "Public Facebook, Instagram, LinkedIn and news links come from Google's index through Google "
            "Alerts emails — Google itself is never scraped. Signed in to Google as the account that owns the "
            f"alerts, create one alert per query below (How often: At most once a day · Sources: Automatic · "
            f"Language: Any · Region: Any · How many: All results · Deliver to: {box['username']}). Then add a "
            f"Gmail filter: from {box['sender_filter']} → apply label '{box['label']}'. The program reads "
            "those emails (read-only) and takes each result's link, title, snippet and date.")).pack(anchor="w")
        text = tk.Text(frame, height=12, width=100, wrap="none", font=("Consolas", 9))
        text.insert("1.0", "\n".join(queries))
        text.configure(state="disabled")
        text.pack(fill="both", expand=True, pady=10)
        status = tk.StringVar(value=f"{len(queries)} alert queries from the approved names and keywords.")
        ttk.Label(frame, textvariable=status, wraplength=700, justify="left").pack(anchor="w")

        def copy() -> None:
            win.clipboard_clear()
            win.clipboard_append("\n".join(queries))
            status.set(f"Copied {len(queries)} queries to the clipboard.")

        def check() -> None:
            cfg = self.cfg
            if cfg is None:
                status.set("Configuration is not available.")
                return
            status.set("Checking the mailbox…")
            self._in_background(
                lambda: check_imap_account(cfg.email, get_secret("NS_IMAP_PASSWORD"),
                                           cfg.settings["http"]["timeout_seconds"]),
                lambda r: status.set(r[1]))

        row = ttk.Frame(frame)
        row.pack(anchor="e", pady=(10, 0))
        ttk.Button(row, text="Copy all queries", command=copy).pack(side="left")
        ttk.Button(row, text="Open Google Alerts", command=lambda: webbrowser.open("https://www.google.com/alerts")
                   ).pack(side="left", padx=8)
        ttk.Button(row, text="Check mailbox for alert emails", command=check).pack(side="left")
        ttk.Button(row, text="Close", command=win.destroy).pack(side="left", padx=(8, 0))

    def open_reports_folder(self) -> None:
        folder = reports_folder(self.project_root, self.cfg)
        folder.mkdir(parents=True, exist_ok=True)
        self._open(folder)

    def run_weekly(self, no_email: bool = False, resend_email: bool = False) -> None:
        import threading
        from tkinter import messagebox
        if self.running:
            return
        if self.config_error:
            messagebox.showerror("Negative Screening", f"Configuration problem:\n{self.config_error}",
                                 parent=self.root)
            return
        if not no_email and not get_secret("NS_SMTP_PASSWORD"):
            if not messagebox.askyesno(
                    "Negative Screening",
                    "Email sending has not been configured, so the report will be saved but NOT emailed.\n\n"
                    "Run the screening anyway? (Set it up in Setup → Email accounts and passwords.)",
                    parent=self.root):
                return
        if resend_email and not messagebox.askyesno(
                "Negative Screening", "Run the screening again and re-send this week's report email, "
                                      "even if it was already sent?", parent=self.root):
            return
        request = staff_run_request(self.today(), no_email=no_email, resend_email=resend_email)
        self.running = True
        for b in (self.run_button, self.latest_button):
            b.state(["disabled"])
        self.progress.configure(mode="indeterminate")
        self.progress.start(12)
        self.status_var.set("Running — this usually takes a minute or two…")
        self._activity(f"--- {now_in_malaysia():%H:%M:%S} screening started"
                       + (" (without email)" if no_email else ""))
        handler = _QueueLogHandler(self.queue)

        def work() -> None:
            try:
                outcome = execute_run(request, self.project_root, handler)
            except Exception as e:  # noqa: BLE001 — shown to staff, never a traceback
                outcome = RunOutcome(2, f"Unexpected error: {_short_error(e)}")
            self.queue.put(("done", outcome))

        threading.Thread(target=work, name="screening-run", daemon=True).start()

    def _poll(self) -> None:
        import queue
        try:
            while True:
                item = self.queue.get_nowait()
                if item[0] == "log":
                    _, level, text = item
                    if level in ("WARNING", "ERROR") or not text.startswith(("source started", "counts:")):
                        self._activity(text)
                    for prefix, stage in self.STAGES:
                        if text.startswith(prefix):
                            self.status_var.set(stage)
                elif item[0] == "done":
                    self._finished(item[1])
                elif item[0] == "call":
                    item[1]()
        except queue.Empty:
            pass
        try:
            self.root.after(150, self._poll)
        except self.tk.TclError:        # window closed
            pass

    def _finished(self, outcome: RunOutcome) -> None:
        from tkinter import messagebox
        self.running = False
        self.progress.stop()
        self.progress.configure(mode="determinate", value=0)
        for b in (self.run_button, self.latest_button):
            b.state(["!disabled"])
        self.refresh_status()
        self.status_var.set(outcome.message)
        lines = [outcome.message]
        if outcome.report_path:
            lines.append(f"\nReport saved:\n{outcome.report_path}")
            self._activity(f"Report saved: {outcome.report_path}")
        if outcome.run_id and self.info and self.info["run"].get("run_id") == outcome.run_id:
            lines.append(f"\nEmail: {self.fields['email'].get()}")
        show = messagebox.showinfo if outcome.exit_code == 0 else (
            messagebox.showwarning if outcome.exit_code in (1, 3) else messagebox.showerror)
        show("Negative Screening", "\n".join(lines), parent=self.root)

    def _in_background(self, fn, done) -> None:
        """Run fn() off the window thread, then done(result) on it."""
        import threading

        def work() -> None:
            try:
                result = fn()
            except Exception as e:  # noqa: BLE001
                result = (False, _short_error(e))
            self.queue.put(("call", lambda: done(result)))
        threading.Thread(target=work, daemon=True).start()

    # ---- setup dialogs -----------------------------------------------------
    def email_setup(self) -> None:
        tk, ttk = self.tk, self.ttk
        from tkinter import messagebox
        if self.cfg is None:
            messagebox.showerror("Negative Screening", f"Configuration problem:\n{self.config_error}",
                                 parent=self.root)
            return
        email_cfg = self.cfg.email
        timeout = self.cfg.settings["http"]["timeout_seconds"]
        win = tk.Toplevel(self.root)
        win.title("Email accounts and passwords")
        win.transient(self.root)
        frame = ttk.Frame(win, padding=16)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, wraplength=620, justify="left", text=(
            "Enter Gmail App Passwords (16 letters, from Google Account → Security → App passwords), not the "
            "normal account password. They are saved in Windows Credential Manager for your Windows user "
            "only — never in files, logs, reports or Git.")).grid(row=0, column=0, columnspan=4, sticky="w")
        entries: dict[str, Any] = {}
        states: dict[str, Any] = {}
        accounts = [("NS_IMAP_PASSWORD", "Google Alerts mailbox",
                     f"{email_cfg['alerts_mailbox']['username']} (label {email_cfg['alerts_mailbox']['label']})"),
                    ("NS_SMTP_PASSWORD", "Report sending account",
                     f"{email_cfg['smtp'].get('username') or email_cfg['sender']} via {email_cfg['smtp']['host']}")]

        def describe(name: str) -> str:
            where = secret_source(name)
            if where == "environment variable":
                return f"Set by the {name} environment variable (it takes precedence)."
            return "Saved." if where else "Not configured."

        for i, (name, title, account) in enumerate(accounts):
            r = 1 + i * 3
            ttk.Label(frame, text=title, style="Key.TLabel").grid(row=r, column=0, sticky="w", pady=(14, 0))
            ttk.Label(frame, text=account).grid(row=r, column=1, columnspan=3, sticky="w", pady=(14, 0))
            ttk.Label(frame, text="App password").grid(row=r + 1, column=0, sticky="w")
            entries[name] = tk.StringVar()
            ttk.Entry(frame, textvariable=entries[name], show="•", width=32).grid(row=r + 1, column=1, sticky="w")
            states[name] = tk.StringVar(value=describe(name))
            ttk.Label(frame, textvariable=states[name]).grid(row=r + 2, column=1, columnspan=3, sticky="w")

        def password(name: str) -> str:
            return normalise_app_password(entries[name].get()) or get_secret(name)

        def show_result(result) -> None:
            ok, message = result
            (messagebox.showinfo if ok else messagebox.showerror)("Negative Screening", message, parent=win)

        ttk.Button(frame, text="Test mailbox", command=lambda: self._in_background(
            lambda: check_imap_account(email_cfg, password("NS_IMAP_PASSWORD"), timeout), show_result)
        ).grid(row=2, column=2, padx=6)
        ttk.Button(frame, text="Test sign-in", command=lambda: self._in_background(
            lambda: check_smtp_account(email_cfg, password("NS_SMTP_PASSWORD"), timeout), show_result)
        ).grid(row=5, column=2, padx=6)
        ttk.Button(frame, text="Send test email", command=lambda: self._in_background(
            lambda: check_smtp_account(email_cfg, password("NS_SMTP_PASSWORD"), timeout,
                                       send_test_to=list(email_cfg["recipients"])), show_result)
        ).grid(row=5, column=3)
        ttk.Label(frame, wraplength=620, justify="left",
                  text="Reports are emailed to: " + ", ".join(email_cfg["recipients"])
                       + "  (recipients and accounts are set in config\\email.json)").grid(
            row=8, column=0, columnspan=4, sticky="w", pady=(14, 0))

        def save() -> None:
            saved = []
            for name, var in entries.items():
                value = normalise_app_password(var.get())
                if value:
                    try:
                        if SECRET_STORE is None:
                            raise OSError("Secure password storage is unavailable.")
                        SECRET_STORE.write(name, value)
                    except (OSError, AttributeError) as e:
                        messagebox.showerror("Negative Screening", f"Could not save the password: {e}", parent=win)
                        return
                    saved.append(SECRET_LABELS[name])
                    var.set("")
                states[name].set(describe(name))
            self.refresh_status()
            messagebox.showinfo("Negative Screening", ("Saved: " + "; ".join(saved)) if saved else
                                "Nothing to save: enter a password first.", parent=win)

        def remove() -> None:
            if not messagebox.askyesno("Negative Screening", "Remove both saved passwords from this computer?",
                                       parent=win):
                return
            for name in entries:
                try:
                    if SECRET_STORE is None:
                        raise OSError("Secure password storage is unavailable.")
                    SECRET_STORE.delete(name)
                except (OSError, AttributeError):
                    pass
                states[name].set(describe(name))
            self.refresh_status()

        row = ttk.Frame(frame)
        row.grid(row=9, column=0, columnspan=4, sticky="e", pady=(16, 0))
        ttk.Button(row, text="Save passwords", command=save).pack(side="left")
        ttk.Button(row, text="Remove saved passwords", command=remove).pack(side="left", padx=8)
        ttk.Button(row, text="Close", command=win.destroy).pack(side="left")

    def schedule_setup(self) -> None:
        tk, ttk = self.tk, self.ttk
        from tkinter import messagebox
        win = tk.Toplevel(self.root)
        win.title("Weekly schedule")
        win.transient(self.root)
        frame = ttk.Frame(win, padding=16)
        frame.pack(fill="both", expand=True)
        status = tk.StringVar(value="Checking…")
        ttk.Label(frame, wraplength=600, justify="left", text=(
            f"Windows Task Scheduler runs '{SCHEDULED_TASK_NAME}' every Monday at the time below for the "
            "previous complete week and emails the report. If the computer is off or asleep then, it runs as "
            "soon as possible afterwards. Only one run at a time. It runs under your Windows account (which "
            "holds the saved passwords), so you must be signed in. Logs go to the logs folder.")
                  ).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(frame, textvariable=status, wraplength=600).grid(row=1, column=0, columnspan=3, sticky="w",
                                                                  pady=10)
        ttk.Label(frame, text="Monday start time (HH:MM)").grid(row=2, column=0, sticky="w")
        at = tk.StringVar(value=SCHEDULED_TASK_DEFAULT_TIME)
        ttk.Entry(frame, textvariable=at, width=8).grid(row=2, column=1, sticky="w")

        def show_status() -> None:
            self._in_background(lambda: (True, self.task_query()), lambda r: status.set(
                f"Installed — next run {r[1]['next_run']}; last run {r[1]['last_run']} "
                f"(result {r[1]['last_result']})." if r[1] else "Not installed."))

        def done(result) -> None:
            ok, message = result
            (messagebox.showinfo if ok else messagebox.showerror)("Negative Screening", message, parent=win)
            show_status()
            self.refresh_status()

        limit = self.cfg.settings["run"]["max_run_minutes"] if self.cfg else 120
        buttons = ttk.Frame(frame)
        buttons.grid(row=3, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="Install / update weekly task", command=lambda: self._in_background(
            lambda: install_scheduled_task(self.project_root, at.get().strip(), limit), done)).pack(side="left")
        ttk.Button(buttons, text="Remove weekly task",
                   command=lambda: self._in_background(uninstall_scheduled_task, done)).pack(side="left", padx=8)
        ttk.Button(buttons, text="Close", command=win.destroy).pack(side="left")
        show_status()

    def about(self) -> None:
        from tkinter import messagebox
        messagebox.showinfo("About this report", (
            f"Negative Screening v{APP_VERSION}\n\n{COVERAGE_STATEMENT}\n\n{SCREENING_NOTE}\n\n"
            "Sources: Google Alerts (public Facebook, Instagram, LinkedIn and web content in Google's index), "
            "the Company website and Securities Commission Malaysia. Screening is rule-based: an approved name "
            "plus an adverse keyword. A FAILED source means it could not be checked — not that nothing was "
            f"found.\n\nProject folder: {self.project_root}"), parent=self.root)

    def close(self) -> None:
        from tkinter import messagebox
        if self.running and not messagebox.askyesno(
                "Negative Screening", "A screening run is still in progress. Close anyway? The run will stop "
                                      "and must be started again.", parent=self.root):
            return
        self.root.destroy()


def run_gui(project_root: Path | None = None) -> int:
    """Open the staff window (NegativeScreening.exe without options, or --gui)."""
    try:
        import tkinter as tk
    except ModuleNotFoundError:
        print("The staff window needs Tkinter (part of the standard Windows Python install).", file=sys.stderr)
        return 2
    if os.name == "nt":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)      # sharp text on scaled displays
        except (AttributeError, OSError):
            pass
    root = tk.Tk()
    StaffApp(root, project_root or PROJECT_ROOT)
    root.mainloop()
    return 0


# =============================================================================
# SELF-TEST HARNESS (no pytest needed)
# Outcomes: PASS · FAIL · PENDING (NotImplementedError) · SKIPPED (missing package)
# =============================================================================

_SELF_TESTS: dict[str, list[Callable[[], None]]] = {}


def selftest(step: str):
    def register(fn: Callable[[], None]):
        _SELF_TESTS.setdefault(step, []).append(fn)
        return fn
    return register


def _run_step_tests(step: str) -> dict[str, int]:
    global SECRET_STORE
    tally = {"PASS": 0, "FAIL": 0, "PENDING": 0, "SKIPPED": 0}
    for fn in _SELF_TESTS.get(step, []):
        name = fn.__name__
        saved_store, SECRET_STORE = SECRET_STORE, MemorySecretStore()   # never the machine's passwords
        try:
            fn()
            outcome, detail = "PASS", ""
        except NotImplementedError as e:
            outcome, detail = "PENDING", f"not implemented: {e}"
        except ModuleNotFoundError as e:
            outcome, detail = "SKIPPED", f"needs {e.name} (pip install -r requirements.txt)"
        except AssertionError as e:
            outcome, detail = "FAIL", str(e) or "assertion failed"
        except Exception as e:  # noqa: BLE001 — report any crash as a failure
            outcome, detail = "FAIL", f"{type(e).__name__}: {e}"
        finally:
            SECRET_STORE = saved_store
        tally[outcome] += 1
        print(f"  [{outcome:7}] {step} {name}" + (f" — {detail}" if detail else ""))
    return tally


def run_self_tests(which: str) -> int:
    steps = list(STEP_STATUS) if which.lower() == "all" else [which.upper()]
    unknown = [s for s in steps if s not in STEP_STATUS]
    if unknown:
        print(f"unknown step: {', '.join(unknown)}")
        return 2
    total = {"PASS": 0, "FAIL": 0, "PENDING": 0, "SKIPPED": 0}
    for s in steps:
        print(f"{s} {STEP_TITLES[s]}  (status: {STEP_STATUS[s]})")
        if not _SELF_TESTS.get(s):
            print("  (no acceptance tests registered)")
        for k, v in _run_step_tests(s).items():
            total[k] += v
    print(f"\nTotal: {total['PASS']} passed, {total['FAIL']} failed, "
          f"{total['PENDING']} pending, {total['SKIPPED']} skipped")
    done_but_not_passing = [s for s in steps if STEP_STATUS[s] == "DONE"
                            and any(True for _ in _SELF_TESTS.get(s, []))
                            and _step_has_non_pass(s)]
    if done_but_not_passing:
        print("WARNING: marked DONE but tests not all passing: " + ", ".join(done_but_not_passing))
    return 1 if total["FAIL"] or done_but_not_passing else 0


def _step_has_non_pass(step: str) -> bool:
    import contextlib
    import io
    with contextlib.redirect_stdout(io.StringIO()):
        t = _run_step_tests(step)
    return bool(t["FAIL"] or t["PENDING"] or t["SKIPPED"])


def print_status() -> None:
    import contextlib
    import io
    print(f"Negative Screening build status (v{APP_VERSION})\n")
    print(f"{'Step':5} {'Title':34} {'Status':12} Tests (pass/fail/pending/skipped)")
    next_step = None
    for s, st in STEP_STATUS.items():
        with contextlib.redirect_stdout(io.StringIO()):
            t = _run_step_tests(s)
        print(f"{s:5} {STEP_TITLES[s]:34} {st:12} "
              f"{t['PASS']}/{t['FAIL']}/{t['PENDING']}/{t['SKIPPED']}")
        if next_step is None and st != "DONE":
            next_step = s
    print()
    if next_step:
        print(f"Next step: {next_step} — {STEP_TITLES[next_step]}. "
              f"Read its section banner, implement, then run --self-test {next_step}.")
    else:
        print("All steps DONE. Run a dry run, then a real run for a past week.")


# =============================================================================
# ACCEPTANCE TESTS — written ahead of the code. Do not weaken; add more freely.
# =============================================================================

def _utc(y, mo, d, h=0, mi=0) -> datetime:
    return datetime(y, mo, d, h, mi, tzinfo=timezone.utc)


def _myt(y, mo, d, h=0, mi=0) -> datetime:
    return datetime(y, mo, d, h, mi, tzinfo=MYT)


def _period(ref: str) -> ReportingPeriod:
    return calculate_previous_completed_week(date.fromisoformat(ref))


def _candidate(url: str, title: str = "", snippet: str = "", source: str = "google_alerts",
               **raw_kwargs) -> Candidate:
    raw = RawItem(source=source, discovery_method="test", discovered_url=url,
                  title=title, snippet=snippet, **raw_kwargs)
    return Candidate(raw=raw, canonical_url=url)


def _test_config() -> Config:
    return Config(settings=json.loads(json.dumps(DEFAULT_SETTINGS)),
                  entities=json.loads(json.dumps(DEFAULT_ENTITIES)),
                  keywords=json.loads(json.dumps(DEFAULT_KEYWORDS)),
                  email=json.loads(json.dumps(DEFAULT_EMAIL)), config_dir=Path("."))


# ---- S01 --------------------------------------------------------------------
@selftest("S01")
def test_models_defaults():
    c = _candidate("https://example.com/a")
    assert c.date_status == DateStatus.UNKNOWN, "dates must default to UNKNOWN"
    assert c.review_status == ReviewStatus.PENDING
    assert c.raw.source_published_at is None and c.raw.alert_received_at is None
    assert PRIORITY_RANK[Priority.CRITICAL] > PRIORITY_RANK[Priority.HIGH] > PRIORITY_RANK[Priority.LOW]


@selftest("S01")
def test_period_contains_requires_aware_datetime():
    p = _period("2026-09-22")
    try:
        p.contains(datetime(2026, 9, 15, 12, 0))
    except ValueError:
        pass
    else:
        raise AssertionError("naive datetimes must be rejected")


# ---- S02 --------------------------------------------------------------------
@selftest("S02")
def test_period_examples_from_brief():
    cases = {
        "2026-09-21": ("2026-09-14", "2026-09-20"),   # Monday
        "2026-09-22": ("2026-09-14", "2026-09-20"),   # Tuesday
        "2026-09-25": ("2026-09-14", "2026-09-20"),   # Friday
        "2026-09-27": ("2026-09-14", "2026-09-20"),   # Sunday
        "2026-09-28": ("2026-09-21", "2026-09-27"),   # next Monday
        "2026-09-29": ("2026-09-21", "2026-09-27"),
        "2027-01-01": ("2026-12-21", "2026-12-27"),
        "2027-01-04": ("2026-12-28", "2027-01-03"),   # year boundary
    }
    for ref, (s, e) in cases.items():
        p = _period(ref)
        assert (p.period_start.isoformat(), p.period_end.isoformat()) == (s, e), \
            f"{ref}: got {p.period_start}–{p.period_end}, expected {s}–{e}"


@selftest("S02")
def test_period_invariants_over_two_years():
    d = date(2026, 1, 1)
    while d < date(2028, 1, 1):
        p = calculate_previous_completed_week(d)
        assert p.period_start.weekday() == 0 and p.period_end.weekday() == 6, d
        assert (p.period_end - p.period_start).days == 6, d
        assert p.period_end < d, f"{d}: period must end before the reference date"
        assert (d - p.period_end).days <= 7, f"{d}: must be the IMMEDIATELY preceding week"
        d += timedelta(days=1)


@selftest("S02")
def test_period_boundaries_in_malaysia_time():
    p = _period("2026-09-25")
    assert p.contains(_myt(2026, 9, 14, 0, 0)), "Monday 00:00 MYT is inside"
    assert p.contains(_myt(2026, 9, 20, 23, 59)), "Sunday 23:59 MYT is inside"
    assert not p.contains(_myt(2026, 9, 21, 0, 0)), "next Monday 00:00 MYT is outside"
    assert not p.contains(_myt(2026, 9, 24, 12, 0)), "current-week content is outside"
    # 2026-09-20 17:00 UTC = 2026-09-21 01:00 MYT -> outside
    assert not p.contains(_utc(2026, 9, 20, 17, 0)), "UTC times must be judged in MYT"
    # 2026-09-13 16:30 UTC = 2026-09-14 00:30 MYT -> inside
    assert p.contains(_utc(2026, 9, 13, 16, 30))


@selftest("S02")
def test_cli_rejects_date_with_scheduled():
    args = build_arg_parser().parse_args(["--scheduled", "--date", "2026-09-22"])
    assert validate_args(args), "--date with --scheduled must be rejected"
    args = build_arg_parser().parse_args(["--date", "2026-09-22"])
    assert validate_args(args) is None
    assert resolve_mode_and_reference(args) == (Mode.MANUAL_TEST_DATE, date(2026, 9, 22))


# ---- S03 --------------------------------------------------------------------
@selftest("S03")
def test_init_and_load_config_roundtrip():
    with tempfile.TemporaryDirectory() as d:
        written = init_config(Path(d))
        assert {p.name for p in written} == set(CONFIG_FILES), written
        assert init_config(Path(d)) == [], "existing files must not be overwritten"
        cfg = load_config(Path(d))
        ids = {e["id"] for e in cfg.entities["entities"]}
        assert ids == {"acme_company", "tan_ah_kow"}, ids
        cats = {c["id"]: c for c in cfg.keywords["categories"]}
        assert len(cats) == 9
        assert cats["sanctions_regulatory"]["priority"] == "CRITICAL"
        terms = {k["term"] for c in cats.values() for k in c["keywords"]}
        assert "arrested" in terms and "rrested" not in terms


@selftest("S03")
def test_unapproved_entities_are_dropped():
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        path = Path(d) / "entities.json"
        data = json.loads(path.read_text(encoding="utf-8"))
        data["entities"].append({"id": "x", "type": "person", "display_name": "Unapproved Person",
                                 "names": ["Unapproved Person"], "variants": [],
                                 "ambiguous_terms": [], "approved": False})
        path.write_text(json.dumps(data), encoding="utf-8")
        cfg = load_config(Path(d))
        assert "x" not in {e["id"] for e in cfg.entities["entities"]}


@selftest("S03")
def test_config_rejects_passwords_and_names_the_file():
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        path = Path(d) / "email.json"
        data = json.loads(path.read_text(encoding="utf-8"))
        data["smtp"]["password"] = "should-not-be-here"
        path.write_text(json.dumps(data), encoding="utf-8")
        try:
            load_config(Path(d))
        except ConfigError as e:
            assert "email.json" in str(e), f"error must name the file: {e}"
        else:
            raise AssertionError("a password in email.json must be rejected")


def _config_error_for(filename: str, mutate: Callable[[dict], object]) -> str:
    """Write default config, apply `mutate` to one file, return load_config's error text."""
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        path = Path(d) / filename
        data = json.loads(path.read_text(encoding="utf-8"))
        mutate(data)
        path.write_text(json.dumps(data), encoding="utf-8")
        try:
            load_config(Path(d))
        except ConfigError as e:
            return str(e)
    raise AssertionError(f"invalid {filename} was accepted")


@selftest("S03")
def test_config_files_match_templates_and_overwrite():
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        cfg = load_config(Path(d))
        assert cfg.settings == DEFAULT_SETTINGS and cfg.keywords == DEFAULT_KEYWORDS
        assert cfg.email == DEFAULT_EMAIL and cfg.entities == DEFAULT_ENTITIES
        assert cfg.config_dir == Path(d)
        (Path(d) / "settings.json").write_text("{}", encoding="utf-8")
        assert len(init_config(Path(d), overwrite=True)) == len(CONFIG_FILES)
        assert load_config(Path(d)).settings == DEFAULT_SETTINGS, "overwrite=True must restore the template"


@selftest("S03")
def test_config_file_level_errors_and_bom():
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        kw = Path(d) / "adverse_keywords.json"
        good = kw.read_text(encoding="utf-8")
        kw.write_text("﻿" + good, encoding="utf-8")
        assert load_config(Path(d)).keywords == DEFAULT_KEYWORDS, "UTF-8 with BOM (Notepad) must load"
        kw.write_text('{"version": "x",', encoding="utf-8")
        try:
            load_config(Path(d))
        except ConfigError as e:
            assert "adverse_keywords.json" in str(e) and "line 1" in str(e), e
        else:
            raise AssertionError("invalid JSON must be rejected")
        kw.write_text(good, encoding="utf-8")
        (Path(d) / "email.json").unlink()
        try:
            load_config(Path(d))
        except ConfigError as e:
            assert "email.json" in str(e) and "not found" in str(e), e
        else:
            raise AssertionError("a missing file must be rejected")


@selftest("S03")
def test_config_validation_names_file_and_field():
    def no_approved(d):
        for e in d["entities"]:
            e["approved"] = False

    cases = [
        ("entities.json", lambda d: d["entities"][0].update(names=[]), "entities[0].names"),
        ("entities.json", lambda d: d["entities"][0].pop("display_name"), "entities[0].display_name"),
        ("entities.json", lambda d: d["entities"][1].update(approved="true"), "entities[1].approved"),
        ("entities.json", lambda d: d["entities"][1].update(id="acme_company"), "entities[1].id"),
        ("entities.json", lambda d: d["entities"][0].update(ambiguous_terms=[
            {"term": "Example", "requires_any_context": ["fund"], "context_window_words": 0}]),
         "entities[0].ambiguous_terms[0].context_window_words"),
        ("entities.json", no_approved, "entities"),
        ("adverse_keywords.json", lambda d: d["categories"][0].update(priority="URGENT"), "categories[0].priority"),
        ("adverse_keywords.json", lambda d: d["categories"][1]["keywords"][0].update(ambiguous="no"),
         "categories[1].keywords[0].ambiguous"),
        ("adverse_keywords.json", lambda d: d["categories"][2]["keywords"].append("fraud"),
         "categories[2].keywords[2]"),
        ("email.json", lambda d: d.pop("sender"), "sender"),
        ("email.json", lambda d: d.update(recipients=[]), "recipients"),
        ("email.json", lambda d: d.update(recipients=["not-an-address"]), "recipients"),
        ("email.json", lambda d: d["alerts_mailbox"].pop("label"), "alerts_mailbox.label"),
        ("settings.json", lambda d: d.update(timezone="UTC"), "timezone"),
        ("settings.json", lambda d: d["screening"].update(keyword_name_window_words=0),
         "screening.keyword_name_window_words"),
        ("settings.json", lambda d: d.update(nlp={"model_name": "any"}), "nlp"),
        ("settings.json", lambda d: d["run"].update(max_run_minutes=True), "run.max_run_minutes"),
        ("settings.json", lambda d: d["sources"]["google_alerts"].update(mode="pop3"), "sources.google_alerts.mode"),
        ("settings.json", lambda d: d["sources"]["sc_malaysia"].update(enabled="yes"),
         "sources.sc_malaysia.enabled"),
    ]
    for filename, mutate, field_name in cases:
        msg = _config_error_for(filename, mutate)
        assert msg.startswith(f"{filename}: {field_name}:"), f"expected '{filename}: {field_name}:' in: {msg}"


@selftest("S03")
def test_config_imap_mailbox_optional_in_rss_mode():
    with tempfile.TemporaryDirectory() as d:
        init_config(Path(d))
        s, e = Path(d) / "settings.json", Path(d) / "email.json"
        settings = json.loads(s.read_text(encoding="utf-8"))
        settings["sources"]["google_alerts"]["mode"] = "rss"
        s.write_text(json.dumps(settings), encoding="utf-8")
        email = json.loads(e.read_text(encoding="utf-8"))
        del email["alerts_mailbox"]
        e.write_text(json.dumps(email), encoding="utf-8")
        load_config(Path(d))   # must not raise: RSS mode needs no mailbox


@selftest("S03")
def test_config_secrets_rejected_in_any_file_without_echoing_value():
    msg = _config_error_for("settings.json", lambda d: d["http"].update(proxy_password="hunter2-value"))
    assert msg.startswith("settings.json: http.proxy_password:"), msg
    assert "hunter2-value" not in msg, "the error message must never echo a secret"
    msg = _config_error_for("entities.json", lambda d: d["entities"][0].update(api_secret="x"))
    assert msg.startswith("entities.json: entities[0].api_secret:"), msg


# ---- S04 --------------------------------------------------------------------
@selftest("S04")
def test_secrets_are_redacted_in_logs():
    secret = "Sup3r-S3cret-AppPass"
    old = os.environ.get("NS_SMTP_PASSWORD")
    os.environ["NS_SMTP_PASSWORD"] = secret
    try:
        with tempfile.TemporaryDirectory() as d:
            logger = setup_logging(Path(d), "testrun")
            logger.info("connecting with %s", secret)
            logger.info("password=%s", "another-secret")
            for h in logger.handlers:
                h.flush()
            text = "".join(p.read_text(encoding="utf-8") for p in Path(d).glob("*.log"))
            assert "connecting with" in text, "log files must be written"
            assert secret not in text and "another-secret" not in text, "secret leaked into logs"
            for h in list(logger.handlers):
                h.close()
                logger.removeHandler(h)
    finally:
        if old is None:
            os.environ.pop("NS_SMTP_PASSWORD", None)
        else:
            os.environ["NS_SMTP_PASSWORD"] = old


def _close_logger(logger: logging.Logger) -> None:
    for h in list(logger.handlers):
        h.close()
        logger.removeHandler(h)


@selftest("S04")
def test_redaction_patterns():
    old = os.environ.get("NS_IMAP_PASSWORD")
    os.environ["NS_IMAP_PASSWORD"] = "abcd efgh ijkl mnop"
    try:
        cases = {
            "login with abcdefghijklmnop": "abcdefghijklmnop",     # app password without spaces
            "login with abcd efgh ijkl mnop": "abcd efgh ijkl mnop",
            '{"password": "p@ss w0rd"}': "p@ss w0rd",
            "GET /x?access_token=AAA111&page=2": "AAA111",
            "api_key: K-999": "K-999",
            "Authorization: Bearer eyJhbGciOi.xyz": "eyJhbGciOi.xyz",
        }
        for text, secret in cases.items():
            out = redact(text)
            assert secret not in out and REDACTED in out, f"not redacted: {text!r} -> {out!r}"
        assert redact("GET /x?access_token=AAA111&page=2").endswith("&page=2"), "over-redacted"
        for harmless in ("max_tokens=512", "tokens used: 12", "Acme Asset Management was fined"):
            assert redact(harmless) == harmless, f"harmless text changed: {harmless!r}"
    finally:
        if old is None:
            os.environ.pop("NS_IMAP_PASSWORD", None)
        else:
            os.environ["NS_IMAP_PASSWORD"] = old


@selftest("S04")
def test_logging_files_idempotence_and_tracebacks():
    secret = "Tr4ceback-S3cret"
    old = os.environ.get("NS_SMTP_PASSWORD")
    os.environ["NS_SMTP_PASSWORD"] = secret
    try:
        with tempfile.TemporaryDirectory() as d:
            setup_logging(Path(d), "run1")
            logger = setup_logging(Path(d), "run1", verbose=True)
            assert len(logger.handlers) == 3, f"duplicate handlers: {logger.handlers}"
            assert logger.propagate is False
            console = [h for h in logger.handlers if type(h) is logging.StreamHandler]
            assert console and console[0].level == logging.DEBUG, "--verbose must show DEBUG on console"
            console[0].setLevel(logging.CRITICAL)   # keep the self-test output readable
            try:
                raise ConnectionError(f"login failed for {secret}")
            except ConnectionError:
                logger.exception("mailbox error")
            logger.info("name: 新加坡 — Tan Ah Kow")
            _close_logger(logger)
            names = {p.name for p in Path(d).glob("*.log")}
            assert names == {"negative_screening.log", "run_run1.log"}, names
            text = (Path(d) / "run_run1.log").read_text(encoding="utf-8")
            assert "mailbox error" in text and "ConnectionError" in text, "traceback must be logged"
            assert secret not in text, "secret leaked through an exception traceback"
            assert "新加坡" in text, "log files must keep non-ASCII text"
            assert "+08:00" in text, "log timestamps must be in Malaysia time"
    finally:
        if old is None:
            os.environ.pop("NS_SMTP_PASSWORD", None)
        else:
            os.environ["NS_SMTP_PASSWORD"] = old


# ---- S05 --------------------------------------------------------------------
def _test_ctx(db_path: Path, ref: str = "2026-09-22") -> RunContext:
    return RunContext(run_id=f"test-{ref}", mode=Mode.MANUAL, period=_period(ref),
                      started_at=_myt(2026, 9, 22, 10, 0), dry_run=False, config=_test_config(),
                      logger=logging.getLogger("ns-test"), db_path=db_path)


@selftest("S05")
def test_upsert_is_rerun_safe_and_keeps_staff_review():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        db_record_run_start(conn, ctx)
        c = _candidate("https://news.example.com/a", title="t", snippet="s", source="google_alerts",
                       discovery_query="q1")
        c.url_key, c.platform, c.content_class = "k1", Platform.PUBLIC_WEB, ContentClass.PUBLIC_WEB
        c.verification, c.section, c.priority = Verification.VERIFIED, Section.MAIN, Priority.HIGH
        c.discovery_queries = {"q1"}
        fid1 = db_upsert_candidate(conn, c, ctx)
        conn.execute("UPDATE finding_periods SET review_status='FALSE POSITIVE' WHERE finding_id=?",
                     (fid1,))
        conn.commit()
        c.discovery_queries = {"q1", "q2"}
        fid2 = db_upsert_candidate(conn, c, ctx)
        assert fid1 == fid2, "same url_key must map to the same finding"
        n = conn.execute("SELECT COUNT(*) FROM findings").fetchone()[0]
        assert n == 1, f"re-run created duplicates: {n} findings"
        rs = conn.execute("SELECT review_status FROM finding_periods WHERE finding_id=?",
                          (fid1,)).fetchone()[0]
        assert rs == "FALSE POSITIVE", "staff review decision was overwritten"
        conn.close()


@selftest("S05")
def test_email_log_detects_already_sent():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        db_record_run_start(conn, ctx)
        assert not db_email_already_sent(conn, ctx.period)
        db_record_email(conn, ctx.run_id, ctx.period, "SENT", 1, None)
        assert db_email_already_sent(conn, ctx.period)
        conn.close()


def _db_candidate(url_key: str = "k1", queries: set[str] | None = None) -> Candidate:
    c = _candidate("https://news.example.com/a", title="t", snippet="s", source="google_alerts")
    c.url_key, c.platform, c.content_class = url_key, Platform.PUBLIC_WEB, ContentClass.PUBLIC_WEB
    c.verification, c.section, c.priority = Verification.VERIFIED, Section.MAIN, Priority.HIGH
    c.discovery_queries = set(queries or ())
    c.entity_matches = [Match("acme_company", "Acme Asset Management", position=0)]
    c.keyword_matches = [Match("", "fraud", "fraud", "Financial Crime", "HIGH", False, 3)]
    c.rules_applied, c.review_status = ["R4"], ReviewStatus.PENDING
    return c


@selftest("S05")
def test_schema_matches_architecture_and_is_idempotent():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "sub" / "t.sqlite"
        db_connect(db).close()
        conn = db_connect(db)            # second open must not fail or duplicate anything
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'"
                                             " AND name NOT LIKE 'sqlite_%'")}
        assert tables == {"runs", "source_status", "findings", "finding_matches", "finding_periods",
                          "discovery_hits", "email_log"}, tables
        assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1
        try:
            conn.execute("INSERT INTO finding_periods (finding_id, period_start, period_end,"
                         " report_section, first_run_id, last_run_id) VALUES (999,'x','y','MAIN','r','r')")
        except sqlite3.IntegrityError:
            pass
        else:
            raise AssertionError("foreign keys must be enforced")
        conn.close()


@selftest("S05")
def test_older_database_gets_date_precision_and_loses_first_seen_dates():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "old.sqlite"
        legacy = sqlite3.connect(db)                 # a findings table from before date_precision existed
        legacy.executescript(SCHEMA_SQL.replace(",\n    date_precision      TEXT", "", 1))
        assert "date_precision" not in {r[1] for r in legacy.execute("PRAGMA table_info(findings)")}
        rows = [("k-year", "2026-09-21T00:00:00+08:00", "FIRST_SEEN",
                 "first found by the run for the week of 21 Sep 2026 – 27 Sep 2026; the SC list gives only the year 2022"),
                ("k-none", "2026-09-21T00:00:00+08:00", "FIRST_SEEN",
                 "first found by the run for the week of 21 Sep 2026 – 27 Sep 2026; the source gives no date"),
                ("k-day", "2026-09-16T09:30:00+08:00", "VERIFIED", "publication date"),
                ("k-unk", None, "UNKNOWN", "no date stated by the source")]
        for key, when, status, basis in rows:
            legacy.execute("INSERT INTO findings (url_key, canonical_url, discovered_url, platform, content_class,"
                           " url_verification, content_date, date_status, date_basis, first_discovered_at,"
                           " last_seen_at) VALUES (?,'u','u','SC Malaysia','regulatory','VERIFIED',?,?,?,'t','t')",
                           (key, when, status, basis))
        legacy.commit()
        legacy.close()
        for _ in range(2):                           # migrates once; opening again changes nothing
            conn = db_connect(db)
            got = {r[0]: tuple(r[1:]) for r in conn.execute(
                "SELECT url_key, content_date, date_precision, date_status FROM findings")}
            conn.close()
            assert got == {"k-year": ("2022-01-01T00:00:00+08:00", "YEAR", "VERIFIED"),
                           "k-none": (None, None, "FIRST_SEEN"),
                           "k-day": ("2026-09-16T09:30:00+08:00", "DAY", "VERIFIED"),
                           "k-unk": (None, None, "UNKNOWN")}, got


@selftest("S05")
def test_rerun_keeps_one_finding_one_match_set_and_no_duplicate_hits():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        db_record_run_start(conn, ctx)
        c = _db_candidate(queries=set())               # no discovery query -> NULL in discovery_hits
        fid = db_upsert_candidate(conn, c, ctx)
        conn.execute("UPDATE findings SET first_discovered_at = '2000-01-01T00:00:00+08:00'")
        conn.commit()
        for _ in range(2):
            assert db_upsert_candidate(conn, c, ctx) == fid
        count = lambda sql: conn.execute(sql).fetchone()[0]  # noqa: E731
        assert count("SELECT COUNT(*) FROM finding_matches") == 2, "matches duplicated on re-run"
        assert count("SELECT COUNT(*) FROM discovery_hits") == 1, "NULL-query hit duplicated on re-run"
        assert count("SELECT first_discovered_at FROM findings") == "2000-01-01T00:00:00+08:00"
        c.keyword_matches = []                          # a later screening found no keyword
        db_upsert_candidate(conn, c, ctx)
        assert count("SELECT COUNT(*) FROM finding_matches WHERE keyword IS NOT NULL") == 0, \
            "stale keyword matches must be replaced"
        conn.close()


@selftest("S05")
def test_review_status_rules_and_second_period():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        db_record_run_start(conn, ctx)
        c = _db_candidate(queries={"q1"})
        fid = db_upsert_candidate(conn, c, ctx)
        c.review_status, c.rules_applied = ReviewStatus.MANUAL_REVIEW, ["R2"]
        db_upsert_candidate(conn, c, ctx)
        row = conn.execute("SELECT review_status, rules_applied FROM finding_periods").fetchone()
        assert tuple(row) == ("MANUAL REVIEW", "R2"), "program-set status must follow the latest rules"
        conn.execute("UPDATE finding_periods SET review_status = 'CONFIRMED ADVERSE'")
        conn.commit()
        c.review_status = ReviewStatus.PENDING
        db_upsert_candidate(conn, c, ctx)
        assert conn.execute("SELECT review_status FROM finding_periods").fetchone()[0] == "CONFIRMED ADVERSE"
        ctx2 = _test_ctx(db, "2026-09-29")              # the next week surfaces the same content
        db_record_run_start(conn, ctx2)
        assert db_upsert_candidate(conn, c, ctx2) == fid
        periods = [r[0] for r in conn.execute("SELECT period_start FROM finding_periods ORDER BY 1")]
        assert periods == ["2026-09-14", "2026-09-21"], periods
        assert conn.execute("SELECT COUNT(*) FROM findings").fetchone()[0] == 1
        conn.close()


@selftest("S05")
def test_report_query_runs_sources_and_email_log():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        db_record_run_start(conn, ctx)
        hi = _db_candidate("k-high", {"q1", "q2"})
        crit = _db_candidate("k-crit", {"q3"})
        crit.priority = Priority.CRITICAL
        db_upsert_candidate(conn, hi, ctx)
        db_upsert_candidate(conn, crit, ctx)
        rows = db_findings_for_period(conn, ctx.period)
        assert [r["url_key"] for r in rows] == ["k-crit", "k-high"], "must be sorted by priority"
        assert rows[1]["query_count"] == 2 and rows[1]["keywords"] == "fraud"
        assert rows[1]["entity_ids"] == "acme_company"
        assert db_findings_for_period(conn, _period("2026-09-29")) == []
        try:
            bad = _db_candidate("k-bad")
            bad.verification = None
            db_upsert_candidate(conn, bad, ctx)
        except ValueError as e:
            assert "verification" in str(e)
        else:
            raise AssertionError("an unprocessed candidate must be refused")
        db_record_source_status(conn, ctx.run_id, [{"source": "SC Malaysia", "status": SourceStatus.FAILED,
                                                    "error": "timeout"}])
        assert tuple(conn.execute("SELECT status, retrieved, error FROM source_status").fetchone()) \
            == ("FAILED", 0, "timeout")
        db_record_run_end(conn, ctx.run_id, 1, Path("reports/x.xlsx"))
        run = conn.execute("SELECT exit_code, report_path, finished_at, config_versions FROM runs").fetchone()
        assert run["exit_code"] == 1 and run["report_path"].endswith("x.xlsx") and run["finished_at"]
        assert json.loads(run["config_versions"])["keywords"] == DEFAULT_KEYWORDS["version"]
        for status in ("SKIPPED", "FAILED"):
            db_record_email(conn, ctx.run_id, ctx.period, status, 0, "test")
        assert not db_email_already_sent(conn, ctx.period), "only a SENT row counts as sent"
        conn.close()


# ---- S06 --------------------------------------------------------------------
@selftest("S06")
def test_google_redirect_decoding():
    target = "https://www.linkedin.com/posts/someone_topic-activity-7240000000000000000-AbCd"
    wrapped = ("https://www.google.com/url?rct=j&sa=t&url=" +
               "https%3A%2F%2Fwww.linkedin.com%2Fposts%2Fsomeone_topic-activity-7240000000000000000-AbCd"
               "&ct=ga&cd=CAEYACoTNz&usg=AOvVaw0")
    assert decode_google_redirect(wrapped) == target
    assert decode_google_redirect("https://example.com/x") == "https://example.com/x"


@selftest("S06")
def test_canonicalisation():
    assert canonicalize_url("https://m.facebook.com/SomePage/posts/123/?fbclid=XYZ&__tn__=R#frag") \
        == "https://facebook.com/SomePage/posts/123"
    assert canonicalize_url("https://www.Example.com/news/a?utm_source=x&b=2&a=1") \
        == "https://example.com/news/a?a=1&b=2"
    assert canonicalize_url("https://my.linkedin.com/in/someone/") == "https://linkedin.com/in/someone"


@selftest("S06")
def test_platform_keys_merge_url_shapes():
    a = canonicalize_url("https://www.linkedin.com/posts/someone_topic-activity-7240000000000000000-AbCd?trk=x")
    b = canonicalize_url("https://www.linkedin.com/feed/update/urn:li:activity:7240000000000000000/")
    assert platform_content_key(a) == platform_content_key(b) == "linkedin:activity:7240000000000000000"
    assert make_url_key(a) == make_url_key(b)
    assert platform_content_key(canonicalize_url("https://www.instagram.com/p/AbC123/")) \
        == platform_content_key(canonicalize_url("https://www.instagram.com/reel/AbC123/?igsh=zz")) \
        == "instagram:media:AbC123"
    assert platform_content_key(canonicalize_url(
        "https://www.facebook.com/permalink.php?story_fbid=111&id=222")) == "facebook:story:222:111"
    assert platform_content_key("https://example.com/news/a") is None
    assert len(make_url_key("https://example.com/news/a")) == 64


@selftest("S06")
def test_redirect_decoding_edge_cases():
    target = "https://news.example.com/a?id=5&page=2"
    enc = "https%3A%2F%2Fnews.example.com%2Fa%3Fid%3D5%26page%3D2"
    assert decode_google_redirect(f"https://www.google.com/url?q={enc}&sa=U") == target, "q= form"
    assert decode_google_redirect(f"https://www.google.com.my/url?url={enc}") == target, "country domain"
    for unchanged in ("https://www.google.com/alerts/remove?source=alertsmail&s=abc",   # footer link
                      "https://www.google.com/url?sa=t&usg=x",                           # no target
                      "https://www.google.com/url?url=javascript:alert(1)",              # not a web URL
                      "https://evil-google.com/url?url=https://x.example",               # not Google
                      "not a url", ""):
        assert decode_google_redirect(unchanged) == unchanged, unchanged


@selftest("S06")
def test_canonicalisation_edge_cases():
    cases = {
        "https://Example.com": "https://example.com",
        "https://example.com/": "https://example.com",
        "HTTPS://WWW.Example.COM/Path/To/Page": "https://example.com/Path/To/Page",   # path case kept
        "https://example.com:443/a": "https://example.com/a",
        "http://example.com:8080/a": "http://example.com:8080/a",
        "https://web.dev/articles/x": "https://web.dev/articles/x",                  # not reduced to 'dev'
        "https://uk.linkedin.com/company/x/": "https://linkedin.com/company/x",
        "https://example.com/a?UTM_Source=x&z=1&si=abc&ref=home": "https://example.com/a?z=1",
        "https://facebook.com/permalink.php?story_fbid=1&id=2&__cft__=x": "https://facebook.com/permalink.php?id=2&story_fbid=1",
        "https://user:pw@example.com/a": "https://example.com/a",                      # credentials dropped
        "  https://example.com/a#section  ": "https://example.com/a",
        "mailto:someone@example.com": "mailto:someone@example.com",
        "": "",
    }
    for raw, expected in cases.items():
        assert canonicalize_url(raw) == expected, f"{raw!r} -> {canonicalize_url(raw)!r}, expected {expected!r}"
    assert canonicalize_url(canonicalize_url("https://m.Example.com/a/?b=2&a=1")) == \
        canonicalize_url("https://m.Example.com/a/?b=2&a=1"), "canonicalisation must be idempotent"


@selftest("S06")
def test_platform_key_variants():
    key = lambda u: platform_content_key(canonicalize_url(u))  # noqa: E731
    assert key("https://www.facebook.com/story.php?story_fbid=111&id=222&comment_id=9") \
        == "facebook:story:222:111:comment:9"
    assert key("https://www.facebook.com/SomePage") is None, "a Page has no content key"
    assert key("https://www.instagram.com/someuser/reel/XyZ_9/") == "instagram:media:XyZ_9"
    assert key("https://www.instagram.com/someuser/") is None, "a profile has no content key"
    assert key("https://www.linkedin.com/posts/a_b-activity-111-x") != key("https://www.linkedin.com/posts/a_b-activity-222-x")
    assert key("https://www.linkedin.com/in/someone") is None


@selftest("S06")
def test_normalize_item():
    item = RawItem("google_alerts", "Google Alerts (search index)",
                   "https://www.google.com/url?rct=j&url=https://www.linkedin.com/feed/update/"
                   "urn:li:activity:7240000000000000000/&ct=ga",
                   title="  Post  about\n Acme ", snippet="line one\n\tline two",
                   discovery_query='site:linkedin.com "Acme Asset Management"')
    c = normalize_item(item)
    assert c.canonical_url == "https://linkedin.com/feed/update/urn:li:activity:7240000000000000000"
    assert c.raw.discovered_url.startswith("https://www.linkedin.com/"), "stored URL must be decoded"
    assert (c.raw.title, c.raw.snippet) == ("Post about Acme", "line one line two")
    assert c.discovery_queries == {'site:linkedin.com "Acme Asset Management"'}
    other = normalize_item(RawItem("google_alerts", "x",
                                   "https://www.linkedin.com/posts/someone_t-activity-7240000000000000000-AbCd"))
    assert c.url_key == other.url_key, "two URL shapes of one LinkedIn post must share a url_key"
    assert other.discovery_queries == set()
    a, b = (normalize_item(RawItem("website", "x", "", title=t)) for t in ("first", "second"))
    assert a.url_key and a.url_key != b.url_key, "URL-less items must not collapse into one"


# ---- S07 --------------------------------------------------------------------
@selftest("S07")
def test_in_run_dedup_unions_queries():
    a = _candidate("https://example.com/a", title="A", discovery_query="q1")
    b = _candidate("https://example.com/a", snippet="snippet from second", discovery_query="q2")
    for c, q in ((a, "q1"), (b, "q2")):
        c.url_key, c.discovery_queries = "same", {q}
    unique, merged = dedupe_in_run([a, b])
    assert len(unique) == 1 and merged == 1
    assert unique[0].discovery_queries == {"q1", "q2"}
    assert unique[0].raw.title == "A" and unique[0].raw.snippet == "snippet from second"


@selftest("S07")
def test_content_hash_normalises_whitespace_and_case():
    assert content_hash("Hello   World") == content_hash("hello world")
    assert content_hash("") is None and content_hash(None) is None


@selftest("S07")
def test_in_run_dedup_keeps_order_and_richest_fields():
    def cand(key, **raw):
        c = _candidate("https://example.com/" + key, **raw)
        c.url_key = key
        return c
    a1 = cand("a", title="Short", content_text="short text")
    b = cand("b", title="B")
    a2 = cand("a", title="A much longer title", content_text="a much longer body text",
              source_published_at=_myt(2026, 9, 15), extra={"sc_alert_list": True})
    unique, merged = dedupe_in_run([a1, b, a2])
    assert [c.url_key for c in unique] == ["a", "b"] and merged == 1
    assert unique[0] is a1, "the first-seen candidate is kept"
    assert a1.raw.title == "A much longer title" and a1.raw.content_text == "a much longer body text"
    assert a1.raw.source_published_at == _myt(2026, 9, 15), "missing facts are filled from the duplicate"
    assert a1.raw.extra.get("sc_alert_list") is True, "flags from a duplicate must not be lost"
    x, y = _candidate("u1"), _candidate("u2")          # url_key '' (never normalised): never merged
    assert dedupe_in_run([x, y]) == ([x, y], 0)


def _dup_case(url, text=None, title="", when=None, publisher=None, query=None):
    c = _candidate(url, title=title, content_text=text, publisher=publisher, discovery_query=query)
    c.discovery_queries = {query} if query else set()
    if when:
        c.content_date, c.date_status = when, DateStatus.VERIFIED
    return c


@selftest("S07")
def test_probable_duplicates_same_site_merge():
    wk = _myt(2026, 9, 17, 10)
    a = _dup_case("https://news.example.com/a", text="Body of the ARTICLE", query="q1")
    b = _dup_case("https://news.example.com/amp/a", text="body of the  article", when=wk, query="q2")
    c = _dup_case("https://blog.example.org/x", title="Acme: fined!", when=wk)
    d = _dup_case("https://blog.example.org/x?page=1", title="acme fined", when=_myt(2026, 9, 17, 22))
    e = _dup_case("https://blog.example.org/y", title="Acme fined", when=_myt(2026, 9, 18, 10))
    items = [a, b, c, d, e]
    mark_probable_duplicates(items)
    assert items == [a, c, e], [i.canonical_url for i in items]
    assert a.discovery_queries == {"q1", "q2"} and any("amp/a" in n for n in a.notes)
    assert a.content_date == wk, "a date found on the duplicate must be kept"
    assert any("same publisher, title and date" in n for n in c.notes)
    assert not any("Also published" in n for n in e.notes), "a different day is a different item"


@selftest("S07")
def test_syndicated_copies_stay_separate_with_notes():
    wk = _myt(2026, 9, 17, 10)
    a = _dup_case("https://thestar.example/a", text="Same syndicated body", publisher="The Star")
    b = _dup_case("https://edge.example/b", text="same  syndicated BODY", publisher="The Edge")
    c = _dup_case("https://site-one.example/p", title="Acme fined", when=wk)
    d = _dup_case("https://site-two.example/q", title="Acme fined", when=wk)
    items = [a, b, c, d]
    mark_probable_duplicates(items)
    assert items == [a, b, c, d], "copies on different sites must not be merged"
    assert "Also published by the edge: https://edge.example/b" in a.notes, a.notes
    assert any(n.startswith("Also published by the star") for n in b.notes), b.notes
    assert any("site-two.example" in n for n in c.notes) and any("site-one.example" in n for n in d.notes)


# ---- S08 --------------------------------------------------------------------
class _BrokenCollector:
    name = "sc_malaysia"

    def collect(self, ctx):
        raise ConnectionError("HTTP 503 after 3 retries")


@selftest("S08")
def test_failed_collector_is_isolated():
    ctx = _test_ctx(Path("unused.sqlite"))
    r = run_collector_safely(_BrokenCollector(), ctx)
    assert r.status == SourceStatus.FAILED and r.items == []
    assert r.error and "503" in r.error


@selftest("S08")
def test_social_rows_inherit_google_alerts_failure():
    results = [SourceResult("google_alerts", SourceStatus.FAILED, error="IMAP login failed"),
               SourceResult("website", SourceStatus.SUCCESS)]
    rows = {r["source"]: r for r in derive_source_health(results, [])}
    for name in ("Google Alerts", "Facebook", "Instagram", "LinkedIn", "Public website"):
        assert rows[name]["status"] == "FAILED", f"{name} must be FAILED, not SUCCESS with 0 results"
    assert rows["Company website"]["status"] == "SUCCESS"
    assert "SC Malaysia" in rows, "every report source needs a row"
    assert "News" not in rows, "there is no separate news source"


class _FakeResponse:
    def __init__(self, status_code=200, text="", headers=None, url=""):
        self.status_code, self.text, self.headers, self.url = status_code, text, headers or {}, url


class _FakeSession:
    """Scripted stand-in for requests.Session: {url: [response or exception, ...]}. No network."""

    def __init__(self, script: dict):
        self.script, self.calls = {k: list(v) for k, v in script.items()}, []

    def get(self, url, **kwargs):
        self.calls.append(url)
        queue = self.script.get(url) or [_FakeResponse(404)]
        step = queue.pop(0) if len(queue) > 1 else queue[0]
        if isinstance(step, BaseException):
            raise step
        step.url = url
        return step


def _fast_settings(retries=2, delay=0.0) -> dict:
    s = json.loads(json.dumps(DEFAULT_SETTINGS))
    s["http"].update(max_retries=retries, backoff_seconds=0, delay_between_requests_seconds=delay)
    return s


_ROBOTS = "https://site.example/robots.txt"


@selftest("S08")
def test_polite_get_retries_and_access_errors():
    log, url = logging.getLogger("ns-test"), "https://site.example/page"
    s = _FakeSession({_ROBOTS: [_FakeResponse(404)],
                      url: [_FakeResponse(503), TimeoutError("slow"), _FakeResponse(200, "ok")]})
    assert polite_get(s, url, _fast_settings(), log).text == "ok"
    assert s.calls.count(url) == 3, "503 and timeouts must be retried"
    for code, exc in ((403, PermissionError), (401, PermissionError), (404, ConnectionError)):
        s = _FakeSession({_ROBOTS: [_FakeResponse(404)], url: [_FakeResponse(code)]})
        try:
            polite_get(s, url, _fast_settings(), log)
        except exc as e:
            assert str(code) in str(e)
        else:
            raise AssertionError(f"HTTP {code} must raise {exc.__name__}")
        assert s.calls.count(url) == 1, f"HTTP {code} must not be retried"
    s = _FakeSession({_ROBOTS: [_FakeResponse(404)], url: [_FakeResponse(503)]})
    try:
        polite_get(s, url, _fast_settings(retries=2), log)
    except ConnectionError as e:
        assert "HTTP 503 after 2 retries" in str(e), e
    else:
        raise AssertionError("a source that keeps failing must raise")
    assert s.calls.count(url) == 3


@selftest("S08")
def test_polite_get_robots_social_and_redirects():
    log = logging.getLogger("ns-test")
    robots = "User-agent: *\nDisallow: /private\n"
    s = _FakeSession({_ROBOTS: [_FakeResponse(200, robots)],
                      "https://site.example/a": [_FakeResponse(200, "a")],
                      "https://site.example/b": [_FakeResponse(200, "b")]})
    polite_get(s, "https://site.example/a", _fast_settings(), log)
    polite_get(s, "https://site.example/b", _fast_settings(), log)
    assert s.calls.count(_ROBOTS) == 1, "robots.txt must be cached per site"
    try:
        polite_get(s, "https://site.example/private/x", _fast_settings(), log)
    except PermissionError:
        pass
    else:
        raise AssertionError("robots.txt Disallow must be respected")
    assert "https://site.example/private/x" not in s.calls
    s = _FakeSession({_ROBOTS: [_FakeResponse(503)], "https://site.example/a": [_FakeResponse(200)]})
    try:
        polite_get(s, "https://site.example/a", _fast_settings(), log)
    except PermissionError:
        assert "https://site.example/a" not in s.calls
    else:
        raise AssertionError("an unreachable robots.txt means 'do not fetch'")
    for social in ("https://www.facebook.com/x", "https://m.facebook.com/x", "https://fb.watch/x",
                   "https://instagram.com/p/x", "https://my.linkedin.com/in/x"):
        s = _FakeSession({})
        try:
            polite_get(s, social, _fast_settings(), log, check_robots=False)
        except PermissionError:
            assert s.calls == [], f"{social} was requested"
        else:
            raise AssertionError(f"{social} must never be fetched")
    s = _FakeSession({_ROBOTS: [_FakeResponse(404)],
                      "https://site.example/r": [_FakeResponse(302, headers={"Location": "/final"})],
                      "https://site.example/final": [_FakeResponse(200, "done")]})
    resp = polite_get(s, "https://site.example/r", _fast_settings(), log)
    assert resp.text == "done" and resp.url == "https://site.example/final"
    s = _FakeSession({_ROBOTS: [_FakeResponse(404)], "https://site.example/r": [
        _FakeResponse(301, headers={"Location": "https://www.facebook.com/SomePage"})]})
    try:
        polite_get(s, "https://site.example/r", _fast_settings(), log)
    except PermissionError:
        assert not any("facebook" in c for c in s.calls), "redirect target on Facebook was fetched"
    else:
        raise AssertionError("a redirect to Facebook must not be followed")


@selftest("S08")
def test_polite_get_waits_between_requests():
    from time import monotonic
    s = _FakeSession({"https://site.example/a": [_FakeResponse(200)]})
    started = monotonic()
    for _ in range(3):
        polite_get(s, "https://site.example/a", _fast_settings(delay=0.05), logging.getLogger("ns-test"),
                   check_robots=False)
    assert monotonic() - started >= 0.09, "requests must be spaced by the configured delay"


@selftest("S08")
def test_http_session_user_agent():
    s = build_http_session(DEFAULT_SETTINGS)
    assert s.headers["User-Agent"] == DEFAULT_SETTINGS["http"]["user_agent"]
    s.close()


@selftest("S08")
def test_run_collector_safely_redacts_and_validates():
    old = os.environ.get("NS_IMAP_PASSWORD")
    os.environ["NS_IMAP_PASSWORD"] = "imap-S3cret-value"
    try:
        class Leaky:
            name = "google_alerts"

            def collect(self, ctx):
                raise RuntimeError("LOGIN failed for user with imap-S3cret-value")

        class Wrong:
            name = "website"

            def collect(self, ctx):
                return None

        class Fine:
            name = "sc_malaysia"

            def collect(self, ctx):
                return SourceResult("sc_malaysia", SourceStatus.SUCCESS, items=[RawItem("sc_malaysia", "x", "u")])

        ctx = _test_ctx(Path("unused.sqlite"))
        r = run_collector_safely(Leaky(), ctx)
        assert r.status == SourceStatus.FAILED and "imap-S3cret-value" not in (r.error or "")
        assert run_collector_safely(Wrong(), ctx).status == SourceStatus.FAILED  # type: ignore[arg-type]  (deliberately broken)
        ok = run_collector_safely(Fine(), ctx)
        assert ok.status == SourceStatus.SUCCESS and len(ok.items) == 1
    finally:
        if old is None:
            os.environ.pop("NS_IMAP_PASSWORD", None)
        else:
            os.environ["NS_IMAP_PASSWORD"] = old


@selftest("S08")
def test_source_health_counts_and_skipped():
    def cand(url, platform, source="google_alerts", entity=True, priority=None, verification=Verification.VERIFIED):
        c = _candidate(url, source=source)
        c.platform, c.verification, c.priority = platform, verification, priority
        c.entity_matches = [Match("acme_company", "x")] if entity else []
        return c
    cands = [cand("https://facebook.com/a", Platform.FACEBOOK, priority=Priority.HIGH),
             cand("https://facebook.com/b", Platform.FACEBOOK, entity=False),
             cand("https://news.example/c", Platform.PUBLIC_WEB, verification=Verification.REJECTED),
             cand("https://acme.example.com/n", Platform.COMPANY_WEBSITE, source="website")]
    items = [RawItem("google_alerts", "x", "u")] * 4           # 4 alert results -> 3 after dedup
    results = [SourceResult("google_alerts", SourceStatus.SUCCESS, items=items),
               SourceResult("website", SourceStatus.SUCCESS, items=[RawItem("website", "x", "u")],
                            warnings=["listing page 2 unreachable"])]
    rows = {r["source"]: r for r in derive_source_health(results, cands)}
    assert [r["source"] for r in derive_source_health(results, cands)] == [
        "Google Alerts", "Facebook", "Instagram", "LinkedIn", "Public website", "Company website", "SC Malaysia"]
    assert rows["Google Alerts"]["retrieved"] == 4 and rows["Google Alerts"]["duplicates"] == 1
    fb = rows["Facebook"]
    assert (fb["status"], fb["retrieved"], fb["flagged"], fb["filtered_out"]) == ("SUCCESS", 2, 1, 1)
    assert rows["Instagram"]["status"] == "SUCCESS" and rows["Instagram"]["retrieved"] == 0
    assert rows["Public website"]["rejected"] == 1
    assert rows["Company website"]["status"] == "SUCCESS" and "unreachable" in rows["Company website"]["error"]
    assert rows["SC Malaysia"]["status"] == "SKIPPED", "a source that did not run must say so"
    none_ran = {r["source"]: r["status"] for r in derive_source_health([], [])}
    assert set(none_ran.values()) == {"SKIPPED"}, none_ran
    failed = {r["source"]: r for r in derive_source_health(
        [SourceResult("google_alerts", SourceStatus.FAILED, error="IMAP login failed")], [])}
    assert failed["Google Alerts"]["error"] == "IMAP login failed"
    for name in ("Facebook", "Instagram", "LinkedIn", "Public website"):
        assert failed[name]["error"] == "Google Alerts unavailable", failed[name]


# ---- S09 --------------------------------------------------------------------
@selftest("S09")
def test_alert_queries():
    qs = generate_alert_queries(DEFAULT_ENTITIES, DEFAULT_KEYWORDS)
    assert 'site:facebook.com "Acme Asset Management"' in qs
    assert 'site:linkedin.com "Tan Ah Kow"' in qs
    assert '"Acme Asset Management"' in qs, "whole-web Tier 1 alert missing"
    assert len(qs) == len(set(qs)), "duplicate queries"
    assert len(qs) <= 60, f"too many alerts: {len(qs)}"
    for q in qs:
        assert any(n in q for n in ("Acme Asset Management", "Tan Ah Kow", "艾克米投资管理", "艾克米投資管理")), \
            f"bare/ambiguous query: {q}"
        assert "陈亚九" not in q and "陳亞九" not in q, "D-29: the common Chinese name is never an alert"
        assert q.count(" OR ") <= 7, f"category chunk too long: {q}"
    assert any('"money laundering"' in q for q in qs), "multi-word keywords must be quoted"


@selftest("S09")
def test_alert_query_plan_shape_and_coverage():
    qs = generate_alert_queries(DEFAULT_ENTITIES, DEFAULT_KEYWORDS)
    tier1 = [q for q in qs if "(" not in q]
    assert qs[:len(tier1)] == tier1, "Tier 1 must come first"
    assert len(tier1) == 20, f"blueprint §2.3 Tier 1 = 20 alerts (5 names x 4 scopes), got {len(tier1)}"
    assert 38 <= len(qs) <= 46, f"blueprint §2.3 expects roughly 40 alerts, got {len(qs)}"
    assert 'site:facebook.com "艾克米投资管理"' in tier1 and '"艾克米投資管理"' in tier1, "D-28: Chinese company alerts"
    assert not any(q.startswith(('"艾克米投资管理" (', '"艾克米投資管理" (')) for q in qs), \
        "Tier 2 pairs English keywords with the shortest ENGLISH name"
    assert '"Acme Asset Management Sdn Bhd"' in tier1
    def terms_of(q: str) -> set[str]:
        return {t.strip('"') for t in q[q.index("(") + 1:-1].split(" OR ")}
    sanctions = {k["term"] for c in DEFAULT_KEYWORDS["categories"] if c["id"] == "sanctions_regulatory"
                 for k in c["keywords"]}
    chunks = [terms_of(q) for q in qs if q.startswith('"Tan Ah Kow" (') and terms_of(q) <= sanctions]
    assert [len(c) for c in chunks] == [8, 8, 7], "the 23 sanctions terms must be split evenly into 3 alerts"
    for name in ("Acme Asset Management", "Tan Ah Kow"):
        tier2 = " ".join(q for q in qs if q.startswith(f'"{name}" ('))
        for cat in DEFAULT_KEYWORDS["categories"]:
            for k in cat["keywords"]:
                t = k["term"]
                assert (f'"{t}"' if " " in t else t) in tier2, f"{t!r} missing from {name} alerts"
    assert not any(q.startswith('"Acme Asset Management Sdn Bhd" (') for q in qs), \
        "Tier 2 uses only the shortest name"


@selftest("S09")
def test_alert_queries_respect_approval_ambiguity_and_limits():
    ents = json.loads(json.dumps(_AMBIGUOUS_TEST_ENTITIES))
    ents["entities"][0]["alert_names"].append("Acme")         # misconfiguration: bare ambiguous term
    ents["entities"].append({"id": "x", "display_name": "Not Approved", "names": ["Not Approved"],
                             "alert_names": ["Not Approved"], "approved": False})
    ents["entities"].append({"id": "y", "display_name": "No Alert Names", "names": ["No Alert Names"],
                             "approved": True})
    qs = generate_alert_queries(ents, DEFAULT_KEYWORDS, max_terms=5)
    assert not any(q in ('"Acme"', 'site:facebook.com "Acme"') or q.startswith('"Acme" (')
                   for q in qs), "a bare 'Acme' query was generated"
    assert not any("Not Approved" in q for q in qs), "unapproved names must not get alerts"
    assert 'site:linkedin.com "No Alert Names"' in qs, "names are used when alert_names is absent"
    assert all(q.count(" OR ") <= 4 for q in qs), "max_terms must be respected"
    fin = [q for q in qs if q.startswith('"Tan Ah Kow" (') and ("fraud" in q or '"money laundering"' in q
                                                                  or "embezzlement" in q)]
    assert [q.count(" OR ") + 1 for q in fin] == [3, 3], f"6 terms must split 3+3, not 5+1: {fin}"


# ---- S10 --------------------------------------------------------------------
_SAMPLE_ALERT_EMAIL = (
    b"Message-ID: <alert-1@google.com>\r\n"
    b"Date: Wed, 16 Sep 2026 09:15:00 +0000\r\n"
    b"From: Google Alerts <googlealerts-noreply@google.com>\r\n"
    b"Subject: Google Alert - site:linkedin.com \"Acme Asset Management\"\r\n"
    b"MIME-Version: 1.0\r\nContent-Type: text/html; charset=UTF-8\r\n\r\n"
    b"<html><body><table><tr><td>"
    b"<a href=\"https://www.google.com/url?rct=j&sa=t&url=https://www.linkedin.com/posts/"
    b"someone_topic-activity-7240000000000000000-AbCd&ct=ga&usg=x\">Example post title about "
    b"<b>Acme Asset Management</b></a>"
    b"<div>LinkedIn</div><div>Illustrative snippet text mentioning Acme Asset Management.</div>"
    b"</td></tr><tr><td>"
    b"<a href=\"https://www.google.com/alerts/remove?source=alertsmail&s=abc\">Delete this alert</a> "
    b"<a href=\"https://www.google.com/alerts?source=alertsmail\">Create another alert</a>"
    b"</td></tr></table></body></html>\r\n"
)


@selftest("S10")
def test_parse_alert_email_synthetic_sample():
    items = parse_alert_email(_SAMPLE_ALERT_EMAIL)
    assert len(items) == 1, f"footer links must be ignored, got {len(items)} items"
    it = items[0]
    assert it.discovered_url.startswith("https://www.linkedin.com/posts/"), it.discovered_url
    assert "google.com" not in it.discovered_url
    assert it.discovery_query == 'site:linkedin.com "Acme Asset Management"'
    assert "Acme Asset Management" in it.title
    assert it.source == "google_alerts" and it.source_message_id == "<alert-1@google.com>"
    assert it.alert_received_at is not None and it.alert_received_at.tzinfo is not None
    assert it.source_published_at is None, "alerts do not state a publication date"


@selftest("S10")
def test_social_results_are_never_fetched():
    class NoNetwork:
        def get(self, *args, **kwargs):
            raise AssertionError("a social-media URL was fetched")

        request = get

    for url in ("https://www.linkedin.com/posts/someone_x-activity-7240000000000000000-AbCd",
                "https://www.facebook.com/SomePage/posts/123",
                "https://www.instagram.com/p/AbC123/"):
        item = RawItem("google_alerts", "Google Alerts (search index)", url, title="t", snippet="s")
        out = enrich_public_web_item(item, NoNetwork(), DEFAULT_SETTINGS, logging.getLogger("ns-test"))
        assert out.content_text is None and out.source_published_at is None


# Synthetic stand-in modelled on the layout of real Google Alerts emails (schema.org itemprop
# markup, NEWS/WEB headings, share + "Flag as irrelevant" links, footer). Replace or complement
# with a real, anonymised alert email once the alerts mailbox has received some.
_REALISTIC_ALERT_HTML = """<html><body><table>
<tr><td><span style="font-size:12px">NEWS</span></td></tr>
<tr><td itemscope itemtype="http://schema.org/Article">
 <a href="https://www.google.com/url?rct=j&amp;sa=t&amp;url=https://www.thestar.example/business/2026/09/16/x&amp;ct=ga&amp;cd=CAIy&amp;usg=AOv"
    itemprop="url"><span itemprop="name"><b>Tan Ah Kow</b> speaks at &quot;fund&quot; forum</span></a>
 <div itemprop="publisher" itemscope itemtype="http://schema.org/Organization"><span itemprop="name">The Star Online</span></div>
 <div itemprop="description">Illustrative snippet: <b>Tan Ah Kow</b> of Acme Asset<br>Management said&nbsp;the fund &#8230;</div>
 <a href="https://www.google.com/alerts/share?hl=en&amp;ss=fb&amp;rt=x">Facebook</a>
 <a href="https://www.google.com/alerts/share?hl=en&amp;ss=tw&amp;rt=x">Twitter</a>
 <a href="https://www.google.com/alerts/feedback?ffu=x&amp;source=alertsmail">Flag as irrelevant</a>
</td></tr>
<tr><td><span>WEB</span></td></tr>
<tr><td itemscope itemtype="http://schema.org/Article">
 <a href="https://www.google.com/url?rct=j&amp;sa=t&amp;url=https://www.facebook.com/permalink.php%3Fstory_fbid%3D111%26id%3D222&amp;ct=ga&amp;usg=x"><img src="https://example.invalid/thumb.png"></a>
 <a href="https://www.google.com/url?rct=j&amp;sa=t&amp;url=https://www.facebook.com/permalink.php%3Fstory_fbid%3D111%26id%3D222&amp;ct=ga&amp;usg=x"
    itemprop="url"><span itemprop="name">Post by an illustrative page</span></a>
 <div itemprop="publisher"><span itemprop="name">Facebook</span></div>
 <div itemprop="description">新加坡 illustrative text mentioning Tan Ah Kow</div>
 <a href="https://www.google.com/alerts/feedback?ffu=y">Flag as irrelevant</a>
</td></tr>
<tr><td><a href="https://www.google.com/alerts?source=alertsmail">See more results</a> |
 <a href="https://www.google.com/alerts/edit?source=alertsmail&amp;s=abc">Edit this alert</a>
 You have received this email because you have subscribed to Google Alerts.
 <a href="https://www.google.com/alerts/remove?source=alertsmail&amp;s=abc">Unsubscribe</a></td></tr>
</table></body></html>"""


def _alert_email_bytes(html_text: str | None = _REALISTIC_ALERT_HTML, subject: str = 'Google Alert – "Tan Ah Kow"',
                       date_header: str = "Wed, 16 Sep 2026 09:15:00 +0000", msg_id: str = "<alert-2@google.com>") -> bytes:
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["Date"], msg["Message-ID"] = (
        subject, "Google Alerts <googlealerts-noreply@google.com>", date_header, msg_id)
    msg.set_content("Plain-text version of the alert.")
    if html_text is not None:
        msg.add_alternative(html_text, subtype="html", cte="quoted-printable")
    return bytes(msg)


@selftest("S10")
def test_parse_realistic_alert_email():
    items = parse_alert_email(_alert_email_bytes())
    assert len(items) == 2, f"expected 2 results (footer/share/flag links and a repeated image link ignored), got {len(items)}"
    news, fb = items
    assert news.discovered_url == "https://www.thestar.example/business/2026/09/16/x"
    assert news.title == 'Tan Ah Kow speaks at "fund" forum', news.title
    assert news.publisher == "The Star Online"
    assert "of Acme Asset Management said the fund" in news.snippet, "a <br> must not cut the snippet: " + news.snippet
    assert not any(w in news.snippet for w in ("Flag", "Twitter", "WEB")), news.snippet
    assert fb.discovered_url == "https://www.facebook.com/permalink.php?story_fbid=111&id=222"
    assert fb.title == "Post by an illustrative page" and fb.publisher == "Facebook"
    assert "新加坡" in fb.snippet, "non-ASCII text must survive the MIME decoding"
    assert {i.discovery_query for i in items} == {'"Tan Ah Kow"'}, "query from an en-dash subject"
    assert news.alert_received_at == _utc(2026, 9, 16, 9, 15) and news.source_message_id == "<alert-2@google.com>"
    assert all(i.source_published_at is None and i.source == "google_alerts" for i in items)
    received = _myt(2026, 9, 16, 17, 20)
    assert parse_alert_email(_alert_email_bytes(), received)[0].alert_received_at == received
    minus_zero = parse_alert_email(_alert_email_bytes(date_header="Wed, 16 Sep 2026 09:15:00 -0000"))
    assert minus_zero[0].alert_received_at == _utc(2026, 9, 16, 9, 15), "a '-0000' Date header is UTC, and aware"
    try:
        parse_alert_email(_alert_email_bytes(html_text=None))
    except ValueError:
        pass
    else:
        raise AssertionError("an alert email without HTML must be reported as unparseable")


@selftest("S10")
def test_parse_plain_layout_alert_email():
    (it,) = parse_alert_email(_SAMPLE_ALERT_EMAIL)
    assert it.publisher == "LinkedIn", "a short first line after the link is the publisher"
    assert it.snippet == "Illustrative snippet text mentioning Acme Asset Management."
    html_text = ('<div><a href="https://www.google.com/url?url=https://a.example/1">First</a></div>'
                 '<div>Site A</div><div>Snippet one.</div>'
                 '<a href="https://www.google.com/alerts/feedback?x=1">Flag as irrelevant</a>'
                 '<div>WEB</div>'
                 '<div><a href="https://www.google.com/url?url=https://b.example/2">Second</a></div>'
                 '<div>Only one line of snippet</div>')
    first, second = parse_alert_email(_alert_email_bytes(html_text))
    assert (first.publisher, first.snippet) == ("Site A", "Snippet one."), "text after a non-result link must be ignored"
    assert (second.publisher, second.snippet) == (None, "Only one line of snippet")
    (desc_only,) = parse_alert_email(_alert_email_bytes(
        '<a href="https://www.google.com/url?url=https://c.example/3">T</a>'
        '<div itemprop="description">Short line one<br>line two</div>'))
    assert desc_only.publisher is None and desc_only.snippet == "Short line one line two", \
        "a description must never be mistaken for the publisher"
    (nested,) = parse_alert_email(_alert_email_bytes(
        '<a href="https://www.google.com/url?url=https://d.example/4">T</a>'
        '<div itemprop="description"><p>Line one.</p><p>Line two.</p></div>'))
    assert (nested.publisher, nested.snippet) == (None, "Line one. Line two."), \
        "itemprop applies to everything inside the element"


@selftest("S10")
def test_parse_iso_datetime_forms():
    assert parse_iso_datetime("2026-09-18T02:15:00Z") == (_utc(2026, 9, 18, 2, 15), False)
    assert parse_iso_datetime("2026-09-18T10:15:00.123456789+0800")[0] == _myt(2026, 9, 18, 10, 15) \
        + timedelta(microseconds=123456)
    value, assumed = parse_iso_datetime("2026-09-18")
    assert assumed and value == _myt(2026, 9, 18), "a date without a zone is read as Malaysia time, flagged"
    assert parse_iso_datetime("18 Sept") == (None, False) and parse_iso_datetime(None) == (None, False)


@selftest("S10")
def test_parse_alert_feed():
    feed = """<?xml version="1.0" encoding="utf-8"?>
<feed xmlns="http://www.w3.org/2005/Atom"><title>Google Alert - "Acme Asset Management"</title>
<entry><id>tag:google.com,2013:googlealerts/feed:1</id>
 <title type="html">&lt;b&gt;Acme Asset Management&lt;/b&gt; fund update</title>
 <link href="https://www.google.com/url?rct=j&amp;sa=t&amp;url=https://news.example/a&amp;ct=ga&amp;usg=x"/>
 <published>2026-09-16T01:00:00Z</published>
 <content type="html">Illustrative &lt;b&gt;snippet&lt;/b&gt; text</content></entry>
</feed>"""
    (item,) = parse_alert_feed(feed)
    assert item.discovered_url == "https://news.example/a" and item.title == "Acme Asset Management fund update"
    assert item.snippet == "Illustrative snippet text"
    assert item.discovery_query == '"Acme Asset Management"'
    assert item.alert_received_at == _utc(2026, 9, 16, 1, 0) and item.source_published_at is None, \
        "the feed's published time is the alert time, not the content date"


_ARTICLE_HTML = """<!doctype html><html><head><title>t</title>
<meta property="article:published_time" content="2026-09-17T11:00:00+08:00">
<script>var hidden = "script text";</script><style>.x{}</style>
<body><h1>Acme Asset Management fined</h1><p>Body text about the asset freeze.</p>
<script>var tracker = "body script text";</script><noscript>enable js</noscript><time datetime="2020-01-01">old</time></body></html>"""


@selftest("S10")
def test_enrich_public_web_item():
    log = logging.getLogger("ns-test")
    settings = _fast_settings()
    page = "https://news.example/article"
    s = _FakeSession({"https://news.example/robots.txt": [_FakeResponse(404)],
                      page: [_FakeResponse(200, _ARTICLE_HTML, {"Content-Type": "text/html; charset=utf-8"})]})
    item = enrich_public_web_item(RawItem("google_alerts", "x", page), s, settings, log)
    assert item.source_published_at == _myt(2026, 9, 17, 11, 0), "article:published_time comes first"
    text = item.content_text or ""
    assert "Body text about the asset freeze." in text and "Acme Asset Management fined" in text
    assert not any(w in text for w in ("script text", "tracker", "enable js", ".x{}")), text
    assert item.final_url == page
    json_ld = ('<html><head><script type="application/ld+json">{"@graph": [{"@type": "WebPage"}, '
               '{"@type": "NewsArticle", "datePublished": "2026-09-15T08:00:00+08:00"}]}</script>'
               '<body><p>Text</p><time datetime="2020-01-01">x</time></body></html>')    # <head> never closed
    s = _FakeSession({"https://news.example/robots.txt": [_FakeResponse(404)],
                      page: [_FakeResponse(200, json_ld, {"Content-Type": "text/html"})]})
    item = enrich_public_web_item(RawItem("google_alerts", "x", page), s, settings, log)
    assert item.source_published_at == _myt(2026, 9, 15, 8, 0) and item.content_text == "Text x"
    for resp in (_FakeResponse(404), _FakeResponse(200, "%PDF-1.7", {"Content-Type": "application/pdf"})):
        s = _FakeSession({"https://news.example/robots.txt": [_FakeResponse(404)], page: [resp]})
        item = enrich_public_web_item(RawItem("google_alerts", "x", page, snippet="kept"), s, settings, log)
        assert item.content_text is None and item.snippet == "kept", "a failed or non-HTML page leaves the item as it was"
    off = _fast_settings()
    off["sources"]["google_alerts"]["fetch_public_web_pages"] = False
    s = _FakeSession({})
    enrich_public_web_item(RawItem("google_alerts", "x", page), s, off, log)
    assert s.calls == [], "no page may be fetched when fetch_public_web_pages is false"


class _FakeImap:
    """Read-only fake IMAP server. It has no store/copy/expunge/delete methods, so any attempt
    to modify the mailbox fails the test."""

    def __init__(self, emails: dict[str, tuple[str, bytes]], select_ok=True):
        self.emails, self.select_ok, self.log = emails, select_ok, []

    def __call__(self, host, port, timeout=None):
        self.log.append(("connect", host, port))
        return self

    def login(self, user, password):
        self.log.append(("login", user))
        return "OK", [b"logged in"]

    def select(self, mailbox, readonly=False):
        self.log.append(("select", mailbox, readonly))
        return ("OK", [str(len(self.emails)).encode()]) if self.select_ok else ("NO", [b"no such mailbox"])

    def search(self, charset, *criteria):
        self.log.append(("search",) + criteria)
        return "OK", [" ".join(self.emails).encode()]

    def fetch(self, msg_id, spec):
        self.log.append(("fetch", msg_id, spec))
        internal, raw = self.emails[msg_id]
        return "OK", [(f'{msg_id} (INTERNALDATE "{internal}" BODY[] {{{len(raw)}}}'.encode(), raw), b")"]

    def logout(self):
        self.log.append(("logout",))
        return "BYE", [b""]


def _alerts_ctx() -> RunContext:
    ctx = _test_ctx(Path("unused.sqlite"))
    ctx.config.settings["http"].update(max_retries=0, backoff_seconds=0, delay_between_requests_seconds=0)
    return ctx


@selftest("S10")
def test_google_alerts_imap_collection():
    old = os.environ.get("NS_IMAP_PASSWORD")
    os.environ["NS_IMAP_PASSWORD"] = "test-app-password"
    try:
        emails = {
            "1": ("16-Sep-2026 17:15:00 +0800", _alert_email_bytes()),
            "2": ("13-Sep-2026 23:50:00 +0800", _alert_email_bytes(msg_id="<old@google.com>")),  # before the period
            "3": ("17-Sep-2026 08:00:00 +0800", _alert_email_bytes(html_text=None, msg_id="<bad@google.com>")),
            "4": ("18-Sep-2026 08:00:00 +0800", _alert_email_bytes(msg_id="<again@google.com>")),  # same results
        }
        imap = _FakeImap(emails)
        page = "https://www.thestar.example/business/2026/09/16/x"
        session = _FakeSession({"https://www.thestar.example/robots.txt": [_FakeResponse(404)],
                                page: [_FakeResponse(200, _ARTICLE_HTML, {"Content-Type": "text/html"})]})
        ctx = _alerts_ctx()
        r = run_collector_safely(GoogleAlertsCollector(imap_factory=imap, session=session), ctx)
        assert r.status == SourceStatus.SUCCESS, r.error
        assert ("select", '"NS-Alerts"', True) in imap.log, "the label must be opened read-only"
        search = next(e for e in imap.log if e[0] == "search")
        assert search[1:] == ("FROM", '"googlealerts-noreply@google.com"', "SINCE", "13-Sep-2026"), search
        assert all("BODY.PEEK[]" in e[2] for e in imap.log if e[0] == "fetch"), "fetch must not mark mail as read"
        assert imap.log[-1] == ("logout",)
        assert len(r.items) == 4, "the email received before the period starts must be left out"
        assert r.stats["emails_read"] == 4 and r.stats["results"] == 4 and r.stats["pages_fetched"] == 1
        assert any("skipped" in w for w in r.warnings), "an unparseable email is a warning, not a failure"
        assert r.items[0].alert_received_at == _myt(2026, 9, 16, 17, 15), "INTERNALDATE is the received time"
        news = [i for i in r.items if "thestar" in i.discovered_url]
        assert len(news) == 2 and all(i.content_text for i in news), "both copies get the page text"
        assert session.calls.count(page) == 1, "a page must be fetched once per run"
        assert all(i.content_text is None for i in r.items if "facebook" in i.discovered_url)
        assert not any("facebook" in c for c in session.calls), "a Facebook result page was fetched"
        imap_missing = _FakeImap(emails, select_ok=False)
        r = run_collector_safely(GoogleAlertsCollector(imap_factory=imap_missing, session=session), ctx)
        assert r.status == SourceStatus.FAILED and "NS-Alerts" in (r.error or "")
        assert imap_missing.log[-1] == ("logout",), "the connection must be closed after a failure"
        os.environ.pop("NS_IMAP_PASSWORD")
        r = run_collector_safely(GoogleAlertsCollector(imap_factory=imap, session=session), ctx)
        assert r.status == SourceStatus.FAILED and "NS_IMAP_PASSWORD" in (r.error or "")
    finally:
        if old is None:
            os.environ.pop("NS_IMAP_PASSWORD", None)
        else:
            os.environ["NS_IMAP_PASSWORD"] = old


@selftest("S10")
def test_google_alerts_rss_collection():
    feed_ok = ('<feed xmlns="http://www.w3.org/2005/Atom"><title>Google Alert - "Tan Ah Kow"</title>'
               '<entry><id>e1</id><title>t</title><link href="https://www.google.com/url?url=https://www.linkedin.com/in/x"/>'
               '<published>2026-09-15T01:00:00Z</published></entry></feed>')
    f1, f2 = "https://www.google.com/alerts/feeds/1/2", "https://www.google.com/alerts/feeds/1/3"
    ctx = _alerts_ctx()
    ga = ctx.config.settings["sources"]["google_alerts"]
    ga.update(mode="rss", rss_feeds=[f1, f2])
    session = _FakeSession({f1: [_FakeResponse(200, feed_ok)], f2: [_FakeResponse(503)]})
    r = run_collector_safely(GoogleAlertsCollector(session=session), ctx)
    assert r.status == SourceStatus.SUCCESS and len(r.items) == 1 and r.warnings, "one feed down = warning"
    assert r.items[0].discovered_url == "https://www.linkedin.com/in/x"
    session = _FakeSession({f1: [_FakeResponse(500)], f2: [_FakeResponse(503)]})
    assert run_collector_safely(GoogleAlertsCollector(session=session), ctx).status == SourceStatus.FAILED
    ga["rss_feeds"] = []
    r = run_collector_safely(GoogleAlertsCollector(session=_FakeSession({})), ctx)
    assert r.status == SourceStatus.FAILED and "rss_feeds" in (r.error or "")


# ---- S11 --------------------------------------------------------------------
# Synthetic feed in the exact shape of a real WordPress RSS feed: channel header and three items
# (an English notice with a PDF link, a news link, and a Chinese title with a percent-encoded slug).
_COMPANY_FEED_FIXTURE = """<?xml version="1.0" encoding="UTF-8"?><rss version="2.0"
	xmlns:content="http://purl.org/rss/1.0/modules/content/"
	xmlns:wfw="http://wellformedweb.org/CommentAPI/"
	xmlns:dc="http://purl.org/dc/elements/1.1/"
	xmlns:atom="http://www.w3.org/2005/Atom"
	xmlns:sy="http://purl.org/rss/1.0/modules/syndication/"
	xmlns:slash="http://purl.org/rss/1.0/modules/slash/"
	>
<channel>
	<title>Acme Asset Management Malaysia</title>
	<atom:link href="https://acme.example.com/feed/" rel="self" type="application/rss+xml" />
	<link>https://acme.example.com</link>
	<description></description>
	<lastBuildDate>Thu, 12 Mar 2026 08:22:16 +0000</lastBuildDate>
	<language>en-US</language>
	<item>
		<title>Notice: Example Regulatory Update (&#8220;ERU&#8221;) for Financial Services</title>
		<link>https://acme.example.com/notice/</link>
		<dc:creator><![CDATA[admin]]></dc:creator>
		<pubDate>Thu, 12 Mar 2026 08:21:13 +0000</pubDate>
				<category><![CDATA[News & Events]]></category>
		<guid isPermaLink="false">https://acme.example.com/?p=1003</guid>
					<description><![CDATA[<p class="date">September 2025</p>
<p><a href="https://acme.example.com/wp-content/uploads/Notice.pdf" target="_blank" rel="noopener">Example Regulatory Update ("ERU") for Financial Services</a></p>]]></description>
										<content:encoded><![CDATA[<p class="date">September 2025</p>
<p><a href="https://acme.example.com/wp-content/uploads/Notice.pdf" target="_blank" rel="noopener">Example Regulatory Update ("ERU") for Financial Services</a></p>]]></content:encoded>
			<slash:comments>0</slash:comments>
	</item>
	<item>
		<title>Industry News: Example article about the fund industry</title>
		<link>https://acme.example.com/industry-news-example-article/</link>
		<dc:creator><![CDATA[user_wp]]></dc:creator>
		<pubDate>Fri, 06 Feb 2026 02:06:48 +0000</pubDate>
				<category><![CDATA[Highlight News]]></category>
		<guid isPermaLink="false">https://acme.example.com/?p=1002</guid>
					<description><![CDATA[<p class="date">5 February, 2026</p>
<p><a href="https://news.example.com/node/1000" target="_blank" rel="noopener">Industry News: Example article about the fund industry</a></p>]]></description>
										<content:encoded><![CDATA[<p class="date">5 February, 2026</p>
<p><a href="https://news.example.com/node/1000" target="_blank" rel="noopener">Industry News: Example article about the fund industry</a></p>]]></content:encoded>
			<slash:comments>0</slash:comments>
	</item>
	<item>
		<title>示例新闻标题 测试中文链接 &#124; 示例节目</title>
		<link>https://acme.example.com/%e7%a4%ba%e4%be%8b%e6%96%b0%e9%97%bb%e6%a0%87%e9%a2%98-%e6%b5%8b%e8%af%95%e4%b8%ad%e6%96%87%e9%93%be%e6%8e%a5-%e7%a4%ba%e4%be%8b%e8%8a%82%e7%9b%ae/</link>
		<dc:creator><![CDATA[admin]]></dc:creator>
		<pubDate>Wed, 21 Jan 2026 09:47:51 +0000</pubDate>
				<category><![CDATA[Highlight News]]></category>
		<guid isPermaLink="false">https://acme.example.com/?p=1001</guid>
					<description><![CDATA[<p class="date">19 January, 2026</p>
<p><a href="https://video.example.com/watch?v=example123&#038;feature=share" target="_blank" rel="noopener">示例新闻标题 测试中文链接 &#124; 示例节目
</a></p>]]></description>
										<content:encoded><![CDATA[<p class="date">19 January, 2026</p>
<p><a href="https://video.example.com/watch?v=example123&#038;feature=share" target="_blank" rel="noopener">示例新闻标题 测试中文链接 &#124; 示例节目
</a></p>]]></content:encoded>
			<slash:comments>0</slash:comments>
	</item>
</channel>
</rss>"""

# Synthetic: the real comments feed was empty on 29 Sep 2026. Shape of a WordPress comment feed item.
_COMPANY_COMMENTS_FIXTURE = """<?xml version="1.0" encoding="UTF-8"?><rss version="2.0"
	xmlns:content="http://purl.org/rss/1.0/modules/content/" xmlns:dc="http://purl.org/dc/elements/1.1/">
<channel><title>Comments for Acme Asset Management Malaysia</title><link>https://acme.example.com</link>
	<item>
		<title>Comment on Industry News: Example article about the fund industry by Example Commenter</title>
		<link>https://acme.example.com/industry-news-example-article/#comment-12</link>
		<dc:creator><![CDATA[Example Commenter]]></dc:creator>
		<pubDate>Tue, 15 Sep 2026 03:00:00 +0000</pubDate>
		<guid isPermaLink="false">https://acme.example.com/industry-news-example-article/#comment-12</guid>
		<description><![CDATA[Illustrative comment text.]]></description>
		<content:encoded><![CDATA[<p>Illustrative comment text.</p>]]></content:encoded>
	</item>
</channel></rss>"""


@selftest("S11")
def test_parse_company_feed():
    items = parse_wordpress_feed(_COMPANY_FEED_FIXTURE)
    assert len(items) == 3
    sst, edge, zh = items
    assert sst.title == "Notice: Example Regulatory Update (“ERU”) for Financial Services"
    assert sst.discovered_url == "https://acme.example.com/notice/"
    assert sst.source_published_at == datetime(2026, 3, 12, 8, 21, 13, tzinfo=timezone.utc), \
        "the feed pubDate (when the company posted it) is the publication date"
    assert sst.extra.get("page_date_text") == "September 2025", "the date shown in the post is kept as a note"
    assert "Example Regulatory Update" in (sst.content_text or "")
    assert (sst.source, sst.discovery_method) == ("website", "Direct public retrieval")
    assert sst.publisher == "Acme Asset Management Malaysia" and sst.source_message_id == "https://acme.example.com/?p=1003"
    assert edge.extra.get("categories") == ["Highlight News"]
    assert zh.title == "示例新闻标题 测试中文链接 | 示例节目"
    assert zh.discovered_url.startswith("https://acme.example.com/%e7%a4%ba")
    assert not any(i.extra.get("is_comment") for i in items)


@selftest("S11")
def test_parse_comment_feed_and_comment_keys():
    (c,) = parse_wordpress_feed(_COMPANY_COMMENTS_FIXTURE, comments=True)
    assert c.extra.get("is_comment") is True and c.content_text == "Illustrative comment text."
    assert c.discovered_url.endswith("#comment-12")
    assert c.source_published_at == _utc(2026, 9, 15, 3, 0)
    post = "https://acme.example.com/industry-news-example-article/"
    assert canonicalize_url(c.discovered_url).endswith("#comment-12"), "a comment anchor identifies content"
    assert make_url_key(canonicalize_url(c.discovered_url)) != make_url_key(canonicalize_url(post)), \
        "a comment must not merge into its post"
    assert canonicalize_url(post + "#respond") == canonicalize_url(post), "other fragments are still removed"


def _wp_feed(entries: list[tuple[str, str]]) -> str:
    """Minimal WordPress-style feed: [(slug, RFC 822 pubDate)], newest first."""
    body = "".join(f"<item><title>Post {slug}</title><link>https://acme.example.com/{slug}/</link>"
                   f"<pubDate>{when}</pubDate><guid>https://acme.example.com/?p={i}</guid>"
                   f"<description>Summary {slug}</description></item>" for i, (slug, when) in enumerate(entries))
    return f'<?xml version="1.0"?><rss version="2.0"><channel><title>Acme</title>{body}</channel></rss>'


def _website_ctx() -> RunContext:
    ctx = _alerts_ctx()                                   # period 14-20 Sep 2026, fast HTTP settings
    ctx.config.settings["sources"]["website"].update(
        feed_url="https://acme.example.com/feed/", comments_feed_url="https://acme.example.com/comments/feed/")
    return ctx


_WP_ROBOTS = {"https://acme.example.com/robots.txt": [_FakeResponse(200, "User-agent: *\nDisallow: /wp-admin/\n")]}


@selftest("S11")
def test_website_collector_window_and_paging():
    feed, comments = "https://acme.example.com/feed/", "https://acme.example.com/comments/feed/"
    page1 = _wp_feed([("current-week", "Thu, 24 Sep 2026 02:00:00 +0000"),
                      ("in-period", "Tue, 15 Sep 2026 02:00:00 +0000"),
                      ("old", "Tue, 01 Sep 2026 02:00:00 +0000")])
    s = _FakeSession({**_WP_ROBOTS, feed: [_FakeResponse(200, page1)],
                      comments: [_FakeResponse(200, _COMPANY_COMMENTS_FIXTURE)]})
    r = run_collector_safely(WebsiteCollector(session=s), _website_ctx())
    assert r.status == SourceStatus.SUCCESS, r.error
    urls = [i.discovered_url for i in r.items]
    assert "https://acme.example.com/in-period/" in urls and "https://acme.example.com/current-week/" in urls, \
        "current-week posts are collected (S14 excludes them from this report)"
    assert "https://acme.example.com/old/" not in urls, "posts older than the period are not collected"
    assert feed + "?paged=2" not in s.calls, "no further page is needed once an older post appears"
    assert comments + "?paged=2" in s.calls, "the comments feed had no older comment, so page 2 is checked"
    assert any(i.extra.get("is_comment") for i in r.items) and r.stats["comments"] == 1
    all_new = _wp_feed([("a", "Sat, 19 Sep 2026 02:00:00 +0000"), ("b", "Fri, 18 Sep 2026 02:00:00 +0000")])
    page2 = _wp_feed([("c", "Thu, 17 Sep 2026 02:00:00 +0000"), ("d", "Mon, 07 Sep 2026 02:00:00 +0000")])
    s = _FakeSession({**_WP_ROBOTS, feed: [_FakeResponse(200, all_new)],
                      feed + "?paged=2": [_FakeResponse(200, page2)], comments: [_FakeResponse(503)]})
    r = run_collector_safely(WebsiteCollector(session=s), _website_ctx())
    assert r.status == SourceStatus.SUCCESS and [i.discovered_url.split("/")[3] for i in r.items] == ["a", "b", "c"]
    assert r.warnings and "comments" in r.warnings[0], "a failed comments feed is a warning"
    s = _FakeSession({**_WP_ROBOTS, feed: [_FakeResponse(200, all_new)], feed + "?paged=2": [_FakeResponse(404)],
                      comments: [_FakeResponse(200, _wp_feed([]))]})
    r = run_collector_safely(WebsiteCollector(session=s), _website_ctx())
    assert r.status == SourceStatus.SUCCESS and len(r.items) == 2, "a 404 past the last page ends paging"


@selftest("S11")
def test_website_collector_failures():
    feed = "https://acme.example.com/feed/"
    s = _FakeSession({**_WP_ROBOTS, feed: [_FakeResponse(503)]})
    r = run_collector_safely(WebsiteCollector(session=s), _website_ctx())
    assert r.status == SourceStatus.FAILED and "503" in (r.error or ""), "the main feed failing fails the source"
    s = _FakeSession({**_WP_ROBOTS, feed: [_FakeResponse(200, "<html>not a feed</html>")]})
    r = run_collector_safely(WebsiteCollector(session=s), _website_ctx())
    assert r.status == SourceStatus.FAILED and "not an RSS feed" in (r.error or ""), r.error
    ctx = _website_ctx()
    ctx.config.settings["sources"]["website"]["feed_url"] = None
    r = run_collector_safely(WebsiteCollector(session=_FakeSession({})), ctx)
    assert r.status == SourceStatus.FAILED and "feed_url" in (r.error or "")
    assert DEFAULT_SETTINGS["sources"]["website"]["feed_url"] == "https://acme.example.com/feed/", \
        "the path chosen after inspecting the site is recorded in the settings template"


# ---- S12 --------------------------------------------------------------------
@selftest("S12")
def test_sc_alert_list_parser_matches_only_entities():
    html = ("<table><tr><th>Name</th><th>Date Added</th></tr>"
            "<tr><td>Some Unrelated Investment Scheme</td><td>10/09/2026</td></tr>"
            "<tr><td>Illustrative Row Acme Asset Management Test</td><td>15/09/2026</td></tr>"
            "</table>")
    items = parse_sc_alert_list(html, ["Acme Asset Management", "Tan Ah Kow"],
                                "https://www.sc.com.my/example-alert-list")
    assert len(items) == 1
    assert items[0].extra.get("sc_alert_list") is True and items[0].source == "sc_malaysia"


# Trimmed from the real SC pages (29 Sep 2026): the page scripts `$X.PMD = {...}` (Investor Alert
# List) and `$X.PPG = {...}` (media releases) with two real entries / four real records, plus one
# SYNTHETIC alert entry naming the monitored company (no such entry exists on the real list).
_SC_ALERT_DATA = {
    "cf1e0250-7566-4738-9fc0-604a002ee200": {
        "5b1925f8-8d1b-4139-a633-450053364988": {
            "name": "Potential Clone Entity - Bursa Coins", "address": "N/A", "website": "https://bursa-coins.my/",
            "date": "2026", "remark": "•\tMisusing the name and credentials of Bursa Malaysia Berhad (an exchange "
            "holding company under the purview of the SC and the Ministry of Finance) to carry out an illegal investment "
            "scheme \r\n•\tCarrying on unlicensed capital market activities of dealing in securities\r\n",
            "group": ",B,", "url": "", "bankInfo": ""},
        "6a847d9d-9c66-4244-a948-0a09c835db24": {
            "name": "Example Dragon BTC Mining-Malaysia", "address": "N/A",
            "website": "https://www.acme-fund.com/ | https://t.me/exampledragonminingma | https://t.me/exampledragonminingmalaysia",
            "date": "2022", "remark": "Offering of securities without SC’s approval", "group": ",H,", "url": "", "bankInfo": ""},
        "synthetic-0001": {
            "name": "Illustrative Clone of ACME ASSET MANAGEMENT SDN. BHD.", "address": "N/A",
            "website": "https://example.invalid/", "date": "2026", "remark": "Synthetic test entry",
            "group": ",I,", "url": "", "bankInfo": ""}},
    "e7290085-09eb-4f22-aa60-c3775195004b": {}}
_SC_RELEASE_DATA = {
    "f3883a8e-7c9e-4192-8caf-5d1c404699d3": {"items": [], "cnt": 3, "pages": {"0": [
        {"Page Name": "ACSR's Mandatory Sustainability Assurance Requirement to Take Effect in 2028",
         "Page Description": "The Advisory Committee on Sustainability Reporting (ACSR) has deferred the commencement.",
         "Page Address": "resources/media/media-release/acsrs-mandatory-sustainability-assurance-requirement-to-take-effect-in-2028",
         "Display Date 2": "17 September 2026", "xDataId": "f6bb99af-9966-40da-b923-1183a2de2fef"},
        {"Page Name": "SC Hosts Inaugural ASEAN Regulatory Sustainability Summit in Collaboration with Durham University",
         "Page Description": "The Securities Commission Malaysia (SC) held the inaugural ASEAN Regulatory Sustainability Summit.",
         "Page Address": "resources/media/media-release/sc-hosts-inaugural-asean-regulatory-sustainability-summit-in-collaboration-with-durham-university",
         "Display Date 2": "11 September 2026", "xDataId": "558d8902-f080-4fe1-bef9-e1057b75b62e"},
        {"Page Name": "SC and Bursa Malaysia Collaborate to Advance Informed Investing Among Retail Investors",
         "Page Description": "The SC and Bursa Malaysia today launched the inaugural InvestSmart x Bursa Marketplace Fair 2026.",
         "Page Address": "resources/media/media-release/sc-and-bursa-malaysia-collaborate-to-advance-informed-investing-among-retail-investors",
         "Display Date 2": "05 September 2026", "xDataId": "b48dd63b-2555-41cb-a36e-c38eefcc2345"}]}},
    "6ccafc50-f1ea-4e42-ae52-3a0e8709a2e9": {"items": [], "cnt": 1, "pages": {"0": [
        {"Page Name": "SC Takes Enforcement Actions on Binance for Illegally Operating in Malaysia",
         "Page Address": "resources/media/media-release/sc-takes-enforcement-actions-on-binance-for-illegally-operating-in-malaysia",
         "Parent Page Name": "Media Releases", "xDataId": "cf3e0c8b-b246-4e7b-8a9e-1649cbfba40b"}]}}}


def _sc_page(variable: str, data: dict) -> str:
    return (f"<!DOCTYPE html><html><head><title>SC</title></head><body><div>layout</div><script>"
            f"$X.{variable} = {json.dumps(data)};$X.OTHER = {{}};</script></body></html>")


_SC_ALERT_URL = "https://www.sc.com.my/investor-alert-list"
_SC_KEY_CLONE = "sc-alert-list:illustrative clone of acme asset management sdn bhd | https example invalid"
_SC_KEY_LOOKALIKE = ("sc-alert-list:example dragon btc mining malaysia | https www acme fund com https t me "
                  "exampledragonminingma https t me exampledragonminingmalaysia")
_SC_RELEASES_URL = "https://www.sc.com.my/resources/media/media-release"


@selftest("S12")
def test_sc_alert_list_embedded_data():
    page = _sc_page("PMD", _SC_ALERT_DATA)
    names = ["Acme Asset Management Sdn Bhd", "Acme Asset Management", "Tan Ah Kow"]
    (hit,) = parse_sc_alert_list(page, names, _SC_ALERT_URL)
    assert hit.title == "SC Investor Alert List: Illustrative Clone of ACME ASSET MANAGEMENT SDN. BHD."
    assert hit.source_published_at is None and hit.extra["sc_listed_year"] == "2026", \
        "the list only states a year: it is kept, never turned into a date"
    assert hit.extra["content_key"] == _SC_KEY_CLONE and hit.discovered_url == _SC_ALERT_URL
    assert (hit.publisher, hit.source) == ("Securities Commission Malaysia", "sc_malaysia")
    (only_full,) = parse_sc_alert_list(page, entity_screening_terms(DEFAULT_ENTITIES), _SC_ALERT_URL)
    assert only_full.extra["content_key"] == _SC_KEY_CLONE, \
        "D-27: with the full company name only, the acme-fund.com row is not even pre-selected"
    with_terms = parse_sc_alert_list(page, entity_screening_terms(_AMBIGUOUS_TEST_ENTITIES), _SC_ALERT_URL)
    assert {i.extra["content_key"] for i in with_terms} == {_SC_KEY_CLONE, _SC_KEY_LOOKALIKE}, \
        "when an ambiguous term IS configured, rows using it are pre-selected for the entity matcher to judge"
    assert parse_sc_alert_list(page, ["Acmeville"], _SC_ALERT_URL) == [], "whole words only"
    a, b = (normalize_item(i) for i in with_terms)
    assert a.url_key != b.url_key and a.canonical_url == b.canonical_url, \
        "entries on one page must not merge: each has its own key"
    group = _SC_ALERT_DATA["cf1e0250-7566-4738-9fc0-604a002ee200"]
    reloaded = {"g": {f"new-random-{i}": e for i, e in enumerate(group.values())}}   # SC re-keys every load
    again = parse_sc_alert_list(_sc_page("PMD", reloaded), entity_screening_terms(_AMBIGUOUS_TEST_ENTITIES),
                                _SC_ALERT_URL)
    assert {i.extra["content_key"] for i in again} == {_SC_KEY_CLONE, _SC_KEY_LOOKALIKE}, \
        "the same entry keeps its identity when the page's random keys change"


@selftest("S12")
def test_sc_alert_table_dates_and_helpers():
    html_text = ("<table><tr><th>Name</th><th>Website</th><th>Date Added</th></tr>"
                 "<tr><td>Tan Ah Kow Illustrative Scheme</td><td>x.example</td><td>15/09/2026</td></tr>"
                 "<tr><td>Other Scheme</td><td>y.example</td><td>16/09/2026</td></tr></table>")
    (item,) = parse_sc_alert_list(html_text, ["Tan Ah Kow"], _SC_ALERT_URL)
    assert item.source_published_at == _myt(2026, 9, 15), "a 'date added' column is the publication date (MYT)"
    assert item.extra["content_key"] == "sc-alert-list:tan ah kow illustrative scheme"
    assert _day_month_year("17 September 2026") == date(2026, 9, 17)
    assert _day_month_year("5 Sept 2026") == date(2026, 9, 5) and _day_month_year("05 Sep, 2026") == date(2026, 9, 5)
    assert _day_month_year("31 February 2026") is None and _day_month_year("2026") is None
    assert _day_month_year("1 Ma 2026") is None, "month abbreviations need at least 3 letters"
    assert _mentions_any("ACME ASSET MANAGEMENT SDN. BHD.", ["Acme Asset Management Sdn Bhd"])
    assert not _mentions_any("Acmeville Asset Management", ["Acme Asset Management"])


@selftest("S12")
def test_sc_media_releases():
    items = parse_sc_media_releases(_sc_page("PPG", _SC_RELEASE_DATA), _SC_RELEASES_URL)
    assert len(items) == 3, "records without a display date (popular links) are not releases"
    first = items[0]
    assert first.discovered_url == ("https://www.sc.com.my/resources/media/media-release/"
                                    "acsrs-mandatory-sustainability-assurance-requirement-to-take-effect-in-2028")
    assert first.source_published_at == _myt(2026, 9, 17) and first.extra["date_precision"] == "day"
    assert first.title.startswith("ACSR's") and first.snippet.startswith("The Advisory Committee")
    try:
        parse_sc_media_releases("<html><body>redesigned page</body></html>", _SC_RELEASES_URL)
    except ValueError as e:
        assert "layout changed" in str(e)
    else:
        raise AssertionError("a page without the release data must be reported, not read as 'no releases'")


def _sc_session(release_page_status: int = 200, alert_status: int = 200, releases_status: int = 200):
    acsr = _SC_RELEASES_URL + "/acsrs-mandatory-sustainability-assurance-requirement-to-take-effect-in-2028"
    return _FakeSession({
        "https://www.sc.com.my/robots.txt": [_FakeResponse(200, "User-agent: *\n\nDisallow:\n")],
        _SC_ALERT_URL: [_FakeResponse(alert_status, _sc_page("PMD", _SC_ALERT_DATA))],
        _SC_RELEASES_URL: [_FakeResponse(releases_status, _sc_page("PPG", _SC_RELEASE_DATA))],
        acsr: [_FakeResponse(release_page_status, "<html><body><h1>ACSR</h1><p>Full release text.</p></body></html>")]})


@selftest("S12")
def test_sc_collector():
    ctx = _alerts_ctx()                                  # period 14-20 Sep 2026
    s = _sc_session()
    r = run_collector_safely(SecuritiesCommissionCollector(session=s), ctx)
    assert r.status == SourceStatus.SUCCESS, r.error
    alerts = [i for i in r.items if i.extra.get("sc_alert_list")]
    releases = [i for i in r.items if not i.extra.get("sc_alert_list")]
    assert len(alerts) == 1, "D-27: only rows naming an approved full name are pre-selected"
    assert all(i.extra["first_seen_period_start"] == "2026-09-14" for i in alerts), \
        "a new undated alert-list hit is dated by this run's period (FIRST_SEEN)"
    assert [i.title[:5] for i in releases] == ["ACSR'"], "only releases from the period start are kept"
    assert releases[0].content_text == "ACSR Full release text.", "the release page is read for its full text"
    assert not any("durham" in c for c in s.calls), "releases before the period are not fetched"
    r = run_collector_safely(SecuritiesCommissionCollector(session=_sc_session(release_page_status=500)), ctx)
    assert r.status == SourceStatus.SUCCESS and r.warnings and r.items, "an unreadable release page is a warning"
    r = run_collector_safely(SecuritiesCommissionCollector(session=_sc_session(releases_status=503)), ctx)
    assert r.status == SourceStatus.SUCCESS and any("announcements" in w for w in r.warnings)
    r = run_collector_safely(SecuritiesCommissionCollector(session=_sc_session(alert_status=503)), ctx)
    assert r.status == SourceStatus.FAILED, "the Investor Alert List failing fails the source"
    s = _sc_session()
    s.script[_SC_ALERT_URL] = [_FakeResponse(200, "<html><body>redesigned</body></html>")]
    r = run_collector_safely(SecuritiesCommissionCollector(session=s), ctx)
    assert r.status == SourceStatus.FAILED and "layout changed" in (r.error or ""), \
        "a changed layout must be FAILED, never 'no alert-list hits'"


@selftest("S12")
def test_sc_first_seen_period_comes_from_the_database():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        earlier = _test_ctx(db, "2026-09-15")            # period 7-13 Sep 2026 found it first
        conn = db_connect(db)
        db_record_run_start(conn, earlier)
        c = _db_candidate(content_key_url_key(_SC_KEY_CLONE))
        db_upsert_candidate(conn, c, earlier)
        conn.close()
        ctx = _alerts_ctx()
        ctx.db_path = db                                 # this run: period 14-20 Sep 2026
        r = run_collector_safely(SecuritiesCommissionCollector(session=_sc_session()), ctx)
        by_key = {i.extra["content_key"]: i for i in r.items if i.extra.get("sc_alert_list")}
        assert by_key[_SC_KEY_CLONE].extra["first_seen_period_start"] == "2026-09-07", \
            "an entry already reported in an earlier week keeps that week (not repeated every week)"
        assert _SC_KEY_LOOKALIKE not in by_key, "D-27: the acme-fund.com row is not a candidate at all"


# ---- S13 --------------------------------------------------------------------
@selftest("S13")
def test_platform_detection_from_hostname_only():
    assert detect_platform("https://linkedin.com/in/x", "google_alerts") == Platform.LINKEDIN
    assert detect_platform("https://facebook.com/page", "google_alerts") == Platform.FACEBOOK
    assert detect_platform("https://instagram.com/p/x", "google_alerts") == Platform.INSTAGRAM
    assert detect_platform("https://sc.com.my/x", "sc_malaysia") == Platform.SC_MALAYSIA
    assert detect_platform("https://acme.example.com/news/a", "website") == Platform.COMPANY_WEBSITE
    # anything that cannot be identified from the hostname is listed as Public website
    assert detect_platform("https://news.example.com/a", "google_alerts") == Platform.PUBLIC_WEB
    # a site:facebook.com alert that returned another domain must NOT become Facebook
    assert detect_platform("https://random-blog.example.org/p", "google_alerts") == Platform.PUBLIC_WEB
    assert Platform.PUBLIC_WEB == "Public website"


@selftest("S13")
def test_url_classification_and_verification():
    ct, ut = classify_url("https://instagram.com/explore/tags/something", Platform.INSTAGRAM)
    assert ut == UrlType.NON_EVIDENCE
    ct, ut = classify_url("https://instagram.com/p/AbC123", Platform.INSTAGRAM)
    assert ut == UrlType.SPECIFIC_CONTENT and "post" in ct.lower()
    ct, ut = classify_url("https://facebook.com/SomePage", Platform.FACEBOOK)
    assert ut == UrlType.CONTAINER
    ct, ut = classify_url("https://linkedin.com/company/some-co", Platform.LINKEDIN)
    assert ut == UrlType.CONTAINER
    ct, ut = classify_url("https://linkedin.com/posts/someone_x-activity-724-AbCd", Platform.LINKEDIN)
    assert ut == UrlType.SPECIFIC_CONTENT

    c = _candidate("https://instagram.com/explore/tags/x")
    c.platform, c.content_class, c.url_type = Platform.INSTAGRAM, ContentClass.SOCIAL, UrlType.NON_EVIDENCE
    assert verify_url(c, True, None)[0] == Verification.REJECTED
    c = _candidate("https://facebook.com/SomePage")
    c.platform, c.content_class, c.url_type = Platform.FACEBOOK, ContentClass.SOCIAL, UrlType.CONTAINER
    assert verify_url(c, True, None)[0] == Verification.PARTIAL, "a Page is never 'the post'"
    c = _candidate("https://linkedin.com/posts/someone_x-activity-724-AbCd")
    c.platform, c.content_class, c.url_type = Platform.LINKEDIN, ContentClass.SOCIAL, UrlType.SPECIFIC_CONTENT
    status, basis = verify_url(c, True, None)
    assert status == Verification.VERIFIED and "snippet" in basis
    assert verify_url(c, False, None)[0] == Verification.PARTIAL


@selftest("S13")
def test_platform_hosts_and_content_classes():
    cases = {"https://business.facebook.com/x": Platform.FACEBOOK, "https://fb.watch/abc": Platform.FACEBOOK,
             "https://fb.com/x": Platform.FACEBOOK, "https://notfacebook.com/x": Platform.PUBLIC_WEB,
             "https://facebook.com.evil.example/x": Platform.PUBLIC_WEB, "https://www.sc.com.my/x": Platform.SC_MALAYSIA,
             "https://xsc.com.my/x": Platform.PUBLIC_WEB, "https://shop.acme.example.com/x": Platform.COMPANY_WEBSITE,
             "https://uk.linkedin.com/in/x": Platform.LINKEDIN, "": Platform.PUBLIC_WEB, "not a url": Platform.PUBLIC_WEB}
    for url, platform in cases.items():
        assert detect_platform(url, "google_alerts") == platform, f"{url!r} -> {detect_platform(url, 'x')}"
    raw = RawItem("website", "x", "u")
    expected = {Platform.FACEBOOK: ContentClass.SOCIAL, Platform.INSTAGRAM: ContentClass.SOCIAL,
                Platform.LINKEDIN: ContentClass.SOCIAL, Platform.SC_MALAYSIA: ContentClass.REGULATORY,
                Platform.COMPANY_WEBSITE: ContentClass.COMPANY, Platform.PUBLIC_WEB: ContentClass.PUBLIC_WEB}
    for platform, cls in expected.items():
        assert content_class_for(platform, raw) == cls
    comment = RawItem("website", "x", "u", extra={"is_comment": True})
    assert content_class_for(Platform.COMPANY_WEBSITE, comment) == ContentClass.SOCIAL, \
        "comments on the company website are judged as social content"


@selftest("S13")
def test_url_pattern_tables():
    # One row per line of SOURCE_ACCESS_BLUEPRINT §3 (URLs already canonicalised).
    fb, ig, li = Platform.FACEBOOK, Platform.INSTAGRAM, Platform.LINKEDIN
    rows = [
        (fb, "https://facebook.com/SomePage/posts/123", "Post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/permalink.php?id=2&story_fbid=1", "Post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/story.php?comment_id=9&id=2&story_fbid=1", "Comment on a post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/SomePage/posts/123?comment_id=9", "Comment on a post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/groups/g1/posts/55", "Group post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/groups/g1/permalink/55", "Group post", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/reel/77", "Video / reel", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/watch?v=77", "Video / reel", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/SomePage/videos/77", "Video / reel", UrlType.SPECIFIC_CONTENT),
        (fb, "https://fb.watch/abc", "Video / reel", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/photo.php?fbid=88", "Photo", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/photo?fbid=88", "Photo", UrlType.SPECIFIC_CONTENT),
        (fb, "https://facebook.com/share/p/xyz", "Share link (not resolved)", UrlType.UNKNOWN),
        (fb, "https://facebook.com/SomePage/about", "Page / profile", UrlType.CONTAINER),
        (fb, "https://facebook.com/profile.php?id=99", "Page / profile", UrlType.CONTAINER),
        (fb, "https://facebook.com/SomePage/reviews", "Page reviews tab (not a specific review)", UrlType.CONTAINER),
        (fb, "https://facebook.com/groups/g1", "Group home", UrlType.CONTAINER),
        (fb, "https://facebook.com/search/top?q=acme", "Search / hashtag / login", UrlType.NON_EVIDENCE),
        (fb, "https://facebook.com/hashtag/acme", "Search / hashtag / login", UrlType.NON_EVIDENCE),
        (fb, "https://facebook.com/login.php", "Search / hashtag / login", UrlType.NON_EVIDENCE),
        (fb, "https://facebook.com", "Home page", UrlType.NON_EVIDENCE),
        (ig, "https://instagram.com/reel/AbC", "Reel", UrlType.SPECIFIC_CONTENT),
        (ig, "https://instagram.com/reels/AbC", "Reel", UrlType.SPECIFIC_CONTENT),
        (ig, "https://instagram.com/tv/AbC", "Video (IGTV)", UrlType.SPECIFIC_CONTENT),
        (ig, "https://instagram.com/someuser/p/AbC", "Post", UrlType.SPECIFIC_CONTENT),
        (ig, "https://instagram.com/someuser", "Profile", UrlType.CONTAINER),
        (ig, "https://instagram.com/stories/someuser/123", "Story (ephemeral)", UrlType.SPECIFIC_CONTENT),
        (ig, "https://instagram.com/explore", "Explore page", UrlType.NON_EVIDENCE),
        (ig, "https://instagram.com/accounts/login", "Login / account page", UrlType.NON_EVIDENCE),
        (li, "https://linkedin.com/feed/update/urn:li:activity:724", "Individual post", UrlType.SPECIFIC_CONTENT),
        (li, "https://linkedin.com/pulse/some-article", "Article (LinkedIn Pulse)", UrlType.SPECIFIC_CONTENT),
        (li, "https://linkedin.com/company/some-co/posts", "Company page", UrlType.CONTAINER),
        (li, "https://linkedin.com/in/someone", "Profile", UrlType.CONTAINER),
        (li, "https://linkedin.com/jobs/view/1", "Other — not evidence", UrlType.NON_EVIDENCE),
        (li, "https://linkedin.com/school/x", "Other — not evidence", UrlType.NON_EVIDENCE),
        (li, "https://linkedin.com/search/results/all?keywords=x", "Other — not evidence", UrlType.NON_EVIDENCE),
        (li, "https://linkedin.com/something-else", "Unknown", UrlType.UNKNOWN),
        (Platform.SC_MALAYSIA, "https://sc.com.my/investor-alert-list", "Investor Alert List", UrlType.CONTAINER),
        (Platform.SC_MALAYSIA, "https://sc.com.my/resources/media/media-release/x", "Media release", UrlType.SPECIFIC_CONTENT),
        (Platform.COMPANY_WEBSITE, "https://acme.example.com/notice", "Article / post", UrlType.SPECIFIC_CONTENT),
        (Platform.COMPANY_WEBSITE, "https://acme.example.com/post#comment-12", "Comment", UrlType.SPECIFIC_CONTENT),
        (Platform.COMPANY_WEBSITE, "https://acme.example.com/category/news", "Listing / tag page", UrlType.CONTAINER),
        (Platform.COMPANY_WEBSITE, "https://acme.example.com", "Home page", UrlType.NON_EVIDENCE),
        (Platform.PUBLIC_WEB, "https://news.example/2026/09/16/story", "Web page", UrlType.SPECIFIC_CONTENT),
        (Platform.PUBLIC_WEB, "https://news.example/files/report.pdf", "Document (PDF)", UrlType.SPECIFIC_CONTENT),
        (Platform.PUBLIC_WEB, "https://cdn.example/img/photo.JPG", "Image", UrlType.NON_EVIDENCE),
        (Platform.PUBLIC_WEB, "https://news.example/search?q=acme", "Search page", UrlType.NON_EVIDENCE),
        (Platform.PUBLIC_WEB, "https://blog.example/?s=acme", "Search page", UrlType.NON_EVIDENCE),
        (Platform.PUBLIC_WEB, "https://news.example/tag/acme", "Listing / tag page", UrlType.CONTAINER),
    ]
    for platform, url, content_type, url_type in rows:
        assert classify_url(url, platform) == (content_type, url_type), f"{url}: {classify_url(url, platform)}"


def _verify_case(url, platform, content_class, text=None, keyword=None, title="", snippet=""):
    c = _candidate(url, title=title, snippet=snippet, content_text=text)
    c.platform, c.content_class = platform, content_class
    c.content_type, c.url_type = classify_url(url, platform)
    if keyword:
        c.keyword_matches = [Match("", keyword, keyword, "cat", "HIGH")]
    return c


@selftest("S13")
def test_verification_branches():
    web = (Platform.PUBLIC_WEB, ContentClass.PUBLIC_WEB)
    page = "https://news.example/2026/09/16/story"
    c = _verify_case(page, *web, text="Acme Asset Management was fined", keyword="fined")
    assert verify_url(c, True, True) == (Verification.VERIFIED, "page text match")
    c = _verify_case(page, *web, text="Acme Asset Management annual report", keyword="fined")
    status, basis = verify_url(c, True, True)
    assert status == Verification.PARTIAL and "keyword" in basis, "name and keyword must both be in the page"
    c = _verify_case(page, *web, text="An unrelated page")
    assert verify_url(c, True, False)[0] == Verification.PARTIAL, "page read but match not re-located"
    c = _verify_case(page, *web)
    status, basis = verify_url(c, True, None)
    assert status == Verification.PARTIAL and "not read" in basis, "a public web page is VERIFIED only when read"
    c = _verify_case("https://facebook.com/share/p/xyz", Platform.FACEBOOK, ContentClass.SOCIAL)
    assert verify_url(c, True, None)[0] == Verification.PARTIAL, "share links are never resolved, so PARTIAL"
    c = _verify_case("https://sc.com.my/investor-alert-list", Platform.SC_MALAYSIA, ContentClass.REGULATORY,
                     text="Illustrative Clone of Acme Asset Management | N/A")
    assert verify_url(c, True, True)[0] == Verification.VERIFIED, "an alert-list entry read from the SC page"
    for bad in ("", "https://www.google.com/url?sa=t&usg=x"):
        c = _verify_case(bad, *web)
        assert verify_url(c, True, True)[0] == Verification.REJECTED, f"{bad!r} must be rejected"
    c = _verify_case("https://cdn.example/img/photo.jpg", *web, text="Acme Asset Management")
    assert verify_url(c, True, True)[0] == Verification.REJECTED, "an image URL is not evidence, even with text"
    before = (c.canonical_url, c.raw.discovered_url, c.raw.final_url)
    verify_url(c, True, True)
    assert (c.canonical_url, c.raw.discovered_url, c.raw.final_url) == before, "verification never alters a URL"


# ---- S14 --------------------------------------------------------------------
@selftest("S14")
def test_alert_email_date_is_never_the_content_date():
    c = _candidate("https://facebook.com/SomePage", alert_received_at=_utc(2026, 9, 16, 9, 15))
    c.platform = Platform.FACEBOOK
    resolve_content_date(c)
    assert c.date_status == DateStatus.UNKNOWN and c.content_date is None
    assert period_position(c, _period("2026-09-22")) == "UNKNOWN"


@selftest("S14")
def test_source_date_and_period_position():
    p = _period("2026-09-25")
    c = _candidate("https://news.example.com/a", source="google_alerts",
                   source_published_at=_utc(2026, 9, 17, 3, 0))
    resolve_content_date(c)
    assert c.date_status == DateStatus.VERIFIED and period_position(c, p) == "IN"
    c = _candidate("https://news.example.com/b", source="google_alerts",
                   source_published_at=_myt(2026, 9, 24, 12, 0))
    resolve_content_date(c)
    assert period_position(c, p) == "AFTER", "current-week content must be excluded"
    c = _candidate("https://news.example.com/c", source="google_alerts",
                   source_published_at=_myt(2026, 9, 2, 12, 0))
    resolve_content_date(c)
    assert period_position(c, p) == "BEFORE"


@selftest("S14")
def test_relative_snippet_date_is_approximate():
    c = _candidate("https://linkedin.com/in/x", snippet="3 days ago · Illustrative text",
                   alert_received_at=_myt(2026, 9, 18, 9, 0))
    resolve_content_date(c)
    assert c.date_status == DateStatus.APPROXIMATE
    assert c.content_date is not None and c.content_date.astimezone(MYT).date() == date(2026, 9, 15)


@selftest("S14")
def test_source_dates_and_their_basis():
    c = _candidate("https://acme.example.com/x", source="website", source_published_at=_utc(2026, 9, 15, 20, 0))
    resolve_content_date(c)
    assert c.content_date == _myt(2026, 9, 16, 4, 0), "16 Sep 04:00 MYT"
    assert c.content_date is not None and c.content_date.tzinfo is MYT, "stored in Malaysia time"
    assert c.date_status == DateStatus.VERIFIED and "feed" in c.date_basis
    c = _candidate("https://news.example/a", content_text="text", source_published_at=_myt(2026, 9, 15),
                   extra={"published_timezone_assumed": "Asia/Kuala_Lumpur"})
    resolve_content_date(c)
    assert "metadata" in c.date_basis and "read as Malaysia time" in c.date_basis
    c = _candidate("https://news.example/b", snippet="Sep 01, 2026 — text", source_published_at=_myt(2026, 9, 16),
                   alert_received_at=_myt(2026, 9, 17))
    resolve_content_date(c)
    assert c.content_date == _myt(2026, 9, 16), "the source's own date beats a snippet date"


@selftest("S14")
def test_snippet_date_forms():
    received = _myt(2026, 9, 18, 9, 0)
    cases = {
        "2 hours ago · text": _myt(2026, 9, 18, 7, 0),
        "1 day ago - text": _myt(2026, 9, 17, 9, 0),
        "a week ago text": _myt(2026, 9, 11, 9, 0),
        "45 mins ago text": _myt(2026, 9, 18, 8, 15),
        "Sep 15, 2026 ... text": _myt(2026, 9, 15),
        "15 Sept 2026 — text": _myt(2026, 9, 15),
        "September 15 2026 text": _myt(2026, 9, 15),
        "2026-09-15 text": _myt(2026, 9, 15),
    }
    for snippet, expected in cases.items():
        c = _candidate("https://news.example/a", snippet=snippet, alert_received_at=received)
        resolve_content_date(c)
        assert (c.date_status, c.content_date) == (DateStatus.APPROXIMATE, expected), f"{snippet!r}: {c.content_date}"
    for snippet in ("The summit was held on 9 September 2026 in KL",      # a date inside the content
                    "Held on 9 Sep 2026, the SC summit",                 # near the start, still content
                    "Sep 30, 2026 upcoming event",                        # after the alert arrived
                    "31 Feb 2026 text", "last week text", ""):
        c = _candidate("https://news.example/a", snippet=snippet, alert_received_at=received)
        resolve_content_date(c)
        assert c.date_status == DateStatus.UNKNOWN and c.content_date is None, f"{snippet!r} gave {c.content_date}"
    c = _candidate("https://news.example/a", snippet="3 days ago text")          # no alert time to count from
    resolve_content_date(c)
    assert c.date_status == DateStatus.UNKNOWN


@selftest("S14")
def test_first_seen_and_period_positions():
    c = _candidate("https://www.sc.com.my/investor-alert-list", source="sc_malaysia",
                   extra={"sc_alert_list": True, "first_seen_period_start": "2026-09-14", "sc_listed_year": "2026"})
    resolve_content_date(c)
    assert (c.date_status, c.date_precision, c.content_date) == \
        (DateStatus.VERIFIED, DatePrecision.YEAR, _myt(2026, 1, 1)), "the listed year, never the week it was found"
    assert "only the year 2026" in c.date_basis and "14 Sep 2026" in c.date_basis
    assert period_position(c, _period("2026-09-22")) == "IN", "reported in the week that first found it"
    assert period_position(c, _period("2026-09-29")) == "BEFORE", "not reported as new in later weeks"
    p = _period("2026-09-25")                                                    # 14-20 Sep 2026
    for when, expected in ((_myt(2026, 9, 14, 0, 0), "IN"), (_myt(2026, 9, 20, 23, 59), "IN"),
                           (_myt(2026, 9, 21, 0, 0), "AFTER"), (_myt(2026, 9, 13, 23, 59), "BEFORE"),
                           (_utc(2026, 9, 13, 16, 0), "IN"), (_utc(2026, 9, 20, 16, 0), "AFTER")):
        c = _candidate("https://news.example/a", source_published_at=when)
        resolve_content_date(c)
        assert period_position(c, p) == expected, f"{when} -> {period_position(c, p)}, expected {expected}"


def _sc_list_candidate(year: str | None, first_seen: str) -> Candidate:
    extra = {"sc_alert_list": True, "first_seen_period_start": first_seen}
    if year is not None:
        extra["sc_listed_year"] = year
    c = _candidate(_SC_ALERT_URL, title="SC Investor Alert List: Example Dragon BTC Mining-Malaysia",
                   source="sc_malaysia", extra=extra)
    c.content_class, c.url_type, c.verification = ContentClass.REGULATORY, UrlType.CONTAINER, Verification.VERIFIED
    c.entity_matches = [Match("acme_company", "Acme", position=0)]
    return c


@selftest("S14")
def test_sc_entry_from_2022_is_historical_but_still_reported():
    p = _period("2026-10-04")                                       # 21-27 Sep 2026, the week it was found
    c = _sc_list_candidate("2022", "2026-09-21")
    resolve_content_date(c)
    assert (c.date_status, c.date_precision, c.content_date) == \
        (DateStatus.VERIFIED, DatePrecision.YEAR, _myt(2022, 1, 1)), "the year SC states is the event date"
    assert "only the year 2022" in c.date_basis and "21 Sep 2026" in c.date_basis
    assert period_position(c, p) == "IN", "listed in the week the system first found it"
    assert period_position(c, _period("2026-10-11")) == "BEFORE", "not reported as new in the following week"
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.MAIN and c.priority == Priority.CRITICAL, \
        "an old SC record is still reported, with the same priority"
    assert historical_status(c.content_date, c.date_precision, p) == HistoricalStatus.HISTORICAL
    assert record_age_text(c.content_date, c.date_precision, p) == "approx. 4 years"
    assert event_date_text(_iso(c.content_date), str(c.date_precision)) == "2022", "never shown as 2022-01-01"
    assert date_status_text(str(c.date_status), str(c.date_precision)) == "VERIFIED (year only)"


@selftest("S14")
def test_year_or_month_only_dates_keep_their_precision():
    p = _period("2026-09-25")                                       # 14-20 Sep 2026
    year_2026, year_2025 = _myt(2026, 1, 1), _myt(2025, 1, 1)
    assert historical_status(year_2026, DatePrecision.YEAR, p) == HistoricalStatus.CURRENT, \
        "a year that includes the reporting week cannot be shown to be older"
    assert historical_status(year_2025, DatePrecision.YEAR, p) == HistoricalStatus.HISTORICAL
    assert record_age_text(year_2026, DatePrecision.YEAR, p) == "", "no age for a CURRENT item"
    assert record_age_text(year_2025, DatePrecision.YEAR, p) == "approx. 1 year"
    assert historical_status(_myt(2026, 9, 1), DatePrecision.MONTH, p) == HistoricalStatus.CURRENT
    assert historical_status(_myt(2026, 8, 1), DatePrecision.MONTH, p) == HistoricalStatus.HISTORICAL
    assert record_age_text(_myt(2026, 8, 1), DatePrecision.MONTH, p) == "approx. 1 month"
    assert record_age_text(_myt(2023, 12, 1), DatePrecision.MONTH, p) == "approx. 3 years", "33 months"
    for iso, precision, shown in ((_iso(year_2025), "YEAR", "2025"), (_iso(_myt(2026, 3, 1)), "MONTH", "2026-03"),
                                  (_iso(_myt(2026, 3, 15)), "DAY", "2026-03-15"), (None, None, "UNKNOWN")):
        assert event_date_text(iso, precision) == shown, (precision, event_date_text(iso, precision))
    # A year / month-only date without a first-seen week is placed only when it lies wholly outside the period.
    c = _candidate("https://news.example/a")
    for when, precision, expected in ((year_2025, DatePrecision.YEAR, "BEFORE"),
                                      (year_2026, DatePrecision.YEAR, "UNKNOWN"),
                                      (_myt(2027, 1, 1), DatePrecision.YEAR, "AFTER"),
                                      (_myt(2026, 8, 1), DatePrecision.MONTH, "BEFORE"),
                                      (_myt(2026, 9, 1), DatePrecision.MONTH, "UNKNOWN")):
        c.content_date, c.date_precision = when, precision
        assert period_position(c, p) == expected, (when, precision, period_position(c, p))
    for text in ("2022", "2022-03", "2022-W10", "2022W10"):       # a year, month or week is not a day
        assert parse_iso_datetime(text) == (None, False), f"{text!r} must not become an exact date"
    assert parse_iso_datetime("2022-W10-1")[0] == _myt(2022, 3, 7), "a week with its day is a full date"


@selftest("S14")
def test_current_and_unknown_dates():
    p = _period("2026-09-25")                                       # 14-20 Sep 2026
    c = _candidate("https://news.example/a", source_published_at=_myt(2026, 9, 16, 9, 30))
    resolve_content_date(c)
    assert (c.date_status, c.date_precision) == (DateStatus.VERIFIED, DatePrecision.DAY)
    assert historical_status(c.content_date, c.date_precision, p) == HistoricalStatus.CURRENT
    assert record_age_text(c.content_date, c.date_precision, p) == ""
    assert historical_status(_myt(2026, 9, 13, 23, 59), DatePrecision.DAY, p) == HistoricalStatus.HISTORICAL
    assert record_age_text(_myt(2026, 9, 8), DatePrecision.DAY, p) == "12 days"
    assert record_age_text(_myt(2026, 3, 21), DatePrecision.DAY, p) == "approx. 5 months"
    assert record_age_text(_myt(2022, 9, 21), DatePrecision.DAY, p) == "approx. 4 years", "47 months"
    c = _candidate("https://facebook.com/SomePage", alert_received_at=_myt(2026, 9, 16, 9, 15))
    resolve_content_date(c)
    assert (c.content_date, c.date_precision, c.date_status) == (None, None, DateStatus.UNKNOWN)
    assert historical_status(c.content_date, c.date_precision, p) == HistoricalStatus.DATE_UNKNOWN
    c = _sc_list_candidate(None, "2026-09-14")                      # an SC entry without even a year
    resolve_content_date(c)
    assert (c.content_date, c.date_status) == (None, DateStatus.FIRST_SEEN), \
        "FIRST_SEEN never puts the week it was found into the event date"
    assert "gives no date" in c.date_basis and period_position(c, p) == "IN"
    assert historical_status(c.content_date, c.date_precision, p) == HistoricalStatus.DATE_UNKNOWN
    assert date_status_text("FIRST_SEEN", None) == "FIRST_SEEN" and date_status_text(None, None) == "UNKNOWN"


# ---- S15 --------------------------------------------------------------------
@selftest("S15")
def test_entity_matching():
    ents = DEFAULT_ENTITIES
    ids = lambda t: {m.entity_id for m in match_entities(t, "", ents)}  # noqa: E731
    assert ids("Acme Asset Management Sdn. Bhd. announced results") == {"acme_company"}
    assert ids("ACME ASSET MANAGEMENT statement") == {"acme_company"}
    assert ids("the acme of good engineering") == set(), "bare 'Acme' matched"
    assert ids("Acme, the fund manager, said") == set(), "D-27: 'Acme' alone is never the company"
    assert ids("Interview with Tan Ah Kow") == {"tan_ah_kow"}
    assert ids("Acmeville summit") == set(), "must match whole words only"
    m = match_entities("Report: Tan Ah Kow interview", "", ents)
    assert m and m[0].position == 1, "position = index of the first matched word"


@selftest("S15")
def test_entity_mentions_positions_and_forms():
    ents = DEFAULT_ENTITIES
    text = ("Acme Asset Management Sdn Bhd issued a statement. " + "filler " * 30 +
            "Later, Acme Asset Management said more.")
    m = [x for x in match_entities(text, "", ents) if x.entity_id == "acme_company"]
    assert [x.position for x in m] == [0, 39], f"every mention, in order: {[x.position for x in m]}"
    assert m[0].matched_text == "Acme Asset Management Sdn Bhd", "the longest overlapping form wins"
    assert [x.matched_text for x in match_entities("Ah Kow Tan spoke", "", ents)] == ["Ah Kow Tan"]
    assert {x.entity_id for x in match_entities("tan-ah-kow and ACME ASSET MANAGEMENT", "", ents)} \
        == {"tan_ah_kow", "acme_company"}
    assert match_entities("Acmeville Asset Managements fund, Tans Ah Kowt", "", ents) == [], \
        "longer words that merely start like a name are not the name"


@selftest("S15")
def test_full_company_name_only():
    """D-27 (owner, 5 Oct 2026): only the full company name (and its approved forms) is screened."""
    ents = DEFAULT_ENTITIES
    ids = lambda t: {m.entity_id for m in match_entities(t, "", ents)}  # noqa: E731
    for text in ["Acme launched a new fund", "Acme unit trust update", "Visit acme.example.com",
                 "Acme Asset", "Asset Management", "Acme Management", "asset management firm fined",
                 "Example Dragon BTC Mining-Malaysia | N/A | https://www.acme-fund.com/",
                 "Acme Asset Holdings Management"]:
        assert ids(text) == set(), f"not the full company name, but matched: {text!r}"
    for text in ["Acme Asset Management", "ACME ASSET MANAGEMENT SDN. BHD.",
                 "acme-asset-management", "艾克米投资管理"]:
        assert ids(text) == {"acme_company"}, f"the full company name was missed: {text!r}"
    ambiguous = [t["term"] for e in ents["entities"] for t in e.get("ambiguous_terms", [])]
    assert ambiguous == ["陈亚九", "陳亞九"], "no single-word terms; only the common Chinese name needs context (D-29)"
    assert entity_screening_terms(ents) == ["Acme Asset Management Sdn Bhd", "Acme Asset Management",
                                            "Acme Asset Management Sdn. Bhd.", "艾克米投资管理", "艾克米投資管理",
                                            "Tan Ah Kow", "Ah Kow Tan", "陈亚九", "陳亞九"]


@selftest("S17")
def test_chinese_names_and_keywords_screen_like_english():
    """Owner's check (5 Oct 2026): 艾克米投资管理 / 陈亚九 and Tan Ah Kow are screened like
    'Acme Asset Management' — full names only — and Chinese keywords work like English ones."""
    def screen(text: str) -> Candidate:
        c = _candidate("https://news.example.com/a", title=text)
        c.content_class, c.url_type, c.verification = ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED
        c.entity_matches = match_entities(text, "", DEFAULT_ENTITIES)
        c.keyword_matches = scan_keywords(text, DEFAULT_KEYWORDS, [m.position for m in c.entity_matches]) \
            if c.entity_matches else []
        apply_flagging_rules(c, 30, approved_relationships(DEFAULT_ENTITIES))
        return c
    ids = lambda t: {m.entity_id for m in screen(t).entity_matches}  # noqa: E731
    # Full names only, in every script — never a part of a name.
    for text in ("Mr Tan said", "Ah Kow said", "Tan Ah said", "Tan AK", "陈先生表示", "亚九", "陈亚",
                 "艾克米", "投资管理公司", "艾", "克米投资管理", "艾克米"):
        assert ids(text) == set(), f"part of a name matched: {text!r}"
    for text, expected in (("Tan Ah Kow said", {"tan_ah_kow"}), ("Ah Kow Tan said", {"tan_ah_kow"}),
                           ("艾克米投资管理发表声明", {"acme_company"}), ("艾克米投資管理發表聲明", {"acme_company"}),
                           ("陈亚九代表艾克米投资管理发言", {"tan_ah_kow", "acme_company"}),
                           ("陳亞九代表艾克米投資管理發言", {"tan_ah_kow", "acme_company"})):
        assert ids(text) == expected, (text, ids(text))
    assert ids("陈亚九被捕") == set() and ids("陳亞九被捕") == set(), "D-29: a common name alone may be someone else"
    # Chinese keywords work like their English terms: same category, priority and ambiguity.
    cases = [("艾克米投资管理涉及洗钱", Priority.HIGH, "R4"), ("艾克米投資管理涉及洗錢", Priority.HIGH, "R4"),
             ("Acme Asset Management involved in money laundering", Priority.HIGH, "R4"),
             ("Tan Ah Kow（陈亚九）被捕", Priority.CRITICAL, "R2"), ("Tan Ah Kow arrested", Priority.CRITICAL, "R2"),
             ("艾克米投资管理被证监会罚款", Priority.HIGH, "R3"), ("Acme Asset Management fined", Priority.HIGH, "R3")]
    for text, priority, rule in cases:
        c = screen(text)
        assert (c.priority, rule in c.rules_applied) == (priority, True), (text, c.priority, c.rules_applied)
    (k,) = screen("艾克米投资管理被证监会罚款").keyword_matches
    assert (k.keyword, k.matched_text, k.ambiguous) == ("fined", "罚款", True), "a Chinese form keeps its English keyword"
    # D-30: a keyword far from the name is REVIEW for every name and script.
    for text in ("Tan Ah Kow " + "word " * 40 + "fraud", "艾克米投资管理" + "的" * 40 + "诈骗"):
        c = screen(text)
        assert (c.priority, c.review_status, c.rules_applied) == (Priority.REVIEW, ReviewStatus.NEEDS_REVIEW, ["R5"]), text


# Test fixture for the ambiguous-term MECHANISM only (D-27: the real configuration has none).
_AMBIGUOUS_TEST_ENTITIES: dict = json.loads(json.dumps(DEFAULT_ENTITIES))
_AMBIGUOUS_TEST_ENTITIES["entities"][0]["ambiguous_terms"] = [
    {"term": "Acme", "requires_any_context": ["asset management", "fund", "unit trust", "investment",
                                                  "Tan Ah Kow", "acme.example.com"],
     "context_window_words": 15}]


@selftest("S15")
def test_ambiguous_term_context_window():
    ents = _AMBIGUOUS_TEST_ENTITIES
    ids = lambda t: {m.entity_id for m in match_entities(t, "", ents)}  # noqa: E731
    near = "Acme " + "word " * 14 + "fund"                   # context 15 words away: inside the window
    far = "Acme " + "word " * 15 + "fund"                    # 16 words away: outside
    assert ids(near) == {"acme_company"} and ids(far) == set()
    assert ids("fund raising at Acme today") == {"acme_company"}, "context may come before the term"
    assert ids("Visit acme.example.com for details") == {"acme_company"}, "the company domain is context"
    (m,) = match_entities("Acme unit trust update", "", ents)
    assert m.matched_text == "Acme (with context 'unit trust')", m.matched_text
    sc_row = "Example Dragon BTC Mining-Malaysia | N/A | https://www.acme-fund.com/ | https://t.me/exampledragonminingma"
    assert ids(sc_row) == {"acme_company"}, \
        "a look-alike 'acme-fund' website on the SC list is still matched, for staff to review (O-12)"
    (m,) = match_entities(sc_row, "", ents)
    assert m.look_alike and LOOK_ALIKE_NOTE in m.matched_text, "…but marked as a look-alike (R9 caps it at REVIEW)"


@selftest("S15")
def test_web_addresses_do_not_make_strong_mentions():
    ents = _AMBIGUOUS_TEST_ENTITIES
    look = lambda t, url="": [m.look_alike for m in match_entities(t, url, ents)]  # noqa: E731
    assert look("Scheme X | acme-fund.com | unlicensed") == [True], "a bare look-alike domain"
    assert look("write to acme@fund.com") == [True], "an email address is a web address too"
    assert look("Acme warns investors http://scam-fund.net/x") == [True], \
        "a context word that only appears in a web address does not confirm the name"
    assert look("nothing here", "https://x.example/acme-fund/abc") == [True], "look-alike in a URL path"
    assert look("Visit acme.example.com for details") == [False], "the company's own domain is a normal mention"
    assert look("email info@acme.example.com now") == [False]
    assert look("Acme launched a new fund") == [False], "plain text is unchanged"
    assert look("clone site acme-asset-management.com") == [False], "full names are never look-alikes"
    assert look("Acme fund news, see https://acme-fund.com") == [False, True], \
        "one plain mention keeps the item a normal finding"
    assert web_address_positions("see https://www.acme-fund.com/ now") == {1, 2, 3, 4, 5}
    assert web_address_positions("Sdn. Bhd. e.g. fund") == set(), "abbreviations are not web addresses"


@selftest("S15")
def test_entity_matching_in_url_and_approval():
    ents = DEFAULT_ENTITIES
    (m,) = match_entities("no name in this text", "https://linkedin.com/company/acme-asset-management", ents)
    assert (m.entity_id, m.position) == ("acme_company", None) and m.matched_text.endswith("(in URL)")
    assert match_entities("", "https://acme.example.com/some-post", ents) == [], "the host is not searched, only the path"
    assert match_entities("Tan Ah Kow", "https://x.example/tan-ah-kow", ents)[0].position == 0, \
        "a text match is preferred over a URL match"
    hidden = json.loads(json.dumps(ents))
    hidden["entities"][1]["approved"] = False
    assert match_entities("Tan Ah Kow", "", hidden) == [], "unapproved entities are never matched"
    assert match_entities("", "", ents) == [] and match_entities("x", "not a url::", ents) == []


# ---- S16 --------------------------------------------------------------------
@selftest("S16")
def test_keyword_boundaries_and_flags():
    kws = DEFAULT_KEYWORDS
    terms = lambda t: {m.keyword for m in scan_keywords(t, kws)}  # noqa: E731
    assert "aml" not in terms("A performance of Hamlet"), "'aml' matched inside a word"
    assert "fined" not in terms("refined oil prices"), "'fined' matched inside a word"
    assert "aml" in terms("AML breach reported")
    assert "money laundering" in terms("alleged Money  Laundering case")
    m = {x.keyword: x for x in scan_keywords("an asset freeze and fees charged", kws)}
    assert m["asset freeze"].priority == "CRITICAL" and not m["asset freeze"].ambiguous
    assert m["charged"].ambiguous, "'charged' must be flagged ambiguous"
    text = "Acme Asset Management was charged yesterday"
    near = {x.keyword: x for x in scan_keywords(text, kws, [0])}
    assert near["charged"].proximity_words is not None and near["charged"].proximity_words <= 4
    assert scan_keywords(text, kws, None)[0].proximity_words is None


@selftest("S16")
def test_keyword_phrases_variants_and_order():
    kws = DEFAULT_KEYWORDS
    terms = lambda t: [m.keyword for m in scan_keywords(t, kws)]  # noqa: E731
    assert terms("A cease-and-desist order; OFAC listed. Then a CYBER-ATTACK.") == \
        ["cease and desist", "ofac", "cyber attack"], "phrases across punctuation, in text order"
    assert terms("sanctioned by the board, a refined fraudulent plan, amlodipine") == [], \
        "inflected forms only count when configured as variants"
    assert terms("fraud, then fraud again") == ["fraud"], "one match per keyword"
    (m,) = scan_keywords("insider trading", kws)
    assert (m.category, m.priority, m.ambiguous, m.entity_id) == ("Market Misconduct", "HIGH", False, "")
    custom = json.loads(json.dumps(kws))
    fraud = next(k for c in custom["categories"] for k in c["keywords"] if k["term"] == "fraud")
    fraud["variants"] = ["fraudulent"]
    (m,) = scan_keywords("a fraudulent plan", custom)
    assert (m.keyword, m.matched_text) == ("fraud", "fraudulent"), "a variant counts as its keyword"


@selftest("S16")
def test_keyword_proximity_is_to_the_nearest_mention():
    kws = DEFAULT_KEYWORDS
    text = "charged " + "x " * 50 + "Acme Asset Management " + "y " * 20 + "was charged"
    (m,) = scan_keywords(text, kws, [51])              # 'charged' at words 0 and 75; name at 51
    assert m.proximity_words == 24, "the nearest occurrence of the keyword counts (75 - 51, not 51 - 0)"
    (m,) = scan_keywords("Tan Ah Kow " + "x " * 40 + "was fined; Tan Ah Kow later", kws, [0, 45])
    assert m.proximity_words == 1, "distance to the nearest name mention ('fined' is word 44, 2nd name at 45)"
    assert scan_keywords("was fined", kws, [])[0].proximity_words is None, "no name position: unknown"


# ---- S17 --------------------------------------------------------------------
def _flag_case(content_class, url_type, verification, entity=True, keyword=None, kw_priority=None,
               ambiguous=False, proximity=None, sc_list=False, date_status=DateStatus.VERIFIED,
               when=None) -> Candidate:
    c = _candidate("https://example.com/x")
    c.content_class, c.url_type, c.verification = content_class, url_type, verification
    if sc_list:
        c.raw.extra["sc_alert_list"] = True
    if entity:
        c.entity_matches = [Match("acme_company", "Acme Asset Management", position=0)]
    if keyword:
        c.keyword_matches = [Match("acme_company", keyword, keyword, "cat", kw_priority,
                                   ambiguous, proximity)]
    c.date_status = date_status
    c.content_date = when if date_status != DateStatus.UNKNOWN else None
    return c


@selftest("S17")
def test_flagging_rules():
    wk = _myt(2026, 9, 17, 12, 0)
    web, spec, ok = ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED
    # R2: unambiguous critical keyword + name -> CRITICAL / MANUAL REVIEW
    c = _flag_case(web, spec, ok, keyword="asset freeze", kw_priority="CRITICAL", proximity=3, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.CRITICAL, ReviewStatus.MANUAL_REVIEW)
    assert "R2" in c.rules_applied
    # R3: ambiguous critical keyword near the name -> HIGH / NEEDS REVIEW
    c = _flag_case(web, spec, ok, keyword="charged", kw_priority="CRITICAL", ambiguous=True,
                   proximity=4, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.HIGH, ReviewStatus.NEEDS_REVIEW), c.priority
    # R5: ambiguous keyword far from the name -> REVIEW / NEEDS REVIEW
    c = _flag_case(web, spec, ok, keyword="charged", kw_priority="CRITICAL", ambiguous=True,
                   proximity=80, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.REVIEW, ReviewStatus.NEEDS_REVIEW), c.priority
    # R4: non-critical keyword -> category priority / PENDING
    c = _flag_case(web, spec, ok, keyword="fraud", kw_priority="HIGH", proximity=2, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.HIGH, ReviewStatus.PENDING)
    c = _flag_case(web, spec, ok, keyword="data breach", kw_priority="MEDIUM", proximity=2, when=wk)
    apply_flagging_rules(c, 30)
    assert c.priority == Priority.MEDIUM
    # highest priority wins
    c = _flag_case(web, spec, ok, keyword="fraud", kw_priority="HIGH", proximity=2, when=wk)
    c.keyword_matches.append(Match("acme_company", "asset freeze", "asset freeze", "cat",
                                   "CRITICAL", False, 5))
    apply_flagging_rules(c, 30)
    assert c.priority == Priority.CRITICAL and c.review_status == ReviewStatus.MANUAL_REVIEW
    # R7: name without any keyword is not reported (no sentiment guessing)
    c = _flag_case(ContentClass.SOCIAL, spec, ok, when=wk)
    apply_flagging_rules(c, 30)
    assert c.priority is None
    # R1: SC Investor Alert List hit
    c = _flag_case(ContentClass.REGULATORY, spec, ok, sc_list=True, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.CRITICAL, ReviewStatus.MANUAL_REVIEW)
    # R6: SC regulatory page naming the company, no keyword
    c = _flag_case(ContentClass.REGULATORY, spec, ok, when=wk)
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.HIGH, ReviewStatus.MANUAL_REVIEW)


@selftest("S17")
def test_report_sections():
    p = _period("2026-09-25")
    wk = _myt(2026, 9, 17, 12, 0)
    c = _flag_case(ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED,
                   keyword="fraud", kw_priority="HIGH", proximity=2, when=wk)
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.MAIN
    c = _flag_case(ContentClass.SOCIAL, UrlType.CONTAINER, Verification.PARTIAL,
                   keyword="fraud", kw_priority="HIGH", proximity=2, date_status=DateStatus.UNKNOWN)
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.DATE_UNKNOWN
    c = _flag_case(ContentClass.SOCIAL, UrlType.NON_EVIDENCE, Verification.REJECTED,
                   keyword="fraud", kw_priority="HIGH", proximity=2, when=wk)
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.UNVERIFIED_URL
    c = _flag_case(ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED,
                   keyword="fraud", kw_priority="HIGH", proximity=2,
                   when=_myt(2026, 9, 24, 12, 0))
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.EXCLUDED_FUTURE, "current-week item leaked into report"
    c = _flag_case(ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED,
                   entity=False, keyword="fraud", kw_priority="HIGH", proximity=2, when=wk)
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.NOT_REPORTED, "keyword without a name must not be reported"


def _rules(keywords: list[tuple[str, str, bool, int | None]], content_class=ContentClass.PUBLIC_WEB,
           sc_list=False, window=30) -> Candidate:
    """Candidate with a name match and keyword matches [(keyword, priority, ambiguous, proximity)]."""
    c = _flag_case(content_class, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED, sc_list=sc_list,
                   when=_myt(2026, 9, 17, 12))
    c.keyword_matches = [Match("", k, k, "cat", pr, amb, prox) for k, pr, amb, prox in keywords]
    apply_flagging_rules(c, window)
    return c


@selftest("S17")
def test_flagging_rule_details():
    P, R = Priority, ReviewStatus
    c = _rules([("charged", "CRITICAL", True, 30)])
    assert (c.priority, c.review_status, c.rules_applied) == (P.HIGH, R.NEEDS_REVIEW, ["R3"]), "30 words = near"
    c = _rules([("charged", "CRITICAL", True, 31)])
    assert (c.priority, c.rules_applied) == (P.REVIEW, ["R5"]), "31 words = not near"
    c = _rules([("charged", "CRITICAL", True, None)])
    assert c.rules_applied == ["R3"], "an unmeasurable distance counts as near"
    c = _rules([("aml", "HIGH", True, 1)])
    assert (c.priority, c.review_status, c.rules_applied) == (P.REVIEW, R.NEEDS_REVIEW, ["R5"]), \
        "an ambiguous keyword from another category is REVIEW even when near"
    c = _rules([("fraud", "HIGH", False, 2), ("phishing", "MEDIUM", True, 2)])
    assert (c.priority, c.review_status, c.rules_applied) == (P.HIGH, R.NEEDS_REVIEW, ["R4", "R5"]), \
        "highest priority wins; review status is the strongest that fired"
    c = _rules([("data breach", "MEDIUM", False, 200)])
    assert (c.priority, c.review_status, c.rules_applied) == (P.REVIEW, R.NEEDS_REVIEW, ["R5"]), \
        "D-30: an unambiguous keyword far from the name is REVIEW, not its category's priority"
    c = _rules([("asset freeze", "CRITICAL", False, 31)])
    assert (c.priority, c.review_status, c.rules_applied) == (P.REVIEW, R.NEEDS_REVIEW, ["R5"]), \
        "D-30: a critical keyword far from the name is not CRITICAL"
    c = _rules([("asset freeze", "CRITICAL", False, 30)])
    assert (c.priority, c.rules_applied) == (P.CRITICAL, ["R2"]), "30 words = near"
    c = _rules([("fraud", "HIGH", False, 2)], content_class=ContentClass.REGULATORY, sc_list=True)
    assert (c.priority, c.review_status, c.rules_applied) == (P.CRITICAL, R.MANUAL_REVIEW, ["R1", "R4"])
    c = _rules([("sanctions", "CRITICAL", True, 90)], content_class=ContentClass.REGULATORY)
    assert c.rules_applied == ["R5"], "R6 applies only when there is no keyword (as documented)"
    c = _rules([], content_class=ContentClass.SOCIAL)
    assert (c.priority, c.rules_applied) == (None, ["R7"]), "a name alone is not reported (no sentiment guessing)"
    orphan = _flag_case(ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED, entity=False,
                        keyword="asset freeze", kw_priority="CRITICAL", when=_myt(2026, 9, 17))
    apply_flagging_rules(orphan, 30)
    assert (orphan.priority, orphan.rules_applied) == (None, ["R7"]), "a keyword without a name is R7"
    for case in (c, _rules([("asset freeze", "CRITICAL", False, 1)])):
        before = (case.priority, case.review_status, list(case.rules_applied))
        apply_flagging_rules(case, 30)
        assert (case.priority, case.review_status, case.rules_applied) == before, "re-applying is stable"
        assert case.review_status not in (R.CONFIRMED_ADVERSE, R.FALSE_POSITIVE)


@selftest("S17")
def test_section_order_and_review():
    p = _period("2026-09-25")                                       # 14-20 Sep 2026
    c = _flag_case(ContentClass.SOCIAL, UrlType.NON_EVIDENCE, Verification.REJECTED,
                   keyword="fraud", kw_priority="HIGH", proximity=2, date_status=DateStatus.UNKNOWN)
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.UNVERIFIED_URL, "a rejected URL goes to manual lookup even undated"
    c = _flag_case(ContentClass.SOCIAL, UrlType.CONTAINER, Verification.PARTIAL,
                   keyword="fraud", kw_priority="HIGH", proximity=2, date_status=DateStatus.UNKNOWN)
    apply_flagging_rules(c, 30)
    assert c.review_status == ReviewStatus.PENDING
    assert assign_section(c, p) == Section.DATE_UNKNOWN and c.review_status == ReviewStatus.MANUAL_REVIEW
    c = _flag_case(ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.REJECTED,
                   keyword="fraud", kw_priority="HIGH", proximity=2, when=_myt(2026, 9, 2))
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.OUT_OF_PERIOD, "dated before the period, even if rejected"
    c = _flag_case(ContentClass.REGULATORY, UrlType.CONTAINER, Verification.VERIFIED, sc_list=True,
                   date_status=DateStatus.FIRST_SEEN)
    c.content_date, c.raw.extra["first_seen_period_start"] = None, "2026-09-14"
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.MAIN, "a first-seen alert-list hit belongs to that week"
    c = _flag_case(ContentClass.SOCIAL, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED, when=_myt(2026, 9, 17))
    apply_flagging_rules(c, 30)
    assert assign_section(c, p) == Section.NOT_REPORTED, "name without keyword"


# ---- S18 --------------------------------------------------------------------
@selftest("S18")
def test_report_path_uses_period_start():
    p = _period("2026-09-25")
    assert report_path_for(p, Path("reports"), False) == Path("reports/Negative_Screening_2026-09-14.xlsx")
    assert report_path_for(p, Path("reports"), True).parent == Path("reports/dry_run")


@selftest("S18")
def test_excel_sheets_and_summary_metadata():
    try:
        from importlib import import_module
        openpyxl = import_module("openpyxl")
    except ImportError:
        return  # Excel report tests require the optional openpyxl dependency.
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        ctx.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
        db_record_run_start(conn, ctx)
        health = [{"source": "Public website", "status": "SUCCESS", "retrieved": 1, "filtered_out": 0,
                   "flagged": 1, "duplicates": 0, "rejected": 0, "error": None},
                  {"source": "SC Malaysia", "status": "FAILED", "retrieved": 0, "filtered_out": 0,
                   "flagged": 0, "duplicates": 0, "rejected": 0, "error": "timeout"}]
        path = generate_excel_report(conn, ctx, health, {"discovered": 0})
        wb = openpyxl.load_workbook(path)
        for sheet in (SHEET_SUMMARY, SHEET_FINDINGS, SHEET_DATE_UNKNOWN, SHEET_UNVERIFIED):
            assert sheet in wb.sheetnames, f"missing sheet {sheet}"
        summary = " ".join(str(v) for row in wb[SHEET_SUMMARY].iter_rows(values_only=True)
                           for v in row if v is not None)
        for needle in ("Execution Mode", "Reference Date", "Reporting Period", "FAILED",
                       "not a complete record"):
            assert needle in summary, f"summary missing {needle!r}"
        header = [c.value for c in wb[SHEET_FINDINGS][1]]
        assert header == FINDINGS_COLUMNS, header
        conn.close()


def _report_candidate(url_key: str, url: str, section: Section, priority: Priority | None,
                      title: str = "t", snippet: str = "s", when: datetime | None = None,
                      date_status: DateStatus = DateStatus.VERIFIED) -> Candidate:
    c = _db_candidate(url_key, {"q1"})
    c.raw.discovered_url, c.canonical_url = url, url
    c.raw.title, c.raw.snippet = title, snippet
    c.section, c.priority = section, priority
    c.content_date, c.date_status = when, date_status
    c.date_precision = DatePrecision.DAY if when is not None else None
    return c


def _report_workbook(d: str, candidates: list[Candidate], out_of_period_sheet: bool = False,
                     dry_run: bool = False):
    """Store the candidates in a fresh database, build the report, return (path, workbook, ctx, conn)."""
    from openpyxl import load_workbook
    db = Path(d) / "t.sqlite"
    conn = db_connect(db)
    ctx = _test_ctx(db)
    ctx.dry_run = dry_run
    ctx.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
    ctx.config.settings["report"]["include_out_of_period_sheet"] = out_of_period_sheet
    try:
        db_record_run_start(conn, ctx)
        for c in candidates:
            db_upsert_candidate(conn, c, ctx)
        path = generate_excel_report(conn, ctx, [], {"discovered": len(candidates)})
        return path, load_workbook(path), ctx, conn
    except BaseException:
        conn.close()
        raise


def _sheet_values(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:] if r[0] is not None and r[1] is not None]


@selftest("S18")
def test_report_url_prefers_the_verified_final_url():
    page = "https://news.example.com/2026/09/16/fraud-story"
    for final, verification, basis, expected in [
        ("https://www.news.example.com/fraud-story-123", "VERIFIED", "page text match",
         "https://www.news.example.com/fraud-story-123"),                  # read + matched: the page's own URL
        ("https://news.example.com/", "VERIFIED", "page text match", page),  # redirected to the home page
        ("https://news.example.com/2026/09", "VERIFIED", "page text match", page),   # up to a parent section
        ("https://news.example.com/other", "PARTIAL", "page read, but the matched text could not be re-located",
         page),                                                             # page not confirmed
        ("https://news.example.com/other", "VERIFIED", "search-index snippet", page),
        ("https://www.facebook.com/x/posts/1", "VERIFIED", "page text match", page),  # never a social final URL
        (None, "VERIFIED", "page text match", page),
    ]:
        assert report_url(page, final, verification, basis) == expected, (final, verification, basis)
    with tempfile.TemporaryDirectory() as d:
        c = _report_candidate("k-f", page, Section.MAIN, Priority.HIGH, when=_myt(2026, 9, 16))
        c.raw.final_url = "https://news.example.com/fraud-story-123"
        c.verification, c.verification_basis = Verification.VERIFIED, "page text match"
        _, wb, _, conn = _report_workbook(d, [c])
        try:
            row = _sheet_values(wb[SHEET_FINDINGS])[0]
            assert row["URL"] == "https://news.example.com/fraud-story-123"
            assert row["Discovered URL"] == page, "the discovered URL stays visible in its own column"
            assert row["Fetched URL"] == "https://news.example.com/fraud-story-123", "the page the program read"
            assert "Discovered as" not in row["Notes"]
            assert row["URL Verification"] == "VERIFIED" and "URL: page text match" in row["Notes"]
        finally:
            conn.close()


@selftest("S18")
def test_report_sections_are_routed_from_sqlite():
    with tempfile.TemporaryDirectory() as d:
        cands = [
            _report_candidate("k-main", "https://news.example.com/main", Section.MAIN, Priority.HIGH,
                              when=_myt(2026, 9, 16, 9, 30)),
            _report_candidate("k-crit", "https://news.example.com/crit", Section.MAIN, Priority.CRITICAL,
                              when=_myt(2026, 9, 17)),
            _report_candidate("k-unk", "https://news.example.com/unk", Section.DATE_UNKNOWN, Priority.HIGH,
                              date_status=DateStatus.UNKNOWN),
            _report_candidate("k-rej", "https://www.facebook.com/hashtag/x", Section.UNVERIFIED_URL,
                              Priority.HIGH, when=_myt(2026, 9, 16)),
            _report_candidate("k-old", "https://news.example.com/old", Section.OUT_OF_PERIOD, Priority.HIGH,
                              when=_myt(2026, 9, 1)),
            _report_candidate("k-none", "https://news.example.com/none", Section.NOT_REPORTED, None),
            _report_candidate("k-fut", "https://news.example.com/fut", Section.EXCLUDED_FUTURE, Priority.HIGH,
                              when=_myt(2026, 9, 22)),
        ]
        path, wb, ctx, conn = _report_workbook(d, cands)
        try:
            assert [r["URL"] for r in _sheet_values(wb[SHEET_OUT_OF_PERIOD])] == ["https://news.example.com/old"], \
                "an item found this week but dated earlier, never reported before, is never hidden (D-24)"
            main = _sheet_values(wb[SHEET_FINDINGS])
            assert [r["URL"] for r in main] == ["https://news.example.com/crit", "https://news.example.com/main"], \
                "main sheet: in-period items only, CRITICAL first"
            assert main[1]["Event / Content Date"] == "2026-09-16 09:30"
            assert main[0]["Event / Content Date"] == "2026-09-17"
            assert main[0]["Historical Status"] == "CURRENT" and not main[0]["Record Age"]
            assert main[0]["Matched Entity"] == "Acme Asset Management Sdn Bhd", "entity shown by display name"
            assert main[0]["Keyword"] == "fraud" and main[0]["Keyword Category"] == "Financial Crime"
            assert main[0]["Report Week"] == "2026-09-14 – 2026-09-20"
            assert "Rules: R4" in main[0]["Notes"] and "q1" in main[0]["Notes"]
            assert "1 query" in main[0]["Discovery Method"]
            unknown = _sheet_values(wb[SHEET_DATE_UNKNOWN])
            assert [r["URL"] for r in unknown] == ["https://news.example.com/unk"]
            assert unknown[0]["Event / Content Date"] == "UNKNOWN" and unknown[0]["Date Status"] == "UNKNOWN"
            assert unknown[0]["Historical Status"] == "DATE UNKNOWN"
            assert [r["URL"] for r in _sheet_values(wb[SHEET_UNVERIFIED])] == ["https://www.facebook.com/hashtag/x"]
            all_urls = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row)
            assert "none" not in all_urls and "fut" not in all_urls, "not-reported / future items stay out"
            ws = wb[SHEET_FINDINGS]
            assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")
            url_col = FINDINGS_COLUMNS.index("URL") + 1
            hyperlink = ws.cell(row=2, column=url_col).hyperlink
            assert hyperlink is not None and hyperlink.target == "https://news.example.com/crit"
            assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith("F8CBAD"), "CRITICAL row highlighted"
            assert not ws.cell(row=3, column=1).fill.fgColor.rgb.endswith("F8CBAD"), "HIGH row not highlighted"
        finally:
            conn.close()
        path2, wb2, _, conn = _report_workbook(str(Path(d) / "b"), cands, out_of_period_sheet=True)
        try:
            assert [r["URL"] for r in _sheet_values(wb2[SHEET_OUT_OF_PERIOD])] == ["https://news.example.com/old"]
        finally:
            conn.close()


@selftest("S18")
def test_item_already_reported_in_an_earlier_week_is_not_repeated():
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        try:
            first = _test_ctx(db, "2026-09-15")                     # week 7–13 Sep: first found, listed
            first.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
            db_record_run_start(conn, first)
            sc = _report_candidate("k-sc", _SC_ALERT_URL, Section.MAIN, Priority.REVIEW,
                                   title="SC Investor Alert List: Example Scheme")
            sc.platform, sc.content_class = Platform.SC_MALAYSIA, ContentClass.REGULATORY
            fid = db_upsert_candidate(conn, sc, first)
            later = _test_ctx(db, "2026-09-22")                     # week 14–20 Sep: found again
            later.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
            db_record_run_start(conn, later)
            sc.section = Section.OUT_OF_PERIOD
            db_upsert_candidate(conn, sc, later)
            assert db_first_reported_period(conn, fid, later.period) == date(2026, 9, 7)
            assert db_first_reported_period(conn, fid, first.period) is None, "only EARLIER weeks count"
            health = [{"source": "SC Malaysia", "status": "SUCCESS", "retrieved": 1, "flagged": 1}]
            from openpyxl import load_workbook
            wb = load_workbook(generate_excel_report(conn, later, health, {"manual_review": 0}))
            assert SHEET_OUT_OF_PERIOD not in wb.sheetnames, "nothing new out of period: no sheet"
            summary = " ".join(str(c.value) for row in wb[SHEET_SUMMARY].iter_rows() for c in row if c.value)
            assert "already reported in the report for 7–13 Sep 2026" in summary, summary
            later.config.settings["report"]["include_out_of_period_sheet"] = True     # audit sheet on
            wb = load_workbook(generate_excel_report(conn, later, health, {"manual_review": 0}))
            (row,) = _sheet_values(wb[SHEET_OUT_OF_PERIOD])
            assert row["Notes"].startswith("Already reported in the report for 7–13 Sep 2026")
        finally:
            conn.close()


@selftest("S17")
def test_look_alike_only_items_are_review_priority():
    web, spec, ok = ContentClass.PUBLIC_WEB, UrlType.SPECIFIC_CONTENT, Verification.VERIFIED
    c = _flag_case(ContentClass.REGULATORY, UrlType.CONTAINER, ok, sc_list=True)
    c.entity_matches = [Match("acme_company", f"Acme ({LOOK_ALIKE_NOTE})", position=3, look_alike=True)]
    apply_flagging_rules(c, 30)
    assert (c.priority, c.review_status) == (Priority.REVIEW, ReviewStatus.NEEDS_REVIEW), (c.priority, c.review_status)
    assert {"R1", "R9"} <= set(c.rules_applied), "R1 still recorded; R9 caps it"
    c = _flag_case(web, spec, ok, keyword="fraud", kw_priority="LOW", proximity=1)
    c.entity_matches[0].look_alike = True
    apply_flagging_rules(c, 30)
    assert c.priority == Priority.LOW, "R9 only lowers a priority, never raises one"
    c = _flag_case(web, spec, ok)                       # look-alike name, no keyword: R7, not reported
    c.entity_matches[0].look_alike = True
    apply_flagging_rules(c, 30)
    assert c.priority is None and c.rules_applied == ["R7"], "R9 never makes an item reported by itself"
    c = _flag_case(ContentClass.REGULATORY, UrlType.CONTAINER, ok, sc_list=True)
    c.entity_matches.append(Match("acme_company", "Acme (web)", position=9, look_alike=True))
    apply_flagging_rules(c, 30)
    assert c.priority == Priority.CRITICAL and "R9" not in c.rules_applied, "a real mention keeps full priority"


@selftest("S08")
def test_no_alert_emails_shows_not_checked():
    warning = NO_ALERT_EMAILS.format(label="NS-Alerts", since="27 Sep 2026", mailbox="m@example.com")
    quiet = [SourceResult("google_alerts", SourceStatus.SUCCESS, warnings=[warning]),
             SourceResult("website", SourceStatus.SUCCESS), SourceResult("sc_malaysia", SourceStatus.SUCCESS)]
    rows = {h["source"]: h for h in derive_source_health(quiet, [])}
    for name in ("Google Alerts", "Facebook", "Instagram", "LinkedIn", "Public website"):
        assert rows[name]["status"] == "NOT CHECKED", (name, rows[name]["status"])
    assert rows["Facebook"]["error"].startswith("No Google Alerts emails arrived")
    assert rows["Company website"]["status"] == rows["SC Malaysia"]["status"] == "SUCCESS"
    subject = build_email_subject(DEFAULT_EMAIL["subject_template"], _period("2026-10-05"), list(rows.values()),
                                  {"critical": 0, "manual_review": 0})
    assert subject.endswith("— social media & news NOT CHECKED (no Google Alerts emails)"), subject
    assert friendly_source_note("NOT CHECKED", rows["Facebook"]["error"]).startswith("Not checked")
    normal = [SourceResult("google_alerts", SourceStatus.SUCCESS), *quiet[1:]]
    rows = {h["source"]: h for h in derive_source_health(normal, [])}
    assert rows["Google Alerts"]["status"] == rows["Facebook"]["status"] == "SUCCESS" and rows["Facebook"]["error"] is None
    assert "NOT CHECKED" not in build_email_subject(DEFAULT_EMAIL["subject_template"], _period("2026-10-05"),
                                                    list(rows.values()), {})


@selftest("S18")
def test_report_separates_event_date_from_discovery_date():
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)                                         # reporting week 14-20 Sep 2026
        ctx.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
        try:
            db_record_run_start(conn, ctx)
            sc = _report_candidate("k-sc", _SC_ALERT_URL, Section.MAIN, Priority.CRITICAL,
                                   title="SC Investor Alert List: Example Dragon BTC Mining-Malaysia")
            sc.platform, sc.content_class = Platform.SC_MALAYSIA, ContentClass.REGULATORY
            sc.content_date, sc.date_precision, sc.date_status = _myt(2022, 1, 1), DatePrecision.YEAR, DateStatus.VERIFIED
            sc.date_basis = "the SC list gives only the year 2022 (no day or month); listed in the week ..."
            news = _report_candidate("k-news", "https://news.example.com/a", Section.MAIN, Priority.HIGH,
                                     when=_myt(2026, 9, 16))
            news.raw.alert_received_at = _myt(2026, 9, 19, 8, 0)   # the alert email is not the content date
            for c in (sc, news):
                db_upsert_candidate(conn, c, ctx)
            with conn:                                              # the runs that first found them
                conn.execute("UPDATE findings SET first_discovered_at = '2026-09-23T10:05:00+08:00'"
                             " WHERE url_key = 'k-sc'")
                conn.execute("UPDATE findings SET first_discovered_at = '2026-09-21T09:00:00+08:00'"
                             " WHERE url_key = 'k-news'")
            db_upsert_candidate(conn, sc, ctx)                      # a re-run never moves the discovery date
            path = generate_excel_report(conn, ctx, [], {"historical": 1})
            wb = load_workbook(path)
            rows = {r["URL"]: r for r in _sheet_values(wb[SHEET_FINDINGS])}
            old, new = rows[_SC_ALERT_URL], rows["https://news.example.com/a"]
            assert (old["Report Week"], old["Event / Content Date"], old["Date Status"]) == \
                ("2026-09-14 – 2026-09-20", "2022", "VERIFIED (year only)"), "still reported, never as 2026"
            assert (old["Historical Status"], old["Record Age"], old["Discovery Date"]) == \
                ("HISTORICAL", "approx. 4 years", "2026-09-23 10:05"), old
            assert old["Priority"] == "CRITICAL", "an old SC record keeps its priority"
            assert (new["Event / Content Date"], new["Historical Status"], new["Discovery Date"]) == \
                ("2026-09-16", "CURRENT", "2026-09-21 09:00"), new
            ws = wb[SHEET_FINDINGS]
            hist_col = FINDINGS_COLUMNS.index("Historical Status") + 1
            row_of = {ws.cell(row=r, column=FINDINGS_COLUMNS.index("URL") + 1).value: r for r in (2, 3)}
            assert ws.cell(row=row_of[_SC_ALERT_URL], column=hist_col).fill.fgColor.rgb.endswith("DDEBF7"), \
                "a HISTORICAL cell stands out, even on a CRITICAL row"
            summary = " ".join(str(v) for row in wb[SHEET_SUMMARY].iter_rows(values_only=True)
                               for v in row if v is not None)
            for needle in ("Dates in this report", "1 CURRENT · 1 HISTORICAL · 0 DATE UNKNOWN",
                           "oldest Event / Content Date: 2022", "NOT new incidents in this reporting week",
                           "Never the date the system found the item", "Not the event date",
                           "Reported items dated before the period (HISTORICAL"):
                assert needle in summary, f"summary missing {needle!r}"
        finally:
            conn.close()


@selftest("S18")
def test_report_text_is_never_a_formula_and_is_cleaned():
    with tempfile.TemporaryDirectory() as d:
        c = _report_candidate("k1", "https://news.example.com/a", Section.MAIN, Priority.HIGH,
                              title="=HYPERLINK(\"http://evil.example\")", snippet="bad\x07char",
                              when=_myt(2026, 9, 16))
        _, wb, _, conn = _report_workbook(d, [c])
        try:
            cell = wb[SHEET_FINDINGS].cell(row=2, column=FINDINGS_COLUMNS.index("Title / Comment / Snippet") + 1)
            assert cell.data_type == "s" and isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"), (cell.data_type, cell.value)
            assert "\x07" not in cell.value and "badchar" in cell.value
            assert len(_excel_text("x" * 40000)) == EXCEL_CELL_LIMIT
        finally:
            conn.close()


@selftest("S18")
def test_report_rerun_archives_previous_and_dry_run_is_separate():
    with tempfile.TemporaryDirectory() as d:
        c = _report_candidate("k1", "https://news.example.com/a", Section.MAIN, Priority.HIGH,
                              when=_myt(2026, 9, 16))
        first, _, ctx, conn = _report_workbook(d, [c])
        try:
            assert first.name == "Negative_Screening_2026-09-14.xlsx"
            second = generate_excel_report(conn, ctx, [], {})
            assert second == first and second.exists()
            archived = list((first.parent / "archive").glob("Negative_Screening_2026-09-14_*.xlsx"))
            assert len(archived) == 1, archived
            third = generate_excel_report(conn, ctx, [], {})
            assert len(list((first.parent / "archive").glob("*.xlsx"))) == 2, "every earlier copy is kept"
            ctx.dry_run = True
            dry = generate_excel_report(conn, ctx, [], {})
            assert dry.parent.name == "dry_run" and dry.exists() and third.exists()
            from openpyxl import load_workbook
            summary = " ".join(str(v) for row in load_workbook(dry)[SHEET_SUMMARY].iter_rows(values_only=True)
                               for v in row if v is not None)
            assert "Dry Run" in summary and "Manual" in summary
        finally:
            conn.close()


@selftest("S18")
def test_summary_source_health_and_counts():
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import load_workbook
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        ctx.mode = Mode.MANUAL_TEST_DATE
        ctx.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
        db_record_run_start(conn, ctx)
        try:
            health = [{"source": "Google Alerts", "status": "SUCCESS", "retrieved": 4, "error": None},
                      {"source": "Company website", "status": "SKIPPED", "error": "not run"},
                      {"source": "SC Malaysia", "status": "FAILED", "error": "timeout"}]
            path = generate_excel_report(conn, ctx, health, {"discovered": 7, "critical": 2})
            rows = list(load_workbook(path)[SHEET_SUMMARY].iter_rows(values_only=True))
            by_label = {r[0]: r[1] for r in rows if r[0]}
            assert by_label["Execution Mode"] == "Manual (test reference date)"
            assert by_label["Sources attempted"] == 2 and by_label["Sources successful"] == 1
            assert by_label["Sources failed"] == 1
            assert by_label["Total discovered candidates"] == 7 and by_label["Critical items"] == 2
            assert by_label["Keywords"] == DEFAULT_KEYWORDS["version"]
            sc = next(r for r in rows if r[0] == "SC Malaysia")
            assert sc[1] == "FAILED" and "timeout" in str(sc[7]) and "does not mean" in str(sc[7])
            ok = next(r for r in rows if r[0] == "Google Alerts")
            assert ok[2] == 4 and not ok[7]
            health = [{"source": "Google Alerts", "status": "FAILED", "error": "LOGIN failed"}] + [
                {"source": str(p), "status": "FAILED", "error": "Google Alerts unavailable"}
                for p in GOOGLE_ALERTS_PLATFORMS]
            path = generate_excel_report(conn, ctx, health, {})
            rows = list(load_workbook(path)[SHEET_SUMMARY].iter_rows(values_only=True))
            by_label = {r[0]: r[1] for r in rows if r[0]}
            assert by_label["Sources failed"] == 1, "platform rows repeat Google Alerts: count it once"
            assert len([r for r in rows if r[1] == "FAILED"]) == 5, "every health row is still shown"
        finally:
            conn.close()


# ---- S19 --------------------------------------------------------------------
@selftest("S19")
def test_email_decision_table():
    dflt = {"send_on_automatic": True, "send_on_manual": False}
    assert decide_email(mode=Mode.AUTOMATIC, dry_run=False, no_email=False,
                        send_email=False, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is True
    assert decide_email(mode=Mode.MANUAL, dry_run=False, no_email=False,
                        send_email=False, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is False
    assert decide_email(mode=Mode.MANUAL, dry_run=False, no_email=False,
                        send_email=True, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is True
    assert decide_email(mode=Mode.MANUAL, dry_run=False, no_email=False,
                        send_email=True, resend_email=False, already_sent=True,
                        email_defaults=dflt)[0] is False
    assert decide_email(mode=Mode.AUTOMATIC, dry_run=False, no_email=False,
                        send_email=False, resend_email=False, already_sent=True,
                        email_defaults=dflt)[0] is False
    assert decide_email(mode=Mode.MANUAL, dry_run=False, no_email=False,
                        send_email=False, resend_email=True, already_sent=True,
                        email_defaults=dflt)[0] is True
    assert decide_email(mode=Mode.MANUAL_TEST_DATE, dry_run=False, no_email=False,
                        send_email=True, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is False
    assert decide_email(mode=Mode.AUTOMATIC, dry_run=True, no_email=False,
                        send_email=False, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is False
    assert decide_email(mode=Mode.AUTOMATIC, dry_run=False, no_email=True,
                        send_email=False, resend_email=False, already_sent=False,
                        email_defaults=dflt)[0] is False


@selftest("S19")
def test_email_decision_reasons_and_config_switches():
    def decide(**kw):
        mode: Mode = kw.get("mode", Mode.MANUAL)
        dry_run: bool = kw.get("dry_run", False)
        no_email: bool = kw.get("no_email", False)
        send_email: bool = kw.get("send_email", False)
        resend_email: bool = kw.get("resend_email", False)
        already_sent: bool = kw.get("already_sent", False)
        email_defaults: dict = kw.get(
            "email_defaults",
            {"send_on_automatic": True, "send_on_manual": False},
        )
        validation: bool = kw.get("validation", False)
        minutes: float | None = kw.get("minutes_since_last_email")
        return decide_email(
            mode=mode, dry_run=dry_run, no_email=no_email, send_email=send_email,
            resend_email=resend_email, already_sent=already_sent,
            email_defaults=email_defaults, validation=validation, minutes_since_last_email=minutes,
        )
    assert decide() == (False, "manual mode: send_on_manual is false")
    assert decide(email_defaults=DEFAULT_SETTINGS["email_defaults"]) == (True, "manual mode default"),         "staff 'Run Weekly Screening' emails the report by default"
    assert decide(mode=Mode.AUTOMATIC, email_defaults=DEFAULT_SETTINGS["email_defaults"]) ==         (True, "automatic mode default")
    assert decide(already_sent=True, mode=Mode.AUTOMATIC) == (False, "already sent for period")
    assert decide(email_defaults=DEFAULT_SETTINGS["email_defaults"], already_sent=True,
                  minutes_since_last_email=2)[0] is False, "a manual re-run inside the cooldown is not emailed"
    assert decide(email_defaults=DEFAULT_SETTINGS["email_defaults"], already_sent=True,
                  minutes_since_last_email=30) == (True, "manual mode default"), "after the cooldown it is"
    assert decide(email_defaults=DEFAULT_SETTINGS["email_defaults"], no_email=True) == (False, "--no-email")
    assert decide(email_defaults=DEFAULT_SETTINGS["email_defaults"], validation=True, resend_email=True)[0]         is False, "a validation run never emails"
    assert decide(email_defaults={"send_on_automatic": False, "send_on_manual": True})[0] is True
    assert decide(mode=Mode.AUTOMATIC, email_defaults={"send_on_automatic": False,
                                                       "send_on_manual": True})[0] is False
    assert decide(dry_run=True, resend_email=True)[0] is False, "a dry run never emails"
    assert decide(mode=Mode.MANUAL_TEST_DATE, resend_email=True)[0] is False, "--date never emails"


_TEST_HEALTH = [
    {"source": "Google Alerts", "status": "FAILED", "retrieved": 0, "flagged": 0, "error": "LOGIN failed"},
    {"source": "Facebook", "status": "FAILED", "error": "Google Alerts unavailable"},
    {"source": "Instagram", "status": "FAILED", "error": "Google Alerts unavailable"},
    {"source": "LinkedIn", "status": "FAILED", "error": "Google Alerts unavailable"},
    {"source": "Public website", "status": "FAILED", "error": "Google Alerts unavailable"},
    {"source": "Company website", "status": "SUCCESS", "retrieved": 3, "flagged": 0, "error": None},
    {"source": "SC Malaysia", "status": "SUCCESS", "retrieved": 2, "flagged": 1, "error": None},
]


@selftest("S19")
def test_email_subject_and_body():
    p = _period("2026-09-25")
    template = DEFAULT_EMAIL["subject_template"]
    counts = {"critical": 2, "manual_review": 5, "discovered": 9}
    ok = [h for h in _TEST_HEALTH if h["source"] in ("Company website", "SC Malaysia")]
    assert build_email_subject(template, p, ok, counts) == \
        "[Negative Screening] 14–20 Sep 2026 — 2 critical, 5 for review"
    assert build_email_subject(template, p, _TEST_HEALTH, counts).endswith("— 1 source failed"), \
        "one mailbox failure is one failed source, not five"
    two = _TEST_HEALTH + [{"source": "SC Malaysia", "status": "FAILED"}]
    assert build_email_subject(template, p, two[:5] + two[7:], counts).endswith("— 2 sources failed")
    assert email_period_text(_period("2026-10-05")) == "28 Sep – 4 Oct 2026"
    assert email_period_text(_period("2027-01-05")) == "28 Dec 2026 – 3 Jan 2027"
    ctx = _test_ctx(Path("unused.sqlite"), "2026-09-25")
    body = build_email_body(ctx, _TEST_HEALTH, counts, "Negative_Screening_2026-09-14.xlsx")
    for needle in ("14 Sep 2026", "20 Sep 2026", "Manual", "Google Alerts: FAILED", "LOGIN failed",
                   "does not mean", "Critical items: 2", "Total discovered candidates: 9",
                   "not a complete record", "Negative_Screening_2026-09-14.xlsx"):
        assert needle in body, f"body missing {needle!r}"


class _FakeSmtp:
    """Stands in for smtplib.SMTP; records what the program did."""
    calls: list = []
    fail_with: BaseException | None = None
    refuse: dict = {}

    def __init__(self, host, port, timeout=None):
        _FakeSmtp.calls.append(("connect", host, port))
        if _FakeSmtp.fail_with is not None:
            raise _FakeSmtp.fail_with

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        _FakeSmtp.calls.append(("quit",))
        return False

    def starttls(self, context=None):
        _FakeSmtp.calls.append(("starttls",))

    def login(self, user, password):
        _FakeSmtp.calls.append(("login", user, password))

    def send_message(self, msg, from_addr=None, to_addrs=None):
        _FakeSmtp.calls.append(("send", msg, from_addr, list(to_addrs or [])))
        return dict(_FakeSmtp.refuse)


def _send_with_fake_smtp(d: str, password: str | None, fail_with=None, refuse=None, recipients=None):
    """Run send_report_email() against _FakeSmtp; return (result, calls, log text)."""
    import io
    import smtplib
    report = Path(d) / "Negative_Screening_2026-09-14.xlsx"
    report.write_bytes(b"PK fake xlsx")
    ctx = _test_ctx(Path(d) / "t.sqlite")
    if recipients:
        ctx.config.email["recipients"] = recipients
    stream = io.StringIO()
    handler = logging.StreamHandler(stream)
    handler.addFilter(RedactingFilter())
    ctx.logger = logging.getLogger("ns-test-email")
    ctx.logger.propagate = False            # keep test output quiet
    ctx.logger.addHandler(handler)
    _FakeSmtp.calls, _FakeSmtp.fail_with, _FakeSmtp.refuse = [], fail_with, refuse or {}
    real_smtp, old_pw = smtplib.SMTP, os.environ.get("NS_SMTP_PASSWORD")
    smtplib.SMTP = _FakeSmtp
    if password is None:
        os.environ.pop("NS_SMTP_PASSWORD", None)
    else:
        os.environ["NS_SMTP_PASSWORD"] = password
    try:
        result = send_report_email(report, ctx, _TEST_HEALTH, {"critical": 1, "manual_review": 2})
    finally:
        smtplib.SMTP = real_smtp
        ctx.logger.removeHandler(handler)
        if old_pw is None:
            os.environ.pop("NS_SMTP_PASSWORD", None)
        else:
            os.environ["NS_SMTP_PASSWORD"] = old_pw
    return result, _FakeSmtp.calls, stream.getvalue()


@selftest("S19")
def test_send_report_email_with_attachment():
    with tempfile.TemporaryDirectory() as d:
        result, calls, log_text = _send_with_fake_smtp(d, "app-pass-1234")
        assert result == ("SENT", 1), result
        names = [c[0] for c in calls]
        assert names == ["connect", "starttls", "login", "send", "quit"], names
        assert calls[0][1:] == ("smtp.gmail.com", 587)
        assert calls[2][1:] == ("operator@example.com", "app-pass-1234"), "password only from the env var"
        msg = calls[3][1]
        assert msg["Subject"].startswith("[Negative Screening] 14–20 Sep 2026 — 1 critical, 2 for review")
        assert msg["To"] == "operator@example.com"
        files = [(a.get_filename(), a.get_content_type(), a.get_content()) for a in msg.iter_attachments()]
        assert files == [("Negative_Screening_2026-09-14.xlsx", "/".join(XLSX_MIME), b"PK fake xlsx")], files
        assert "not a complete record" in msg.get_body(("plain",)).get_content()
        assert "app-pass-1234" not in log_text and "app-pass-1234" not in msg.as_string()


@selftest("S19")
def test_send_report_email_failures_are_reported_not_raised():
    import smtplib
    with tempfile.TemporaryDirectory() as d:
        result, calls, log_text = _send_with_fake_smtp(d, None)
        assert result == ("FAILED", 0) and calls == [], "no password -> no connection attempt"
        assert "NS_SMTP_PASSWORD" in log_text
        result, _, log_text = _send_with_fake_smtp(d, "secret-pw-9",
                                                   fail_with=smtplib.SMTPAuthenticationError(535, b"bad creds"))
        assert result == ("FAILED", 0) and "bad creds" in log_text and "secret-pw-9" not in log_text
        result, _, _ = _send_with_fake_smtp(d, "pw", fail_with=OSError("network unreachable"))
        assert result == ("FAILED", 0)
        result, _, log_text = _send_with_fake_smtp(d, "pw", refuse={"b@example.com": (550, b"no such user")},
                                                   recipients=["a@example.com", "b@example.com"])
        assert result == ("SENT", 1) and "b@example.com" in log_text, "partly delivered: count the delivered"


class _SavedEnv:
    """Temporarily remove / set secret environment variables for a test."""

    def __init__(self, **values: str | None) -> None:
        self.values = {name: values.get(name) for name in SECRET_ENV_VARS}

    def __enter__(self):
        self.saved = {name: os.environ.get(name) for name in SECRET_ENV_VARS}
        for name, value in self.values.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
        return self

    def __exit__(self, *exc):
        for name, value in self.saved.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
        return False


@selftest("S19")
def test_secrets_from_environment_or_credential_store():
    with _SavedEnv():
        assert get_secret("NS_SMTP_PASSWORD") == "" and secret_source("NS_SMTP_PASSWORD") is None
        assert SECRET_STORE is not None
        SECRET_STORE.write("NS_SMTP_PASSWORD", "stored-Pw-77")
        assert get_secret("NS_SMTP_PASSWORD") == "stored-Pw-77", "staff setup: Credential Manager"
        assert secret_source("NS_SMTP_PASSWORD") == "Windows Credential Manager"
        assert redact("login with stored-Pw-77 failed") == "login with *** failed", "stored secrets are redacted"
    with _SavedEnv(NS_SMTP_PASSWORD="env-Pw-11"):
        assert get_secret("NS_SMTP_PASSWORD") == "env-Pw-11", "an environment variable takes precedence"
        assert secret_source("NS_SMTP_PASSWORD") == "environment variable"
    assert normalise_app_password(" abcd efgh ijkl mnop ") == "abcdefghijklmnop"
    assert normalise_app_password("pass word 1") == "pass word 1", "only the Gmail 4x4 form loses spaces"
    with _SavedEnv():
        r = run_collector_safely(GoogleAlertsCollector(), _test_ctx(Path(tempfile.gettempdir()) / "unused.sqlite"))
        assert r.status == SourceStatus.FAILED and r.error and "not configured" in r.error, r.error
        assert "NS_IMAP_PASSWORD" in r.error, "the failure names what to configure"
    source = inspect_source_text()
    for name in SECRET_ENV_VARS:
        assert f'{name} = "' not in source and f"{name}='" not in source, "no secret value in source"


def inspect_source_text() -> str:
    return Path(__file__).read_text(encoding="utf-8") if not IS_FROZEN else ""


@selftest("S19")
def test_windows_credential_store_round_trip():
    if os.name != "nt":
        raise ModuleNotFoundError("Windows Credential Manager", name="Windows Credential Manager")
    store, name = WindowsCredentialStore(), f"NS_SELFTEST_{os.urandom(4).hex()}"
    try:
        assert store.read(name) is None
        store.write(name, "Zz-ünï-9 x")
        assert store.read(name) == "Zz-ünï-9 x", "the real Credential Manager round trip"
    finally:
        store.delete(name)
    assert store.read(name) is None, "deleted"


@selftest("S19")
def test_account_checks_for_setup():
    import imaplib
    import smtplib
    cfg = copy_json(DEFAULT_EMAIL)
    _FakeSmtp.calls, _FakeSmtp.fail_with, _FakeSmtp.refuse = [], None, {}
    ok, message = check_smtp_account(cfg, "pw-1", smtp_factory=_FakeSmtp)
    assert ok and "Signed in" in message and [c[0] for c in _FakeSmtp.calls] == ["connect", "starttls", "login", "quit"]
    _FakeSmtp.calls = []
    ok, message = check_smtp_account(cfg, "pw-1", smtp_factory=_FakeSmtp, send_test_to=["a@example.com"])
    assert ok and "Test email sent to 1 of 1" in message and _FakeSmtp.calls[3][0] == "send"
    assert _FakeSmtp.calls[3][1].get_payload() and not _FakeSmtp.calls[3][1].is_multipart(), "no attachment"
    _FakeSmtp.fail_with = smtplib.SMTPAuthenticationError(535, b"bad creds pw-1")
    ok, message = check_smtp_account(cfg, "pw-1", smtp_factory=_FakeSmtp)
    _FakeSmtp.fail_with = None
    assert not ok and "rejected" in message
    assert check_smtp_account(cfg, "")[0] is False

    class FakeImap:
        label_ok = True
        alerts = True

        def __init__(self, host, port, timeout=None):
            self.host = host

        def login(self, user, password):
            if password != "good":
                raise imaplib.IMAP4.error("AUTHENTICATIONFAILED")

        def select(self, box, readonly=False):
            assert readonly, "the mailbox is only ever read"
            return ("OK" if FakeImap.label_ok else "NO"), [b""]

        def search(self, charset, *criteria):
            assert criteria[:2] == ("FROM", '"googlealerts-noreply@google.com"'), criteria
            return "OK", [b"1 2 3" if FakeImap.alerts else b""]

        def logout(self):
            pass
    FakeImap.alerts = True
    ok, message = check_imap_account(cfg, "good", imap_factory=FakeImap)
    assert ok and "3 Google Alerts email(s) in the label, 3 in the last 14 days" in message, message
    FakeImap.alerts = False
    ok, message = check_imap_account(cfg, "good", imap_factory=FakeImap)
    assert ok and "NO Google Alerts emails yet" in message and "Setup → Google Alerts" in message, message
    assert check_imap_account(cfg, "bad", imap_factory=FakeImap) == (False, "Could not sign in: error: AUTHENTICATIONFAILED")
    FakeImap.label_ok = False
    ok, message = check_imap_account(cfg, "good", imap_factory=FakeImap)
    assert not ok and "NS-Alerts" in message


def copy_json(data: dict) -> dict:
    return json.loads(json.dumps(data))


# ---- S20 --------------------------------------------------------------------
@selftest("S20")
def test_run_lock():
    with tempfile.TemporaryDirectory() as d:
        lock = Path(d) / "run.lock"
        assert acquire_run_lock(lock, 120) is True
        assert acquire_run_lock(lock, 120) is False, "second run must be refused"
        release_run_lock(lock)
        assert acquire_run_lock(lock, 120) is True
        release_run_lock(lock)


@selftest("S20")
def test_run_lock_stale_and_ownership():
    with tempfile.TemporaryDirectory() as d:
        lock = Path(d) / "sub" / "run.lock"
        assert acquire_run_lock(lock, 120) is True, "the lock folder is created"
        info = json.loads(lock.read_text(encoding="utf-8"))
        assert info["pid"] == os.getpid() and info["token"] and info["created"] and "host" in info
        release_run_lock(lock)
        assert not lock.exists(), "released lock is removed"

        other = json.dumps({"pid": 999999, "host": "other", "token": "theirs", "created": "x"})
        lock.write_text(other, encoding="utf-8")
        release_run_lock(lock)
        assert lock.exists(), "never delete a lock this process did not take"
        assert acquire_run_lock(lock, 120) is False, "a fresh foreign lock is respected"
        assert lock.read_text(encoding="utf-8") == other, "a refused run leaves the lock unchanged"

        old = datetime.now(timezone.utc).timestamp() - 121 * 60
        os.utime(lock, (old, old))
        assert acquire_run_lock(lock, 120) is True, "a lock older than max_age_minutes is stale"
        assert json.loads(lock.read_text(encoding="utf-8"))["pid"] == os.getpid(), "stale lock replaced"
        assert [p.name for p in lock.parent.iterdir()] == ["run.lock"], "no side files left behind"

        lock.write_text(other, encoding="utf-8")   # our lock was taken over (e.g. replaced as stale)
        release_run_lock(lock)
        assert lock.read_text(encoding="utf-8") == other, "release must not delete the new owner's lock"
        lock.unlink()

        lock.write_text("not json", encoding="utf-8")   # a half-written lock still counts while fresh
        assert acquire_run_lock(lock, 120) is False
        os.utime(lock, (old, old))
        assert acquire_run_lock(lock, 120) is True
        release_run_lock(lock)


@selftest("S20")
def test_scheduler_wrapper():
    with tempfile.TemporaryDirectory() as d:
        path = write_scheduler_files(Path(d))
        assert path == Path(d) / "run_weekly.bat"
        raw = path.read_bytes()
        text = raw.decode("utf-8")
        assert "--scheduled" in text and "--date" not in text
        assert "\r\n" in text, "Windows batch files need CRLF line endings"
        assert raw.count(b"\n") == raw.count(b"\r\n"), "no bare LF line endings"
        fallback = "\\".join(VENV_PYTHON.parts) if IS_FROZEN else sys.executable     # the exe is not Python
        assert f'"{fallback}" negative_screening.py --scheduled' in text, "no .venv: this interpreter"
        assert 'cd /d "%~dp0"' in text, "runs from the project folder whatever Task Scheduler's start folder"
        assert text.index('cd /d "%~dp0"') < text.index("--scheduled"), "changes folder before starting"
        assert '"NegativeScreening.exe" --scheduled' in text, "the packaged app when present"

        venv_python = Path(d) / VENV_PYTHON
        venv_python.parent.mkdir(parents=True)
        venv_python.write_bytes(b"")
        text = write_scheduler_files(Path(d)).read_bytes().decode("utf-8")
        assert '".venv\\Scripts\\python.exe" negative_screening.py --scheduled\r\n' in text, "project .venv"


@selftest("S20")
def test_pipeline_gate_includes_s20():
    assert "S20" in PIPELINE_REQUIRED_STEPS and "S03" in PIPELINE_REQUIRED_STEPS
    saved = STEP_STATUS["S20"]
    try:
        STEP_STATUS["S20"] = "TODO"
        assert pipeline_ready() == ["S20"], "a run must not reach an unfinished S20 (run lock)"
    finally:
        STEP_STATUS["S20"] = saved


@selftest("S20")
def test_project_root_for_script_and_exe():
    with tempfile.TemporaryDirectory() as d:
        root = Path(d).resolve()
        (root / "config").mkdir()
        (root / "dist").mkdir()
        assert resolve_project_root(False, "x", str(root / "negative_screening.py")) == root
        assert resolve_project_root(True, str(root / EXE_NAME), "") == root, "exe copied into the project"
        assert resolve_project_root(True, str(root / "dist" / EXE_NAME), "") == root, "fresh build in dist\\"
        other = root / "standalone" / "app"
        other.mkdir(parents=True)
        assert resolve_project_root(True, str(other / EXE_NAME), "") == other, "a standalone copy keeps its folder"


@selftest("S20")
def test_run_ids_are_unique_for_same_second_runs():
    with tempfile.TemporaryDirectory() as d:
        db, started = Path(d) / "t.sqlite", _myt(2026, 9, 22, 10, 30).replace(second=12)
        assert unique_run_id(db, started) == "20260922T103012+0800", "no database yet"
        conn = db_connect(db)
        try:
            with conn:
                for rid in ("20260922T103012+0800", "20260922T103012+0800-2"):
                    conn.execute("INSERT INTO runs (run_id, started_at, execution_mode, reference_date,"
                                 " period_start, period_end) VALUES (?, 'x', 'MANUAL', 'x', 'x', 'x')", (rid,))
        finally:
            conn.close()
        assert unique_run_id(db, started) == "20260922T103012+0800-3"

@selftest("S20")
def test_scheduled_task_definition():
    import xml.etree.ElementTree as ET
    assert next_weekday_at(date(2026, 9, 30), 0, "09:00") == datetime(2026, 10, 5, 9, 0), "next Monday"
    assert next_weekday_at(date(2026, 9, 28), 0, "09:00") == datetime(2026, 9, 28, 9, 0), "today is Monday"
    xml = scheduled_task_xml(Path(r"C:\NS\NegativeScreening.exe"), "--scheduled", Path(r"C:\NS"),
                             Path(r"C:\NS\logs"), "PC\\staff", datetime(2026, 10, 5, 9, 0), 120)
    ns = {"t": "http://schemas.microsoft.com/windows/2004/02/mit/task"}
    root = ET.fromstring(xml.split("?>", 1)[1])
    triggers = root.find("t:Triggers", ns)
    assert triggers is not None and [c.tag.split("}")[1] for c in triggers] == ["CalendarTrigger"], "one weekly trigger, no boot/logon trigger"
    week = triggers.find("t:CalendarTrigger/t:ScheduleByWeek", ns)
    assert week is not None and week.find("t:DaysOfWeek/t:Monday", ns) is not None and week.findtext("t:WeeksInterval", None, ns) == "1"
    assert triggers.findtext("t:CalendarTrigger/t:StartBoundary", None, ns) == "2026-10-05T09:00:00"
    settings = root.find("t:Settings", ns)
    assert settings is not None
    assert settings.findtext("t:StartWhenAvailable", None, ns) == "true", "a missed Monday runs when available"
    assert settings.findtext("t:MultipleInstancesPolicy", None, ns) == "IgnoreNew", "no overlapping runs"
    assert settings.findtext("t:ExecutionTimeLimit", None, ns) == "PT120M"
    assert settings.findtext("t:DisallowStartIfOnBatteries", None, ns) == "false"
    action = root.find("t:Actions/t:Exec", ns)
    assert action is not None
    assert action.findtext("t:Command", None, ns) == r"C:\NS\NegativeScreening.exe"
    assert action.findtext("t:Arguments", None, ns) == "--scheduled" and "--date" not in xml
    assert action.findtext("t:WorkingDirectory", None, ns) == r"C:\NS", "starts in the project folder"
    assert root.findtext("t:Principals/t:Principal/t:UserId", None, ns) == "PC\\staff"


@selftest("S20")
def test_install_and_remove_task_with_fake_schtasks():
    calls: list[list[str]] = []
    installed = {"yes": False}
    header = '"HostName","TaskName","Next Run Time","Status","Last Run Time","Last Result","Task To Run"'

    def runner(args):
        calls.append(args)
        if args[0] == "/Create":
            installed["yes"] = True
            return 0, "SUCCESS: The scheduled task was successfully created."
        if args[0] == "/Query":
            if not installed["yes"]:
                return 1, "ERROR: The system cannot find the file specified."
            return 0, header + '\n"PC","\\Negative Screening - Weekly","05/10/2026 09:00:00","Ready","N/A","267011","x.exe --scheduled"'
        if args[0] == "/Delete":
            installed["yes"] = False
            return 0, "SUCCESS"
        return 1, "unexpected"
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        ok, message = install_scheduled_task(root, "9:00", runner=runner)
        assert not ok and calls == [], "a bad time is refused before Task Scheduler is called"
        ok, message = install_scheduled_task(root, "09:00", runner=runner)
        assert ok and "every Monday at 09:00" in message, message
        assert calls[-1][:3] == ["/Create", "/TN", SCHEDULED_TASK_NAME] and calls[-1][-1] == "/F"
        xml = Path(calls[-1][4]).read_bytes()
        assert xml[:2] in (b"\xff\xfe", b"\xfe\xff"), "schtasks /XML needs a UTF-16 file"
        text = xml.decode("utf-16")
        if IS_FROZEN:
            assert f"<Command>{Path(sys.executable).resolve()}</Command>" in text, "the running exe itself"
        else:
            assert "run_weekly.bat" in text and (root / "run_weekly.bat").exists(), "no exe yet: the launcher"
        (root / EXE_NAME).write_bytes(b"MZ")
        install_scheduled_task(root, "08:30", runner=runner)
        text = Path(calls[-1][4]).read_bytes().decode("utf-16")
        assert f"<Command>{(root / EXE_NAME).resolve()}</Command>" in text and "<Arguments>--scheduled</Arguments>" in text
        assert "T08:30:00" in text
        info = query_scheduled_task(runner)
        assert info is not None and info["next_run"] == "05/10/2026 09:00:00" and info["status"] == "Ready"
        assert uninstall_scheduled_task(runner)[0] is True and calls[-1][0] == "/Delete"
        assert query_scheduled_task(runner) is None
        n = len(calls)
        assert uninstall_scheduled_task(runner) == (True, f"'{SCHEDULED_TASK_NAME}' is not installed; nothing to remove.")
        assert all(c[0] != "/Delete" for c in calls[n:]), "nothing deleted when not installed"
        ok, message = install_scheduled_task(root, "09:00", runner=lambda a: (1, "ERROR: Access is denied."))
        assert not ok and "Access is denied" in message, "a refused install is reported, not claimed"


@selftest("S20")
def test_staff_status_texts_and_report_opening():
    assert "not been configured" in friendly_source_note("FAILED", IMAP_NOT_CONFIGURED)
    assert "Traceback" not in friendly_source_note("FAILED", "RuntimeError: " + IMAP_NOT_CONFIGURED)
    assert friendly_source_note("FAILED", "Google Alerts unavailable").startswith("NOT CHECKED")
    assert friendly_source_note("SUCCESS", None) == "Checked"
    assert "not been configured" in friendly_email_status({"status": "FAILED", "reason": SMTP_NOT_CONFIGURED}, None)
    assert "already emailed" in friendly_email_status(
        {"status": "SKIPPED", "reason": "already sent for period"}, {"sent_at": "2026-09-28T09:01:00+08:00"})
    assert friendly_email_status({"status": "SENT", "recipient_count": 2, "sent_at": "2026-09-28T09:01:00+08:00"},
                                 None).startswith("Sent to 2 recipient(s), Mon 28 Sep 2026 09:01")
    assert describe_last_run(None)["when"] == "No screening has been run yet."
    for day, start in [(22, 14), (25, 14), (27, 14), (28, 21)]:        # Tue, Fri, Sun, Mon
        req = staff_run_request(date(2026, 9, day))
        assert req.mode == Mode.MANUAL and not req.no_email and not req.dry_run and not req.validation
        assert calculate_previous_completed_week(req.reference).period_start == date(2026, 9, start)
    with tempfile.TemporaryDirectory() as d:
        folder = Path(d)
        assert latest_report(folder) is None
        old, new = folder / "Negative_Screening_2026-09-07.xlsx", folder / "Negative_Screening_2026-09-14.xlsx"
        for i, f in enumerate([old, new]):
            f.write_bytes(b"PK")
            os.utime(f, (1_800_000_000 + i, 1_800_000_000 + i))
        (folder / "dry_run").mkdir()
        (folder / "dry_run" / "Negative_Screening_2026-09-21.xlsx").write_bytes(b"PK")
        assert latest_report(folder) == new, "newest report in reports\\ itself, not dry_run\\"
        assert latest_report(folder, {"run": {"report_path": str(old), "dry_run": 0}}) == old, "last run's report"
        opened: list[str] = []
        open_path(new, opener=opened.append)
        assert opened == [str(new)]


def _e2e_project(d: str) -> Path:
    root = Path(d)
    init_config(root / "config")
    return root


class _FakeWebsiteCollector:
    name = "website"

    def collect(self, ctx: RunContext) -> SourceResult:
        item = RawItem(source="website", discovery_method="Direct public retrieval",
                       discovered_url="https://acme.example.com/news/acme-asset-management-fraud-probe",
                       title="Acme Asset Management named in fraud probe",
                       snippet="Acme Asset Management Sdn Bhd was named in a fraud probe.",
                       source_published_at=datetime(2026, 9, 16, 10, 0, tzinfo=MYT))
        return SourceResult(self.name, SourceStatus.SUCCESS, items=[item])


def _e2e_run(root: Path, req: RunRequest) -> RunOutcome:
    import contextlib
    import io
    with contextlib.redirect_stderr(io.StringIO()):
        outcome = execute_run(req, root)
    logger = logging.getLogger(LOGGER_NAME)
    for h in list(logger.handlers):       # release the log files so the temp folder can go
        logger.removeHandler(h)
        h.close()
    return outcome


@selftest("S20")
def test_end_to_end_runs_manual_any_day_scheduled_and_validation():
    import smtplib
    from openpyxl import load_workbook
    global build_collectors
    real_collectors, real_smtp = build_collectors, smtplib.SMTP
    build_collectors = lambda ctx: [GoogleAlertsCollector(), _FakeWebsiteCollector()]   # noqa: E731
    smtplib.SMTP = _FakeSmtp
    _FakeSmtp.calls, _FakeSmtp.fail_with, _FakeSmtp.refuse = [], None, {}
    try:
        with _SavedEnv(), tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as d:
            root = _e2e_project(d)
            assert SECRET_STORE is not None, "secret store is unavailable"
            SECRET_STORE.write("NS_SMTP_PASSWORD", "app-pw-e2e")
            prod_db = root / "data" / "screening.sqlite"

            # Tuesday: manual run, previous complete week, report saved and emailed by default
            out = _e2e_run(root, staff_run_request(date(2026, 9, 22)))
            assert out.exit_code == 1, "Google Alerts not configured -> completed with a FAILED source"
            report = root / "reports" / "Negative_Screening_2026-09-14.xlsx"
            assert out.report_path == report and report.is_file(), out.report_path
            sends = [c for c in _FakeSmtp.calls if c[0] == "send"]
            assert len(sends) == 1 and sends[0][1].get_filename() is None
            assert any(part.get_filename() == report.name for part in sends[0][1].iter_attachments())
            assert not (root / "data" / "run.lock").exists(), "the lock is released"
            d1 = describe_last_run(read_last_run(prod_db))
            assert d1["email"].startswith("Sent to 1 recipient(s)") and d1["alerts_not_configured"]
            assert d1["period"] == "Monday 14 Sep 2026 – Sunday 20 Sep 2026"
            ga = [s for s in d1["sources"] if s[0] == "Google Alerts"][0]
            assert ga[1] == "FAILED" and "has not been configured" in ga[4]
            wb = load_workbook(report)
            assert "Acme Asset Management named in fraud probe" in " ".join(
                str(c.value) for row in wb["Findings"].iter_rows() for c in row), "finding reported"

            # Friday: same week, report regenerated (old one archived); seconds later = inside the
            # 10-minute manual cooldown, so not emailed again
            out = _e2e_run(root, staff_run_request(date(2026, 9, 25)))
            assert out.report_path == report and len([c for c in _FakeSmtp.calls if c[0] == "send"]) == 1
            assert describe_last_run(read_last_run(prod_db))["email"].startswith("Not sent: an email was sent 0 min ago")
            assert len(list((root / "reports" / "archive").glob("Negative_Screening_2026-09-14_*.xlsx"))) == 1

            # Sunday: Run without email (possibly in the same second as the Friday run)
            out = _e2e_run(root, staff_run_request(date(2026, 9, 27), no_email=True))
            assert out.exit_code in (0, 1), out
            assert describe_last_run(read_last_run(prod_db))["email"] == "Not sent (run without email)."

            # Monday: scheduled run = new week, emailed by default
            out = _e2e_run(root, RunRequest(mode=Mode.AUTOMATIC, reference=date(2026, 9, 28)))
            assert out.report_path == root / "reports" / "Negative_Screening_2026-09-21.xlsx"
            assert len([c for c in _FakeSmtp.calls if c[0] == "send"]) == 2
            d2 = describe_last_run(read_last_run(prod_db))
            assert d2["mode"] == "Automatic (weekly schedule)" and d2["email"].startswith("Sent to 1")

            # Validation: staging database, reports\validation\, email preview only
            last_run = read_last_run(prod_db)
            assert last_run is not None, "expected a production run before validation"
            before = last_run["run"]["run_id"]
            out = _e2e_run(root, RunRequest(mode=Mode.MANUAL, reference=date(2026, 9, 22), validation=True))
            assert out.db_path == root / "data" / "screening_validation.sqlite"
            assert out.report_path == root / "reports" / "validation" / "Negative_Screening_2026-09-14.xlsx"
            assert (root / "reports" / "validation" / "Negative_Screening_2026-09-14_email_preview.eml").is_file()
            assert len([c for c in _FakeSmtp.calls if c[0] == "send"]) == 2, "validation never sends"
            last_run_after_validation = read_last_run(prod_db)
            assert last_run_after_validation is not None, "validation must not remove production run data"
            assert last_run_after_validation["run"]["run_id"] == before, "production data untouched"
            assert out.report_path is not None
            summary = " ".join(str(c.value) for row in load_workbook(out.report_path)["Summary"].iter_rows()
                               for c in row)
            assert "VALIDATION DATA" in summary

            # Another run holds the lock -> refused, nothing written
            lock = root / "data" / "run.lock"
            lock.write_text('{"pid": 1, "host": "other", "token": "t"}', encoding="utf-8")
            out = _e2e_run(root, staff_run_request(date(2026, 9, 29)))
            assert out.exit_code == 3 and "in progress" in out.message and lock.exists()
            lock.unlink()

            # Email not configured: report kept, failure explained
            SECRET_STORE.delete("NS_SMTP_PASSWORD")
            out = _e2e_run(root, staff_run_request(date(2026, 10, 6), resend_email=True))   # past the cooldown
            assert out.exit_code == 1 and out.report_path is not None and out.report_path.is_file()
            assert "not been configured" in describe_last_run(read_last_run(prod_db))["email"]
    finally:
        build_collectors, smtplib.SMTP = real_collectors, real_smtp


@selftest("S20")
def test_reports_folder_outside_the_project():
    with tempfile.TemporaryDirectory() as d:
        project, outside = Path(d) / "project", Path(d) / "Negative_Screening_Reports"
        ensure_project_dirs(project)
        cfg = _test_config()
        cfg.settings["paths"]["reports_dir"] = str(outside)
        assert ensure_report_dirs(project, cfg) == outside == reports_folder(project, cfg)
        assert (outside / "archive").is_dir() and (outside / "dry_run").is_dir()
        assert not (project / "reports").exists(), "no reports folder is made inside the project"
        period = calculate_previous_completed_week(date(2026, 9, 22))
        weekly = report_path_for(period, outside, False)
        assert report_path_for(period, outside, True) == outside / "dry_run" / weekly.name
        weekly.write_bytes(b"PK")
        archived = _archive_existing_report(weekly)
        assert archived is not None and archived.parent == outside / "archive", archived


@selftest("S20")
def test_gui_launch():
    import tkinter as tk
    try:
        root = tk.Tk()
    except tk.TclError as e:
        raise ModuleNotFoundError(f"a display for the window ({e})", name="a display") from e
    try:
        root.withdraw()
        with _SavedEnv(), tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as d:
            app = StaffApp(root, Path(d), task_query=lambda: None, today=lambda: date(2026, 9, 25))
            root.update()
            assert (Path(d) / "config" / "settings.json").exists(), "missing config is created"
            for sub in ("data", "logs", "reports", "reports/archive", "reports/dry_run"):
                assert (Path(d) / sub).is_dir(), sub
            texts = [app.run_button.cget("text"), app.latest_button.cget("text"), app.folder_button.cget("text"),
                     app.print_button.cget("text")]
            assert texts == ["Run Weekly Screening", "Open Latest Report", "Open Reports Folder", "Print Report"]
            assert app.fields["targets"].get().startswith("Acme Asset Management Sdn Bhd — Company; searched as:")
            assert "Tan Ah Kow — Person (Founder)" in app.fields["targets"].get()
            assert "Friday 14 Sep" not in app.period_var.get() and "Monday 14 Sep 2026 – Sunday 20 Sep 2026" in \
                app.period_var.get(), "a Friday click scans the previous complete week"
            assert "Google Alerts email access has not been configured" in app.notice_var.get()
            assert app.fields["when"].get() == "No screening has been run yet."
            assert app.fields["schedule"].get().startswith("Not installed")
    finally:
        root.destroy()


@selftest("S10")
def test_google_alerts_empty_mailbox_is_a_warning_not_silence():
    with _SavedEnv(NS_IMAP_PASSWORD="test-app-password"):
        ctx = _alerts_ctx()
        r = run_collector_safely(GoogleAlertsCollector(imap_factory=_FakeImap({}), session=object()), ctx)
    assert r.status == SourceStatus.SUCCESS and r.items == []
    assert any("no Google Alerts emails arrived in the 'NS-Alerts' label since 13 Sep 2026" in w for w in r.warnings), \
        r.warnings
    health = {h["source"]: h for h in derive_source_health([r], [])}
    assert health["Google Alerts"]["error"].startswith("Warnings: no Google Alerts emails arrived")
    for platform in ("Facebook", "Instagram", "LinkedIn", "Public website"):
        assert "No Google Alerts emails arrived" in health[platform]["error"], "never a silent 0 results"
        assert friendly_source_note("SUCCESS", health[platform]["error"]).startswith("Not checked")
    info = {"run": {"exit_code": 0, "period_start": "2026-09-14", "period_end": "2026-09-20",
                    "started_at": "2026-09-22T10:00:00+08:00", "execution_mode": "MANUAL"},
            "sources": list(health.values()), "email": None, "period_sent": None, "findings": []}
    assert describe_last_run(info)["result"].startswith("Completed — but no Google Alerts emails arrived"),         "not 'every source was checked' when Google's index was not read"
    note = friendly_source_note("SUCCESS", health["Google Alerts"]["error"])
    assert "NO Google Alerts emails arrived" in note and "Setup → Google Alerts" in note
    assert friendly_source_note("SUCCESS", "Warnings: comments feed not read") == "Checked — comments feed not read"


@selftest("S18")
def test_report_shows_targets_and_prints_on_a4():
    with tempfile.TemporaryDirectory() as d:
        c = _report_candidate("k-p", "https://news.example.com/p", Section.MAIN, Priority.HIGH, when=_myt(2026, 9, 16))
        _, wb, ctx, conn = _report_workbook(d, [c])
        try:
            summary = {r[0]: r[1] for r in wb[SHEET_SUMMARY].iter_rows(values_only=True) if r[0]}
            assert "Screened targets" in summary
            for target in screened_targets(ctx.config.entities):
                target_summary = summary.get(target["name"])
                assert isinstance(target_summary, str) and target_summary.startswith(target["kind"]), target
            assert "searched as: Acme Asset Management Sdn Bhd, Acme Asset Management" in str(\
                summary["Acme Asset Management Sdn Bhd"])
            assert str(summary["Tan Ah Kow"]).startswith("Person (Founder)")
            assert isinstance(summary["Sources attempted"], int), "counts stay numbers in Excel"
            for ws in wb.worksheets:
                assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9, "A4"
                assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
                assert ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage, "fitted to one page wide"
                header_center = ws.oddHeader.center if ws.oddHeader is not None else None
                assert header_center is not None and isinstance(header_center.text, str) and "14 Sep 2026 – 20 Sep 2026" in header_center.text
                assert ws.oddFooter is not None and ws.oddFooter.right.text == "Page &P of &N"
            assert wb[SHEET_FINDINGS].print_title_rows in ("1:1", "$1:$1"), "column headings repeat on every page"
        finally:
            conn.close()
    body = build_email_body(_test_ctx(Path("unused.sqlite")), _TEST_HEALTH, {"critical": 1}, "r.xlsx")
    assert "Screened targets: Acme Asset Management Sdn Bhd (Company); Tan Ah Kow (Person (Founder))" in body


@selftest("S20")
def test_printable_report_view():
    import smtplib
    global build_collectors
    real_collectors, real_smtp = build_collectors, smtplib.SMTP

    class Hostile(_FakeWebsiteCollector):
        def collect(self, ctx):
            r = super().collect(ctx)
            r.items[0].title = "Acme Asset Management fraud <script>alert(1)</script>"
            return r
    build_collectors = lambda ctx: [Hostile()]   # noqa: E731
    smtplib.SMTP = _FakeSmtp
    _FakeSmtp.calls, _FakeSmtp.fail_with, _FakeSmtp.refuse = [], None, {}
    try:
        with _SavedEnv(), tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as d:
            root = _e2e_project(d)
            cfg = load_config(root / "config")
            db = root / "data" / "screening.sqlite"
            assert write_print_view(root, cfg, db) is None, "no run yet: nothing to print"
            _e2e_run(root, staff_run_request(date(2026, 9, 22), no_email=True))
            path = write_print_view(root, cfg, db)
            assert path is not None, "a completed run should produce a printable report"
            assert path == root / "reports" / "print" / "Negative_Screening_2026-09-14.html"
            html_text = path.read_text(encoding="utf-8")
            assert "Monday 14 Sep 2026 – Sunday 20 Sep 2026" in html_text
            assert "<b>Acme Asset Management Sdn Bhd</b> — Company" in html_text
            assert "<b>Tan Ah Kow</b> — Person (Founder)" in html_text, "the screened targets are shown"
            assert '<a href="https://acme.example.com/news/acme-asset-management-fraud-probe">' in html_text
            assert "<script>alert(1)" not in html_text and "&lt;script&gt;alert(1)" in html_text, "web text escaped"
            assert "size: A4 landscape" in html_text and "window.print()" in html_text
            assert "1 in Findings" in html_text and "Not sent (run without email)." in html_text
            info = read_last_run(db)
            assert info is not None, "the completed run should be available"
            assert describe_findings(info["findings"]).startswith("1 in Findings (1 ")
    finally:
        build_collectors, smtplib.SMTP = real_collectors, real_smtp
    assert describe_findings([]) == "No reportable findings in this period from the sources that were checked."
    assert describe_findings([("MAIN", "CRITICAL", 2), ("MAIN", "HIGH", 1), ("DATE_UNKNOWN", "HIGH", 1),
                              ("NOT_REPORTED", None, 9)]) == \
        "3 in Findings (2 critical, 1 high) · 1 date unknown · 0 unverified URLs"


@selftest("S19")
def test_manual_email_cooldown():
    dflt = {"send_on_automatic": True, "send_on_manual": True, "manual_email_cooldown_minutes": 10}

    def decide(**kw: Any) -> tuple[bool, str]:
        args: dict[str, Any] = dict(mode=Mode.MANUAL, dry_run=False, no_email=False, send_email=False,
                                    resend_email=False, already_sent=True, email_defaults=dflt,
                                    minutes_since_last_email=None)
        args.update(kw)
        return decide_email(**args)
    assert DEFAULT_SETTINGS["email_defaults"]["manual_email_cooldown_minutes"] == 10
    assert decide() == (True, "manual mode default"), "manual scan of an emailed week: sent again"
    send, reason = decide(minutes_since_last_email=3.7)
    assert not send and reason == "cooldown: an email was sent 3 min ago; manual scans email at most once every 10 min"
    assert decide(minutes_since_last_email=10)[0] is True, "cooldown over at 10 minutes"
    assert decide(minutes_since_last_email=3, resend_email=True)[0] is True, "re-send overrides the cooldown"
    assert decide(minutes_since_last_email=3, send_email=True)[0] is False, "the cooldown also applies to --send-email"
    assert decide(minutes_since_last_email=3, no_email=True) == (False, "--no-email")
    assert decide(mode=Mode.AUTOMATIC, minutes_since_last_email=500) == (False, "already sent for period"), \
        "the Monday task still emails a week only once"
    assert decide(email_defaults={**dflt, "manual_email_cooldown_minutes": 5}, minutes_since_last_email=6)[0] is True
    no_cooldown = {"send_on_automatic": True, "send_on_manual": True}
    assert decide(email_defaults=no_cooldown) == (False, "already sent for period"), "setting absent: old rule"
    assert "Not sent: an email was sent 3 min ago" in friendly_email_status(
        {"status": "SKIPPED", "reason": reason.replace("3 min", "3 min")}, None)
    with tempfile.TemporaryDirectory() as d:
        conn = db_connect(Path(d) / "t.sqlite")
        try:
            ctx = _test_ctx(Path(d) / "t.sqlite")
            db_record_run_start(conn, ctx)
            now = _myt(2026, 9, 22, 10, 30)
            assert db_minutes_since_last_email(conn, now) is None
            db_record_email(conn, ctx.run_id, ctx.period, "FAILED", 0, "x")
            assert db_minutes_since_last_email(conn, now) is None, "a failed email starts no cooldown"
            with conn:
                conn.execute("INSERT INTO email_log (run_id, period_start, sent_at, recipient_count, status)"
                             " VALUES (?, '2026-09-14', ?, 1, 'SENT')", (ctx.run_id, "2026-09-22T10:24:00+08:00"))
            minutes_since_last_email = db_minutes_since_last_email(conn, now)
            assert minutes_since_last_email is not None
            assert abs(minutes_since_last_email - 6.0) < 1e-6
        finally:
            conn.close()
    bad = copy_json(DEFAULT_SETTINGS)
    bad["email_defaults"]["manual_email_cooldown_minutes"] = "ten"
    try:
        _validate_settings(bad)
        raise AssertionError("a non-number cooldown must be rejected")
    except ConfigError as e:
        assert "manual_email_cooldown_minutes" in str(e)


@selftest("S18")
def test_report_records_discovered_and_fetched_links():
    with tempfile.TemporaryDirectory() as d:
        social = _report_candidate("k-s", "https://www.facebook.com/acme/posts/123", Section.MAIN,
                                   Priority.HIGH, when=_myt(2026, 9, 16))
        web = _report_candidate("k-w", "https://news.example.com/a?utm_source=alert", Section.MAIN,
                                Priority.HIGH, when=_myt(2026, 9, 17))
        web.raw.final_url = "https://news.example.com/a"
        feed = _report_candidate("k-f", "https://acme.example.com/news/x", Section.MAIN, Priority.HIGH,
                                 when=_myt(2026, 9, 18))
        _, wb, _, conn = _report_workbook(d, [social, web, feed])
        try:
            rows = {r["Discovered URL"]: r for r in _sheet_values(wb[SHEET_FINDINGS])}
            assert rows["https://www.facebook.com/acme/posts/123"]["Fetched URL"] == NOT_FETCHED_SOCIAL
            assert rows["https://news.example.com/a?utm_source=alert"]["Fetched URL"] == "https://news.example.com/a"
            assert rows["https://acme.example.com/news/x"]["Fetched URL"] == NOT_FETCHED_OTHER
            ws = wb[SHEET_FINDINGS]
            col = FINDINGS_COLUMNS.index("Fetched URL") + 1
            linked = [ws.cell(row=r, column=col).hyperlink for r in range(2, 5)]
            assert sum(1 for h in linked if h is not None) == 1, "only real fetched URLs are links"
        finally:
            conn.close()


_MEDIA_FB_SEMINAR = "https://www.facebook.com/AcmeAM/posts/pfbid0AaAaA111"
_MEDIA_FB_ADVERSE = "https://www.facebook.com/AcmeAM/posts/pfbid0BbBbB222"
_MEDIA_LINKEDIN = "https://www.linkedin.com/posts/tan-ah-kow_update-activity-7250000000000000000-AbCd"
_MEDIA_NEWS = "https://www.thestar.com.my/business/2026/09/22/acme-asset-management-probed-for-fraud"


def _media_alert_html() -> str:
    """A Google Alerts email in the real layout: every result is a google.com/url wrapper. Google
    titles both Facebook posts with the page name, so they share publisher + title + day."""
    from urllib.parse import quote

    def result(url: str, title: str, publisher: str, snippet: str) -> str:
        href = f"https://www.google.com/url?rct=j&amp;sa=t&amp;url={quote(url, safe='')}&amp;ct=ga&amp;cd=CAIy&amp;usg=AOv"
        return (f'<tr><td itemscope itemtype="http://schema.org/Article"><a href="{href}" itemprop="url">'
                f'<span itemprop="name">{title}</span></a><div itemprop="publisher"><span itemprop="name">'
                f'{publisher}</span></div><div itemprop="description">{snippet}</div>'
                f'<a href="https://www.google.com/alerts/feedback?ffu=x">Flag as irrelevant</a></td></tr>')
    return "<html><body><table>" + "".join([
        result(_MEDIA_FB_SEMINAR, "Acme Asset Management | Facebook", "Facebook",
               "1 day ago ... Acme Asset Management thanks investors for joining our seminar."),
        result(_MEDIA_FB_ADVERSE, "Acme Asset Management | Facebook", "Facebook",
               "1 day ago ... Investors say <b>Acme Asset Management</b> is a fraud after returns stopped."),
        result(_MEDIA_LINKEDIN, "Tan Ah Kow on LinkedIn: market update", "LinkedIn",
               "2 days ago ... Tan Ah Kow of Acme Asset Management questioned in fraud investigation."),
        result(_MEDIA_NEWS, "Acme Asset Management probed for fraud", "The Star",
               "Acme Asset Management Sdn Bhd is being probed for alleged fraud."),
    ]) + '<tr><td><a href="https://www.google.com/alerts/remove?s=abc">Unsubscribe</a></td></tr></table></body></html>'


@selftest("S07")
def test_different_social_posts_with_the_same_title_are_not_merged():
    items = parse_alert_email(_alert_email_bytes(_media_alert_html(), msg_id="<media@google.com>"),
                              _myt(2026, 9, 23, 9, 15))
    cands, _ = dedupe_in_run([normalize_item(i) for i in items])
    for c in cands:
        c.platform = detect_platform(c.canonical_url, c.raw.source)
        c.content_class = content_class_for(c.platform, c.raw)
        resolve_content_date(c)
    fb = [c for c in cands if c.platform == Platform.FACEBOOK]
    assert len(fb) == 2 and fb[0].content_date == fb[1].content_date and fb[0].raw.title == fb[1].raw.title
    mark_probable_duplicates(cands)
    assert [c.raw.discovered_url for c in cands] == [_MEDIA_FB_SEMINAR, _MEDIA_FB_ADVERSE, _MEDIA_LINKEDIN, _MEDIA_NEWS]
    adverse = cands[1]
    assert "fraud" in adverse.raw.snippet and "seminar" not in adverse.raw.snippet, "texts must not be swapped"
    assert "seminar" in cands[0].raw.snippet and not cands[0].notes


@selftest("S20")
def test_flagged_media_urls_reach_sqlite_and_excel():
    """Google Alerts email -> collector -> screening -> SQLite -> saved .xlsx (read back from disk)."""
    import contextlib
    from openpyxl import load_workbook
    global build_collectors
    real_collectors = build_collectors
    article = ('<html><head><meta property="article:published_time" content="2026-09-22T10:00:00+08:00">'
               '</head><body><p>Acme Asset Management Sdn Bhd is being probed for alleged fraud.</p></body></html>')
    session = _FakeSession({"https://www.thestar.com.my/robots.txt": [_FakeResponse(404)],
                            _MEDIA_NEWS: [_FakeResponse(200, article, {"Content-Type": "text/html"})]})
    imap = _FakeImap({"1": ("23-Sep-2026 09:15:00 +0800",
                            _alert_email_bytes(_media_alert_html(), subject='Google Alert – "Acme Asset Management"',
                                               msg_id="<media@google.com>"))})
    build_collectors = lambda ctx: [GoogleAlertsCollector(imap_factory=imap, session=session)]   # noqa: E731
    try:
        with _SavedEnv(NS_IMAP_PASSWORD="test-app-password"), \
                tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as d:
            root = _e2e_project(d)
            settings_file = root / "config" / "settings.json"
            settings = json.loads(settings_file.read_text(encoding="utf-8"))
            settings["http"].update(max_retries=0, backoff_seconds=0, delay_between_requests_seconds=0)
            settings["paths"]["reports_dir"] = str(Path(d) / "Negative_Screening_Reports")
            settings_file.write_text(json.dumps(settings), encoding="utf-8")
            out = _e2e_run(root, staff_run_request(date(2026, 10, 1), no_email=True))
            assert out.report_path is not None and out.report_path.is_file(), out.message
            assert not any("facebook.com" in u or "linkedin.com" in u for u in session.calls), \
                f"a social page was fetched: {session.calls}"

            # SQLite: the exact URL each source supplied; social posts were never fetched
            assert out.db_path is not None
            with contextlib.closing(sqlite3.connect(out.db_path)) as conn:
                conn.row_factory = sqlite3.Row
                db = {r["discovered_url"]: r for r in conn.execute(
                    "SELECT f.discovered_url, f.final_url, fp.report_section, fp.priority, fp.review_status"
                    " FROM findings f JOIN finding_periods fp USING (finding_id)")}
            assert set(db) == {_MEDIA_FB_SEMINAR, _MEDIA_FB_ADVERSE, _MEDIA_LINKEDIN, _MEDIA_NEWS}, set(db)
            # The LinkedIn snippet names Tan Ah Kow AND the company: R8 sends it to human review.
            review = {_MEDIA_FB_ADVERSE: "PENDING", _MEDIA_LINKEDIN: "NEEDS REVIEW", _MEDIA_NEWS: "PENDING"}
            for url in (_MEDIA_FB_ADVERSE, _MEDIA_LINKEDIN):
                assert db[url]["final_url"] is None and db[url]["report_section"] == "MAIN", dict(db[url])
                assert (db[url]["priority"], db[url]["review_status"]) == ("HIGH", review[url]), \
                    "the screening flag (priority) and the staff review status are separate fields"
            assert db[_MEDIA_NEWS]["final_url"] == _MEDIA_NEWS and db[_MEDIA_NEWS]["priority"] == "HIGH"
            assert db[_MEDIA_FB_SEMINAR]["report_section"] == "NOT_REPORTED", "name without keyword (R7)"

            # Excel, read back from the saved file: current columns, same URL, real hyperlink
            wb = load_workbook(out.report_path)
            ws = wb[SHEET_FINDINGS]
            assert [c.value for c in ws[1]] == FINDINGS_COLUMNS, "runtime workbook = current FINDINGS_COLUMNS"
            col = {name: i + 1 for i, name in enumerate(FINDINGS_COLUMNS)}
            rows = {ws.cell(row=r, column=col["Discovered URL"]).value: r for r in range(2, ws.max_row + 1)}
            assert set(rows) == {_MEDIA_FB_ADVERSE, _MEDIA_LINKEDIN, _MEDIA_NEWS}, set(rows)
            for url, r in rows.items():
                assert isinstance(url, str)
                cell = ws.cell(row=r, column=col["URL"])
                assert cell.value == url and cell.hyperlink is not None and cell.hyperlink.target == url, \
                    (url, cell.value, cell.hyperlink)
                assert ws.cell(row=r, column=col["Priority"]).value == "HIGH"
                assert ws.cell(row=r, column=col["Manual Review Status"]).value == review[url]
            fb = rows[_MEDIA_FB_ADVERSE]
            assert ws.cell(row=fb, column=col["Platform"]).value == "Facebook"
            assert ws.cell(row=fb, column=col["Fetched URL"]).value == NOT_FETCHED_SOCIAL
            title = ws.cell(row=fb, column=col["Title / Comment / Snippet"]).value
            assert isinstance(title, str) and "fraud" in title
            fetched_url_cell = ws.cell(row=rows[_MEDIA_NEWS], column=col["Fetched URL"])
            assert fetched_url_cell.hyperlink is not None and fetched_url_cell.hyperlink.target == _MEDIA_NEWS

            # Run log answers: was a URL discovered, what entered SQLite, what was written to Excel
            log = next((root / "logs").glob("run_*.log")).read_text(encoding="utf-8")
            for stage in ("collected", "stored", "excel"):
                assert f"url trace [{stage}]" in log and _MEDIA_FB_ADVERSE in log.split(f"url trace [{stage}]", 1)[1], stage
    finally:
        build_collectors = real_collectors


@selftest("S12")
def test_sc_item_file_is_labelled_and_the_url_stays_the_list_page():
    data = json.loads(json.dumps(_SC_ALERT_DATA))
    entries = data["cf1e0250-7566-4738-9fc0-604a002ee200"]
    evidence = "https://www.sc.com.my/clients/Security_Commission_X/contentms/img/fake-cert/fake-cert-acme.jpg"
    entries["synthetic-0001"]["url"] = evidence
    entries["6a847d9d-9c66-4244-a948-0a09c835db24"]["name"] += " (Acme Asset Management)"
    items = {i.title: i for i in parse_sc_alert_list(_sc_page("PMD", data), ["Acme Asset Management"], _SC_ALERT_URL)}
    with_file = items["SC Investor Alert List: Illustrative Clone of ACME ASSET MANAGEMENT SDN. BHD."]
    without = items["SC Investor Alert List: Example Dragon BTC Mining-Malaysia (Acme Asset Management)"]
    for item in (with_file, without):
        c = normalize_item(item)
        c.platform = detect_platform(c.canonical_url, c.raw.source)
        assert item.discovered_url == _SC_ALERT_URL, "the item URL is the list page the source gave"
        assert classify_url(c.canonical_url, c.platform)[1] == UrlType.CONTAINER
    assert f"{SC_ITEM_FILE_LABEL}{evidence}" in with_file.snippet, with_file.snippet
    assert SC_ITEM_FILE_LABEL not in without.snippet, "no item file is invented when SC gives none"
    assert "https://www.acme-fund.com/" in without.snippet, "the entry's website stays as listed text"


def _alias_entities(**aliases: list[str]) -> dict:
    """The default approved entities with exactly the given aliases per entity id (others: none)."""
    entities = json.loads(json.dumps(DEFAULT_ENTITIES))
    for e in entities["entities"]:
        e["aliases"] = list(aliases.get(e["id"], []))
    return entities


def _keywords_with(term: str) -> dict:
    """The default keywords plus `term`, unless the built-in dictionary already has it (D-28: the
    Chinese forms such as 欺诈 are now variants of their English keyword)."""
    keywords = json.loads(json.dumps(DEFAULT_KEYWORDS))
    known = {f for c in keywords["categories"] for k in c["keywords"] for f in [k["term"], *k.get("variants", [])]}
    if term not in known:
        keywords["categories"][0]["keywords"].append({"term": term, "ambiguous": False})
    return keywords


@selftest("S15")
def test_chinese_aliases_match_only_when_approved():
    configured = _alias_entities(tan_ah_kow=["陈亚九"], acme_company=["艾克米投资管理"])
    person = match_entities("据报道，陈亚九昨日接受调查。", "", configured)
    assert [(m.entity_id, m.matched_text, m.position) for m in person] == [("tan_ah_kow", "陈亚九", 3)], \
        "据 报 道 are tokens 0-2; the Chinese comma is punctuation, not a token"
    company = match_entities("艾克米投资管理有限公司发表声明", "", configured)
    assert [(m.entity_id, m.matched_text) for m in company] == [("acme_company", "艾克米投资管理")], company
    mixed = match_entities("Report: 陈亚九 and Acme Asset Management", "", configured)
    assert {(m.entity_id, m.matched_text) for m in mixed} == {
        ("tan_ah_kow", "陈亚九"), ("acme_company", "Acme Asset Management")}
    assert match_entities("陈亚 attended", "", configured) == [], "part of an alias is not a match"
    assert match_entities("陳亞九", "", configured) == [], "traditional characters only when listed too"
    for text in ("Chen Yajiu was questioned in a fraud probe", "Chen Ya Jiu", "chenyajiu"):
        assert match_entities(text, "", configured) == [], f"{text!r} is not an approved alias"
    assert match_entities("据报道，陈亚九昨日接受调查。", "", _alias_entities()) == [], \
        "without the approved alias in entities.json nothing is inferred or translated"
    assert match_entities("Chen Yajiu", "", _alias_entities(tan_ah_kow=["Chen Yajiu"]))[0].entity_id \
        == "tan_ah_kow", "a romanised alias works only once it is approved and listed"
    # The alias position feeds keyword proximity and the flagging rules like an English name.
    for text, keywords in (("陈亚九被指控欺诈", _keywords_with("欺诈")),
                           ("陈亚九 was named in a fraud investigation", DEFAULT_KEYWORDS)):
        c = _candidate("https://news.example.com/a", title=text)
        c.entity_matches = match_entities(text, "", configured)
        c.keyword_matches = scan_keywords(text, keywords, [m.position for m in c.entity_matches])
        apply_flagging_rules(c, 15)
        assert c.priority == Priority.HIGH and c.rules_applied == ["R4"], (text, c.priority, c.rules_applied)
        proximity = c.keyword_matches[0].proximity_words
        assert proximity is not None and proximity <= 15, (text, proximity)   # within the name window
    assert tokenize_words("Tan 陈亚九, fraud!") == ["tan", "陈", "亚", "九", "fraud"]


@selftest("S15")
def test_alias_config_validation_and_unchanged_sc_keys():
    assert DEFAULT_ENTITIES["entities"][0]["aliases"] == ["艾克米投资管理", "艾克米投資管理"]
    assert DEFAULT_ENTITIES["entities"][1]["aliases"] == [], "D-29: 陈亚九 is an ambiguous term, not an alias"
    assert [t["term"] for t in DEFAULT_ENTITIES["entities"][1]["ambiguous_terms"]] == ["陈亚九", "陳亞九"]
    _validate_entities(_alias_entities(tan_ah_kow=["陈亚九"]))
    for bad in ([""], "陈亚九", ["！！"]):
        entities = _alias_entities()
        entities["entities"][1]["aliases"] = bad
        try:
            _validate_entities(entities)
        except ConfigError:
            continue
        raise AssertionError(f"aliases {bad!r} must be rejected")
    assert "陈亚九" in screened_targets(_alias_entities(tan_ah_kow=["陈亚九"]))[1]["names"]
    assert "陈亚九" in entity_screening_terms(_alias_entities(tan_ah_kow=["陈亚九"]))
    entry = {"name": "艾克米投资管理 Clone Fund", "website": "https://clone.example/"}
    assert sc_alert_content_key(entry) == "sc-alert-list:clone fund | https clone example", \
        "SC identity keys keep the ASCII-only rule, so stored entries are not reported again as new"


# The five relationship cases (client spec, 1 Oct 2026). "Lim Ah Kow" / 陈亚强 are other people.
_RELATIONSHIP_CASES = {
    "case1_company": ["Acme Asset Management was named in a fraud complaint."],
    "case2_founder": ["Tan Ah Kow was named in a fraud probe.", "Tan Ah Kow（陈亚九）被指控欺诈"],
    "case3_linked": ["Tan Ah Kow of Acme Asset Management was named in a fraud probe.",
                     "陈亚九与艾克米投资管理涉嫌欺诈"],
    "case4_other_person_and_company": ["Lim Ah Kow, a director at Acme Asset Management, was named in a fraud probe.",
                                       "李明在艾克米投资管理涉嫌欺诈"],
    "case5_other_person_only": ["Lim Ah Kow was named in a fraud probe.", "陈亚强涉嫌欺诈"],
    # D-29 (owner, 5 Oct 2026): 陈亚九 is a common name, so on its own it may be someone else.
    "case6_common_chinese_name_alone": ["陈亚九被指控欺诈", "陳亞九被指控欺詐"],
}


def _screen_case(text: str, entities: dict | None = None, url: str = "https://news.example.com/a") -> Candidate:
    """Pipeline steps 5–6 (S15–S17) for one item, with the default approved entities and
    relationships and the default keywords (which include the Chinese forms such as 欺诈, D-28)."""
    entities = entities if entities is not None else DEFAULT_ENTITIES
    c = _candidate(url, title=text)
    c.content_class, c.verification = ContentClass.PUBLIC_WEB, Verification.VERIFIED
    c.content_date, c.date_status = _myt(2026, 9, 16), DateStatus.VERIFIED
    c.entity_matches = match_entities(text, "", entities)
    if c.entity_matches:
        c.keyword_matches = scan_keywords(text, _keywords_with("欺诈"), [m.position for m in c.entity_matches])
    apply_flagging_rules(c, 15, approved_relationships(entities))
    c.section = assign_section(c, _period("2026-09-22"))
    return c


@selftest("S17")
def test_entity_relationship_cases():
    tan, company = "tan_ah_kow", "acme_company"
    for case, texts in _RELATIONSHIP_CASES.items():
        for text in texts:
            c = _screen_case(text)
            ids = {m.entity_id for m in c.entity_matches}
            link = describe_entity_link(ids, DEFAULT_ENTITIES)
            if case == "case1_company":
                assert ids == {company} and link["type"] == "Company finding" and not link["associated"], (text, link)
                assert (c.priority, c.review_status, "R8" in c.rules_applied) == (Priority.HIGH, ReviewStatus.PENDING, False)
            elif case == "case2_founder":
                assert ids == {tan} and link["type"] == "Founder finding" and not link["relationship"], (text, link)
                assert (c.priority, c.review_status) == (Priority.HIGH, ReviewStatus.PENDING), text
            elif case == "case3_linked":
                assert ids == {tan, company} and link["type"] == "Relationship-linked finding", (text, link)
                assert link["associated"] == "Acme Asset Management Sdn Bhd"
                assert link["relationship"].startswith(
                    "Tan Ah Kow → Founder of → Acme Asset Management Sdn Bhd (approved relationship)")
                assert "needs human review" in link["relationship"], "never claims Tan Ah Kow is responsible"
                assert c.priority == Priority.HIGH, "R8 never raises the priority"
                assert c.review_status == ReviewStatus.NEEDS_REVIEW and "R8" in c.rules_applied, (text, c.rules_applied)
            elif case == "case4_other_person_and_company":
                assert ids == {company}, f"the other person must not become {tan}: {text} -> {ids}"
                assert link["type"] == "Company finding" and not link["associated"] and not link["relationship"]
                assert c.review_status == ReviewStatus.PENDING and "R8" not in c.rules_applied
            else:
                assert ids == set() and c.priority is None and c.section == Section.NOT_REPORTED, (text, ids)
            assert all(m.entity_id == "" for m in c.keyword_matches), "keywords are not attributed to a person"
    # The relationship is used only when it is listed and approved: never inferred from the text.
    unlisted = json.loads(json.dumps(DEFAULT_ENTITIES))
    unlisted["relationships"] = []
    not_approved = json.loads(json.dumps(DEFAULT_ENTITIES))
    not_approved["relationships"][0]["approved"] = False
    for entities in (unlisted, not_approved):
        c = _screen_case(_RELATIONSHIP_CASES["case3_linked"][0], entities)
        link = describe_entity_link({m.entity_id for m in c.entity_matches}, entities)
        assert link["type"] == "Several entities — no approved relationship" and not link["relationship"]
        assert "R8" not in c.rules_applied and c.review_status == ReviewStatus.PENDING
    bad = json.loads(json.dumps(DEFAULT_ENTITIES))
    bad["relationships"][0]["object"] = "someone_unknown"
    try:
        _validate_entities(bad)
        raise AssertionError("a relationship to an unknown entity id must be rejected")
    except ConfigError:
        pass


@selftest("S18")
def test_report_relationship_columns():
    cands, expected = [], {}
    for case, texts in _RELATIONSHIP_CASES.items():
        url = f"https://news.example.com/{case}"
        c = _screen_case(texts[-1], url=url)              # the Chinese wording (English for case 1)
        assert c.section is not None
        r = _report_candidate(case, url, c.section, c.priority, title=texts[-1], snippet="", when=_myt(2026, 9, 16))
        r.entity_matches, r.keyword_matches = c.entity_matches, c.keyword_matches
        r.review_status, r.rules_applied = c.review_status, c.rules_applied
        cands.append(r)
        expected[url] = case
    with tempfile.TemporaryDirectory() as d:
        path, _, _, conn = _report_workbook(d, cands)
        conn.close()
        from openpyxl import load_workbook
        wb = load_workbook(path)                              # read back from disk
        assert [c.value for c in wb[SHEET_FINDINGS][1]] == FINDINGS_COLUMNS
        for name in ("Matched Entity", "Matched Name / Alias", "Associated Entity", "Relationship", "Keyword", "URL"):
            assert name in FINDINGS_COLUMNS, name
        rows = {r["URL"]: r for r in _sheet_values(wb[SHEET_FINDINGS])}
        case = lambda n: rows[f"https://news.example.com/{n}"]     # noqa: E731
        assert set(rows) == {u for u, n in expected.items()
                             if n not in ("case5_other_person_only", "case6_common_chinese_name_alone")}, set(rows)
        c1, c2, c3, c4 = (case("case1_company"), case("case2_founder"), case("case3_linked"),
                          case("case4_other_person_and_company"))
        assert (c1["Finding Type"], c1["Matched Entity"], c1["Associated Entity"]) == \
            ("Company finding", "Acme Asset Management Sdn Bhd", None)
        assert (c2["Finding Type"], c2["Matched Entity"], c2["Relationship"]) == \
            ("Founder finding", "Tan Ah Kow", None)
        assert c2["Matched Name / Alias"] == "Tan Ah Kow; 陈亚九 (with context 'Tan Ah Kow')", c2["Matched Name / Alias"]
        assert c3["Finding Type"] == "Relationship-linked finding"
        assert c3["Matched Name / Alias"] == ("Acme Asset Management Sdn Bhd: 艾克米投资管理; "
                                              "Tan Ah Kow: 陈亚九 (with context '艾克米投资管理')"), c3["Matched Name / Alias"]
        assert c3["Associated Entity"] == "Acme Asset Management Sdn Bhd" and "needs human review" in c3["Relationship"]
        assert c3["Manual Review Status"] == "NEEDS REVIEW" and c3["Keyword"] == "fraud (欺诈)", c3["Keyword"]
        assert (c4["Matched Entity"], c4["Matched Name / Alias"], c4["Associated Entity"]) == \
            ("Acme Asset Management Sdn Bhd", "艾克米投资管理", None), "the other person is not Tan Ah Kow"
        summary = " ".join(str(v) for row in wb[SHEET_SUMMARY].iter_rows(values_only=True) for v in row if v)
        assert "Tan Ah Kow → Founder of → Acme Asset Management Sdn Bhd (associated_with)" in summary


@selftest("S18")
def test_chinese_alias_reported_with_english_entity_name():
    title = "陈亚九被指控欺诈"
    snippet = "艾克米投资管理与陈亚九的声明（2026年9月16日）"
    keywords = _keywords_with("欺诈")
    configured = _alias_entities(tan_ah_kow=["陈亚九"], acme_company=["艾克米投资管理"])
    person = _report_candidate("k-zh", "https://news.example.com/zh/%E5%BC%A0", Section.MAIN, Priority.HIGH,
                               title=title, snippet="欺诈调查", when=_myt(2026, 9, 16))
    both = _report_candidate("k-both", "https://news.example.com/zh/both", Section.MAIN, Priority.HIGH,
                             title=title, snippet=snippet, when=_myt(2026, 9, 17))
    for c in (person, both):
        text = f"{c.raw.title} {c.raw.snippet}"
        c.entity_matches = match_entities(text, "", configured)
        c.keyword_matches = scan_keywords(text, keywords, [m.position for m in c.entity_matches])
    with tempfile.TemporaryDirectory() as d:
        path, _, _, conn = _report_workbook(d, [person, both])
        try:
            stored = conn.execute("SELECT title, snippet FROM findings WHERE url_key = 'k-zh'").fetchone()
            assert tuple(stored) == (title, "欺诈调查"), stored
            matched = {r[0] for r in conn.execute("SELECT matched_text FROM finding_matches WHERE keyword IS NULL")}
            assert matched == {"陈亚九", "艾克米投资管理"}, matched
        finally:
            conn.close()
        from openpyxl import load_workbook
        rows = {r["URL"]: r for r in _sheet_values(load_workbook(path)[SHEET_FINDINGS])}   # read back from disk
        zh = rows["https://news.example.com/zh/%E5%BC%A0"]
        assert zh["Matched Entity"] == "Tan Ah Kow" and zh["Matched Name / Alias"] == "陈亚九", zh
        assert zh["Title / Comment / Snippet"] == f"{title} — 欺诈调查" and zh["Keyword"] == "fraud (欺诈)", zh["Keyword"]
        two = rows["https://news.example.com/zh/both"]
        assert two["Matched Entity"] == "Acme Asset Management Sdn Bhd, Tan Ah Kow"
        assert two["Matched Name / Alias"] == "Acme Asset Management Sdn Bhd: 艾克米投资管理; Tan Ah Kow: 陈亚九"
        assert snippet in two["Title / Comment / Snippet"]


@selftest("S18")
def test_no_link_reasons():
    assert link_or_reason("https://a.example/x") == "https://a.example/x"
    assert link_or_reason("") == link_or_reason(None) == NO_LINK_EMPTY
    assert link_or_reason("www.a.example").startswith(NO_LINK) and "www.a.example" in link_or_reason("www.a.example")
    web = "https://news.example.com/a"
    assert fetched_url_text("https://news.example.com/b", web) == "https://news.example.com/b"
    assert fetched_url_text(None, "https://www.instagram.com/p/x", via_alerts=True) == NOT_FETCHED_SOCIAL
    assert fetched_url_text(None, web, via_alerts=True, fetch_public_web=True) == NOT_FETCHED_UNREAD
    assert fetched_url_text(None, web, via_alerts=True, fetch_public_web=False) == NOT_FETCHED_SWITCHED_OFF
    assert fetched_url_text(None, web) == NOT_FETCHED_OTHER
    assert fetched_url_text(None, "").startswith(NO_LINK)

    def health(**kw):
        h = {"source": "SC Malaysia", "status": "SUCCESS", "retrieved": 0, "filtered_out": 0,
             "flagged": 0, "duplicates": 0, "rejected": 0, "error": None}
        h.update(kw)
        return source_link_status(h, 0, 0, False)
    assert source_link_status({"status": "SUCCESS"}, 2, 2, False) == "2 link(s) in this report"
    assert "none has a web address" in source_link_status({"status": "SUCCESS"}, 0, 1, False)
    expected = {
        health(status="FAILED", error="timeout"): "the source failed — timeout",
        health(status="SKIPPED", error="not run (disabled or not selected)"): "the source was not run",
        health(error="Warnings: no Google Alerts emails arrived"): "nothing was retrieved — Warnings: no Google Alerts",
        health(): "nothing was retrieved in this run",
        health(retrieved=3, filtered_out=3): "3 item(s) retrieved, but none named a screened target",
        health(retrieved=4, duplicates=1, filtered_out=3): "4 item(s) retrieved, but none named a screened target",
        health(retrieved=3, filtered_out=2): "1 named a screened target, but none matched the adverse-keyword",
        health(retrieved=3, filtered_out=1, flagged=1): "1 item(s) flagged, but none is in this report's sheets",
    }
    for text, needle in expected.items():
        assert text.startswith(NO_LINK) and needle in text, (needle, text)


@selftest("S18")
def test_report_explains_missing_links():
    with tempfile.TemporaryDirectory() as d:
        sc = _report_candidate("k-sc", "https://www.sc.com.my/investor-alert-list", Section.MAIN, Priority.CRITICAL,
                               when=_myt(2026, 9, 16))
        sc.platform = Platform.SC_MALAYSIA
        blank = _report_candidate("k-none", "", Section.MAIN, Priority.HIGH, when=_myt(2026, 9, 17))
        db = Path(d) / "t.sqlite"
        conn = db_connect(db)
        ctx = _test_ctx(db)
        ctx.config.settings["paths"]["reports_dir"] = str(Path(d) / "reports")
        unmatched = _candidate("https://acme.example.com/a", source="website")     # names no target
        health = derive_source_health(
            [SourceResult("google_alerts", SourceStatus.SUCCESS,
                          warnings=["no Google Alerts emails arrived in the 'NS-Alerts' label"]),
             SourceResult("website", SourceStatus.SUCCESS, items=[unmatched.raw]),
             SourceResult("sc_malaysia", SourceStatus.SUCCESS, items=[sc.raw])], [sc, unmatched])
        try:
            db_record_run_start(conn, ctx)
            for c in (sc, blank):
                db_upsert_candidate(conn, c, ctx)
            from openpyxl import load_workbook
            wb = load_workbook(generate_excel_report(conn, ctx, health, {}))
            links = {r[0]: r[1] for r in wb[SHEET_SUMMARY].iter_rows(values_only=True) if r[0]}
            assert links["Clickable links (URL column)"] == 1, links["Clickable links (URL column)"]
            assert links["SC Malaysia"] == "1 link(s) in this report"
            assert str(links["Facebook"]).startswith(NO_LINK + "nothing was retrieved — No Google Alerts emails arrived")
            assert links["Google Alerts"] == links["Public website"], "Google Alerts covers the Public website item"
            assert links["Company website"] == NO_LINK + "1 item(s) retrieved, but none named a screened target"
            assert links["Public website"] == NO_LINK + "1 item(s) listed, but none has a web address — " \
                                                        "see the URL column for each item's reason"
            rows = {r["Platform"]: r for r in _sheet_values(wb[SHEET_FINDINGS])}
            assert rows["Public website"]["URL"] == rows["Public website"]["Discovered URL"] == NO_LINK_EMPTY
            assert rows["Public website"]["Fetched URL"].startswith(NO_LINK)
            ws = wb[SHEET_FINDINGS]
            url_col = FINDINGS_COLUMNS.index("URL") + 1
            linked = {ws.cell(row=r, column=FINDINGS_COLUMNS.index("Platform") + 1).value:
                      ws.cell(row=r, column=url_col).hyperlink for r in (2, 3)}
            assert linked["SC Malaysia"] is not None and linked["Public website"] is None
            assert wb[SHEET_DATE_UNKNOWN].cell(row=3, column=1).value == NO_ITEMS_IN_SHEET
        finally:
            conn.close()


if __name__ == "__main__":
    sys.exit(main())
